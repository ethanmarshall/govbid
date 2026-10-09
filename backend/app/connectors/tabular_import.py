"""Import opportunities from files you download yourself.

DIBBS blocks automated access, so DIBBS RFQs come in by file:
  - a CSV/XLSX export of RFQ search results, or
  - the daily RFQ index text file (parsed best-effort by pattern matching).
Agency forecasts (Acquisition Gateway, agency OSDBU pages) come in as CSV/XLSX.

Headers are matched loosely, so most exports work without editing.
"""
from __future__ import annotations

import csv
import io
import re
from pathlib import Path

ALIASES: dict[str, list[str]] = {
    "solicitation_number": ["solicitation", "solicitation number", "solicitation #", "sol #", "rfq", "rfq number", "contract number", "listing id", "id"],
    "title": ["title", "requirement", "requirement title", "nomenclature", "item", "item description", "description of requirement", "project title"],
    "description": ["description", "requirement description", "details", "summary"],
    "agency": ["agency", "organization", "department", "buying office", "contracting office", "office", "buyer"],
    "naics": ["naics", "naics code", "primary naics"],
    "psc": ["psc", "psc code", "product service code", "fsc"],
    "nsn": ["nsn", "national stock number", "niin"],
    "quantity": ["qty", "quantity"],
    "set_aside_desc": ["set aside", "set-aside", "setaside", "small business set aside", "acquisition strategy", "socioeconomic"],
    "response_deadline": ["return by", "return by date", "response date", "due date", "closing date", "response deadline", "anticipated solicitation date", "estimated solicitation release", "solicitation date"],
    "posted_date": ["issue date", "posted", "posted date", "date posted", "publish date"],
    "estimated_value": ["estimated value", "value", "dollar range", "estimated contract value", "contract value", "estimated award amount"],
    "place_of_performance": ["place of performance", "location", "pop"],
    "url": ["url", "link", "record url"],
}

# Text description -> SAM.gov set-aside code
SET_ASIDE_TEXT = [
    (r"service[- ]disabled|sdvosb|sdvo", "SDVOSBC"),
    (r"\bedwosb\b|economically disadvantaged women", "EDWOSB"),
    (r"women[- ]owned|\bwosb\b", "WOSB"),
    (r"hubzone|hub zone", "HZC"),
    (r"8\s*\(a\)|\b8a\b", "8A"),
    (r"veteran[- ]owned|\bvosb\b", "VSA"),
    (r"partial small business", "SBP"),
    (r"small business|total small|\bsb\b set", "SBA"),
    (r"unrestricted|full and open|none|n/a", ""),
]


def set_aside_code_from_text(text: str) -> str:
    t = (text or "").strip().lower()
    if not t:
        return ""
    for pattern, code in SET_ASIDE_TEXT:
        if re.search(pattern, t):
            if code == "SDVOSBC" and "sole" in t:
                return "SDVOSBS"
            return code
    return ""


def _norm(h: str) -> str:
    return re.sub(r"[^a-z0-9#() ]+", " ", (h or "").lower()).strip()


def _map_headers(headers: list[str]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    normed = [_norm(h) for h in headers]
    for field, aliases in ALIASES.items():
        for alias in aliases:
            for i, h in enumerate(normed):
                if i not in mapping and h == alias:
                    mapping[i] = field
                    break
            else:
                continue
            break
    # Second pass: substring matches for anything still unmapped
    for field, aliases in ALIASES.items():
        if field in mapping.values():
            continue
        for i, h in enumerate(normed):
            if i in mapping:
                continue
            if any(len(a) > 3 and a in h for a in aliases):
                mapping[i] = field
                break
    return mapping


def _rows_from_file(path: Path) -> list[list[str]]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        return [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
    text = path.read_text(errors="replace")
    first_line = text.splitlines()[0] if text.strip() else ""
    dialect = csv.excel_tab if "\t" in first_line else csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _parse_money(v: str) -> float | None:
    m = re.findall(r"[\d,.]+", v or "")
    nums = []
    for x in m:
        try:
            nums.append(float(x.replace(",", "")))
        except ValueError:
            pass
    if not nums:
        return None
    low = (v or "").lower()
    if re.search(r"\d\s*(m|mm|mil|million)\b", low):
        mult = 1_000_000
    elif re.search(r"\d\s*(b|billion)\b", low):
        mult = 1_000_000_000
    elif re.search(r"\d\s*(k|thousand)\b", low):
        mult = 1000
    else:
        mult = 1
    return max(nums) * mult


def parse_table(path: Path, source: str) -> list[dict]:
    rows = _rows_from_file(path)
    # Find the header row: first row with at least 2 recognizable headers
    header_idx, mapping = 0, {}
    for i, row in enumerate(rows[:15]):
        m = _map_headers(row)
        if len(m) >= 2:
            header_idx, mapping = i, m
            break
    if not mapping:
        return []
    out = []
    for n, row in enumerate(rows[header_idx + 1 :]):
        if not any(c.strip() for c in row):
            continue
        rec: dict = {}
        for idx, field in mapping.items():
            if idx < len(row):
                rec[field] = row[idx].strip()
        if not (rec.get("title") or rec.get("solicitation_number") or rec.get("nsn")):
            continue
        out.append(_finish(rec, source, n))
    return out


SOL_RE = re.compile(r"\b(SP[0-9A-Z]{4}-?\d{2}-?[A-Z]-?[0-9A-Z]{4})\b")
NSN_RE = re.compile(r"\b(\d{4}-?\d{2}-?\d{3}-?\d{4})\b")
DATE_RE = re.compile(r"\b(\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}-\d{2}-\d{2})\b")


def parse_dibbs_index(path: Path) -> list[dict]:
    """Best-effort parser for the DIBBS daily RFQ index text file.

    The file is fixed-width and its exact layout is not published, so this pulls out
    the solicitation number, NSN and dates by pattern and keeps the remaining text
    as the item name. Always spot-check results against DIBBS.
    """
    out = []
    for n, line in enumerate(path.read_text(errors="replace").splitlines()):
        sol = SOL_RE.search(line)
        if not sol:
            continue
        nsn = NSN_RE.search(line)
        dates = DATE_RE.findall(line)
        rest = line
        for frag in filter(None, [sol.group(1), nsn.group(1) if nsn else None, *dates]):
            rest = rest.replace(frag, " ")
        name = re.sub(r"\s{2,}", " ", re.sub(r"[^A-Za-z0-9,()/&.\- ]", " ", rest)).strip()
        name = re.sub(r"^\d+\s+|\s+\d+$", "", name)  # strip stray leading/trailing numbers (qty, PR)
        rec = {
            "solicitation_number": sol.group(1),
            "nsn": nsn.group(1) if nsn else "",
            "title": name[:200] or sol.group(1),
            "response_deadline": dates[-1] if dates else "",
            "posted_date": dates[0] if len(dates) > 1 else "",
            "agency": "DLA",
        }
        out.append(_finish(rec, "dibbs", n))
    return out


def _format_nsn(nsn: str) -> str:
    digits = re.sub(r"\D", "", nsn or "")
    if len(digits) == 13:
        return f"{digits[:4]}-{digits[4:6]}-{digits[6:9]}-{digits[9:]}"
    return nsn or ""


def _finish(rec: dict, source: str, n: int) -> dict:
    nsn = _format_nsn(rec.get("nsn", ""))
    sol = rec.get("solicitation_number", "")
    sa_text = rec.get("set_aside_desc", "")
    sol_compact = sol.replace("-", "")
    url = rec.get("url", "")
    if source == "dibbs" and sol and not url:
        url = f"https://www.dibbs.bsm.dla.mil/rfq/rfqrec.aspx?sn={sol_compact}"
    psc = rec.get("psc", "") or (nsn[:4] if nsn else "")
    return {
        "source": source,
        "external_id": sol_compact or f"{source}-{nsn or n}-{rec.get('title', '')[:40]}",
        "solicitation_number": sol,
        "title": rec.get("title") or rec.get("description", "")[:200] or sol,
        "agency": rec.get("agency") or ("DLA" if source == "dibbs" else ""),
        "notice_type": {"dibbs": "RFQ", "forecast": "Forecast"}.get(source, "Imported"),
        "set_aside_code": set_aside_code_from_text(sa_text),
        "set_aside_desc": sa_text,
        "naics": re.sub(r"\D", "", rec.get("naics", ""))[:6],
        "psc": psc,
        "nsn": nsn,
        "quantity": rec.get("quantity", ""),
        "posted_date": rec.get("posted_date", ""),
        "response_deadline": rec.get("response_deadline", ""),
        "place_of_performance": rec.get("place_of_performance", ""),
        "url": url,
        "description": rec.get("description", ""),
        "estimated_value": _parse_money(rec.get("estimated_value", "")),
        "raw": rec,
    }


def filter_watchlist(records: list[dict], watchlist: list[str]) -> list[dict]:
    """Keep records whose NSN or FSC starts with any watchlist entry. Empty watchlist keeps everything."""
    keys = [re.sub(r"\D", "", w) for w in watchlist if w and re.sub(r"\D", "", w)]
    if not keys:
        return records
    kept = []
    for r in records:
        digits = re.sub(r"\D", "", r.get("nsn", "")) or re.sub(r"\D", "", r.get("psc", ""))
        if digits and any(digits.startswith(k) for k in keys):
            kept.append(r)
    return kept
