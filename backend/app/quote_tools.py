"""Quote tools for saved part quotes: customer quote documents, vendor RFQs and win/loss insights.

Customer quote: a PDF or Word document built from the saved quote result (prices and lead times only: internal
cost and margin never appear). Each part quote gets a customer quote number Q-YYYY-NNNN the first time one is made.

Vendor RFQ: one record per vendor with an email draft, a mailto: link (mailto cannot carry attachments) and a zip
package of the STEP model, drawing PDF and cut list or BOM. If the drawing or the solicitation is marked export
controlled or carries distribution statement B through F, files are left out of the package and the draft says why.
A recorded vendor response becomes a VendorQuote, so make-or-buy picks it up.

Win/loss insights: groups quotes marked won, lost or submitted by kind, process, material and FSC and reports the
win rate, our price against the award price (when we lost and the NSN history has an award price), and margins.
"""
from __future__ import annotations

import csv
import io
import json
import re
import statistics
import urllib.parse
import zipfile
from datetime import date, datetime, timedelta
from xml.sax.saxutils import escape

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .models import PartQuote
from .models_crm import Organization, VendorQuote
from .models_quote_tools import RFQ_STATUSES, CustomerQuoteDoc, QuoteToolSettings, VendorRFQ
from .services import get_profile


class QuoteToolError(ValueError):
    pass


class NotFound(LookupError):
    pass


PACKAGING_TEXT = {
    "commercial": "Commercial packaging",
    "mil_std_2073": "Military packaging per MIL-STD-2073-1, codes per the solicitation",
}
SETTINGS_FIELDS = ("validity_days", "payment_terms", "fob", "shipping", "inspection_acceptance", "address", "footer_text")


def get_quote(db: Session, quote_id: int) -> PartQuote:
    pq = db.get(PartQuote, quote_id)
    if not pq:
        raise NotFound(f"Quote {quote_id} not found")
    return pq


# ---------------------------------------------------------------- settings
def get_settings(db: Session) -> QuoteToolSettings:
    s = db.get(QuoteToolSettings, 1)
    if s is None:
        s = QuoteToolSettings(id=1)
        db.add(s)
        db.commit()
    return s


def settings_dict(s: QuoteToolSettings) -> dict:
    return {k: getattr(s, k) for k in SETTINGS_FIELDS}


def update_settings(db: Session, changes: dict) -> dict:
    s = get_settings(db)
    for k, v in (changes or {}).items():
        if k not in SETTINGS_FIELDS or v is None:
            continue
        if k == "validity_days":
            try:
                v = int(v)
            except (TypeError, ValueError):
                raise QuoteToolError("validity_days must be a whole number")
            if not 1 <= v <= 365:
                raise QuoteToolError("validity_days must be between 1 and 365")
        else:
            v = str(v)
        setattr(s, k, v)
    db.commit()
    return settings_dict(s)


# ---------------------------------------------------------------- quote numbers
def ensure_number(db: Session, pq: PartQuote) -> CustomerQuoteDoc:
    doc = db.scalars(select(CustomerQuoteDoc).where(CustomerQuoteDoc.part_quote_id == pq.id)).first()
    if doc and pq.created_at and doc.created_at and doc.created_at < pq.created_at:
        st = get_settings(db)  # left over from a deleted quote whose id SQLite reused; retire its number
        if (doc.year, doc.seq) > (st.last_number_year or 0, st.last_number_seq or 0):
            st.last_number_year, st.last_number_seq = doc.year, doc.seq
        db.delete(doc)
        db.flush()
        doc = None
    if doc:
        return doc
    year = date.today().year
    s = get_settings(db)
    issued = (s.last_number_seq or 0) if s.last_number_year == year else 0
    seq = max(db.scalar(select(func.max(CustomerQuoteDoc.seq)).where(CustomerQuoteDoc.year == year)) or 0, issued) + 1
    doc = CustomerQuoteDoc(part_quote_id=pq.id, number=f"Q-{year}-{seq:04d}", year=year, seq=seq, options={})
    s.last_number_year, s.last_number_seq = year, seq
    db.add(doc)
    db.commit()
    return doc


# ---------------------------------------------------------------- shared part facts
def _revision(spec: dict) -> str:
    return spec.get("revision") or (spec.get("drawing") or {}).get("revision") or ""


def _drawing_id(spec: dict) -> str:
    return ((spec.get("drawing") or {}).get("drawing_id")
            or ((spec.get("cad") or {}).get("options") or {}).get("drawing_id")
            or (spec.get("options") or {}).get("drawing_id")
            or spec.get("drawing_id") or "")


def _kind(pq: PartQuote) -> str:
    from .quotes import quote_dict
    return quote_dict(pq, full=False)["kind"]


def _process(spec: dict) -> str:
    if spec.get("process"):
        return str(spec["process"])
    ops = spec.get("operations") or []
    if ops and isinstance(ops[0], dict) and ops[0].get("type"):
        return str(ops[0]["type"])
    return spec.get("kind") or ""


def _finishes(spec: dict) -> list[str]:
    out = []
    for f in spec.get("finishes") or []:
        name = f.get("type") if isinstance(f, dict) else f
        if name:
            out.append(str(name))
    return out


def part_lines(pq: PartQuote) -> list[str]:
    """Customer-safe description lines for the part."""
    spec = pq.spec or {}
    lines = [pq.name or spec.get("name") or "Part"]
    ident = []
    if pq.part_number:
        ident.append(f"P/N {pq.part_number}" + (f" Rev {_revision(spec)}" if _revision(spec) else ""))
    if pq.nsn:
        ident.append(f"NSN {pq.nsn}")
    if ident:
        lines.append(", ".join(ident))
    if spec.get("kind") == "box_build":
        lines.append(box_summary(spec))
    mat = spec.get("material")
    fin = _finishes(spec)
    if mat or fin:
        lines.append(", ".join(x for x in [f"Material: {mat}" if mat else "", f"Finish: {', '.join(fin)}" if fin else ""] if x))
    return lines


def box_summary(spec: dict) -> str:
    """One customer-safe line saying what a box build includes (no costs)."""
    parts = []
    enc = spec.get("enclosure") or {}
    if enc.get("source") in ("catalog", "custom"):
        parts.append("enclosure")
    elif enc.get("source") == "customer":
        parts.append("customer-furnished enclosure")
    boards = sum(float(b.get("qty_per") or 1) for b in spec.get("pcbs") or [])
    if boards:
        parts.append(f"{boards:g} circuit card assembl{'y' if boards == 1 else 'ies'}")
    n = sum(float(x.get("qty") or 0) for x in spec.get("lines") or [] if x.get("type") != "hardware")
    if n:
        parts.append(f"{n:g} panel and internal components")
    if any((c.get("spec") or {}).get("kind") == "harness" for c in spec.get("children") or []) or (spec.get("wiring") or {}).get("wires"):
        parts.append("internal wiring")
    if spec.get("peripherals"):
        parts.append(f"{len(spec['peripherals'])} peripheral item(s)")
    lab = spec.get("labor") or {}
    tests = [t for k, t in (("functional_test_minutes", "functional test"), ("burn_in_hours", "burn-in"), ("hipot", "safety test")) if lab.get(k)]
    if tests:
        parts.append(", ".join(tests))
    return ("Assembled unit: " + ", ".join(parts)) if parts else "Assembled unit"


# ---------------------------------------------------------------- customer quote
def customer_quote_data(db: Session, quote_id: int, opts: dict | None = None, remember: bool = True) -> dict:
    """Everything the customer quote shows. `opts` overrides the saved defaults and is remembered for next time."""
    pq = get_quote(db, quote_id)
    spec = pq.spec or {}
    breaks = (pq.result or {}).get("price_breaks") or []
    if not breaks:
        raise QuoteToolError("This quote has no prices yet. Save it with at least one quantity.")
    s = get_settings(db)
    doc = ensure_number(db, pq)
    o = {**(doc.options or {}), **{k: v for k, v in (opts or {}).items() if v not in (None, "")}}
    try:
        validity = int(o.get("validity_days") or s.validity_days or 30)
    except (TypeError, ValueError):
        raise QuoteToolError("validity_days must be a whole number")
    wanted = o.get("quantities") or []
    if isinstance(wanted, str):
        wanted = [x for x in re.split(r"[,\s]+", wanted) if x]
    try:
        wanted = {int(x) for x in wanted}
    except (TypeError, ValueError):
        raise QuoteToolError("quantities must be whole numbers")
    rows = [b for b in breaks if not wanted or b["quantity"] in wanted]
    if not rows:
        raise QuoteToolError(f"None of those quantities are in the quote. Available: {', '.join(str(b['quantity']) for b in breaks)}")
    if remember:
        keep = {k: o.get(k) for k in ("customer_name", "attn", "validity_days", "fob", "payment_terms", "notes", "shipping",
                                      "inspection_acceptance", "packaging")}
        keep["quantities"] = sorted(wanted)
        doc.options = keep
        db.commit()

    profile = get_profile(db)
    from .models_writing import CapabilitySettings
    cap = db.get(CapabilitySettings, 1)
    certs = profile.certifications or {}
    status_line = []
    if certs.get("SDVOSB") == "certified":
        status_line.append("Service-Disabled Veteran-Owned Small Business (SBA VetCert)")
    elif certs.get("VOSB") == "certified":
        status_line.append("Veteran-Owned Small Business (SBA VetCert)")
    elif certs.get("SB") == "certified":
        status_line.append("Small Business")
    opp = pq.opportunity
    today = date.today()
    pk_level = (spec.get("packaging") or {}).get("level") or ""
    return {
        "number": doc.number,
        "date": today.isoformat(),
        "valid_until": (today + timedelta(days=validity)).isoformat(),
        "validity_days": validity,
        "company": {
            "name": profile.name or "", "uei": profile.uei or "", "cage": profile.cage or "",
            "address": [ln.strip() for ln in (s.address or "").splitlines() if ln.strip()],
            "phone": cap.contact_phone if cap else "", "email": cap.contact_email if cap else "",
            "website": cap.website if cap else "", "contact_name": cap.contact_name if cap else "",
            "contact_title": cap.contact_title if cap else "", "status": status_line,
        },
        "customer": {"name": o.get("customer_name") or (opp.agency if opp else "") or "", "attn": o.get("attn") or ""},
        "reference": {
            "solicitation_number": opp.solicitation_number if opp else "",
            "opportunity_title": opp.title if opp else "",
            "nsn": pq.nsn or "", "part_number": pq.part_number or "", "revision": _revision(spec),
        },
        "description": part_lines(pq),
        "lines": [{"quantity": b["quantity"], "unit_price": b["unit_price"],
                   "extended_price": round(b.get("total_price") if b.get("total_price") is not None else b["unit_price"] * b["quantity"], 2),
                   "lead_time": f"{b['lead_time_days']} days ARO" if b.get("lead_time_days") else "To be confirmed"}
                  for b in rows],
        "terms": {
            "fob": o.get("fob") or s.fob, "payment_terms": o.get("payment_terms") or s.payment_terms,
            "shipping": o.get("shipping") or s.shipping,
            "packaging": o.get("packaging") or PACKAGING_TEXT.get(pk_level, pk_level.replace("_", " ").capitalize() if pk_level else "Commercial packaging"),
            "inspection_acceptance": o.get("inspection_acceptance") or s.inspection_acceptance,
        },
        "notes": o.get("notes") or "",
        "footer": s.footer_text or "",
        "options": doc.options,
    }


def _filename(d: dict, ext: str) -> str:
    ref = re.sub(r"[^A-Za-z0-9._-]+", "_", d["reference"]["part_number"] or d["reference"]["nsn"] or "")[:40]
    return f"{d['number']}{('_' + ref) if ref else ''}.{ext}"


def _money(v: float) -> str:
    return f"${v:,.2f}"


def _term_rows(d: dict) -> list[tuple[str, str]]:
    t = d["terms"]
    return [("Quote valid", f"{d['validity_days']} days (until {d['valid_until']})"), ("FOB", t["fob"]),
            ("Payment terms", t["payment_terms"]), ("Shipping", t["shipping"]), ("Packaging", t["packaging"]),
            ("Inspection and acceptance", t["inspection_acceptance"])]


def _ref_rows(d: dict) -> list[tuple[str, str]]:
    r = d["reference"]
    rows = [("Solicitation", r["solicitation_number"]), ("NSN", r["nsn"]), ("Part number", r["part_number"]), ("Revision", r["revision"])]
    return [(k, v) for k, v in rows if v]


def render_pdf(d: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=0.75 * inch, rightMargin=0.75 * inch,
                            topMargin=0.6 * inch, bottomMargin=0.6 * inch, title=f"Quote {d['number']}")
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontSize=9.5, leading=12.5)
    small = ParagraphStyle("s", parent=body, fontSize=8.5, leading=11, textColor=colors.HexColor("#444444"))
    big = ParagraphStyle("h", parent=ss["Title"], fontSize=16, leading=19, alignment=0, spaceAfter=2)
    right = ParagraphStyle("r", parent=body, alignment=2)
    P = lambda t, st=body: Paragraph(escape(str(t)).replace("\n", "<br/>"), st)  # noqa: E731

    c = d["company"]
    left = [P(c["name"] or "Your company name (set it in Profile)", big)]
    for ln in c["address"]:
        left.append(P(ln, small))
    contact = " | ".join(x for x in [c["phone"], c["email"], c["website"]] if x)
    if contact:
        left.append(P(contact, small))
    ids = " | ".join(x for x in [f"UEI {c['uei']}" if c["uei"] else "", f"CAGE {c['cage']}" if c["cage"] else ""] if x)
    if ids:
        left.append(P(ids, small))
    for st in c["status"]:
        left.append(P(st, small))
    rt = [Paragraph("<b>QUOTATION</b>", ParagraphStyle("q", parent=right, fontSize=14, leading=18)),
          P(f"Quote no. {d['number']}", right), P(f"Date {d['date']}", right), P(f"Valid until {d['valid_until']}", right)]
    head = Table([[left, rt]], colWidths=[4.4 * inch, 2.6 * inch])
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, 0), 1, colors.HexColor("#333333")),
                              ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    story = [head, Spacer(1, 10)]

    cust = d["customer"]
    to_lines = [x for x in [cust["name"], f"Attn: {cust['attn']}" if cust["attn"] else ""] if x]
    ref = [f"{k}: {v}" for k, v in _ref_rows(d)]
    t = Table([[P("To", small), P("Reference", small)],
               [P("\n".join(to_lines) or "-"), P("\n".join(ref) or "-")]], colWidths=[3.5 * inch, 3.5 * inch])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("TEXTCOLOR", (0, 0), (-1, 0), colors.grey)]))
    story += [t, Spacer(1, 10)]

    story.append(Paragraph("<b>Item</b>", body))
    story.append(P("\n".join(d["description"])))
    story.append(Spacer(1, 6))
    rows = [["Quantity", "Unit price", "Extended price", "Lead time"]]
    for ln in d["lines"]:
        rows.append([f"{ln['quantity']:,}", _money(ln["unit_price"]), _money(ln["extended_price"]), ln["lead_time"]])
    tbl = Table(rows, colWidths=[1.3 * inch, 1.6 * inch, 1.8 * inch, 2.3 * inch])
    tbl.setStyle(TableStyle([
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 9.5), ("FONT", (0, 1), (-1, -1), "Helvetica", 9.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")), ("ALIGN", (0, 0), (2, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#bbbbbb")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story += [tbl, Spacer(1, 6), P("Each row is a separate price for that total quantity.", small), Spacer(1, 10)]

    terms = Table([[P(k, small), P(v)] for k, v in _term_rows(d)], colWidths=[1.9 * inch, 5.1 * inch])
    terms.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#dddddd"))]))
    story += [Paragraph("<b>Terms</b>", body), terms, Spacer(1, 10)]
    if d["notes"]:
        story += [Paragraph("<b>Notes</b>", body), P(d["notes"]), Spacer(1, 10)]

    sig = [P(f"{c['contact_name'] or ''}{', ' + c['contact_title'] if c['contact_title'] else ''}"), P(c["name"])]
    story += [Spacer(1, 14), P("Authorized signature: ______________________________    Date: ____________"), *sig]
    if d["footer"]:
        story += [Spacer(1, 14), P(d["footer"], small)]
    doc.build(story)
    return buf.getvalue()


def render_docx(d: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)
    c = d["company"]
    h = doc.add_paragraph()
    r = h.add_run(c["name"] or "Your company name (set it in Profile)")
    r.bold, r.font.size = True, Pt(16)
    for ln in c["address"] + [" | ".join(x for x in [c["phone"], c["email"], c["website"]] if x),
                              " | ".join(x for x in [f"UEI {c['uei']}" if c["uei"] else "", f"CAGE {c['cage']}" if c["cage"] else ""] if x)] + c["status"]:
        if ln:
            doc.add_paragraph(ln).paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rr = p.add_run(f"QUOTATION\nQuote no. {d['number']}\nDate {d['date']}\nValid until {d['valid_until']}")
    rr.bold = True

    cust = d["customer"]
    doc.add_heading("To", level=3)
    doc.add_paragraph("\n".join(x for x in [cust["name"], f"Attn: {cust['attn']}" if cust["attn"] else ""] if x) or "-")
    refs = _ref_rows(d)
    if refs:
        doc.add_heading("Reference", level=3)
        doc.add_paragraph("\n".join(f"{k}: {v}" for k, v in refs))
    doc.add_heading("Item", level=3)
    doc.add_paragraph("\n".join(d["description"]))
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, ["Quantity", "Unit price", "Extended price", "Lead time"]):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
    for ln in d["lines"]:
        cells = t.add_row().cells
        for cell, text in zip(cells, [f"{ln['quantity']:,}", _money(ln["unit_price"]), _money(ln["extended_price"]), ln["lead_time"]]):
            cell.text = text
    doc.add_paragraph("Each row is a separate price for that total quantity.")
    doc.add_heading("Terms", level=3)
    tt = doc.add_table(rows=0, cols=2)
    tt.style = "Table Grid"
    for k, v in _term_rows(d):
        cells = tt.add_row().cells
        cells[0].text, cells[1].text = k, v
    if d["notes"]:
        doc.add_heading("Notes", level=3)
        doc.add_paragraph(d["notes"])
    doc.add_paragraph("")
    doc.add_paragraph("Authorized signature: ______________________________    Date: ____________")
    who = f"{c['contact_name'] or ''}{', ' + c['contact_title'] if c['contact_title'] else ''}"
    if who:
        doc.add_paragraph(who)
    doc.add_paragraph(c["name"])
    if d["footer"]:
        doc.add_paragraph(d["footer"])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- export-control guard
def export_check(db: Session, pq: PartQuote) -> dict:
    """Is the technical data for this quote export controlled or limited distribution? Returns {flagged, reasons}."""
    spec = pq.spec or {}
    reasons: list[str] = []
    did = _drawing_id(spec)
    if did:
        from .drawings_api import DRAWING_DIR
        try:
            meta = json.loads((DRAWING_DIR / f"{did}.json").read_text())
        except (OSError, ValueError):
            meta = {}
        read = meta.get("read") or {}
        if read.get("export_controlled"):
            reasons.append("The drawing carries an export-control (ITAR/EAR) marking.")
        letter = (read.get("distribution") or {}).get("letter") or ""
        if letter and letter != "A":
            reasons.append(f"The drawing is marked Distribution Statement {letter}.")
    dr = spec.get("drawing") or {}
    if dr.get("export_controlled") and not any("export-control" in r for r in reasons):
        reasons.append("The drawing carries an export-control (ITAR/EAR) marking.")
    letter = dr.get("distribution") if isinstance(dr.get("distribution"), str) else ""
    if letter and letter.upper() != "A" and not any("Distribution Statement" in r for r in reasons):
        reasons.append(f"The drawing is marked Distribution Statement {letter.upper()}.")
    if pq.opportunity is not None:
        from .bidding import export_flags_for
        flags = export_flags_for(pq.opportunity)
        if flags.get("flagged"):
            reasons.append("The solicitation is flagged export controlled" + (f" ({', '.join(flags['hits'][:4])})." if flags.get("hits") else "."))
    return {"flagged": bool(reasons), "reasons": reasons}


EXPORT_WARNING = ("Files were not included. Share controlled or limited-distribution technical data only with vendors "
                  "that hold an active JCP certification (DD Form 2345) and are authorized recipients, keep the export-control "
                  "and distribution markings on every copy, and send it through an approved channel.")


# ---------------------------------------------------------------- RFQ package
def _package_files(db: Session, pq: PartQuote) -> list[tuple[str, object]]:
    """(name, source) for the technical files of a quote. Source is a Path or bytes."""
    spec = pq.spec or {}
    out: list[tuple[str, object]] = []
    cad = spec.get("cad") or {}
    if cad.get("file_id"):
        try:
            from .cad_quote import step_path
            p = step_path(cad["file_id"])
            if p.exists():
                out.append((cad.get("filename") or "part.step", p))
        except Exception:  # noqa: BLE001  (a missing model only drops the file)
            pass
    did = _drawing_id(spec)
    if did and re.fullmatch(r"[0-9a-f]{32}", did):
        from .drawings_api import DRAWING_DIR
        p = DRAWING_DIR / f"{did}.pdf"
        if p.exists():
            name = (spec.get("drawing") or {}).get("filename") or "drawing.pdf"
            out.append((name if name.lower().endswith(".pdf") else name + ".pdf", p))
    if spec.get("kind") == "extrusion_build" and spec.get("lines"):
        try:
            from . import extrusion, quotes
            qtys = spec.get("quantities") or [1]
            out.append(("cut_list.xlsx", extrusion.cut_list_xlsx(spec["lines"], quotes.get_config(db), int(qtys[0]), pq.name or "")))
        except Exception:  # noqa: BLE001
            pass
    else:
        bom = spec.get("bom") if isinstance(spec.get("bom"), list) else spec.get("lines") if isinstance(spec.get("lines"), list) else None
        if bom and all(isinstance(r, dict) for r in bom):
            out.append(("bom.csv", _csv(bom)))
    return out


def _csv(rows: list[dict]) -> bytes:
    cols: list[str] = []
    for r in rows:
        for k, v in r.items():
            if k not in cols and not isinstance(v, (dict, list)):
                cols.append(k)
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=cols, extrasaction="ignore")
    w.writeheader()
    for r in rows:
        w.writerow({k: r.get(k, "") for k in cols})
    return buf.getvalue().encode()


def rfq_package(db: Session, rfq_id: int) -> tuple[str, bytes]:
    rfq = _rfq(db, rfq_id)
    pq = get_quote(db, rfq.part_quote_id)
    buf = io.BytesIO()
    names: list[str] = []
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("RFQ.txt", f"Subject: {rfq.subject}\n\n{rfq.body}\n")
        if rfq.include_files and not rfq.files_blocked:
            for name, src in _package_files(db, pq):
                base = name
                n = 2
                while name in names:
                    name = f"{n}_{base}"
                    n += 1
                names.append(name)
                if isinstance(src, (bytes, bytearray)):
                    z.writestr(name, bytes(src))
                else:
                    z.write(src, name)
        elif rfq.files_blocked:
            z.writestr("FILES_NOT_INCLUDED.txt", EXPORT_WARNING + "\n\n" + "\n".join(rfq.warnings or []) + "\n")
    slug = re.sub(r"[^A-Za-z0-9]+", "_", rfq.vendor_name or "vendor").strip("_")[:30]
    return f"RFQ_{pq.id}_{slug}.zip", buf.getvalue()


# ---------------------------------------------------------------- RFQ drafts
def requirements(pq: PartQuote, quantities: list[int]) -> list[str]:
    spec = pq.spec or {}
    req = []
    if spec.get("process"):
        req.append(f"Process: {str(spec['process']).replace('_', ' ')}")
    if spec.get("material"):
        req.append(f"Material: {spec['material']}" + (" (material certifications required)" if spec.get("material_certs_required") else ""))
    elif spec.get("material_certs_required"):
        req.append("Material certifications required")
    fin = _finishes(spec)
    if fin:
        req.append(f"Finish: {', '.join(fin)}")
    tol = spec.get("tolerance")
    if tol:
        req.append("Tolerances: per drawing" + (f" ({tol} tolerance work)" if tol != "standard" else ""))
    insp = spec.get("inspection") or {}
    if insp.get("first_article"):
        req.append("First article inspection and report required")
    if insp.get("certificate_of_conformance", True):
        req.append("Certificate of conformance required with each shipment")
    lvl = (spec.get("packaging") or {}).get("level")
    if lvl:
        req.append(f"Packaging: {PACKAGING_TEXT.get(lvl, lvl)}")
    if spec.get("kind") == "extrusion_build":
        req.append("Cut list and parts list attached")
    return req


def _first_contact(org: Organization | None) -> tuple[str, str]:
    if not org:
        return "", ""
    for c in org.contacts or []:
        if c.email:
            return c.name or "", c.email
    return ((org.contacts or [None])[0].name if org.contacts else ""), ""


def draft_email(db: Session, pq: PartQuote, vendor_name: str, contact_name: str, quantities: list[int], due_date: str,
                message: str, files_note: str) -> tuple[str, str]:
    spec = pq.spec or {}
    profile = get_profile(db)
    from .models_writing import CapabilitySettings
    cap = db.get(CapabilitySettings, 1)
    ident = ", ".join(x for x in [f"P/N {pq.part_number}" + (f" Rev {_revision(spec)}" if _revision(spec) else "") if pq.part_number else "",
                                   f"NSN {pq.nsn}" if pq.nsn else ""] if x)
    subject = f"RFQ: {pq.name or 'part'}" + (f" ({ident})" if ident else "") + (f", quote by {due_date}" if due_date else "")
    first = (contact_name or "").split(" ")[0]
    lines = [f"Hello {first}," if first else f"Hello {vendor_name} team,", ""]
    if message.strip():
        lines += [message.strip(), ""]
    lines.append("Please quote the following:")
    lines.append("")
    lines.append(f"Part: {pq.name or 'part'}" + (f", {ident}" if ident else ""))
    lines.append(f"Quantities: {', '.join(f'{q:,}' for q in quantities)} (unit price at each quantity)")
    for r in requirements(pq, quantities):
        lines.append(r)
    sol = pq.opportunity.solicitation_number if pq.opportunity else ""
    if sol:
        lines.append(f"Government solicitation: {sol}")
    if due_date:
        lines.append(f"Please reply by: {due_date}")
    lines += ["", "In your quote please include: unit price at each quantity, lead time in days after receipt of order, "
              "any tooling or setup charges, freight, how long the quote is valid, and your business size and whether "
              "you manufacture the part yourself (with the country of manufacture).", "", files_note, "", "Thank you,"]
    for x in [cap.contact_name if cap else "", cap.contact_title if cap else "", profile.name or "",
              cap.contact_phone if cap else "", cap.contact_email if cap else ""]:
        if x:
            lines.append(x)
    return subject, "\n".join(lines).strip() + "\n"


def mailto(email: str, subject: str, body: str) -> str:
    q = urllib.parse.urlencode({"subject": subject, "body": body}, quote_via=urllib.parse.quote)
    return f"mailto:{urllib.parse.quote(email or '', safe='@.+-_')}?{q}"


def _rfq(db: Session, rfq_id: int) -> VendorRFQ:
    r = db.get(VendorRFQ, rfq_id)
    if not r:
        raise NotFound(f"RFQ {rfq_id} not found")
    return r


def rfq_dict(db: Session, r: VendorRFQ) -> dict:
    vq = db.get(VendorQuote, r.vendor_quote_id) if r.vendor_quote_id else None
    files = []
    if r.include_files and not r.files_blocked:
        pq = db.get(PartQuote, r.part_quote_id)
        files = [n for n, _ in _package_files(db, pq)] if pq else []
    overdue = r.status in ("sent", "awaiting") and r.due_date and r.due_date < date.today().isoformat()
    return {"id": r.id, "part_quote_id": r.part_quote_id, "organization_id": r.organization_id, "vendor_name": r.vendor_name,
            "vendor_email": r.vendor_email, "status": r.status, "due_date": r.due_date, "overdue": bool(overdue),
            "message": r.message, "quantities": r.quantities or [], "include_files": r.include_files,
            "files_blocked": r.files_blocked, "warnings": r.warnings or [], "subject": r.subject, "body": r.body,
            "mailto": mailto(r.vendor_email, r.subject, r.body), "package_files": ["RFQ.txt"] + files,
            "vendor_quote_id": r.vendor_quote_id if vq else None,
            "response": {"prices": vq.prices, "lead_days": vq.lead_days, "tooling_charge": vq.tooling_charge, "freight": vq.freight,
                         "quote_ref": vq.quote_ref, "valid_until": vq.valid_until, "notes": vq.notes} if vq else None,
            "response_notes": r.response_notes,
            "responded_at": r.responded_at.isoformat() if r.responded_at else None,
            "created_at": r.created_at.isoformat() if r.created_at else None}


def _quantities(pq: PartQuote, quantities) -> list[int]:
    try:
        qs = sorted({int(q) for q in (quantities or []) if int(q) > 0})
    except (TypeError, ValueError):
        raise QuoteToolError("quantities must be whole numbers")
    if not qs:
        qs = [b["quantity"] for b in (pq.result or {}).get("price_breaks") or []] or ([pq.quoted_quantity] if pq.quoted_quantity else [])
    if not qs:
        raise QuoteToolError("Give at least one quantity.")
    return qs


def create_rfqs(db: Session, quote_id: int, vendor_ids: list[int], due_date: str = "", message: str = "",
                include_files: bool = True, quantities: list[int] | None = None) -> dict:
    pq = get_quote(db, quote_id)
    if not vendor_ids:
        raise QuoteToolError("Pick at least one vendor.")
    if due_date and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", due_date):
        raise QuoteToolError("due_date must be YYYY-MM-DD")
    qs = _quantities(pq, quantities)
    check = export_check(db, pq) if include_files else {"flagged": False, "reasons": []}
    blocked = bool(include_files and check["flagged"])
    warnings: list[str] = []
    if blocked:
        warnings = check["reasons"] + [EXPORT_WARNING]
        from .bidding import jcp_status_note
        note = jcp_status_note(get_profile(db), True)
        if note:
            warnings.append(note)
    available = [n for n, _ in _package_files(db, pq)] if include_files and not blocked else []
    if blocked:
        files_note = ("Technical data is controlled or limited distribution, so it is not attached. If you are JCP-certified "
                      "and authorized, reply and we will send it through an approved channel.")
    elif include_files and available:
        files_note = f"Attached: {', '.join(available)}."
    elif include_files:
        files_note = "No drawing or model is on file; reply if you need more detail."
    else:
        files_note = "Drawings and models are available on request."
    created = []
    for vid in vendor_ids:
        org = db.get(Organization, int(vid))
        if not org:
            raise NotFound(f"Vendor {vid} not found")
        contact, email = _first_contact(org)
        subject, body = draft_email(db, pq, org.name, contact, qs, due_date, message, files_note)
        r = VendorRFQ(part_quote_id=pq.id, organization_id=org.id, vendor_name=org.name, vendor_email=email, status="sent",
                      due_date=due_date, message=message, quantities=qs, include_files=bool(include_files),
                      files_blocked=blocked, warnings=warnings, subject=subject, body=body)
        db.add(r)
        created.append(r)
    db.commit()
    return {"rfqs": [rfq_dict(db, r) for r in created], "warnings": warnings, "files_blocked": blocked,
            "mailto_note": "Email links open a draft in your mail app. A mailto link cannot attach files: download the package and attach it yourself."}


def list_rfqs(db: Session, quote_id: int) -> list[dict]:
    rows = db.scalars(select(VendorRFQ).where(VendorRFQ.part_quote_id == quote_id).order_by(VendorRFQ.id)).all()
    return [rfq_dict(db, r) for r in rows]


def update_rfq(db: Session, rfq_id: int, changes: dict) -> dict:
    r = _rfq(db, rfq_id)
    if "status" in changes and changes["status"] is not None:
        if changes["status"] not in RFQ_STATUSES:
            raise QuoteToolError(f"status must be one of {RFQ_STATUSES}")
        r.status = changes["status"]
    if changes.get("due_date") is not None:
        if changes["due_date"] and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", changes["due_date"]):
            raise QuoteToolError("due_date must be YYYY-MM-DD")
        r.due_date = changes["due_date"]
    if changes.get("vendor_email") is not None:
        r.vendor_email = changes["vendor_email"]
    db.commit()
    return rfq_dict(db, r)


def decline_rfq(db: Session, rfq_id: int, notes: str = "") -> dict:
    r = _rfq(db, rfq_id)
    r.status = "declined"
    r.response_notes = notes or r.response_notes
    r.responded_at = datetime.utcnow()
    db.commit()
    return rfq_dict(db, r)


def delete_rfq(db: Session, rfq_id: int) -> None:
    r = _rfq(db, rfq_id)
    db.delete(r)
    db.commit()


def record_response(db: Session, rfq_id: int, data: dict) -> dict:
    """Save the vendor's prices as a VendorQuote (created or updated) so make-or-buy compares it."""
    r = _rfq(db, rfq_id)
    prices = []
    for p in data.get("prices") or []:
        try:
            q, u = int(p.get("quantity")), float(p.get("unit_price"))
        except (TypeError, ValueError):
            continue
        if q > 0 and u >= 0:
            prices.append({"quantity": q, "unit_price": u})
    if not prices:
        raise QuoteToolError("Enter at least one quantity with a unit price.")
    vq = db.get(VendorQuote, r.vendor_quote_id) if r.vendor_quote_id else None
    if vq is None or vq.part_quote_id != r.part_quote_id:
        vq = VendorQuote(part_quote_id=r.part_quote_id, organization_id=r.organization_id, vendor_name=r.vendor_name)
        db.add(vq)
    vq.prices = sorted(prices, key=lambda p: p["quantity"])
    lead = data.get("lead_days")
    vq.lead_days = int(lead) if lead not in (None, "") else None
    for k in ("tooling_charge", "freight"):
        if data.get(k) not in (None, ""):
            setattr(vq, k, float(data[k]))
    for k in ("quote_ref", "valid_until"):
        if data.get(k) is not None:
            setattr(vq, k, str(data[k]))
    vq.notes = str(data.get("notes") or "") or f"From RFQ #{r.id}"
    db.flush()
    r.vendor_quote_id = vq.id
    r.status = "responded"
    r.response_notes = str(data.get("notes") or "")
    r.responded_at = datetime.utcnow()
    db.commit()
    return rfq_dict(db, r)


TAB_BY_KIND = {"extrusion_build": "extrusion", "harness": "harness", "panel": "panel", "labels": "labels", "flat_dxf": "flat", "box_build": "box"}


def quote_url(pq: PartQuote | None) -> str:
    """App link that reopens a saved quote on its own tab."""
    if pq is None:
        return "/part-quotes?tab=saved"
    spec = pq.spec or {}
    tab = TAB_BY_KIND.get(_kind(pq)) or ("instant" if (spec.get("cad") or {}).get("filename") else "quote")
    return f"/part-quotes?{'' if tab == 'instant' else 'tab=' + tab + '&'}id={pq.id}"


def rfq_calendar_items(db: Session) -> list[dict]:
    out = []
    for r in db.scalars(select(VendorRFQ).where(VendorRFQ.status.in_(("sent", "awaiting")), VendorRFQ.due_date != "")).all():
        try:
            d = date.fromisoformat(r.due_date)
        except ValueError:
            continue
        out.append({"uid": f"vendor-rfq-{r.id}@govbid", "date": d, "summary": f"Vendor quote due: {r.vendor_name}",
                    "description": f"{r.subject}", "url": quote_url(db.get(PartQuote, r.part_quote_id))})
    return out


def rfq_dashboard_items(db: Session) -> list[str]:
    today = date.today().isoformat()
    rows = db.scalars(select(VendorRFQ).where(VendorRFQ.status.in_(("sent", "awaiting")), VendorRFQ.due_date != "",
                                              VendorRFQ.due_date < today)).all()
    return [f"Vendor quote from {r.vendor_name} was due {r.due_date} (quote #{r.part_quote_id}). Follow up or mark it declined." for r in rows]


# ---------------------------------------------------------------- win/loss insights
MARGIN_BANDS = [(0, 10), (10, 20), (20, 30), (30, 40), (40, None)]
MIN_POINTS = 5


def _band(m: float | None) -> int | None:
    if m is None:
        return None
    for i, (lo, hi) in enumerate(MARGIN_BANDS):
        if m >= lo and (hi is None or m < hi):
            return i
    return 0  # negative margin goes in the lowest band


def _band_label(i: int) -> str:
    lo, hi = MARGIN_BANDS[i]
    return f"{lo}%+" if hi is None else f"{lo} to {hi}%"


def _award_price(db: Session, pq: PartQuote) -> tuple[float | None, str]:
    """The award price that beat a lost quote: the first award on or after the quote date, else the latest award."""
    if not pq.nsn:
        return None, ""
    from . import nsn_history
    try:
        h = nsn_history.history(db, pq.nsn)
    except ValueError:
        return None, ""
    awards = [r for r in h["records"] if r.get("source") not in ("quote_lost", "quote_won") and r.get("unit_price") is not None]
    if not awards:
        return None, ""
    since = (pq.created_at or datetime.utcnow()).date().isoformat()
    after = sorted([a for a in awards if (a.get("award_date") or "") >= since], key=lambda a: a.get("award_date") or "")
    pick = after[0] if after else awards[0]  # records are newest first
    return float(pick["unit_price"]), pick.get("award_date") or ""


def quote_point(db: Session, pq: PartQuote) -> dict:
    spec, result = pq.spec or {}, pq.result or {}
    breaks = result.get("price_breaks") or []
    qty = pq.quoted_quantity or (breaks[0]["quantity"] if breaks else None)
    br = None
    if breaks and qty:
        br = min(breaks, key=lambda b: (abs(b["quantity"] - qty), b["quantity"]))
    price = pq.quoted_unit_price if pq.quoted_unit_price is not None else (br["unit_price"] if br else None)
    margin = None
    if br and price:
        if br.get("unit_cost") is not None:
            margin = round((price - br["unit_cost"]) / price * 100, 1)
        elif br.get("margin_pct") is not None:
            margin = br["margin_pct"]
    award, award_date = (None, "")
    if pq.status == "lost":
        award, award_date = _award_price(db, pq)
    nsn_digits = re.sub(r"\D", "", pq.nsn or "")
    return {
        "id": pq.id, "name": pq.name, "status": pq.status, "nsn": pq.nsn, "quantity": qty, "unit_price": price,
        "margin_pct": margin, "award_price": award, "award_date": award_date,
        "price_ratio": round(price / award, 3) if (award and price) else None,
        "kind": _kind(pq), "process": _process(spec) or "(none)", "material": spec.get("material") or "(none)",
        "fsc": nsn_digits[:4] if len(nsn_digits) == 13 else "(no NSN)",
        "solicitation_number": pq.opportunity.solicitation_number if pq.opportunity else "",
    }


def _mean(xs: list[float]) -> float | None:
    return round(statistics.mean(xs), 1) if xs else None


def summarize(points: list[dict]) -> dict:
    won = [p for p in points if p["status"] == "won"]
    lost = [p for p in points if p["status"] == "lost"]
    decided = len(won) + len(lost)
    ratios = [p["price_ratio"] for p in lost if p["price_ratio"] is not None]
    bands = [{"label": _band_label(i), "won": 0, "lost": 0} for i in range(len(MARGIN_BANDS))]
    for p in won + lost:
        b = _band(p["margin_pct"])
        if b is not None:
            bands[b]["won" if p["status"] == "won" else "lost"] += 1
    if decided < MIN_POINTS:
        suggestion = {"enough": False, "low": None, "high": None,
                      "text": f"Not enough data: {decided} won or lost quote{'s' if decided != 1 else ''} (need {MIN_POINTS})."}
    else:
        win_bands = [i for i, b in enumerate(bands) if b["won"]]
        if not win_bands:
            suggestion = {"enough": True, "low": None, "high": None,
                          "text": f"No wins yet in {decided} decided quotes. Price lower or check fit before bidding more of these."}
        else:
            i = max(win_bands)
            lo, hi = MARGIN_BANDS[i]
            b = bands[i]
            suggestion = {"enough": True, "low": lo, "high": hi,
                          "text": f"Target {_band_label(i)} margin: the highest band that still won "
                                  f"({b['won']} won, {b['lost']} lost in that band; {decided} decided quotes in this group)."}
    return {
        "quotes": len(points), "won": len(won), "lost": len(lost), "submitted": sum(p["status"] == "submitted" for p in points),
        "win_rate": round(len(won) / decided * 100, 1) if decided else None,
        "avg_price_ratio": round(statistics.mean(ratios), 3) if ratios else None, "ratio_count": len(ratios),
        "avg_margin_won": _mean([p["margin_pct"] for p in won if p["margin_pct"] is not None]),
        "avg_margin_lost": _mean([p["margin_pct"] for p in lost if p["margin_pct"] is not None]),
        "bands": bands, "suggestion": suggestion,
    }


def insights(db: Session) -> dict:
    rows = db.scalars(select(PartQuote).where(PartQuote.status.in_(("won", "lost", "submitted"))).order_by(PartQuote.id)).all()
    points = [quote_point(db, pq) for pq in rows]
    groups = {}
    for dim in ("kind", "process", "material", "fsc"):
        vals: dict[str, list[dict]] = {}
        for p in points:
            vals.setdefault(p[dim], []).append(p)
        groups[dim] = sorted(({"value": v, **summarize(ps)} for v, ps in vals.items()), key=lambda g: (-g["quotes"], g["value"]))
    return {"overall": summarize(points), "groups": groups, "points": points, "min_points": MIN_POINTS,
            "bands": [_band_label(i) for i in range(len(MARGIN_BANDS))]}
