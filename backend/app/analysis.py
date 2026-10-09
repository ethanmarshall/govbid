"""Requirements breakdown and compliance matrix for a solicitation.

With ANTHROPIC_API_KEY set, Claude reads the solicitation text and returns a structured
breakdown. Without it, a rule-based pass still extracts "shall/must" requirements,
key FAR/DFARS clauses, deadlines, page limits and the evaluation method.
"""
from __future__ import annotations

import json
import re
from html import unescape
from pathlib import Path

from .config import ANTHROPIC_API_KEY, ANTHROPIC_MODEL

MAX_CHARS_FOR_MODEL = 180_000

# Clauses worth flagging, with a plain-language note on why each matters.
KEY_CLAUSES = {
    "52.219-14": "Limitations on Subcontracting: caps what you can pay non-similarly-situated subcontractors (50% services, 85% general construction, 75% specialty trades).",
    "52.219-27": "Notice of SDVOSB Set-Aside: only SBA-certified SDVOSBs may receive the award.",
    "52.219-1": "Small Business Program Representations: confirm your size under the solicitation NAICS.",
    "52.219-6": "Notice of Total Small Business Set-Aside.",
    "52.204-7": "System for Award Management: you must be registered and active in SAM.",
    "52.204-24": "Covered telecom representation (Section 889): required representation about Huawei/ZTE etc. equipment.",
    "52.204-25": "Prohibition on covered telecom equipment (Section 889).",
    "52.204-26": "Covered telecom equipment representation.",
    "52.204-21": "Basic safeguarding of covered contractor information systems.",
    "252.204-7012": "DFARS cyber incident reporting and NIST SP 800-171 controls for CUI.",
    "252.204-7019": "NIST SP 800-171 assessment score must be posted in SPRS.",
    "252.204-7021": "CMMC requirement: you need the stated CMMC level before award.",
    "252.204-7025": "CMMC level requirement notice.",
    "52.212-1": "Instructions to Offerors (Commercial): submission format and content.",
    "52.212-2": "Evaluation (Commercial): how offers are evaluated.",
    "52.212-3": "Offeror Representations and Certifications (Commercial).",
    "52.225-1": "Buy American (supplies): domestic end product rules.",
    "52.225-2": "Buy American Certificate.",
    "252.225-7001": "DFARS Buy American / balance of payments.",
    "52.222-41": "Service Contract Labor Standards: wage determination applies to service employees.",
    "52.222-6": "Construction Wage Rate Requirements (Davis-Bacon).",
    "52.228-15": "Performance and payment bonds (construction).",
    "52.232-33": "Payment by Electronic Funds Transfer (SAM).",
    "252.211-7003": "Item unique identification (IUID) marking.",
    "252.246-7008": "Sources of electronic parts (counterfeit avoidance).",
    "52.246-2": "Inspection of supplies (fixed price).",
    "252.232-7006": "Wide Area WorkFlow (WAWF/PIEE) invoicing.",
}

# Standards and data items cited in solicitation text. Order matters: first match wins for the type label.
STANDARD_PATTERNS = [
    ("Military standard", r"\bMIL-(?:STD|HDBK|PRF|DTL|SPEC|[A-Z])-\d{1,6}(?:-\d{1,2})?[A-Z]?\b"),
    ("Data item description", r"\bDI-[A-Z]{4}-\d{5}[A-Z]?\b"),
    ("ASME", r"\bASME\s+Y14\.\d+(?:\.\d+)?M?\b|\bASME\s+B\d+(?:\.\d+)*\b"),
    ("ISA", r"\b(?:ANSI/)?ISA-?\d+(?:\.\d+)+\b"),
    ("ISO", r"\bISO(?:/IEC)?\s?\d{4,5}(?::\d{4})?\b"),
    ("SAE aerospace", r"\bAS\s?9\d{3}[A-D]?\b"),
    ("IPC", r"\bIPC(?:/WHMA)?-[A-Z]?-?\d{3,4}[A-Z]?\b|\bJ-STD-\d{3}[A-Z]?\b"),
    ("NIST", r"\bNIST\s+SP\s+800-\d+[A-Z]?\b"),
    ("SAE/EIA", r"\b(?:SAE\s+)?EIA-\d{3,4}[A-Z]?\b"),
    ("Federal standard", r"\bFED-STD-\d+[A-Z]?\b|\bA-A-\d{4,5}[A-Z]?\b"),
]


def cited_standards(text: str) -> list[dict]:
    """Find standards, specs and DIDs the solicitation references, with how often each appears."""
    found: dict[str, dict] = {}
    for label, pattern in STANDARD_PATTERNS:
        for m in re.finditer(pattern, text, re.I):
            key = re.sub(r"\s+", " ", m.group(0)).upper().replace("ISO/IEC", "ISO/IEC ").replace("  ", " ").strip()
            if key not in found:
                free = label in ("Military standard", "Data item description", "NIST", "Federal standard")
                found[key] = {"standard": key, "type": label, "count": 0, "free": free}
            found[key]["count"] += 1
    return sorted(found.values(), key=lambda x: (-x["count"], x["standard"]))


# A sentence may contain dots inside paragraph numbers (C.3.1) or decimals (4.5), so only a dot that is
# not followed by a digit ends it.
REQ_RE = re.compile(r"(?:[^.\n]|\.(?=\d))*\b(shall|must|is required to|are required to|will be required to)\b(?:[^.\n]|\.(?=\d))*(?:\.|\n|$)", re.I)
LEAD_LABEL_RE = re.compile(r"^\s*((?:[A-Z]\.)?\d+(?:\.\d+)*|[A-Z]\.\d+(?:\.\d+)*|\([a-z0-9]+\))\s+(?=[A-Z(])")
SECTION_RE = re.compile(r"\b(?:section|sec\.?)\s+([A-M])\b", re.I)
PARA_RE = re.compile(r"^\s*((?:\d+\.){1,4}\d*|[A-Z]\.\d+(?:\.\d+)*|\([a-z0-9]+\))\s+", re.M)
DATE_PAT = r"(?:(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+\d{1,2},?\s+\d{4}|\d{1,2}/\d{1,2}/\d{2,4}|\d{4}-\d{2}-\d{2}|\d{1,2}\s+(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\s+\d{4})"


# ---------------------------------------------------------------- text extraction
def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            return "\n".join((p.extract_text() or "") for p in reader.pages)
        if suffix == ".docx":
            import docx

            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)
        if suffix in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook

            wb = load_workbook(str(path), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"## Sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    if any(v is not None for v in row):
                        lines.append(" | ".join("" if v is None else str(v) for v in row))
            return "\n".join(lines)
        raw = path.read_text(errors="replace")
        if suffix in (".html", ".htm") or "<html" in raw[:500].lower():
            return html_to_text(raw)
        return raw
    except Exception as exc:  # unreadable or scanned file
        return f"[Could not extract text from {path.name}: {exc}]"


def html_to_text(html: str) -> str:
    html = re.sub(r"(?is)<(script|style).*?</\1>", " ", html)
    html = re.sub(r"(?i)<br\s*/?>|</p>|</li>|</h\d>|</tr>", "\n", html)
    text = unescape(re.sub(r"<[^>]+>", " ", html))
    return re.sub(r"[ \t]+", " ", re.sub(r"\n\s*\n+", "\n\n", text)).strip()


# ---------------------------------------------------------------- rule-based analysis
def heuristic_analysis(opp: dict, docs: list[tuple[str, str]]) -> dict:
    corpus = "\n\n".join(t for _, t in docs)
    low = corpus.lower()

    clauses = [
        {"clause": c, "note": note}
        for c, note in KEY_CLAUSES.items()
        if re.search(rf"(?<![\d.]){re.escape(c)}(?![\d])", corpus)
    ]

    if re.search(r"lowest price technically acceptable|\blpta\b", low):
        method = "Lowest Price Technically Acceptable (LPTA): price decides among offers that pass the technical review."
    elif re.search(r"best value|trade-?off", low):
        method = "Best value tradeoff: non-price factors can outweigh a lower price."
    elif re.search(r"lowest (evaluated )?price|lowest priced", low):
        method = "Lowest price."
    else:
        method = "Not stated in the text reviewed. Check Section M or the evaluation provision (52.212-2)."

    factors = []
    for m in re.finditer(r"factor\s*(\d+)\s*[:\-–]?\s*([A-Za-z][A-Za-z /&,\-]{3,60})", corpus, re.I):
        f = f"Factor {m.group(1)}: {m.group(2).strip()}"
        if f not in factors:
            factors.append(f)

    page_limits = sorted({m.group(0).strip() for m in re.finditer(r"[^.\n]{0,60}\b(?:not (?:to )?exceed|limited to|maximum of|no more than)\s+\d+\s+pages?[^.\n]{0,40}", corpus, re.I)})[:10]

    deadlines = []
    for m in re.finditer(rf"[^.\n]{{0,80}}\b(?:due|no later than|deadline|closing|submit(?:ted)? by|questions)\b[^.\n]{{0,80}}?{DATE_PAT}[^.\n]{{0,40}}", corpus, re.I):
        s = re.sub(r"\s+", " ", m.group(0)).strip()
        if s not in deadlines:
            deadlines.append(s)

    matrix = []
    seen = set()
    for name, text in docs:
        for m in REQ_RE.finditer(text):
            sentence = re.sub(r"\s+", " ", m.group(0)).strip()
            own = LEAD_LABEL_RE.match(sentence)
            if own:
                sentence = sentence[own.end():].strip()
            key = sentence.lower()[:120]
            if len(sentence) < 25 or key in seen:
                continue
            seen.add(key)
            before = text[: m.start()]
            sec = SECTION_RE.findall(before[-4000:])
            para = [own.group(1)] if own else PARA_RE.findall(before[-600:])
            if own and sec and own.group(1).upper().startswith(sec[-1].upper() + "."):
                sec = []  # "C.3.1" already names Section C
            matrix.append(
                {
                    "id": len(matrix) + 1,
                    "requirement": sentence[:600],
                    "reference": " ".join(x for x in (f"Sec. {sec[-1].upper()}" if sec else "", para[-1] if para else "") if x),
                    "source": name,
                    "category": _categorize(sentence),
                    "response_location": "",
                    "status": "open",
                }
            )
            if len(matrix) >= 400:
                break

    red_flags = []
    if "52.219-14" in [c["clause"] for c in clauses]:
        red_flags.append("Limitations on subcontracting applies. Plan how much work you will self-perform.")
    if re.search(r"252\.204-7021|cmmc", low):
        red_flags.append("CMMC requirement present. Confirm the level and whether you hold it before investing in a bid.")
    if re.search(r"bond", low) and re.search(r"52\.228", corpus):
        red_flags.append("Bonding is required.")
    if re.search(r"security clearance|facility clearance|\bfcl\b|secret clearance", low):
        red_flags.append("Clearance requirements mentioned.")
    if re.search(r"past performance", low) and re.search(r"(three|3)\s+(relevant|recent)", low):
        red_flags.append("Specific number of past performance references requested.")
    if re.search(r"site visit", low):
        red_flags.append("Site visit mentioned. Check whether attendance is mandatory.")

    return {
        "summary": _quick_summary(opp, corpus),
        "breakdown": {
            "scope": _first_paragraph_after(corpus, ["scope of work", "statement of work", "performance work statement", "background", "requirement"]),
            "evaluation_method": method,
            "evaluation_factors": factors[:10],
            "submission_instructions": _first_paragraph_after(corpus, ["instructions to offerors", "proposal submission", "submission instructions", "quotes shall be submitted", "submission of quotes"]),
            "deadlines": deadlines[:10],
            "page_limits": page_limits,
            "key_clauses": clauses,
            "cited_standards": cited_standards(corpus),
            "eligibility_notes": [],
            "red_flags": red_flags,
            "questions_to_ask": [],
        },
        "compliance_matrix": matrix,
    }


def _categorize(s: str) -> str:
    l = s.lower()
    if re.search(r"page|font|format|volume|submit|copies|electronic|email|label", l):
        return "Submission"
    if re.search(r"price|pricing|cost|clin|invoice|payment", l):
        return "Pricing"
    if re.search(r"past performance|reference|experience", l):
        return "Past performance"
    if re.search(r"deliver|ship|f\.?o\.?b|packag|marking|iuid", l):
        return "Delivery"
    if re.search(r"certif|represent|sam|register|insurance|license|bond", l):
        return "Eligibility / reps"
    if re.search(r"report|meeting|status|deliverable", l):
        return "Reporting"
    return "Technical / performance"


def _first_paragraph_after(text: str, headings: list[str]) -> str:
    low = text.lower()
    for h in headings:
        i = low.find(h)
        if i >= 0:
            chunk = text[i : i + 1200]
            return re.sub(r"\s+", " ", chunk).strip()[:900]
    return ""


def format_deadline(s: str) -> str:
    from datetime import datetime

    try:
        dt = datetime.fromisoformat(s)
        if "T" in s:
            tz = dt.strftime("%z")
            tz = f" (UTC{tz[:3]}:{tz[3:]})" if tz else ""
            return dt.strftime("%B %d, %Y at %I:%M %p").replace(" 0", " ") + tz
        return dt.strftime("%B %d, %Y").replace(" 0", " ")
    except ValueError:
        return s


def _quick_summary(opp: dict, corpus: str) -> str:
    bits = [opp.get("title") or "Untitled opportunity"]
    if opp.get("agency"):
        bits.append(f"from {opp['agency']}")
    s = " ".join(bits) + "."
    if opp.get("set_aside_desc"):
        s += f" Set-aside: {opp['set_aside_desc']}."
    if opp.get("naics"):
        s += f" NAICS {opp['naics']}."
    if opp.get("response_deadline"):
        s += f" Responses due {format_deadline(opp['response_deadline'])}."
    if not corpus.strip():
        s += " No solicitation text was available, so only notice metadata was used. Upload the solicitation documents for a full breakdown."
    return s


# ---------------------------------------------------------------- Claude analysis
SYSTEM_PROMPT = """You are a federal contracting analyst helping a small business decide whether and how to bid.
Read the solicitation material and return ONLY a JSON object, no prose, with this shape:
{
  "summary": "3-5 plain sentences: what is being bought, by whom, contract type, period of performance, and whether it suits a small business",
  "breakdown": {
    "scope": "what the contractor must deliver or perform",
    "contract_type": "e.g. Firm-Fixed-Price, IDIQ, BPA; empty if not stated",
    "period_of_performance": "",
    "place_of_performance": "",
    "evaluation_method": "LPTA, best value tradeoff, lowest price, etc., with a one-line explanation",
    "evaluation_factors": ["Factor name and relative importance"],
    "submission_instructions": "format, volumes, delivery method, where to submit",
    "deadlines": ["Questions due ...", "Proposals due ... (include time zone)"],
    "page_limits": ["Volume I Technical: 10 pages"],
    "key_clauses": [{"clause": "52.219-14", "note": "why it matters to this bidder"}],
    "eligibility_notes": ["set-aside, size standard, certifications, licenses, clearances, CMMC level, bonding"],
    "red_flags": ["anything that could disqualify a small bidder or make the bid risky"],
    "questions_to_ask": ["questions worth submitting to the contracting officer"]
  },
  "compliance_matrix": [
    {"id": 1, "requirement": "verbatim or tightly paraphrased requirement", "reference": "section/paragraph, e.g. L.3.2 or PWS 4.1", "source": "file name", "category": "Submission | Technical / performance | Pricing | Past performance | Delivery | Eligibility / reps | Reporting", "response_location": "", "status": "open"}
  ]
}
Include every shall/must requirement from the instructions, evaluation criteria and statement of work in the compliance matrix.
Use empty strings or empty lists when information is not present. Do not invent dates or numbers."""


def claude_analysis(opp: dict, docs: list[tuple[str, str]], eligibility: dict | None = None) -> dict:
    import anthropic

    client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)
    header = {k: opp.get(k) for k in ("title", "solicitation_number", "agency", "notice_type", "set_aside_desc", "naics", "psc", "response_deadline", "place_of_performance")}
    material, used = [], 0
    for name, text in docs:
        room = MAX_CHARS_FOR_MODEL - used
        if room <= 0:
            break
        chunk = text[:room]
        used += len(chunk)
        material.append(f'<document name="{name}">\n{chunk}\n</document>')
    user = (
        f"Notice metadata:\n{json.dumps(header, indent=2)}\n\n"
        + (f"Bidder eligibility check: {json.dumps(eligibility)}\n\n" if eligibility else "")
        + ("\n\n".join(material) if material else "No documents attached. Work from the metadata and say what is missing.")
    )
    msg = client.messages.create(
        model=ANTHROPIC_MODEL,
        max_tokens=16000,
        system=SYSTEM_PROMPT,
        messages=[{"role": "user", "content": user}],
    )
    text = "".join(b.text for b in msg.content if getattr(b, "type", "") == "text")
    data = _parse_json(text)
    matrix = data.get("compliance_matrix") or []
    for i, row in enumerate(matrix, 1):
        row.setdefault("id", i)
        row.setdefault("status", "open")
        row.setdefault("response_location", "")
    return {"summary": data.get("summary", ""), "breakdown": data.get("breakdown", {}), "compliance_matrix": matrix}


def _parse_json(text: str) -> dict:
    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start >= 0 and end > start:
            return json.loads(text[start : end + 1])
        raise


def analyze(opp: dict, docs: list[tuple[str, str]], eligibility: dict | None = None) -> tuple[str, dict]:
    """Returns (method, result)."""
    if ANTHROPIC_API_KEY:
        try:
            result = claude_analysis(opp, docs, eligibility)
            # Merge in rule-based clause detection so known clauses are never missed.
            h = heuristic_analysis(opp, docs)
            have = {c.get("clause") for c in result["breakdown"].get("key_clauses", [])}
            result["breakdown"].setdefault("key_clauses", [])
            result["breakdown"]["key_clauses"] += [c for c in h["breakdown"]["key_clauses"] if c["clause"] not in have]
            result["breakdown"]["cited_standards"] = h["breakdown"]["cited_standards"]
            return "claude", result
        except Exception as exc:
            result = heuristic_analysis(opp, docs)
            result["breakdown"]["red_flags"].insert(0, f"AI analysis failed ({exc.__class__.__name__}); showing rule-based results.")
            return "heuristic", result
    return "heuristic", heuristic_analysis(opp, docs)


# ---------------------------------------------------------------- export
def matrix_to_xlsx(opp: dict, analysis: dict, out_path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Compliance Matrix"
    ws["A1"] = f"{opp.get('solicitation_number') or ''} {opp.get('title') or ''}".strip()
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"{opp.get('agency') or ''} | Due: {opp.get('response_deadline') or 'n/a'}"
    headers = ["#", "Requirement", "Reference", "Source", "Category", "Proposal location", "Status", "Owner", "Notes"]
    ws.append([])
    ws.append(headers)
    hdr_row = ws.max_row
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=hdr_row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
    for r in analysis.get("compliance_matrix", []):
        ws.append([r.get("id"), r.get("requirement"), r.get("reference"), r.get("source"), r.get("category"), r.get("response_location"), r.get("status"), "", ""])
    widths = [5, 80, 14, 22, 20, 20, 10, 14, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=hdr_row + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)
    ws.auto_filter.ref = f"A{hdr_row}:{get_column_letter(len(headers))}{ws.max_row}"

    b = analysis.get("breakdown", {})
    ws2 = wb.create_sheet("Breakdown")
    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["B"].width = 110
    ws2.append(["Summary", analysis.get("summary", "")])
    for key, val in b.items():
        if isinstance(val, list):
            def _line(x):
                if isinstance(x, dict) and "clause" in x:
                    return f"{x.get('clause')}: {x.get('note')}"
                if isinstance(x, dict) and "standard" in x:
                    return f"{x.get('standard')} ({x.get('type')}, cited {x.get('count')}x)"
                return str(x)
            val = "\n".join(_line(x) for x in val)
        elif isinstance(val, dict):
            if key == "export_control":
                val = ("Yes: " + ", ".join(val.get("hits", [])) + "\n" + "\n".join(val.get("snippets", []))) if val.get("flagged") else "No markings found"
            else:
                val = "\n".join(f"{k}: {v}" for k, v in val.items())
        elif val is not None and not isinstance(val, (str, int, float)):
            val = str(val)
        ws2.append([key.replace("_", " ").title(), val])
    for row in ws2.iter_rows():
        row[0].font = Font(bold=True)
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(out_path)
    return out_path
