"""Clause flowdown checker: which prime contract clauses go into a subcontract or vendor PO.

The clause table, conditions, thresholds and sources live in app/flowdown_data.py (verified on
acquisition.gov, 2026-10-08). This module adds clause extraction from a stored solicitation
analysis and a DOCX PO terms attachment.
"""
from __future__ import annotations

import io
import json
import re
from datetime import date

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import flowdown_data as fd
from .db import get_db
from .models import Analysis, Opportunity

router = APIRouter(prefix="/api/flowdown")

# FAR 52.xxx-x(x) and DFARS 252.xxx-xxxx. The lookbehind keeps "252.204-7012" from also matching as FAR.
FAR_RE = re.compile(r"(?<![\d.])52\.\d{3}-\d{1,3}(?!\d)")
DFARS_RE = re.compile(r"(?<![\d.])252\.\d{3}-\d{4}(?!\d)")


def extract_clause_numbers(text: str) -> list[str]:
    """All FAR (52.xxx-xx) and DFARS (252.xxx-xxxx) clause numbers in the text, in first-seen order."""
    out: list[str] = []
    found = [(m.start(), m.group(0)) for m in FAR_RE.finditer(text or "")]
    found += [(m.start(), m.group(0)) for m in DFARS_RE.finditer(text or "")]
    for _, n in sorted(found):
        if n not in out:
            out.append(n)
    return out


def opportunity_clauses(db: Session, opp_id: int) -> list[str]:
    """Clause numbers for an opportunity: analysis breakdown key_clauses plus a regex scan of the
    stored analysis summary, breakdown and compliance matrix. Raises LookupError if not found."""
    opp = db.get(Opportunity, opp_id)
    if not opp:
        raise LookupError("Opportunity not found")
    a = db.scalars(select(Analysis).where(Analysis.opportunity_id == opp_id)).first()
    if not a:
        return []
    nums: list[str] = []
    for c in (a.breakdown or {}).get("key_clauses") or []:
        raw = c.get("clause", "") if isinstance(c, dict) else str(c)
        nums += extract_clause_numbers(str(raw))
    blob = "\n".join([a.summary or "", json.dumps(a.breakdown or {}), json.dumps(a.compliance_matrix or [])])
    nums += extract_clause_numbers(blob)
    seen: list[str] = []
    for n in nums:
        if n not in seen:
            seen.append(n)
    return seen


class Subcontract(BaseModel):
    value: float = 0
    commercial: bool = False
    cots: bool = False
    prime_commercial: bool = False
    services: bool = False
    supplies: bool = True
    construction: bool = False
    maintenance_repair: bool = False
    involves_fci: bool = False
    involves_cui: bool = False
    operationally_critical: bool = False
    international: bool = False
    us_work: bool = True
    performance_days: int = 0
    small_business: bool = False
    further_subcontracting: bool = False
    pii: bool = False
    electronic_parts: bool = False
    original_manufacturer: bool = False
    specialty_metals: bool = False
    iuid: bool = False
    ocean_shipping: bool = False
    resale_no_value_added: bool = False
    contingency_support: bool = False
    sca_covered: bool = False


class CheckIn(BaseModel):
    clauses: list[str] = Field(default_factory=list)
    clause_text: str = ""  # pasted text; clause numbers are pulled out with the regex
    opportunity_id: int | None = None
    subcontract: Subcontract = Field(default_factory=Subcontract)


def run_check(db: Session, body: CheckIn) -> dict:
    nums: list[str] = []
    for c in body.clauses:
        found = extract_clause_numbers(c)
        nums += found or ([c.strip()] if c.strip() else [])
    nums += extract_clause_numbers(body.clause_text)
    opp_title = ""
    if body.opportunity_id:
        try:
            nums += opportunity_clauses(db, body.opportunity_id)
        except LookupError as exc:
            raise HTTPException(404, str(exc))
        o = db.get(Opportunity, body.opportunity_id)
        opp_title = (o.solicitation_number + " " if o.solicitation_number else "") + (o.title or "")
    results = fd.check(nums, body.subcontract.model_dump())
    return {
        "clauses_in": [r["number"] for r in results],
        "results": results,
        "flow_down": [r for r in results if r["flow_down"] == "yes"],
        "check": [r for r in results if r["flow_down"] == "check"],
        "not_required": [r for r in results if r["flow_down"] == "no"],
        "unknown": [r["number"] for r in results if r["status"] == "unknown"],
        "opportunity": opp_title.strip(),
        "commercial_note": fd.COMMERCIAL_NOTE,
    }


@router.get("/clauses")
def list_clauses():
    return {"clauses": [fd.public(c) for c in fd.CLAUSES], "thresholds": fd.THRESHOLDS,
            "commercial_note": fd.COMMERCIAL_NOTE, "subcontract_defaults": fd.SUBCONTRACT_DEFAULTS,
            "verified_on": "2026-10-08"}


@router.get("/opportunity/{opp_id}/clauses")
def get_opportunity_clauses(opp_id: int, db: Session = Depends(get_db)):
    try:
        return {"clauses": opportunity_clauses(db, opp_id)}
    except LookupError as exc:
        raise HTTPException(404, str(exc))


@router.post("/check")
def check(body: CheckIn, db: Session = Depends(get_db)):
    return run_check(db, body)


class AttachmentIn(CheckIn):
    vendor: str = ""
    po_number: str = ""
    prime_contract: str = ""
    buyer: str = ""
    include: list[str] = Field(default_factory=list)  # extra clause numbers to add (e.g. "check" items you decided to flow)
    exclude: list[str] = Field(default_factory=list)


def build_attachment(rows: list[dict], vendor: str = "", po_number: str = "", prime_contract: str = "", buyer: str = "") -> bytes:
    """DOCX PO terms attachment listing the clauses incorporated by reference."""
    import docx
    from docx.shared import Pt

    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    d.add_heading("Purchase Order Terms: Flowdown Clauses", level=1)
    meta = [("Purchase order", po_number), ("Seller", vendor), ("Buyer", buyer), ("Prime contract", prime_contract),
            ("Date", date.today().isoformat())]
    for k, v in meta:
        if v:
            p = d.add_paragraph()
            p.add_run(f"{k}: ").bold = True
            p.add_run(v)
    d.add_paragraph(
        "The following clauses are incorporated by reference in this purchase order with the same force and effect "
        "as if they were given in full text. The full text of FAR clauses is available at https://www.acquisition.gov/far "
        "and DFARS clauses at https://www.acquisition.gov/dfars. In these clauses, \"Contractor\" means the Seller and "
        "\"Contracting Officer\" means the Buyer's purchasing representative, except where the clause requires notice "
        "to or action by the Government. Each clause applies in the version cited in the Buyer's prime contract."
    )
    for fam in ("FAR", "DFARS"):
        sub = [r for r in rows if r.get("family") == fam]
        if not sub:
            continue
        d.add_heading(f"{'Federal Acquisition Regulation (48 CFR Chapter 1)' if fam == 'FAR' else 'Defense FAR Supplement (48 CFR Chapter 2)'}", level=2)
        t = d.add_table(rows=1, cols=3)
        t.style = "Table Grid"
        hdr = t.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Clause", "Title", "Date / extent"
        for r in sub:
            cells = t.add_row().cells
            cells[0].text = r["number"]
            cells[1].text = r.get("title") or ""
            ext = r.get("date") or ""
            if r.get("portion"):
                ext = (ext + ", " if ext else "") + r["portion"]
            cells[2].text = ext
    d.add_paragraph()
    d.add_paragraph("The Seller shall include these clauses in its lower-tier subcontracts to the extent each clause requires.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@router.post("/po-attachment")
def po_attachment(body: AttachmentIn, db: Session = Depends(get_db)):
    res = run_check(db, body)
    rows = [r for r in res["results"] if r["flow_down"] == "yes" and r["number"] not in body.exclude]
    have = {r["number"] for r in rows}
    for n in body.include:
        if n not in have:
            r = fd.evaluate(n, body.subcontract.model_dump())
            rows.append(r)
            have.add(n)
    if not rows:
        raise HTTPException(400, "No clauses to flow down for this subcontract.")
    order = {c["number"]: i for i, c in enumerate(fd.CLAUSES)}
    rows.sort(key=lambda r: (order.get(r["number"], 10_000), r["number"]))
    data = build_attachment(rows, body.vendor, body.po_number, body.prime_contract, body.buyer)
    name = re.sub(r"[^A-Za-z0-9_-]+", "_", f"PO_{body.po_number or 'terms'}_flowdown") + ".docx"
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})
