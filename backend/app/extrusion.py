"""Quote T-slot aluminum extrusion builds (frames, carts, enclosures, test stands, trainer frames).

Input comes from a drawing PDF with a BOM and cut list, a CSV/XLSX BOM, or lines typed by hand.
Each line is matched to the catalog in app/extrusion_catalog.py (profiles, hardware, panels) and
priced with the "extrusion" section of the shop-rate config:

  per build = profile material + cuts + machining ops + hardware + panels + assembly + inspection + packaging
  per lot   = kitting + supplier order charge + freight + certificate of conformance (+ first article)
  price(q)  = (per build x q + per lot) x (1 + G&A) x (1 + profit), at least the minimum lot charge

Profile material is priced one of two ways (config extrusion.pricing_mode):
  cut_to_length: the supplier cuts each piece; price per inch of part length plus a cut charge per piece.
  stock: you buy full sticks and cut in house. Pieces for all builds in the lot are nested into sticks
         first-fit decreasing, so material per build drops as the lot grows.

Catalog prices are placeholders until the owner enters distributor prices. Parsing is best effort:
lines it cannot match come back as "unmatched" for the user to map.
"""
from __future__ import annotations

import base64
import csv
import io
import json
import math
import re
from collections import Counter, defaultdict
from pathlib import Path

from . import pricing
from .config import ANTHROPIC_API_KEY, ANTHROPIC_MODEL
from .extrusion_catalog import FAMILIES, MACHINING_OPS, PANEL_MATERIALS, catalog

MM_PER_IN = 25.4


class ExtrusionError(ValueError):
    pass


# ---------------------------------------------------------------- units
_NUM = r"\d+(?:\.\d+)?|\.\d+"
_FRAC = r"(\d+)\s+(\d+)/(\d+)|(\d+)/(\d+)|(\d+(?:\.\d+)?|\.\d+)"
LEN_MM = re.compile(rf"(?<![\w.])({_NUM})\s*MM\b", re.I)
LEN_CM = re.compile(rf"(?<![\w.])({_NUM})\s*CM\b", re.I)
LEN_M = re.compile(rf"(?<![\w.])({_NUM})\s*M\b(?![\w-])", re.I)
LEN_FT = re.compile(rf"(?<![\w.])({_NUM})\s*(?:FT\b|FEET\b|FOOT\b|')(?:\s*-?\s*({_NUM})\s*(?:IN\b|INCH(?:ES)?\b|\")?)?", re.I)
LEN_IN = re.compile(r"(?<![\w./])(\d+\s+\d+/\d+|\d+/\d+|\d+(?:\.\d+)?|\.\d+)\s*(?:IN\b|INCH(?:ES)?\b|\"|'')", re.I)
BARE_DEC = re.compile(r"(?<![\w./-])(\d+\.\d+)(?![\w./-])")
BARE_INT = re.compile(r"(?<![\w./-])(\d{1,4})(?![\w./-])")
METRIC_DOC = re.compile(r"DIMENSIONS\s+(?:ARE\s+)?IN\s+(?:MILLIMETERS|MM)|ALL\s+DIMENSIONS\s+(?:ARE\s+)?IN\s+MM|UNITS\s*[:=]?\s*(?:MM|MILLIMETERS)", re.I)


def _frac(s: str) -> float:
    s = s.strip()
    m = re.fullmatch(r"(\d+)\s+(\d+)/(\d+)", s)
    if m:
        return int(m.group(1)) + int(m.group(2)) / int(m.group(3))
    m = re.fullmatch(r"(\d+)/(\d+)", s)
    if m:
        return int(m.group(1)) / int(m.group(2))
    return float(s)


def find_length(text: str, default_unit: str = "in", allow_bare_int: bool = False) -> tuple[float | None, str, tuple[int, int] | None]:
    """First length in `text` -> (inches, unit seen, span). Explicit units win over bare numbers."""
    t = text or ""
    best: tuple[int, float, str, tuple[int, int]] | None = None
    for rx, unit in ((LEN_MM, "mm"), (LEN_CM, "cm"), (LEN_M, "m"), (LEN_FT, "ft"), (LEN_IN, "in")):
        m = rx.search(t)
        if not m:
            continue
        if unit == "mm":
            v = float(m.group(1)) / MM_PER_IN
        elif unit == "cm":
            v = float(m.group(1)) * 10 / MM_PER_IN
        elif unit == "m":
            v = float(m.group(1)) * 1000 / MM_PER_IN
        elif unit == "ft":
            v = float(m.group(1)) * 12 + (float(m.group(2)) if m.group(2) else 0.0)
        else:
            v = _frac(m.group(1))
        if best is None or m.start() < best[0]:
            best = (m.start(), v, unit, m.span())
    if best:
        return round(best[1], 4), best[2], best[3]
    m = BARE_DEC.search(t) or (BARE_INT.search(t) if allow_bare_int else None)
    if m:
        v = float(m.group(1))
        return round(v / MM_PER_IN if default_unit == "mm" else v, 4), default_unit, m.span()
    return None, "", None


def parse_length(text: str, default_unit: str = "in") -> float | None:
    """Length in inches from text like 24.000, 24", 24 1/2 IN, 610 mm, 2 ft, 2'-6". Bare numbers use default_unit."""
    return find_length(str(text or ""), default_unit, allow_bare_int=True)[0]


# ---------------------------------------------------------------- machining callouts
OP_PATTERNS: list[tuple[str, str]] = [
    ("drill_thru_cbore", r"DRILL\s*(?:THRU|THROUGH)\s*(?:&|AND|\+|/)\s*(?:C\s*'?\s*BORE|COUNTER\s*-?\s*BORE)"),
    ("central_connector_cbore", r"CENTRAL\s*CONNECTOR"),
    ("miter_cbore", r"MITER(?:CUT)?\s*(?:C\s*'?\s*BORE|COUNTER\s*-?\s*BORE)|MITRE\s*COUNTER"),
    ("end_tap", r"END\s*-?\s*TAP|TAP(?:PED)?\s*(?:BOTH|EACH|ONE)\s*ENDS?|TAP\s*END"),
    ("access_hole", r"ACCESS\s*HOLES?"),
    ("counterbore", r"C\s*'?\s*BORE|COUNTER\s*-?\s*BORE|\bCBORE\b"),
    ("miter_cut", r"MITER(?:\s*CUT)?|MITRE|MITERCUT|\b\d{1,2}(?:\.\d+)?\s*(?:°|DEG\b|DEGREE)"),
    ("drill_thru", r"DRILL\s*(?:THRU|THROUGH)|THRU\s*HOLE|\bDRILL\b"),
]
OP_START = re.compile("|".join(p for _, p in OP_PATTERNS), re.I)
OP_LABELS = {k: v[0] for k, v in MACHINING_OPS.items()}


def parse_machining(text: str) -> list[dict]:
    """Machining callouts -> [{op, count}]. 'END TAP BOTH ENDS' = 2, '2X ACCESS HOLE' = 2, '(4) CBORE' = 4."""
    ops: Counter = Counter()
    for seg in re.split(r"[,;\n]|\s+&\s+(?!C\s*'?\s*BORE|COUNTER)|\bAND\b(?!\s*(?:C\s*'?\s*BORE|COUNTER))", text or "", flags=re.I):
        seg = seg.strip()
        if not seg:
            continue
        for op, pat in OP_PATTERNS:
            m = re.search(pat, seg, re.I)
            if not m:
                continue
            n = 1
            before = seg[: m.start()]
            after = seg[m.end():]
            c = re.search(r"(?:\(\s*(\d{1,3})\s*\)|\b(\d{1,3})\s*[X×]\s*)$", before.strip() + " ", re.I) or re.search(r"(?:\(\s*(\d{1,3})\s*\)|\b(\d{1,3})\s*[X×])\s*$", before, re.I)
            if c:
                n = int(c.group(1) or c.group(2))
            elif re.search(r"BOTH\s*ENDS|EACH\s*END|2\s*ENDS", seg, re.I):
                n = 2
            else:
                k = re.search(r"^\s*(?:\(\s*(\d{1,3})\s*\)|(\d{1,3})\s*(?:PL(?:ACE)?S?|PLCS)\b)", after, re.I)
                if k:
                    n = int(k.group(1) or k.group(2))
            ops[op] += n
            break  # one operation per segment; the more specific patterns come first
    return [{"op": k, "count": v} for k, v in ops.items()]


def machining_text(ops: list[dict]) -> str:
    return ", ".join(f"{o['count']}X {OP_LABELS.get(o['op'], o['op'])}" if o.get("count", 1) != 1 else OP_LABELS.get(o["op"], o["op"])
                     for o in ops or [])


# ---------------------------------------------------------------- catalog matching
HW_RULES: list[tuple[str, str]] = [
    ("ANCHOR", r"ANCHOR"),
    ("TNUT-DROP", r"DROP[- ]?IN|ROLL[- ]?IN|T[- ]?NUT.{0,20}DROP"),
    ("TNUT-ECON", r"ECON(?:OMY|\.)?\s*(?:T[- ]?NUT|TEE)"),
    ("TNUT-SLIDE", r"SLIDE[- ]?IN|\bT[- ]?NUTS?\b|TEE\s*NUT"),
    ("BRKT-GUSSET", r"GUSSET"),
    ("BRKT-5H", r"5[- ]?HOLE.{0,20}(?:CORNER|BRACKET)|(?:CORNER|BRACKET).{0,20}5[- ]?HOLE"),
    ("BRKT-4H", r"CORNER\s*BRACKET|INSIDE\s*CORNER|\bBRACKET"),
    ("PLATE-T", r"\bT[- ]?(?:JOINING\s*)?PLATE|TEE\s*(?:JOINING\s*)?PLATE|JOINING\s*PLATE,?\s*T\b"),
    ("PLATE-L", r"\bL[- ]?(?:JOINING\s*)?PLATE|JOINING\s*PLATE,?\s*L\b|CORNER\s*PLATE"),
    ("PLATE-FLAT", r"JOINING\s*(?:PLATE|STRIP)|FLAT\s*PLATE|STRAIGHT\s*PLATE"),
    ("ENDCAP", r"END\s*CAP"),
    ("CASTER-LOCK", r"CASTER.{0,30}(?:LOCK|BRAKE)|(?:LOCKING|BRAKE).{0,20}CASTER"),
    ("CASTER", r"CASTER"),
    ("FOOT", r"LEVEL+ING\s*(?:FOOT|FEET|PAD|MOUNT)|\bFOOT\b|\bFEET\b"),
    ("RETAINER", r"RETAIN|GASKET"),
    ("HINGE", r"HINGE"),
    ("HANDLE", r"HANDLE|\bPULL\b"),
    ("SCREW-BHSCS", r"BHSCS|SHCS|FHSCS|BUTTON\s*HEAD|SOCKET\s*(?:HEAD|CAP)|\bSCREWS?\b|\bBOLTS?\b"),
]
FAMILY_HINTS: list[tuple[str, str]] = [
    (r"\b10\s*SERIES|\b1/4\s*-\s*20\b", "10"),
    (r"\b15\s*SERIES|\b5/16\s*-\s*18\b", "15"),
    (r"\b(?:20|25|30)\s*(?:MM\s*)?SERIES|\bM[45]\b|\bM6\b|\bHFS[56]\b", "MS"),
    (r"\b(?:40|45)\s*(?:MM\s*)?SERIES|\bM8\b|\bHFS8\b", "ML"),
]
PANEL_WORDS = [("PC", r"POLY\s*-?\s*CARB|POLYCARBONATE|LEXAN|MAKROLON|\bPC\b"),
               ("ACR", r"ACRYLIC|PLEXI|PMMA|PERSPEX"),
               ("ACM", r"\bACM\b|DIBOND|ALUMINUM\s*COMPOSITE|ALUM(?:INUM)?\.?\s*COMP")]
DIMS = re.compile(rf"((?:\d+\s+\d+/\d+|\d+/\d+|{_NUM}))\s*(?:\"|IN\b|MM\b)?\s*[X×]\s*((?:\d+\s+\d+/\d+|\d+/\d+|{_NUM}))\s*(?:\"|IN\b|MM\b)?(?:\s*[X×]\s*((?:\d+\s+\d+/\d+|\d+/\d+|{_NUM}))\s*(?:\"|IN\b|MM\b)?)?", re.I)


def _profile_index(cat: dict) -> list[tuple[re.Pattern, str]]:
    ids = sorted((p["id"] for p in cat["profiles"]), key=len, reverse=True)
    out = []
    for pid in ids:
        pat = re.escape(pid).replace(r"\-", r"[-\s]?")
        out.append((re.compile(rf"(?<![\w.-]){pat}(?![\w.]|-[A-Z0-9])", re.I), pid))
    return out


def match_profile(text: str, cat: dict, metric_doc: bool = False) -> tuple[str | None, float, tuple[int, int] | None, str]:
    """Find a profile part number in text -> (catalog id, confidence, span, note)."""
    ids = {p["id"] for p in cat["profiles"]}
    for rx, pid in _profile_index(cat):
        m = rx.search(text)
        if not m:
            continue
        note = ""
        conf = 0.95
        # Bare 2020 / 3030 exist as fractional and metric profiles. A metric drawing probably means the metric one.
        metric_twin = f"{pid[:2]}-{pid}" if re.fullmatch(r"\d{4}", pid) else None
        if metric_twin in ids and metric_doc:
            return metric_twin, 0.6, m.span(), f"'{pid}' on a metric drawing read as {metric_twin}; check it."
        if metric_twin in ids:
            conf, note = 0.75, f"'{pid}' read as the fractional {pid}; use {metric_twin} if it is metric."
        return pid, conf, m.span(), note
    # Bare metric numbers without the series prefix, e.g. 4040 -> 40-4040
    for m in re.finditer(r"(?<![\w.-])(\d{2})(\d{2})(?![\w.])", text):
        cand = f"{m.group(1)}-{m.group(1)}{m.group(2)}"
        if cand in ids:
            return cand, 0.7, m.span(), f"'{m.group(0)}' read as {cand}."
    # Size callouts such as "40 X 40 EXTRUSION" or "T-SLOT 30x30"
    if re.search(r"EXTRU|PROFILE|T[- ]?SLOT|FRAMING", text, re.I):
        m = re.search(r"(?<![\d.])(20|25|30|40|45)\s*(?:MM)?\s*[X×]\s*(20|25|30|40|45|60|80|90)\s*(?:MM)?(?![\d.])", text, re.I)
        if m:
            cand = f"{m.group(1)}-{m.group(1)}{m.group(2)}"
            if cand in ids:
                return cand, 0.5, m.span(), f"Profile size {m.group(0)} read as {cand}; confirm the series."
    return None, 0.0, None, ""


def _family_hint(text: str) -> str | None:
    for pat, fam in FAMILY_HINTS:
        if re.search(pat, text, re.I):
            return fam
    return None


def match_hardware(text: str, family: str | None, hint_text: str = "") -> tuple[str | None, float, str]:
    for base, pat in HW_RULES:
        if re.search(pat, text, re.I):
            hint = _family_hint(text) or _family_hint(hint_text)
            fam = hint or family or "15"
            conf = 0.85 if hint else (0.65 if family else 0.4)
            note = "" if hint else (f"Series taken from the build's profiles ({FAMILIES[fam].split(' (')[0]})." if family
                                    else "Series unknown; assumed 15 Series. Pick the right one.")
            return f"{base}-{fam}", conf, note
    return None, 0.0, ""


def _dims_in(m: re.Match, mm: bool) -> list[float]:
    vals = [_frac(g) for g in m.groups() if g]
    return [v / MM_PER_IN if mm else v for v in vals]


def match_panel(text: str, cat: dict, default_unit: str = "in") -> tuple[str | None, float, dict, str]:
    mat = next((k for k, pat in PANEL_WORDS if re.search(pat, text, re.I)), None)
    if not mat and not re.search(r"\bPANEL\b|\bSHEET\b", text, re.I):
        return None, 0.0, {}, ""
    m = DIMS.search(text)
    dims: dict = {}
    mm = bool(re.search(r"\bMM\b", text, re.I)) or default_unit == "mm"
    thk = None
    if m:
        vals = _dims_in(m, mm)
        if len(vals) == 3:
            vals.sort()
            thk, dims = vals[0], {"width_in": round(vals[1], 3), "height_in": round(vals[2], 3)}
        else:
            dims = {"width_in": round(vals[0], 3), "height_in": round(vals[1], 3)}
    t = re.search(rf"((?:\d+/\d+|{_NUM}))\s*(\"|IN\b|MM\b)?\s*(?:THK|THICK)", text, re.I)
    if t:
        thk = _frac(t.group(1)) / (MM_PER_IN if (t.group(2) or "").upper() == "MM" else 1)
    if not mat:
        return None, 0.0, dims, "Panel material not recognized."
    options = [p for p in cat["panels"] if p["material"] == mat and p["thickness_in"]]
    if not options:
        return None, 0.0, dims, ""
    if thk is None:
        best = options[0]
        note = f"Thickness not given; assumed {best['name']}."
        conf = 0.5
    else:
        best = min(options, key=lambda p: abs(p["thickness_in"] - thk))
        conf = 0.9 if abs(best["thickness_in"] - thk) < 0.02 else 0.6
        note = "" if conf > 0.8 else f"{PANEL_MATERIALS[mat]} {thk:.3f} in matched to nearest {best['name']}."
    if not dims:
        conf = min(conf, 0.5)
        note = (note + " Panel size missing: enter width and height.").strip()
    return best["id"], conf, dims, note


# ---------------------------------------------------------------- line building
def _blank_line(**kw) -> dict:
    base = {"item": "", "qty": 1, "kind": "unmatched", "catalog_id": "", "part_number": "", "description": "",
            "length_in": None, "width_in": None, "height_in": None, "machining": [], "notes": "", "confidence": 0.0,
            "unit_price": None, "source": "", "raw": ""}
    base.update(kw)
    return base


def make_line(*, item: str = "", qty: int | None = 1, part: str = "", desc: str = "", length: str | float | None = None,
              machining: str = "", notes: str = "", default_unit: str = "in", cat: dict, metric_doc: bool = False,
              family: str | None = None, source: str = "", raw: str = "") -> dict:
    """Match one BOM/cut-list row to the catalog."""
    text = f"{part} {desc}".strip()
    line = _blank_line(item=str(item or ""), qty=int(qty or 1), part_number=part.strip(), description=desc.strip(),
                       machining=parse_machining(machining), notes=notes.strip(), source=source, raw=raw or text)
    if isinstance(length, (int, float)):
        line["length_in"] = round(float(length), 4)
    elif length not in (None, ""):
        line["length_in"] = parse_length(str(length), default_unit)
    pid, conf, _, note = match_profile(text, cat, metric_doc)
    if pid:
        line.update(kind="profile", catalog_id=pid, confidence=conf)
        if note:
            line["notes"] = (line["notes"] + " " + note).strip()
        if not line["length_in"]:
            line["confidence"] = min(conf, 0.5)
        return line
    hid, conf, note = match_hardware(text, family, notes)
    if hid:
        line.update(kind="hardware", catalog_id=hid, confidence=conf, length_in=None)
        if note:
            line["notes"] = (line["notes"] + " " + note).strip()
        return line
    pan, conf, dims, note = match_panel(text, cat, default_unit)
    if pan:
        line.update(kind="panel", catalog_id=pan, confidence=conf, length_in=None, **dims)
        if note:
            line["notes"] = (line["notes"] + " " + note).strip()
        return line
    if dims:
        line.update(dims)
    if note:
        line["notes"] = (line["notes"] + " " + note).strip()
    return line


def _resolve_families(lines: list[dict], cat: dict, raw_rows: list[dict]) -> None:
    """Second pass: hardware with no series hint takes the build's main profile family."""
    fams = Counter()
    by_id = {p["id"]: p for p in cat["profiles"]}
    for ln in lines:
        if ln["kind"] == "profile" and by_id.get(ln["catalog_id"], {}).get("family"):
            fams[by_id[ln["catalog_id"]]["family"]] += ln["qty"]
    if not fams:
        return
    main = fams.most_common(1)[0][0]
    for ln, row in zip(lines, raw_rows):
        if ln["kind"] == "hardware" and ln["confidence"] < 0.8:
            hid, conf, note = match_hardware(f"{row.get('part', '')} {row.get('desc', '')}", main, ln.get("notes", ""))
            if hid:
                ln["catalog_id"], ln["confidence"] = hid, conf
                ln["notes"] = re.sub(r"Series (?:unknown|taken)[^.]*\.(?: Pick the right one\.)?", "", ln["notes"]).strip()
                ln["notes"] = (ln["notes"] + " " + note).strip()


# ---------------------------------------------------------------- text (drawing PDF) parsing
HEADER_WORDS = {
    "item": r"\bITEM(?:\s*NO\.?)?\b|\bFIND\s*NO\b|\bBALLOON\b|\bMARK\b",
    "qty": r"\bQTY\.?\b|\bQUANTITY\b|\bQ'?TY\b",
    "part": r"\bPART\s*(?:NO\.?|NUMBER|#)?|\bP/N\b",
    "desc": r"\bDESCRIPTION\b|\bDESC\.?\b",
    "length": r"\bLENGTH\b|\bCUT\s*LENGTH\b|\bLEN\.?\b",
    "machining": r"\bMACHINING\b|\bOPERATIONS?\b|\bFAB(?:RICATION)?\b|\bNOTES?\b",
}
STOP = re.compile(r"^\s*(?:NOTES?\s*:?$|GENERAL\s+NOTES|UNLESS\s+OTHERWISE|DRAWN\b|CHECKED\b|TITLE\b|SCALE\b|SHEET\b|DWG\b|SIZE\s+CAGE|DISTRIBUTION|WARNING)", re.I)
ROW = re.compile(r"^\s*([A-Z]{1,2}|\d{1,3})\s+(.+)$", re.I)  # "1. TEXT" is a note, not a row


def _header_cols(line: str) -> list[tuple[int, str]] | None:
    """Header row -> [(column start, key)] in order, or None when the line is not a BOM/cut list header."""
    found = []
    for key, pat in HEADER_WORDS.items():
        m = re.search(pat, line, re.I)
        if m:
            found.append((m.start(), key))
    found.sort()
    keys = [k for _, k in found]
    if "qty" in keys and len(keys) >= 3 and ("part" in keys or "desc" in keys):
        return found
    return None


def _header(line: str) -> list[str] | None:
    cols = _header_cols(line)
    return [k for _, k in cols] if cols else None


def _columnar_row(raw: str, cols: list[tuple[int, str]], default_unit: str) -> dict | None:
    """Split a layout-preserved row into cells (runs separated by 2+ spaces) under the header columns."""
    cells = [(m.start(), m.group(0)) for m in re.finditer(r"\S+(?: \S+)*", raw)]
    if len(cells) < 2 or len(cols) < 2:
        return None
    out: dict[str, str] = {}
    for start, text in cells:
        key = None
        for cstart, k in cols:
            if cstart <= start + 2:
                key = k
        if key is None:
            return None
        out[key] = (out.get(key, "") + " " + text).strip()
    q = out.get("qty", "")
    if q and not re.fullmatch(r"\d{1,5}", q):
        return None
    if not (out.get("part") or out.get("desc")):
        return None
    length = parse_length(out["length"], default_unit) if out.get("length") else None
    if out.get("length") and length is None:
        return None
    return {"item": out.get("item", ""), "part": out.get("part", ""), "desc": out.get("desc", ""), "qty": int(q) if q else 1,
            "length": length, "machining": out.get("machining", "")}


def _strip_spans(text: str, spans: list[tuple[int, int]]) -> str:
    out = list(text)
    for a, b in spans:
        for i in range(a, b):
            out[i] = " "
    return "".join(out)


def _parse_row(body: str, order: list[str], cat: dict, default_unit: str, metric_doc: bool) -> dict:
    """Split one table row (item number already removed) into part, desc, qty, length and machining text."""
    m = OP_START.search(body)
    pre, mach = (body[: m.start()], body[m.start():]) if m else (body, "")
    # A count right before the first op ("2X ACCESS HOLE") belongs to the op
    c = re.search(r"(?:\(\s*\d{1,3}\s*\)|\b\d{1,3}\s*[X×])\s*$", pre, re.I)
    if c and mach:
        mach, pre = pre[c.start():] + mach, pre[: c.start()]
    qty = None
    spans: list[tuple[int, int]] = []
    if [k for k in order if k != "item"][:1] == ["qty"]:  # ITEM QTY PART ... order
        q = re.match(r"\s*(\d{1,4})(?![\w./])", pre)
        if q:
            qty = int(q.group(1))
            spans.append(q.span())
    pre2 = _strip_spans(pre, spans)
    pid, _, pspan, _ = match_profile(pre2, cat, metric_doc)
    part = ""
    if pspan:
        part = pre2[pspan[0]: pspan[1]].strip()
        spans.append(pspan)
    else:
        tok = re.search(r"(?<!\S)([A-Z0-9]+(?:-[A-Z0-9]+)+)(?!\S)", pre2, re.I)
        if tok and re.search(r"[A-Z]", tok.group(1), re.I) and re.search(r"\d", tok.group(1)):
            part = tok.group(1)
            spans.append(tok.span(1))
    pre3 = _strip_spans(pre, spans)
    # Panel/size dimensions stay in the description and are not lengths or quantities
    dim_spans = [d.span() for d in DIMS.finditer(pre3)]
    search = _strip_spans(pre3, dim_spans)
    # Nominal sizes such as 1.5" X 1.5" were removed above; machining locations live in `mach`
    has_len_col = "length" in order
    length, _, lspan = find_length(search, default_unit)
    if lspan:
        spans.append(lspan)
        search = _strip_spans(search, [lspan])
    if qty is None:
        q = re.search(r"\bQTY\.?\s*:?\s*(\d{1,4})\b|\b(\d{1,4})\s*(?:PCS|EA|X)\b", search, re.I)
        if q:
            qty = int(q.group(1) or q.group(2))
            spans.append(q.span())
            search = _strip_spans(search, [q.span()])
        else:
            ints = list(BARE_INT.finditer(search))
            if ints:
                q = ints[-1]
                if has_len_col and length is None and len(ints) >= 2 and order.index("length") > order.index("qty"):
                    # ... QTY LENGTH with a bare integer length
                    length = parse_length(q.group(1), default_unit)
                    spans.append(q.span())
                    q = ints[-2]
                qty = int(q.group(1))
                spans.append(q.span())
    if length is None and has_len_col:  # a LENGTH column with a bare whole number (common in mm)
        ints = list(BARE_INT.finditer(_strip_spans(search, spans)))
        if ints:
            length = parse_length(ints[-1].group(1), default_unit)
            spans.append(ints[-1].span())
    desc =re.sub(r"\s+", " ", _strip_spans(pre, spans)).strip(" ,;-")
    return {"part": part, "desc": desc, "qty": qty or 1, "length": length, "machining": mach.strip()}


def parse_bom_text(text: str, config: dict | None = None) -> dict:
    """Find the BOM / parts list and cut list in drawing text and match each row to the catalog.

    Returns {lines, unmatched, warnings, units, metric_doc}. Cut list rows replace BOM profile rows that
    have no length; quantities are cross-checked.
    """
    return _parse_text_with(text, None, config)


def _parse_text_with(text: str, cat: dict, config: dict | None = None) -> dict:
    cat = cat or catalog(config)
    t = text or ""
    metric_doc = bool(METRIC_DOC.search(t))
    default_unit = "mm" if metric_doc else "in"
    lines = t.splitlines()
    bom_rows: list[dict] = []
    cut_rows: list[dict] = []
    warnings: list[str] = []
    section = None
    order: list[str] = []
    cols: list[tuple[int, str]] = []
    misses = 0
    current: list[dict] | None = None
    for ln in lines:
        s = ln.strip()
        if not s:
            continue
        if re.search(r"\bCUT\s*LIST\b|\bCUTTING\s*LIST\b", s, re.I):
            section, current, order, cols, misses = "cut", cut_rows, [], [], 0
            h = _header_cols(ln)
            if h:
                cols, order = h, [k for _, k in h]
            continue
        if re.fullmatch(r"(?:BILL\s+OF\s+MATERIALS?|PARTS?\s+LIST|BOM)\s*:?", s, re.I):
            section, current, order, cols, misses = "bom", bom_rows, [], [], 0
            continue
        h = _header_cols(ln)
        if h:
            if section != "cut":
                section, current = "bom", bom_rows
            cols, order, misses = h, [k for _, k in h], 0
            continue
        if section is None:
            continue
        if STOP.search(s):
            section, current = None, None
            continue
        m = ROW.match(s)
        label, body = "", s
        # Cut list rows may start with a mark (A, B, 1, 2) or straight with the profile part number
        if m and (m.group(1).isdigit() or section == "cut"):
            label, body = m.group(1), m.group(2)
        if section == "cut":
            is_row = bool(match_profile(body, cat, metric_doc)[0])  # cut lists only carry profiles
        else:
            is_row = bool(label) and label.isdigit() and bool(re.search(r"[A-Z]{2,}|\d{4}", body, re.I))
        if is_row:
            row = _columnar_row(ln, cols, default_unit) if cols else None
            if row is None:
                row = _parse_row(body, order, cat, default_unit, metric_doc)
                row["item"] = label
            row.update(raw=s, section=section)
            current.append(row)
            misses = 0
            continue
        # A wrapped continuation line with machining callouts belongs to the row above
        if current and OP_START.search(s) and not ROW.match(s):
            current[-1]["machining"] = (current[-1]["machining"] + ", " + s).strip(", ")
            continue
        misses += 1
        if misses >= 3:
            section, current = None, None

    # Merge cut list into BOM
    rows = list(bom_rows)
    if cut_rows:
        cut_profiles = Counter()
        for r in cut_rows:
            pid = match_profile(f"{r['part']} {r['desc']}", cat, metric_doc)[0]
            cut_profiles[pid] += r["qty"]
        kept = []
        for r in rows:
            pid = match_profile(f"{r['part']} {r['desc']}", cat, metric_doc)[0]
            if pid and pid in cut_profiles and r["length"] is None:
                if r["qty"] != cut_profiles[pid]:
                    warnings.append(f"BOM lists {r['qty']} of {pid} but the cut list has {cut_profiles[pid]} pieces. The cut list was used.")
                continue
            kept.append(r)
        rows = kept + cut_rows
    out_lines = [make_line(item=r["item"], qty=r["qty"], part=r["part"], desc=r["desc"], length=r["length"], machining=r["machining"],
                           default_unit=default_unit, cat=cat, metric_doc=metric_doc, source=r["section"], raw=r["raw"]) for r in rows]
    _resolve_families(out_lines, cat, rows)
    if not rows:
        warnings.append("No BOM or cut list table was found in the text. Enter the lines by hand or upload the BOM as a spreadsheet.")
    if metric_doc:
        warnings.append("The drawing says dimensions are in millimeters: bare lengths were read as mm and converted to inches.")
    warnings += _line_warnings(out_lines, cat)
    return {"lines": out_lines, "unmatched": [l for l in out_lines if l["kind"] == "unmatched"], "warnings": warnings,
            "units": default_unit, "metric_doc": metric_doc}


def _line_warnings(lines: list[dict], cat: dict) -> list[str]:
    w = []
    prof = {p["id"]: p for p in cat["profiles"]}
    for ln in lines:
        if ln["kind"] == "profile":
            if not ln["length_in"]:
                w.append(f"{ln['catalog_id']} (item {ln['item'] or '?'}) has no cut length.")
            elif prof.get(ln["catalog_id"], {}).get("metric") is False and ln["length_in"] > 250:
                w.append(f"{ln['catalog_id']} length {ln['length_in']} in is longer than a full stick. Was it in mm?")
            elif prof.get(ln["catalog_id"], {}).get("metric") and ln["length_in"] > 250:
                w.append(f"{ln['catalog_id']} length {ln['length_in']} in is longer than a full stick. Check the units.")
    n = sum(1 for l in lines if l["kind"] == "unmatched")
    if n:
        w.append(f"{n} line(s) did not match the catalog. Map them or enter a unit price, or they are left out of the price.")
    low = [l for l in lines if l["kind"] != "unmatched" and l["confidence"] < 0.7]
    if low:
        w.append(f"{len(low)} line(s) matched with low confidence. Check them.")
    return w


# ---------------------------------------------------------------- spreadsheet parsing
COL_PATTERNS = [
    ("length", r"^(?:CUT\s*)?LENGTH|^LEN\b|^CUT\s*LEN|^SIZE$|^CUT$"),
    ("qty", r"^QTY|^QUANTITY|^Q'?TY|^COUNT$|^PCS$"),
    ("part", r"^PART|^P/?N\b|^MFG\s*(?:PART|P/?N)|^CATALOG|^SKU|^ITEM\s*(?:NUMBER|#)$|^MODEL"),
    ("desc", r"^DESC|^NAME$|^TITLE$"),
    ("machining", r"^MACHIN|^OPERATION|^OPS?$|^FAB"),
    ("notes", r"^NOTE|^REMARK|^COMMENT"),
    ("item", r"^ITEM|^FIND|^NO\.?$|^#$|^MARK|^LINE$"),
    ("unit", r"^UNITS?$|^UOM$"),
]


def _map_header(cells: list) -> dict[str, int]:
    found: dict[str, int] = {}
    for i, c in enumerate(cells):
        name = str(c or "").strip().upper()
        if not name:
            continue
        for key, pat in COL_PATTERNS:
            if key not in found and re.search(pat, name):
                found[key] = i
                if key == "length" and re.search(r"\bMM\b|\(MM\)|MILLIM", name):
                    found["_mm"] = 1
                break
    return found


def _read_rows(filename: str, data: bytes) -> list[list]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise ExtrusionError(f"Could not read the spreadsheet: {exc}") from exc
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if name.endswith((".csv", ".txt", ".tsv")):
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.excel_tab if name.endswith(".tsv") or ("\t" in text.split("\n", 1)[0] and "," not in text.split("\n", 1)[0]) else csv.excel
        return [r for r in csv.reader(io.StringIO(text), dialect)]
    raise ExtrusionError("Upload a PDF, CSV or XLSX file.")


def parse_bom_table(file: str | Path | bytes, filename: str = "", config: dict | None = None) -> dict:
    """Parse a CSV/XLSX BOM or cut list with flexible headers (Qty, Part, Description, Length, Cut length, Machining, Notes)."""
    if isinstance(file, (str, Path)):
        filename = filename or Path(file).name
        data = Path(file).read_bytes()
    else:
        data = file
    rows = _read_rows(filename, data)
    cat = catalog(config)
    hdr_i, cols = None, {}
    for i, r in enumerate(rows[:30]):
        c = _map_header(r)
        if len([k for k in c if not k.startswith("_")]) >= 2 and ("qty" in c or "length" in c) and ("part" in c or "desc" in c):
            hdr_i, cols = i, c
            break
    if hdr_i is None:
        raise ExtrusionError("No header row found. The sheet needs columns such as Qty, Part, Description and Length.")
    head_text = " ".join(str(x or "") for r in rows[:hdr_i] for x in r)
    metric_doc = bool(cols.get("_mm")) or bool(METRIC_DOC.search(head_text))
    default_unit = "mm" if metric_doc else "in"

    def cell(r: list, key: str) -> str:
        i = cols.get(key)
        if i is None or i >= len(r) or r[i] is None:
            return ""
        v = r[i]
        if isinstance(v, float) and v.is_integer() and key in ("qty", "item"):
            v = int(v)
        return str(v).strip()

    raw_rows, lines, warnings = [], [], []
    for r in rows[hdr_i + 1:]:
        if not any(str(x or "").strip() for x in r):
            continue
        part, desc = cell(r, "part"), cell(r, "desc")
        if not (part or desc):
            continue
        q = cell(r, "qty")
        try:
            qty = int(float(q)) if q else 1
        except ValueError:
            qty = 1
            warnings.append(f"Quantity '{q}' for '{part or desc}' is not a number; used 1.")
        unit = cell(r, "unit").lower()
        row_unit = "mm" if unit in ("mm", "millimeter", "millimeters") else ("in" if unit in ("in", "inch", "inches") else default_unit)
        length = cell(r, "length")
        rr = {"part": part, "desc": desc}
        raw_rows.append(rr)
        lines.append(make_line(item=cell(r, "item"), qty=max(qty, 1), part=part, desc=desc, length=length or None,
                               machining=cell(r, "machining") + ("," + cell(r, "notes") if OP_START.search(cell(r, "notes")) else ""),
                               notes=cell(r, "notes"), default_unit=row_unit, cat=cat, metric_doc=metric_doc, source="table",
                               raw=" | ".join(str(x) for x in r if x not in (None, ""))))
    _resolve_families(lines, cat, raw_rows)
    if metric_doc:
        warnings.append("Lengths were read as millimeters and converted to inches.")
    warnings += _line_warnings(lines, cat)
    return {"lines": lines, "unmatched": [l for l in lines if l["kind"] == "unmatched"], "warnings": warnings,
            "units": default_unit, "metric_doc": metric_doc}


# ---------------------------------------------------------------- AI read
AI_PROMPT = """This PDF is a drawing of a T-slot aluminum extrusion build (frame, cart, enclosure or stand).
Transcribe its bill of materials and cut list. Return only JSON:
{"units": "in" or "mm",
 "lines": [{"item": "", "qty": 1, "part_number": "", "description": "", "length": "cut length as written with unit, empty for hardware",
            "machining": "machining callouts as written (end taps, counterbores, access holes, miters)", "notes": ""}]}
One entry per BOM row; when a cut list gives lengths for a profile, give one entry per cut length instead of the BOM row.
Use empty strings when something is not on the drawing. Do not guess."""


def _claude_read_bom(data: bytes) -> dict:
    """Send the PDF to Claude and return its JSON transcription. Tests replace this function."""
    import anthropic

    client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)
    msg = client.messages.create(
        model=ANTHROPIC_MODEL, max_tokens=8000,
        messages=[{"role": "user", "content": [
            {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": base64.b64encode(data).decode()}},
            {"type": "text", "text": AI_PROMPT},
        ]}],
    )
    text = "".join(b.text for b in msg.content if getattr(b, "type", "") == "text")
    m = re.search(r"\{.*\}", text, re.S)
    if not m:
        raise ExtrusionError("Claude did not return JSON for the drawing.")
    return json.loads(m.group(0))


def _lines_from_ai(ai: dict, cat: dict) -> list[dict]:
    unit = "mm" if str(ai.get("units") or "").lower().startswith("mm") else "in"
    rows, out = [], []
    for r in ai.get("lines") or []:
        try:
            qty = max(int(float(r.get("qty") or 1)), 1)
        except (TypeError, ValueError):
            qty = 1
        rows.append({"part": str(r.get("part_number") or ""), "desc": str(r.get("description") or "")})
        out.append(make_line(item=str(r.get("item") or ""), qty=qty, part=rows[-1]["part"], desc=rows[-1]["desc"],
                             length=str(r.get("length") or "") or None, machining=str(r.get("machining") or ""),
                             notes=str(r.get("notes") or ""), default_unit=unit, cat=cat, metric_doc=unit == "mm", source="ai"))
    _resolve_families(out, cat, rows)
    return out


def ai_block_reason(text: str) -> str | None:
    """Reason a drawing must not go to an outside service (export-controlled or limited distribution), or None."""
    from .drawing import parse_text
    from .drawing_quote import ai_allowed

    return ai_allowed(parse_text(text or ""))


def pdf_text(path: str | Path) -> tuple[str, int]:
    """PDF text with the page layout kept, so table columns stay on one line. Falls back to plain extraction."""
    from .drawing import extract_pdf_text

    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        parts = []
        for p in reader.pages:
            try:
                parts.append(p.extract_text(extraction_mode="layout") or "")
            except Exception:  # noqa: BLE001  (older pypdf or an odd page)
                parts.append(p.extract_text() or "")
        text = "\n".join(parts)
        if len(re.findall(r"[A-Za-z0-9]", text)) >= 20:
            return text, len(reader.pages)
    except Exception:  # noqa: BLE001
        pass
    return extract_pdf_text(Path(path))


def parse_file(path: str | Path, filename: str = "", use_ai: bool = False, force: bool = False, config: dict | None = None) -> dict:
    """Parse an uploaded PDF drawing, CSV or XLSX into BOM lines. Raises ExtrusionError when the AI read is refused."""
    path = Path(path)
    filename = filename or path.name
    name = filename.lower()
    cat = catalog(config)
    if not name.endswith(".pdf"):
        out = parse_bom_table(path.read_bytes(), filename, config)
        out["source"] = {"filename": filename, "type": "table", "ai_used": False}
        out["drawing"] = {}
        return out
    from .drawing import DrawingError, parse_text

    try:
        text, pages = pdf_text(path)
    except DrawingError as exc:
        raise ExtrusionError(str(exc)) from exc
    has_text = len(re.findall(r"[A-Za-z0-9]", text)) >= 20
    out = _parse_text_with(text, cat, config) if has_text else {"lines": [], "unmatched": [], "warnings": [], "units": "in", "metric_doc": False}
    read = parse_text(text) if has_text else {}
    out["drawing"] = {k: read.get(k) for k in ("part_number", "drawing_number", "revision", "title")} if read else {}
    out["drawing"]["export_controlled"] = bool(read.get("export_controlled"))
    out["drawing"]["distribution"] = (read.get("distribution") or {}).get("letter") or ""
    out["source"] = {"filename": filename, "type": "pdf", "pages": pages, "text_found": has_text, "ai_used": False}
    if read:
        out["warnings"] = [w for w in read.get("warnings", []) if "Distribution" in w or "Export" in w] + out["warnings"]
    parsed_ok = any(l["kind"] != "unmatched" for l in out["lines"])
    if not parsed_ok and use_ai:
        blocked = ai_block_reason(text)
        if blocked and not force:
            raise ExtrusionError(blocked + " Enter the lines by hand, or set force only if you are sure you may share it.")
        if not ANTHROPIC_API_KEY:
            out["warnings"].insert(0, "Set ANTHROPIC_API_KEY to read this drawing with Claude. Enter the lines by hand for now.")
            return out
        try:
            ai = _claude_read_bom(path.read_bytes())
        except Exception as exc:  # noqa: BLE001  (network, quota, bad JSON)
            out["warnings"].insert(0, f"The AI read failed ({exc}). Enter the lines by hand.")
            return out
        lines = _lines_from_ai(ai, cat)
        out.update(lines=lines, unmatched=[l for l in lines if l["kind"] == "unmatched"])
        out["warnings"] = ["Lines read by Claude. Check every quantity and length against the drawing."] + _line_warnings(lines, cat)
        out["source"]["ai_used"] = True
    elif not parsed_ok and not has_text:
        out["warnings"].insert(0, "The drawing is scanned (no text layer). Enter the lines by hand"
                               + (", or read it with Claude (only for drawings you may share)." if ANTHROPIC_API_KEY else "."))
    return out


# ---------------------------------------------------------------- nesting
def nest_ffd(pieces: list[float], stock_len: float, kerf: float = 0.0, end_trim: float = 0.0) -> dict:
    """First-fit decreasing: pack cut lengths (inches) into sticks of stock_len.

    Each stick loses end_trim before cutting. A piece fits if it is no longer than what is left;
    every cut then also consumes one kerf. Pieces longer than the usable stick are returned in
    `oversize` (special order). Returns sticks [{pieces, used, drop}], counts and waste.
    """
    usable = stock_len - end_trim
    if usable <= 0:
        raise ExtrusionError("Stock length must be longer than the end trim.")
    sticks: list[dict] = []
    oversize: list[float] = []
    for p in sorted((float(x) for x in pieces), reverse=True):
        if p <= 0:
            continue
        if p > usable + 1e-9:
            oversize.append(p)
            continue
        for s in sticks:
            if s["left"] + 1e-9 >= p:
                s["pieces"].append(p)
                s["left"] = max(s["left"] - p - kerf, 0.0)
                break
        else:
            sticks.append({"pieces": [p], "left": max(usable - p - kerf, 0.0)})
    used = sum(sum(s["pieces"]) for s in sticks)
    bought = len(sticks) * stock_len
    out_sticks = [{"pieces": [round(x, 4) for x in s["pieces"]], "used_in": round(sum(s["pieces"]), 4),
                   "drop_in": round(s["left"], 4)} for s in sticks]
    return {
        "stock_length_in": stock_len, "sticks": len(sticks), "stick_detail": out_sticks,
        "bought_in": round(bought, 3), "used_in": round(used, 3),
        "waste_in": round(bought - used, 3), "waste_pct": round((bought - used) / bought * 100, 1) if bought else 0.0,
        "drop_in": round(sum(s["left"] for s in sticks), 3), "oversize": oversize,
    }


# ---------------------------------------------------------------- pricing
def _line_obj(category: str, item: str, cost: float, *, basis: str = "per_build", hours: float | None = None,
              rate: float | None = None, note: str = "") -> dict:
    return {"category": category, "item": item, "basis": basis, "hours": None if hours is None else round(hours, 3),
            "rate": rate, "cost": round(cost, 2), "note": note}


def _num(v, default: float = 0.0) -> float:
    try:
        return float(v if v not in (None, "") else default)
    except (TypeError, ValueError):
        raise ExtrusionError(f"Expected a number, got {v!r}")


def normalize_lines(lines: list[dict]) -> list[dict]:
    out = []
    for i, ln in enumerate(lines or []):
        ln = dict(_blank_line(), **(ln or {}))
        try:
            ln["qty"] = int(float(ln.get("qty") or 0))
        except (TypeError, ValueError):
            raise ExtrusionError(f"Line {i + 1}: quantity must be a whole number")
        for k in ("length_in", "width_in", "height_in", "unit_price"):
            if ln.get(k) in ("", None):
                ln[k] = None
            else:
                ln[k] = _num(ln[k])
        ops = []
        for o in ln.get("machining") or []:
            if isinstance(o, str):
                o = {"op": o, "count": 1}
            if o.get("op") and int(_num(o.get("count"), 1)) > 0:
                ops.append({"op": o["op"], "count": int(_num(o.get("count"), 1))})
        ln["machining"] = ops
        out.append(ln)
    return out


def price_build(lines: list[dict], config: dict | None = None, quantities: list[int] | None = None, options: dict | None = None) -> dict:
    """Price a T-slot build. Returns the pricing.estimate shape (per_part_lines are per build) plus nesting,
    cut list and purchase list data."""
    cfg = pricing.merged_config(config)  # saved overrides or a full config; both merge onto the defaults
    ext = cfg["extrusion"]
    rates = cfg["rates"]
    opts = dict(options or {})
    cat = catalog(cfg)
    prof = {p["id"]: p for p in cat["profiles"]}
    hw = {h["id"]: h for h in cat["hardware"]}
    pan = {p["id"]: p for p in cat["panels"]}
    mach_price = {m["id"]: m["price"] for m in cat["machining"]}
    lines = normalize_lines(lines)
    try:
        qtys = sorted({int(q) for q in (quantities or [1]) if int(q) > 0})
    except (TypeError, ValueError):
        raise ExtrusionError("Build quantities must be positive whole numbers")
    if not qtys:
        raise ExtrusionError("Give at least one build quantity")
    mode = opts.get("pricing_mode") or ext.get("pricing_mode", "cut_to_length")
    if mode not in ("cut_to_length", "stock"):
        raise ExtrusionError("pricing_mode must be cut_to_length or stock")
    warnings: list[str] = []
    assumptions: list[str] = [ext.get("note") or ""]
    per_build: list[dict] = []
    per_lot: list[dict] = []

    # ---- profiles
    pieces_by_profile: dict[str, list[tuple[float, int, dict]]] = defaultdict(list)
    cut_rows = []
    n_pieces = 0
    weight = 0.0
    for ln in lines:
        if ln["kind"] != "profile" or ln["qty"] <= 0:
            continue
        p = prof.get(ln["catalog_id"])
        if not p:
            warnings.append(f"Profile '{ln['catalog_id']}' is not in the catalog; line left out.")
            continue
        if not ln["length_in"] or ln["length_in"] <= 0:
            warnings.append(f"{p['id']} line {ln['item'] or ''} has no length; left out.".replace("  ", " "))
            continue
        pieces_by_profile[p["id"]].append((ln["length_in"], ln["qty"], ln))
        n_pieces += ln["qty"]
        if p.get("weight_lb_per_in"):
            weight += p["weight_lb_per_in"] * ln["length_in"] * ln["qty"]
        cut_rows.append(ln)

    stock_len = ext["stock_length_in"]
    kerf, trim = ext["kerf_in"], ext["end_trim_in"]

    def material_for(q: int) -> tuple[float, dict]:
        """Profile material cost for q builds and the nesting per profile."""
        total = 0.0
        nest = {}
        for pid, items in pieces_by_profile.items():
            p = prof[pid]
            sl = stock_len["metric" if p.get("metric") else "fractional"]
            lengths = [L for L, n, _ in items for _ in range(n * q)]
            nres = nest_ffd(lengths, sl, kerf, trim)
            nest[pid] = nres
            if mode == "stock":
                stick_price = sl * p["price_per_in"] * ext["full_stick_price_factor"]
                total += nres["sticks"] * stick_price + sum(nres["oversize"]) * p["price_per_in"]
            else:
                total += sum(lengths) * p["price_per_in"]
        return total, nest

    mat0, nest0 = material_for(qtys[0])
    for pid, items in pieces_by_profile.items():
        p = prof[pid]
        inches = sum(L * n for L, n, _ in items)
        if mode == "stock":
            nres = nest0[pid]
            sl = nres["stock_length_in"]
            cost = (nres["sticks"] * sl * p["price_per_in"] * ext["full_stick_price_factor"] + sum(nres["oversize"]) * p["price_per_in"]) / qtys[0]
            note = f"{nres['sticks']} stick(s) of {sl:g} in for {qtys[0]} build(s), {nres['waste_pct']}% drop and kerf"
        else:
            cost = inches * p["price_per_in"]
            note = f"{inches:.2f} in cut to length"
        per_build.append(_line_obj("material", f"{pid} profile", cost, rate=p["price_per_in"], note=note))
        if any(nest0[pid]["oversize"]):
            warnings.append(f"{pid}: {len(nest0[pid]['oversize'])} piece(s) longer than a {nest0[pid]['stock_length_in']:g} in stick. Special order or splice.")
    if mode == "stock" and len(qtys) > 1:
        assumptions.append("Stock mode: pieces for the whole lot are nested into full sticks, so material per build changes with quantity. "
                           f"The material lines show quantity {qtys[0]}; price breaks use each quantity's nesting.")

    if n_pieces:
        if mode == "stock":
            h = ext["inhouse_cut_minutes"] * n_pieces / 60
            per_build.append(_line_obj("cutting", f"saw cuts in house, {n_pieces} piece(s)", h * rates["fabrication"], hours=h, rate=rates["fabrication"]))
        else:
            per_build.append(_line_obj("cutting", f"supplier cut charge, {n_pieces} cut(s)", n_pieces * ext["cut_charge"], rate=ext["cut_charge"]))

    # ---- machining
    op_counts: Counter = Counter()
    for ln in cut_rows:
        for o in ln["machining"]:
            op_counts[o["op"]] += o["count"] * ln["qty"]
    for op, n in sorted(op_counts.items()):
        if op not in mach_price:
            warnings.append(f"Machining operation '{op}' has no price; left out.")
            continue
        per_build.append(_line_obj("machining", f"{OP_LABELS.get(op, op)} x {n}", n * mach_price[op], rate=mach_price[op]))

    # ---- hardware
    counts = Counter()
    hw_cost = 0.0
    for ln in lines:
        if ln["kind"] != "hardware" or ln["qty"] <= 0:
            continue
        h = hw.get(ln["catalog_id"])
        if not h:
            warnings.append(f"Hardware '{ln['catalog_id']}' is not in the catalog; line left out.")
            continue
        price = ln["unit_price"] if ln["unit_price"] is not None else h["unit_price"]
        hw_cost += ln["qty"] * price
        counts[h["assembly"]] += ln["qty"]
        per_build.append(_line_obj("hardware", f"{h['name']} x {ln['qty']}", ln["qty"] * price, rate=price))

    # ---- panels
    n_panels = 0
    for ln in lines:
        if ln["kind"] != "panel" or ln["qty"] <= 0:
            continue
        p = pan.get(ln["catalog_id"])
        if not p:
            warnings.append(f"Panel '{ln['catalog_id']}' is not in the catalog; line left out.")
            continue
        if not (ln["width_in"] and ln["height_in"]):
            warnings.append(f"Panel {p['name']} has no width and height; left out.")
            continue
        sqft = ln["width_in"] * ln["height_in"] / 144 * ln["qty"]
        cost = sqft * (1 + ext["panel_waste_factor"]) * p["price_per_sqft"]
        n_panels += ln["qty"]
        per_build.append(_line_obj("panels", f"{p['name']}, {ln['qty']} at {ln['width_in']:g} x {ln['height_in']:g} in", cost,
                                   rate=p["price_per_sqft"], note=f"{sqft:.2f} sq ft + {ext['panel_waste_factor']:.0%} waste"))

    # ---- custom priced lines (unmatched lines with a unit price)
    for ln in lines:
        if ln["kind"] in ("unmatched", "custom") and ln["unit_price"] is not None and ln["qty"] > 0:
            per_build.append(_line_obj("hardware", f"{ln['description'] or ln['part_number'] or 'custom item'} x {ln['qty']}",
                                       ln["qty"] * ln["unit_price"], rate=ln["unit_price"], note="price entered by hand"))
    skipped = [l for l in lines if l["kind"] == "unmatched" and l["unit_price"] is None and l["qty"] > 0]
    if skipped:
        warnings.append(f"{len(skipped)} unmatched line(s) are not priced: " + ", ".join((l["part_number"] or l["description"] or "?")[:30] for l in skipped[:5]))

    # ---- assembly
    am = ext["assembly_minutes"]
    joints = int(_num(opts.get("joints"))) if opts.get("joints") not in (None, "") else counts["joint"]
    if joints == 0 and n_pieces > 1 and opts.get("joints") in (None, ""):
        joints = max(n_pieces - 1, 0)
        assumptions.append(f"No joining hardware listed: assumed {joints} joint(s) (pieces minus one). Set joints to override.")
    minutes = am["per_build"] + joints * am["per_joint"] + counts["fastener"] * am["per_fastener"] + n_panels * am["per_panel"] + counts["accessory"] * am["per_accessory"]
    if lines:
        h = minutes / 60
        per_build.append(_line_obj("assembly", f"assembly: {joints} joint(s), {counts['fastener']} fastener(s), {n_panels} panel(s), {counts['accessory']} accessory item(s)",
                                   h * rates["assembly"], hours=h, rate=rates["assembly"]))
        h = ext["inspection_minutes_per_build"] / 60
        per_build.append(_line_obj("inspection", "build inspection (squareness, fasteners, dimensions)", h * rates["inspection"], hours=h, rate=rates["inspection"]))

    # ---- packaging
    pk_level = opts.get("packaging_level") or "commercial"
    pk = cfg["packaging"].get(pk_level)
    if pk is None:
        raise ExtrusionError(f"packaging_level must be one of {list(cfg['packaging'])}")
    per_build.append(_line_obj("packaging", f"{pk_level} packaging and crating", pk["per_part"] + ext["crate_per_build"],
                               note=f"${pk['per_part']:.2f} packaging + ${ext['crate_per_build']:.2f} crate"))

    # ---- per lot
    h = ext["kitting_hours_per_lot"]
    per_lot.append(_line_obj("setup", "cut list, kitting and staging", h * rates["fabrication"], basis="per_lot", hours=h, rate=rates["fabrication"]))
    if ext["supplier_order_charge"]:
        per_lot.append(_line_obj("material", "extrusion supplier order (inbound freight, handling)", ext["supplier_order_charge"], basis="per_lot"))
    if pk["per_lot"]:
        per_lot.append(_line_obj("packaging", f"{pk_level} lot labels and marking", pk["per_lot"], basis="per_lot"))
    ins = cfg["inspection"]
    lead_extra = 0
    if opts.get("first_article"):
        h = ins["first_article_hours"]
        per_lot.append(_line_obj("inspection", "first article inspection and report", h * rates["inspection"], basis="per_lot", hours=h, rate=rates["inspection"]))
        lead_extra += cfg["lead_time"]["first_article_days"]
    if opts.get("certificate_of_conformance", True):
        per_lot.append(_line_obj("inspection", "certificate of conformance", ins["cert_per_lot"], basis="per_lot"))
    freight = _num(opts.get("freight_per_lot"), cfg["default_freight_per_lot"])
    per_lot.append(_line_obj("freight", "outbound freight", freight, basis="per_lot"))

    # ---- totals
    ga = _num(opts.get("ga_rate"), cfg["ga_rate"])
    profit = _num(opts.get("profit_rate"), cfg["profit_rate"])
    build_cost_ex_mat = sum(l["cost"] for l in per_build if not (l["category"] == "material"))
    lot_cost = sum(l["cost"] for l in per_lot)
    breaks = []
    nest_by_q = {}
    for q in qtys:
        mat, nest = (mat0, nest0) if q == qtys[0] else material_for(q)
        nest_by_q[q] = nest
        cost = (build_cost_ex_mat * q) + mat + lot_cost
        price = max(cost * (1 + ga) * (1 + profit), cfg["min_lot_charge"])
        days = cfg["lead_time"]["base_days"] + int(ext["supplier_lead_days"]) + lead_extra + math.ceil(q / max(ext["builds_per_day"], 1))
        breaks.append({"quantity": q, "total_cost": round(cost, 2), "unit_cost": round(cost / q, 2),
                       "unit_price": round(price / q, 2), "total_price": round(price, 2),
                       "margin_pct": round((price - cost) / price * 100, 1) if price else 0.0, "lead_time_days": days})
    if not pieces_by_profile:
        warnings.append("No priced extrusion profiles in this build.")
    assumptions.append(f"Profile pricing: {'full sticks nested first-fit decreasing, cut in house' if mode == 'stock' else 'supplier cuts to length, price per inch plus a cut charge'}.")
    assumptions.append("Catalog prices are placeholders until you enter your distributor's prices in the Extrusion rates panel.")
    build_qty = int(opts.get("cut_list_builds") or qtys[0])
    return {
        "kind": "extrusion_build",
        "part": {k: opts.get(k) for k in ("name", "part_number", "nsn") if opts.get(k)},
        "material": "T-slot aluminum extrusion build",
        "pricing_mode": mode,
        "part_weight_lb": round(weight, 2),
        "stock_volume_in3": 0.0,
        "per_part_lines": per_build,
        "per_lot_lines": per_lot,
        "per_part_cost": round(sum(l["cost"] for l in per_build), 2),
        "per_lot_cost": round(lot_cost, 2),
        "ga_rate": ga,
        "profit_rate": profit,
        "price_breaks": breaks,
        "assumptions": [a for a in assumptions if a],
        "warnings": warnings,
        "config_note": ext.get("note", ""),
        "nesting": {"quantity": qtys[0], "profiles": {pid: {k: v for k, v in n.items()} for pid, n in nest0.items()}},
        "nesting_by_quantity": {str(q): {pid: {"sticks": n["sticks"], "waste_pct": n["waste_pct"]} for pid, n in nest.items()} for q, nest in nest_by_q.items()},
        "purchase_list": purchase_list(lines, cfg, build_qty, mode=mode),
        "counts": {"pieces": n_pieces, "joints": joints, "fasteners": counts["fastener"], "panels": n_panels, "accessories": counts["accessory"]},
    }


def estimate_spec(spec: dict, config: dict | None = None) -> dict:
    """Price a saved extrusion_build spec (used by quotes.save_quote)."""
    opts = dict(spec.get("options") or {})
    for k in ("name", "part_number", "nsn"):
        if spec.get(k):
            opts.setdefault(k, spec[k])
    try:
        return price_build(spec.get("lines") or [], config, spec.get("quantities") or [1], opts)
    except ExtrusionError as exc:
        raise pricing.SpecError(str(exc)) from exc


# ---------------------------------------------------------------- purchase list and cut list
def purchase_list(lines: list[dict], config: dict | None = None, builds: int = 1, mode: str | None = None) -> list[dict]:
    """What to order for `builds` builds, by supplier part number."""
    cfg = pricing.merged_config(config)
    ext = cfg["extrusion"]
    mode = mode or ext["pricing_mode"]
    cat = catalog(cfg)
    prof = {p["id"]: p for p in cat["profiles"]}
    hw = {h["id"]: h for h in cat["hardware"]}
    pan = {p["id"]: p for p in cat["panels"]}
    lines = normalize_lines(lines)
    rows: list[dict] = []
    by_prof: dict[str, list[dict]] = defaultdict(list)
    for ln in lines:
        if ln["kind"] == "profile" and ln["catalog_id"] in prof and ln["length_in"]:
            by_prof[ln["catalog_id"]].append(ln)
    for pid, lns in by_prof.items():
        p = prof[pid]
        if mode == "stock":
            sl = ext["stock_length_in"]["metric" if p.get("metric") else "fractional"]
            n = nest_ffd([l["length_in"] for l in lns for _ in range(l["qty"] * builds)], sl, ext["kerf_in"], ext["end_trim_in"])
            price = sl * p["price_per_in"] * ext["full_stick_price_factor"]
            rows.append({"supplier": p["vendor"], "part_number": pid, "description": f"{p['name']}, full stick {sl:g} in",
                         "qty": n["sticks"], "unit": "stick", "unit_price": round(price, 2), "extended": round(n["sticks"] * price, 2), "machining": ""})
            for L in n["oversize"]:
                rows.append({"supplier": p["vendor"], "part_number": pid, "description": f"{p['name']}, special length {L:g} in",
                             "qty": 1, "unit": "piece", "unit_price": round(L * p["price_per_in"], 2), "extended": round(L * p["price_per_in"], 2), "machining": ""})
        else:
            for l in lns:
                q = l["qty"] * builds
                price = l["length_in"] * p["price_per_in"] + ext["cut_charge"]
                rows.append({"supplier": p["vendor"], "part_number": pid,
                             "description": f"{p['name']}, cut to {l['length_in']:g} in ({l['length_in'] * MM_PER_IN:.1f} mm)",
                             "qty": q, "unit": "piece", "unit_price": round(price, 2), "extended": round(q * price, 2),
                             "machining": machining_text(l["machining"])})
                for o in l["machining"]:
                    mp = ext["machining"].get(o["op"], 0.0)
                    rows.append({"supplier": p["vendor"], "part_number": f"{pid} op", "description": f"{OP_LABELS.get(o['op'], o['op'])} on {pid} @ {l['length_in']:g} in",
                                 "qty": o["count"] * q, "unit": "op", "unit_price": round(mp, 2), "extended": round(o["count"] * q * mp, 2), "machining": ""})
    agg: dict[str, dict] = {}
    for ln in lines:
        if ln["kind"] == "hardware" and ln["catalog_id"] in hw:
            h = hw[ln["catalog_id"]]
            key = ln["catalog_id"] + "|" + (ln["part_number"] or "")
            price = ln["unit_price"] if ln["unit_price"] is not None else h["unit_price"]
            r = agg.setdefault(key, {"supplier": "", "part_number": ln["part_number"] or h["id"], "description": h["name"],
                                     "qty": 0, "unit": "each", "unit_price": round(price, 2), "extended": 0.0, "machining": ""})
            r["qty"] += ln["qty"] * builds
            r["extended"] = round(r["qty"] * price, 2)
    rows += list(agg.values())
    for ln in lines:
        if ln["kind"] == "panel" and ln["catalog_id"] in pan and ln["width_in"] and ln["height_in"]:
            p = pan[ln["catalog_id"]]
            sqft = ln["width_in"] * ln["height_in"] / 144
            rows.append({"supplier": "", "part_number": ln["part_number"] or p["id"], "description": f"{p['name']}, {ln['width_in']:g} x {ln['height_in']:g} in",
                         "qty": ln["qty"] * builds, "unit": "panel", "unit_price": round(sqft * p["price_per_sqft"], 2),
                         "extended": round(sqft * p["price_per_sqft"] * ln["qty"] * builds, 2), "machining": ""})
        elif ln["kind"] in ("unmatched", "custom") and ln["unit_price"] is not None:
            rows.append({"supplier": "", "part_number": ln["part_number"], "description": ln["description"], "qty": ln["qty"] * builds,
                         "unit": "each", "unit_price": ln["unit_price"], "extended": round(ln["unit_price"] * ln["qty"] * builds, 2), "machining": ""})
    return rows


def cut_list(lines: list[dict], builds: int = 1) -> list[dict]:
    rows = []
    marks = iter("ABCDEFGHIJKLMNOPQRSTUVWXYZ")
    for i, ln in enumerate(normalize_lines(lines)):
        if ln["kind"] != "profile" or not ln["length_in"]:
            continue
        mark = ln["item"] or next(marks, str(i + 1))
        rows.append({"mark": mark, "profile": ln["catalog_id"], "length_in": round(ln["length_in"], 3),
                     "length_mm": round(ln["length_in"] * MM_PER_IN, 1), "qty_per_build": ln["qty"], "qty_total": ln["qty"] * builds,
                     "machining": machining_text(ln["machining"]), "notes": ln["notes"]})
    return rows


def _xlsx(sheets: list[tuple[str, list[str], list[list]]], title: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for name, header, rows in sheets:
        ws = wb.create_sheet(name[:31])
        start = 1
        if title:
            ws.cell(row=1, column=1, value=title).font = Font(bold=True)
            start = 3
        for j, h in enumerate(header, 1):
            ws.cell(row=start, column=j, value=h).font = Font(bold=True)
        for i, r in enumerate(rows, start + 1):
            for j, v in enumerate(r, 1):
                ws.cell(row=i, column=j, value=v)
        for j, h in enumerate(header, 1):
            width = max([len(str(h))] + [len(str(r[j - 1])) for r in rows if j - 1 < len(r) and r[j - 1] is not None]) + 2
            ws.column_dimensions[ws.cell(row=start, column=j).column_letter].width = min(width, 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def cut_list_xlsx(lines: list[dict], config: dict | None = None, builds: int = 1, title: str = "") -> bytes:
    """Workbook with a Cut list sheet and a Stick nesting sheet for `builds` builds."""
    cfg = pricing.merged_config(config)
    ext = cfg["extrusion"]
    cat = catalog(cfg)
    prof = {p["id"]: p for p in cat["profiles"]}
    rows = cut_list(lines, builds)
    cl = [[r["mark"], r["profile"], r["length_in"], r["length_mm"], r["qty_per_build"], r["qty_total"], r["machining"], r["notes"]] for r in rows]
    nest_rows = []
    by_prof: dict[str, list[float]] = defaultdict(list)
    for r in rows:
        by_prof[r["profile"]] += [r["length_in"]] * r["qty_total"]
    for pid, lengths in by_prof.items():
        p = prof.get(pid, {})
        sl = ext["stock_length_in"]["metric" if p.get("metric") else "fractional"]
        n = nest_ffd(lengths, sl, ext["kerf_in"], ext["end_trim_in"])
        for i, s in enumerate(n["stick_detail"], 1):
            nest_rows.append([pid, i, sl, ", ".join(f"{x:g}" for x in s["pieces"]), s["used_in"], s["drop_in"]])
        for L in n["oversize"]:
            nest_rows.append([pid, "special", L, f"{L:g}", L, 0])
    head = f"{title + ': ' if title else ''}cut list for {builds} build(s). Lengths in inches (mm shown)."
    return _xlsx([
        ("Cut list", ["Mark", "Profile", "Length (in)", "Length (mm)", "Qty per build", f"Qty for {builds}", "Machining", "Notes"], cl),
        ("Stick nesting", ["Profile", "Stick", "Stock length (in)", "Pieces (in)", "Used (in)", "Drop (in)"], nest_rows),
    ], head)


def purchase_list_xlsx(lines: list[dict], config: dict | None = None, builds: int = 1, title: str = "", mode: str | None = None) -> bytes:
    rows = purchase_list(lines, config, builds, mode)
    data = [[r["supplier"], r["part_number"], r["description"], r["qty"], r["unit"], r["unit_price"], r["extended"], r["machining"]] for r in rows]
    total = round(sum(r["extended"] for r in rows), 2)
    data.append(["", "", "Total (placeholder prices until you enter supplier prices)", "", "", "", total, ""])
    return _xlsx([("Purchase list", ["Supplier", "Part number", "Description", "Qty", "Unit", "Unit price", "Extended", "Machining"], data)],
                 f"{title + ': ' if title else ''}purchase list for {builds} build(s)")
