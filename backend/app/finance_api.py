"""Finance: invoices tracked against PIEE (WAWF), Prompt Payment due dates and interest, an aging
report, bookkeeping CSV exports, and loans (SBA and others) with draws, payments and amortization.

The app does not connect to PIEE (it needs a CAC or PIEE login). The owner records what was
submitted there and this module works out when payment is due.

Rules used, with sources (checked 2026-10-08):
- FAR 52.232-25(a)(1)(i), Prompt Payment: the due date is the later of the 30th day after the
  designated billing office receives a proper invoice, or the 30th day after Government acceptance.
  (a)(5)(i): for computing interest only, acceptance is deemed to occur on the 7th day after delivery
  unless the contract says otherwise or there is a disagreement over quantity, quality or compliance.
  (a)(4): if the due date falls on a weekend or legal holiday, payment may be made on the next working
  day with no interest. https://www.acquisition.gov/far/52.232-25
- 5 CFR 1315.4(b): an invoice is deemed received on the later of (i) the date a readable electronic
  transmission is received, or (ii) the 7th day after delivery or completion of services, unless the
  agency accepted earlier (acceptance date substitutes) or the contract sets a longer acceptance period
  (actual acceptance or the end of that period substitutes). 1315.4(f)/(g): payment is due 30 days after
  the start of the payment period, unless the contract sets a date. 1315.4(h): weekend/holiday rule.
  https://www.ecfr.gov/current/title-5/chapter-III/subchapter-B/part-1315/section-1315.4
- 5 CFR 1315.10(a): interest runs from the day after the due date through the payment date at the
  rate in effect on the day after the due date; unpaid interest is added to principal at the end of
  each 30-day period for up to one year; it stops accruing after one year; under $1 need not be paid;
  360-day year. https://www.ecfr.gov/current/title-5/chapter-III/subchapter-B/part-1315/section-1315.10
- Treasury publishes the Prompt Payment interest rate every six months:
  https://fiscal.treasury.gov/prompt-payment/rates.html (4.75% for July 1 to December 31, 2026, as read
  on 2026-10-08). The rate is stored in finance settings so the owner can add new periods.
- Accelerated payment: FAR 32.009-1(a)(2) says DoD shall provide accelerated payments to the fullest
  extent permitted by law, with a goal of 15 days after receipt of a proper invoice and all other
  required documentation (FAR 32.009-1(a)(1) gives the same goal for other agencies paying small
  business contractors). DFARS 232.906(a)(ii): the FAR 32.906 bar on early payment does not apply to
  small business invoices, and no interest is owed if the Government does not pay early.
  https://www.acquisition.gov/far/32.009-1 , https://www.acquisition.gov/dfars/232.906-making-payments.
- WAWF document types: DFARS 252.232-7006 (JAN 2023) Wide Area WorkFlow Payment Instructions, paragraph
  (f): cost voucher for cost-type, labor-hour and T&M items; the invoice and receiving report the
  contracting officer specifies for fixed-price items that ship (a "combo" document creates both in one
  step); Invoice 2in1 for fixed-price services without shipment; progress payment, performance based
  payment and commercial financing requests. https://www.acquisition.gov/dfars/252.232-7006-wide-area-workflow-payment-instructions.
  DLA Aviation vendor hub: "Invoice and Receiving Report (Combo)" is the common document; a stand-alone
  invoice is used for service CLINs such as first article test.
  https://www.dla.mil/Aviation/Business/Vendor-Information-Hub/Table-of-Contents/Details/Article/3411088/combo-invoicing-process/
- QuickBooks Online invoice import (Intuit help, "Import multiple invoices"): required fields are invoice
  number, customer, invoice date, due date and item amount; each line needs its own copy of those;
  up to 1,000 rows and 100 invoices per import; not available when sales tax is set up; columns are
  mapped during import. Intuit's help page does not print the exact sample headers, so the header names
  here follow the sample file's style and should be checked against Settings > Import data > Invoices >
  Download sample csv.
- Wave: Wave's help documents CSV upload of transactions (you choose the date, description and amount
  columns; too many columns cause errors). No CSV invoice import is documented, so Wave gets a 3-column
  payments file and the invoice export is a generic CSV.
  https://support.waveapps.com/hc/en-us/articles/208621626-Upload-a-bank-or-credit-card-statement-in-csv-format
"""
from __future__ import annotations

import calendar
import csv
import io
import statistics
from collections import defaultdict
from datetime import date, datetime, timedelta
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from pydantic import BaseModel, ConfigDict
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import services
from .db import get_db
from .models_finance import INVOICE_STATUSES, LOAN_PROGRAMS, FinanceSettings, Invoice, Loan, LoanDraw, LoanPayment
from .models_jobs import Job

router = APIRouter(prefix="/api/finance", tags=["finance"])

DEFAULT_RATES = [{"start": "2026-07-01", "end": "2026-12-31", "rate": 4.75,
                  "source": "fiscal.treasury.gov/prompt-payment/rates.html"}]

DOC_TYPES = [
    {"key": "combo", "label": "Invoice and Receiving Report (Combo)",
     "when": "Fixed-price supplies that ship. Creates the invoice and the receiving report in one step. DLA calls it the common document.",
     "source": "DFARS 252.232-7006(f)(1)(ii)(A); DLA Aviation vendor hub"},
    {"key": "2in1", "label": "Invoice 2in1",
     "when": "Fixed-price services with no shipment, when the contract's WAWF instructions call for it. Meets the requirements for both the invoice and the receiving report.",
     "source": "DFARS 252.232-7006(f)(1)(ii)(B)"},
    {"key": "invoice", "label": "Invoice (stand-alone)",
     "when": "Payment request only, when the contract's WAWF instructions specify it. DLA Aviation uses it for service CLINs such as first article test.",
     "source": "DLA Aviation vendor hub"},
    {"key": "receiving_report", "label": "Receiving Report (stand-alone)",
     "when": "Shipment and acceptance record without a payment request, when the contract's WAWF instructions specify it.",
     "source": "DFARS 252.232-7006(f)(1)(ii)(A)"},
    {"key": "cost_voucher", "label": "Cost Voucher",
     "when": "Cost-type line items, including labor-hour and time-and-materials.", "source": "DFARS 252.232-7006(f)(1)(i)"},
    {"key": "progress_payment", "label": "Progress Payment Request",
     "when": "Customary progress payments based on costs incurred.", "source": "DFARS 252.232-7006(f)(1)(iii)"},
    {"key": "performance_based", "label": "Performance Based Payment Request",
     "when": "Performance-based payments.", "source": "DFARS 252.232-7006(f)(1)(iv)"},
    {"key": "commercial_financing", "label": "Commercial Financing Request",
     "when": "Commercial financing.", "source": "DFARS 252.232-7006(f)(1)(v)"},
    {"key": "emailed", "label": "Emailed invoice (not PIEE)",
     "when": "Primes and commercial customers who take an emailed invoice. Use the generated invoice document.", "source": ""},
]
DOC_TYPE_LABELS = {d["key"]: d["label"] for d in DOC_TYPES}

PROGRAM_LABELS = {"sba_7a": "SBA 7(a)", "sba_express": "SBA Express", "sba_504": "SBA 504", "sba_microloan": "SBA Microloan",
                  "va_other": "VA or other", "line_of_credit": "Line of credit", "other": "Other"}

FUND_PURPOSES = ["Working capital", "Equipment", "Inventory and materials", "Real estate", "Software and IT",
                 "Refinance debt", "Payroll", "Other"]

OPEN_STATUSES = ("submitted", "accepted")


# ------------------------------------------------------------------ date helpers
def _d(s: str | None) -> date | None:
    if not s:
        return None
    try:
        return date.fromisoformat(str(s)[:10])
    except ValueError:
        return None


def _nth_weekday(year: int, month: int, weekday: int, n: int) -> date:
    first = date(year, month, 1)
    return first + timedelta(days=(weekday - first.weekday()) % 7 + 7 * (n - 1))


def _last_weekday(year: int, month: int, weekday: int) -> date:
    last = date(year, month, calendar.monthrange(year, month)[1])
    return last - timedelta(days=(last.weekday() - weekday) % 7)


def federal_holidays(year: int) -> set[date]:
    """Legal public holidays (5 U.S.C. 6103) with the Saturday-to-Friday / Sunday-to-Monday observance rule."""
    fixed = [date(year, 1, 1), date(year, 6, 19), date(year, 7, 4), date(year, 11, 11), date(year, 12, 25)]
    out = {_nth_weekday(year, 1, 0, 3), _nth_weekday(year, 2, 0, 3), _last_weekday(year, 5, 0),
           _nth_weekday(year, 9, 0, 1), _nth_weekday(year, 10, 0, 2), _nth_weekday(year, 11, 3, 4)}
    for d in fixed:
        out.add(d - timedelta(days=1) if d.weekday() == 5 else d + timedelta(days=1) if d.weekday() == 6 else d)
    # New Year's Day of the next year observed on Friday Dec 31
    if date(year + 1, 1, 1).weekday() == 5:
        out.add(date(year, 12, 31))
    return out


def next_business_day(d: date) -> date:
    """d itself if it is a working day, otherwise the next working day."""
    while d.weekday() >= 5 or d in federal_holidays(d.year):
        d += timedelta(days=1)
    return d


def add_months(d: date, n: int) -> date:
    m = d.month - 1 + n
    y, m = d.year + m // 12, m % 12 + 1
    return date(y, m, min(d.day, calendar.monthrange(y, m)[1]))


# ------------------------------------------------------------------ prompt payment
def payment_due(submitted: str, ship_date: str = "", acceptance_date: str = "", acceptance_period_days: int = 7,
                disputed: bool = False, due_override: str = "") -> dict:
    """Prompt Payment due date for an invoice (FAR 52.232-25, 5 CFR 1315.4).

    Returns {"due_date", "penalty_free_through", "deemed_receipt", "acceptance_basis", "acceptance_used",
    "accelerated_goal", "estimate", "basis"} with ISO date strings ("" when unknown)."""
    receipt = _d(submitted)
    out = {"due_date": "", "penalty_free_through": "", "deemed_receipt": "", "acceptance_used": "",
           "acceptance_basis": "", "accelerated_goal": "", "estimate": False, "basis": ""}
    if receipt is None:
        out["basis"] = "Not submitted yet."
        return out
    delivery, accepted = _d(ship_date), _d(acceptance_date)
    period = max(int(acceptance_period_days or 7), 1)
    if delivery and not disputed:
        constructive = delivery + timedelta(days=period)
        if accepted and accepted < constructive:
            acc, why = accepted, "actual acceptance (before the acceptance period ended)"
        else:
            acc = constructive
            why = (f"constructive acceptance, day {period} after delivery" +
                   (" (actual acceptance came later, but for Prompt Payment it is deemed on this day)" if accepted and accepted > constructive else ""))
            out["estimate"] = accepted is None
    elif accepted:
        acc, why = accepted, "actual acceptance" + (" (disputed, so no constructive acceptance)" if disputed else "")
    else:
        acc, why = None, ("disputed and not yet accepted" if disputed else "no ship or acceptance date entered")
        out["estimate"] = True
    deemed = max(receipt, acc) if acc else receipt
    if _d(due_override):
        due, basis = _d(due_override), "Payment date set by the contract."
    else:
        due = deemed + timedelta(days=30)
        basis = (f"30 days after {deemed.isoformat()}, the later of PIEE receipt ({receipt.isoformat()}) and acceptance"
                 + (f" ({acc.isoformat()}, {why})." if acc else f" ({why}; using receipt for now)."))
    goal_from = max(receipt, accepted) if accepted else (max(receipt, acc) if acc else receipt)
    out.update({"due_date": due.isoformat(), "penalty_free_through": next_business_day(due).isoformat(),
                "deemed_receipt": deemed.isoformat(), "acceptance_used": acc.isoformat() if acc else "",
                "acceptance_basis": why, "accelerated_goal": (goal_from + timedelta(days=15)).isoformat(), "basis": basis})
    return out


def rate_on(rates: list[dict], d: date) -> tuple[float | None, str]:
    """Annual percent in effect on d from the stored Treasury rate periods, plus a note."""
    good = [r for r in (rates or []) if _d(r.get("start")) and r.get("rate") not in (None, "")]
    for r in good:
        end = _d(r.get("end")) or date.max
        if _d(r["start"]) <= d <= end:
            return float(r["rate"]), f"{r['rate']}% (Treasury rate for {r['start']} to {r.get('end') or 'open'})"
    if not good:
        return None, "No Prompt Payment rate stored. Add the current rate from fiscal.treasury.gov in Finance settings."
    latest = max(good, key=lambda r: r["start"])
    return float(latest["rate"]), (f"{latest['rate']}% (latest stored rate; no stored period covers {d.isoformat()}, "
                                   "so check fiscal.treasury.gov and add it in Finance settings)")


def late_interest(amount: float, due_date: str, pay_date: str, rates: list[dict]) -> dict:
    """Prompt Payment late interest (5 CFR 1315.10): from the day after the due date through the
    payment date, compounded every 30 days, 360-day year, accrual capped at one year, under $1 not owed.
    No interest when payment lands on or before the next working day after a weekend/holiday due date."""
    due, paid = _d(due_date), _d(pay_date)
    out = {"days_late": 0, "interest": 0.0, "payable": 0.0, "rate": None, "rate_note": "", "note": ""}
    if not due or not paid or not amount:
        return out
    if paid <= next_business_day(due):
        return out
    days = (paid - due).days
    out["days_late"] = days
    rate, note = rate_on(rates, due + timedelta(days=1))
    out["rate"], out["rate_note"] = rate, note
    if rate is None:
        out["note"] = note
        return out
    accrual = min(days, 365)
    r = rate / 100.0
    periods, rem = divmod(accrual, 30)
    total = amount * (1 + r * 30 / 360) ** periods * (1 + r * rem / 360)
    interest = round(total - amount, 2)
    out["interest"] = interest
    out["payable"] = interest if interest >= 1 else 0.0
    notes = []
    if days > 365:
        notes.append("Interest stops accruing after one year.")
    if 0 < interest < 1:
        notes.append("Interest under $1 need not be paid.")
    out["note"] = " ".join(notes)
    return out


# ------------------------------------------------------------------ invoice helpers
def _num(v) -> float | None:
    try:
        return None if v in (None, "") else float(v)
    except (TypeError, ValueError):
        return None


def clean_lines(lines: list[dict]) -> list[dict]:
    out = []
    for ln in lines or []:
        q, p = _num(ln.get("quantity")), _num(ln.get("unit_price"))
        amt = round(q * p, 2) if q is not None and p is not None else (_num(ln.get("amount")) or 0.0)
        out.append({"clin": str(ln.get("clin") or ""), "description": str(ln.get("description") or ""),
                    "quantity": q, "unit": str(ln.get("unit") or ""), "unit_price": p, "amount": amt})
    return out


def invoice_total(inv: Invoice) -> float:
    return round(sum(float(ln.get("amount") or 0) for ln in (inv.lines or [])), 2)


def get_settings(db: Session) -> FinanceSettings:
    s = db.get(FinanceSettings, 1)
    if s is None:
        s = FinanceSettings(id=1, invoice_prefix="INV-", next_invoice_seq=1, seq_width=4,
                            prompt_pay_rates=list(DEFAULT_RATES), payment_terms="Net 30")
        db.add(s)
        db.commit()
    return s


def _next_number(db: Session, s: FinanceSettings) -> str:
    while True:
        n = f"{s.invoice_prefix or ''}{int(s.next_invoice_seq or 1):0{int(s.seq_width or 4)}d}"
        s.next_invoice_seq = int(s.next_invoice_seq or 1) + 1
        if not db.scalar(select(Invoice.id).where(Invoice.number == n)):
            return n


def invoice_view(inv: Invoice, rates: list[dict], today: date | None = None) -> dict:
    today = today or date.today()
    total = invoice_total(inv)
    pp = payment_due(inv.submitted_date, inv.ship_date, inv.acceptance_date, inv.acceptance_period_days,
                     bool(inv.disputed), inv.due_date_override)
    if inv.status == "paid":
        interest = late_interest(total, pp["due_date"], inv.paid_date, rates)
    elif inv.status in OPEN_STATUSES:
        interest = late_interest(total, pp["due_date"], today.isoformat(), rates)
        interest["note"] = ("Accrued to today if paid now. " + interest["note"]).strip() if interest["days_late"] else interest["note"]
    else:
        interest = late_interest(0, "", "", rates)
    paid = inv.amount_paid if inv.amount_paid is not None else (total if inv.status == "paid" else 0.0)
    sub = _d(inv.submitted_date)
    end = _d(inv.paid_date) if inv.status == "paid" else today
    return {
        "id": inv.id, "number": inv.number, "job_id": inv.job_id, "customer": inv.customer,
        "contract_number": inv.contract_number, "delivery_order": inv.delivery_order,
        "doc_type": inv.doc_type, "doc_type_label": DOC_TYPE_LABELS.get(inv.doc_type, inv.doc_type),
        "lines": inv.lines or [], "shipment_number": inv.shipment_number, "ship_date": inv.ship_date,
        "invoice_date": inv.invoice_date, "submitted_date": inv.submitted_date, "status": inv.status,
        "acceptance_date": inv.acceptance_date, "acceptance_period_days": inv.acceptance_period_days,
        "disputed": bool(inv.disputed), "due_date_override": inv.due_date_override,
        "paid_date": inv.paid_date, "amount_paid": inv.amount_paid, "interest_paid": inv.interest_paid,
        "rejection_reason": inv.rejection_reason, "notes": inv.notes, "history": inv.history or [],
        "total": total, "balance": round(total - (paid or 0.0), 2) if inv.status != "draft" else total,
        "days_since_submission": (end - sub).days if sub and end else None,
        "prompt_payment": pp, "interest": interest,
    }


def sync_job_status(db: Session, job_id: int | None) -> None:
    """Set the job to "invoiced" once something is submitted and to "paid" when every submitted invoice is paid."""
    if not job_id:
        return
    job = db.get(Job, job_id)
    if job is None or job.status in ("closed", "cancelled"):
        return
    statuses = [s for s in db.scalars(select(Invoice.status).where(Invoice.job_id == job_id)).all() if s != "draft"]
    if not statuses:
        return
    if all(s == "paid" for s in statuses):
        job.status = "paid"
    elif job.status in ("awarded", "in_work", "inspection", "shipped", "paid"):
        job.status = "invoiced"


# ------------------------------------------------------------------ schemas
class SettingsIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    invoice_prefix: str | None = None
    next_invoice_seq: int | None = None
    seq_width: int | None = None
    prompt_pay_rates: list[dict[str, Any]] | None = None
    remit_to: str | None = None
    payment_terms: str | None = None
    invoice_footer: str | None = None


class InvoiceIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    number: str | None = None
    job_id: int | None = None
    customer: str | None = None
    contract_number: str | None = None
    delivery_order: str | None = None
    doc_type: str | None = None
    lines: list[dict[str, Any]] | None = None
    shipment_number: str | None = None
    ship_date: str | None = None
    invoice_date: str | None = None
    submitted_date: str | None = None
    status: str | None = None
    acceptance_date: str | None = None
    acceptance_period_days: int | None = None
    disputed: bool | None = None
    due_date_override: str | None = None
    paid_date: str | None = None
    amount_paid: float | None = None
    interest_paid: float | None = None
    rejection_reason: str | None = None
    notes: str | None = None


class StatusIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    status: str
    date: str | None = None
    note: str | None = None
    rejection_reason: str | None = None
    amount_paid: float | None = None
    interest_paid: float | None = None


class LoanIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    lender: str | None = None
    program: str | None = None
    principal: float | None = None
    rate_type: str | None = None
    rate: float | None = None
    rate_base: str | None = None
    term_months: int | None = None
    start_date: str | None = None
    first_payment_date: str | None = None
    amortizing: bool | None = None
    payment_amount: float | None = None
    balance_basis: str | None = None
    day_count: str | None = None
    status: str | None = None
    notes: str | None = None


class DrawIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    date: str
    amount: float
    purpose: str = ""


class LoanPaymentIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    date: str
    amount: float
    principal_override: float | None = None
    interest_override: float | None = None
    note: str = ""


class CalcIn(BaseModel):
    principal: float
    rate: float
    term_months: int


# ------------------------------------------------------------------ meta and settings
def _settings_view(s: FinanceSettings) -> dict:
    return {"invoice_prefix": s.invoice_prefix, "next_invoice_seq": s.next_invoice_seq, "seq_width": s.seq_width,
            "prompt_pay_rates": s.prompt_pay_rates or [], "remit_to": s.remit_to, "payment_terms": s.payment_terms,
            "invoice_footer": s.invoice_footer,
            "next_number_preview": f"{s.invoice_prefix or ''}{int(s.next_invoice_seq or 1):0{int(s.seq_width or 4)}d}"}


@router.get("/meta")
def meta(db: Session = Depends(get_db)):
    s = get_settings(db)
    today = date.today()
    rate, note = rate_on(s.prompt_pay_rates or [], today)
    return {
        "doc_types": DOC_TYPES, "statuses": INVOICE_STATUSES,
        "programs": [{"key": k, "label": PROGRAM_LABELS[k]} for k in LOAN_PROGRAMS],
        "fund_purposes": FUND_PURPOSES, "settings": _settings_view(s),
        "current_rate": rate, "current_rate_note": note,
        "rate_source_url": "https://fiscal.treasury.gov/prompt-payment/rates.html",
        "rules": [
            "Due date: 30 days after the later of PIEE receipt of a proper invoice and acceptance (FAR 52.232-25(a)(1)(i)).",
            "If not accepted sooner, acceptance is deemed to occur on the 7th day after delivery or service completion, unless the contract sets a longer period or there is a dispute (FAR 52.232-25(a)(5)(i), 5 CFR 1315.4(b)).",
            "A due date on a weekend or federal holiday can be paid the next working day without interest (5 CFR 1315.4(h)).",
            "Late interest is paid automatically: from the day after the due date through payment, Treasury rate in effect the day after the due date, compounded every 30 days, 360-day year, up to one year; under $1 is not owed (5 CFR 1315.10).",
            "Accelerated payment: DoD has a goal of paying small business contractors within 15 days of a proper invoice and all required documents (FAR 32.009-1(a)). It is a goal, not a deadline; no interest is owed if it is missed (DFARS 232.906(a)(ii)).",
            "The contract's DFARS 252.232-7006 Wide Area WorkFlow Payment Instructions name the document type to use. Follow the contract when it differs from the guide here.",
        ],
    }


@router.get("/settings")
def read_settings(db: Session = Depends(get_db)):
    return _settings_view(get_settings(db))


@router.put("/settings")
def update_settings(body: SettingsIn, db: Session = Depends(get_db)):
    s = get_settings(db)
    data = body.model_dump(exclude_none=True)
    if "prompt_pay_rates" in data:
        rates = []
        for r in data["prompt_pay_rates"]:
            start, rate = _d(r.get("start")), _num(r.get("rate"))
            if not start or rate is None:
                raise HTTPException(400, "Each rate period needs a start date and a rate.")
            end = _d(r.get("end"))
            rates.append({"start": start.isoformat(), "end": end.isoformat() if end else "", "rate": rate,
                          "source": str(r.get("source") or "")})
        data["prompt_pay_rates"] = sorted(rates, key=lambda r: r["start"])
    for k, v in data.items():
        setattr(s, k, v)
    db.commit()
    return _settings_view(s)


@router.get("/jobs")
def job_choices(db: Session = Depends(get_db)):
    """Jobs to invoice against (short list for pickers)."""
    rows = db.scalars(select(Job).order_by(Job.id.desc())).all()
    return [{"id": j.id, "title": j.title, "customer": j.customer, "contract_number": j.contract_number,
             "delivery_order": j.delivery_order, "status": j.status, "shipped_date": j.shipped_date} for j in rows]


# ------------------------------------------------------------------ invoices
def _get_invoice(db: Session, inv_id: int) -> Invoice:
    inv = db.get(Invoice, inv_id)
    if inv is None:
        raise HTTPException(404, "Invoice not found")
    return inv


def _apply(inv: Invoice, data: dict) -> None:
    if "doc_type" in data and data["doc_type"] not in DOC_TYPE_LABELS:
        raise HTTPException(400, f"Unknown document type {data['doc_type']}")
    if "status" in data and data["status"] not in INVOICE_STATUSES:
        raise HTTPException(400, f"Unknown status {data['status']}")
    if "lines" in data:
        data["lines"] = clean_lines(data["lines"])
    for k, v in data.items():
        if hasattr(inv, k) and k not in ("id", "history"):
            setattr(inv, k, v)


@router.get("/invoices")
def list_invoices(status: str = "", job_id: int | None = None, db: Session = Depends(get_db)):
    q = select(Invoice).order_by(Invoice.id.desc())
    if status:
        q = q.where(Invoice.status == status)
    if job_id:
        q = q.where(Invoice.job_id == job_id)
    rates = get_settings(db).prompt_pay_rates or []
    return [invoice_view(i, rates) for i in db.scalars(q).all()]


@router.post("/invoices")
def create_invoice(body: InvoiceIn, db: Session = Depends(get_db)):
    """Create an invoice. With job_id, contract number, delivery order, customer, CLINs and ship date are copied from the job."""
    s = get_settings(db)
    data = body.model_dump(exclude_none=True)
    inv = Invoice(doc_type="combo", status="draft", lines=[], history=[], acceptance_period_days=7,
                  invoice_date=date.today().isoformat())
    if body.job_id:
        job = db.get(Job, body.job_id)
        if job is None:
            raise HTTPException(404, "Job not found")
        inv.job_id = job.id
        inv.customer, inv.contract_number, inv.delivery_order = job.customer, job.contract_number, job.delivery_order
        inv.ship_date = job.shipped_date or ""
        inv.lines = clean_lines([{"clin": c.get("clin"), "description": c.get("description") or c.get("part_number") or "",
                                  "quantity": c.get("quantity"), "unit": c.get("unit") or "EA", "unit_price": c.get("unit_price")}
                                 for c in (job.clins or [])])
    number = (data.pop("number", "") or "").strip()
    if number and db.scalar(select(Invoice.id).where(Invoice.number == number)):
        raise HTTPException(409, f"Invoice number {number} is already used")
    inv.number = number or _next_number(db, s)
    _apply(inv, data)
    inv.history = [{"date": date.today().isoformat(), "status": inv.status, "note": "Created"}]
    db.add(inv)
    db.flush()
    sync_job_status(db, inv.job_id)
    db.commit()
    return invoice_view(inv, s.prompt_pay_rates or [])


@router.get("/invoices/{inv_id}")
def get_invoice(inv_id: int, db: Session = Depends(get_db)):
    return invoice_view(_get_invoice(db, inv_id), get_settings(db).prompt_pay_rates or [])


@router.put("/invoices/{inv_id}")
def update_invoice(inv_id: int, body: InvoiceIn, db: Session = Depends(get_db)):
    inv = _get_invoice(db, inv_id)
    data = body.model_dump(exclude_unset=True)
    for k in ("amount_paid", "interest_paid", "job_id"):
        if k in data and data[k] is None:
            setattr(inv, k, None)
            data.pop(k)
    data = {k: v for k, v in data.items() if v is not None}
    if "number" in data:
        n = data["number"].strip()
        if not n:
            raise HTTPException(400, "Invoice number cannot be blank")
        if db.scalar(select(Invoice.id).where(Invoice.number == n, Invoice.id != inv.id)):
            raise HTTPException(409, f"Invoice number {n} is already used")
        data["number"] = n
    old_status, old_job = inv.status, inv.job_id
    _apply(inv, data)
    if inv.status != old_status:
        inv.history = [*(inv.history or []), {"date": date.today().isoformat(), "status": inv.status, "note": "Edited"}]
    db.flush()
    sync_job_status(db, inv.job_id)
    if old_job != inv.job_id:
        sync_job_status(db, old_job)
    db.commit()
    return invoice_view(inv, get_settings(db).prompt_pay_rates or [])


@router.post("/invoices/{inv_id}/status")
def set_invoice_status(inv_id: int, body: StatusIn, db: Session = Depends(get_db)):
    """Record a PIEE status change. The date goes to the matching field (submitted, acceptance or paid date)."""
    inv = _get_invoice(db, inv_id)
    if body.status not in INVOICE_STATUSES:
        raise HTTPException(400, f"Unknown status {body.status}")
    when = (_d(body.date) or date.today()).isoformat()
    if body.status == "submitted":
        inv.submitted_date = when  # a resubmission after rejection restarts the clock
        inv.rejection_reason = "" if inv.status == "rejected" else inv.rejection_reason
    elif body.status == "rejected":
        inv.rejection_reason = body.rejection_reason or body.note or inv.rejection_reason
    elif body.status == "accepted":
        inv.acceptance_date = when
    elif body.status == "paid":
        inv.paid_date = when
        inv.amount_paid = body.amount_paid if body.amount_paid is not None else invoice_total(inv)
        if body.interest_paid is not None:
            inv.interest_paid = body.interest_paid
    inv.status = body.status
    inv.history = [*(inv.history or []), {"date": when, "status": body.status,
                                          "note": body.note or body.rejection_reason or ""}]
    db.flush()
    sync_job_status(db, inv.job_id)
    db.commit()
    return invoice_view(inv, get_settings(db).prompt_pay_rates or [])


@router.delete("/invoices/{inv_id}")
def delete_invoice(inv_id: int, db: Session = Depends(get_db)):
    inv = _get_invoice(db, inv_id)
    job_id = inv.job_id
    db.delete(inv)
    db.flush()
    sync_job_status(db, job_id)
    db.commit()
    return {"ok": True}


def build_invoice_docx(db: Session, inv: Invoice) -> bytes:
    """Invoice document (DOCX) for records and for customers who take emailed invoices."""
    p = services.get_profile(db)
    s = get_settings(db)
    v = invoice_view(inv, s.prompt_pay_rates or [])
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    h = doc.add_paragraph()
    r = h.add_run(p.name or "Your company")
    r.bold, r.font.size = True, Pt(16)
    ids = [f"UEI {p.uei}" if p.uei else "", f"CAGE {p.cage}" if p.cage else ""]
    addr = s.remit_to or "\n".join(str(getattr(p, k)) for k in ("address", "city", "state", "zip") if getattr(p, k, ""))
    if addr:
        doc.add_paragraph(addr)
    if any(ids):
        doc.add_paragraph("   ".join(x for x in ids if x))
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tr = t.add_run("INVOICE")
    tr.bold, tr.font.size = True, Pt(20)

    pp = v["prompt_payment"]
    info = [("Invoice number", inv.number), ("Invoice date", inv.invoice_date), ("Bill to", inv.customer),
            ("Contract number", inv.contract_number), ("Delivery / task order", inv.delivery_order),
            ("Shipment number", inv.shipment_number), ("Ship / completion date", inv.ship_date),
            ("Payment terms", s.payment_terms), ("Payment due", pp["due_date"] if pp["due_date"] else "")]
    info = [(a, b) for a, b in info if b]
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    for a, b in info:
        row = tbl.add_row().cells
        row[0].text, row[1].text = a, str(b)
        row[0].paragraphs[0].runs[0].bold = True
    doc.add_paragraph()

    lt = doc.add_table(rows=1, cols=6)
    lt.style = "Table Grid"
    for c, txt in zip(lt.rows[0].cells, ["CLIN", "Description", "Qty", "Unit", "Unit price", "Amount"]):
        c.text = txt
        c.paragraphs[0].runs[0].bold = True
    for ln in inv.lines or []:
        c = lt.add_row().cells
        q = ln.get("quantity")
        c[0].text = str(ln.get("clin") or "")
        c[1].text = str(ln.get("description") or "")
        c[2].text = "" if q is None else f"{q:g}"
        c[3].text = str(ln.get("unit") or "")
        c[4].text = "" if ln.get("unit_price") is None else f"${ln['unit_price']:,.2f}"
        c[5].text = f"${float(ln.get('amount') or 0):,.2f}"
    tot = lt.add_row().cells
    tot[4].text = "Total"
    tot[4].paragraphs[0].runs[0].bold = True
    tot[5].text = f"${v['total']:,.2f}"
    tot[5].paragraphs[0].runs[0].bold = True

    if inv.doc_type != "emailed":
        doc.add_paragraph(f"Submitted in PIEE (WAWF) as: {v['doc_type_label']}"
                          + (f" on {inv.submitted_date}" if inv.submitted_date else "") + ". This copy is for records.")
    if inv.notes:
        doc.add_paragraph(inv.notes)
    if s.invoice_footer:
        doc.add_paragraph(s.invoice_footer)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@router.get("/invoices/{inv_id}/document.docx")
def invoice_document(inv_id: int, db: Session = Depends(get_db)):
    inv = _get_invoice(db, inv_id)
    data = build_invoice_docx(db, inv)
    safe = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in inv.number)
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers={"Content-Disposition": f'attachment; filename="invoice-{safe}.docx"'})


# ------------------------------------------------------------------ aging
SUBMISSION_BUCKETS = [("0-30", 0, 30), ("31-60", 31, 60), ("61-90", 61, 90), ("90+", 91, 10**9)]
PAST_DUE_BUCKETS = [("not_due", None, 0), ("1-15", 1, 15), ("16-30", 16, 30), ("31-60", 31, 60), ("60+", 61, 10**9)]


def aging_report(db: Session, today: date | None = None) -> dict:
    """Outstanding invoices bucketed by days since PIEE submission and days past the payment due date,
    plus payment-days statistics for paid invoices."""
    today = today or date.today()
    rates = get_settings(db).prompt_pay_rates or []
    invs = db.scalars(select(Invoice).order_by(Invoice.submitted_date)).all()
    rows, by_sub, by_due = [], {b[0]: {"count": 0, "amount": 0.0} for b in SUBMISSION_BUCKETS}, \
        {b[0]: {"count": 0, "amount": 0.0} for b in PAST_DUE_BUCKETS}
    rejected = []
    for inv in invs:
        v = invoice_view(inv, rates, today)
        if inv.status == "rejected":
            rejected.append(v)
        if inv.status not in OPEN_STATUSES:
            continue
        age = v["days_since_submission"] or 0
        due = _d(v["prompt_payment"]["due_date"])
        past = (today - due).days if due and today > next_business_day(due) else 0
        v["days_past_due"] = past
        bal = v["balance"]
        sb = next(b[0] for b in SUBMISSION_BUCKETS if b[1] <= age <= b[2])
        db_ = "not_due" if past <= 0 else next(b[0] for b in PAST_DUE_BUCKETS[1:] if b[1] <= past <= b[2])
        v["submission_bucket"], v["past_due_bucket"] = sb, db_
        for bucket, key in ((by_sub, sb), (by_due, db_)):
            bucket[key]["count"] += 1
            bucket[key]["amount"] = round(bucket[key]["amount"] + bal, 2)
        rows.append(v)
    paid = [invoice_view(i, rates, today) for i in invs if i.status == "paid" and _d(i.submitted_date) and _d(i.paid_date)]
    days = [(_d(p["paid_date"]) - _d(p["submitted_date"])).days for p in paid]
    on_time = sum(1 for p in paid if p["interest"]["days_late"] == 0)
    stats = {"count": len(days)}
    if days:
        stats.update({"average_days": round(sum(days) / len(days), 1), "median_days": statistics.median(days),
                      "min_days": min(days), "max_days": max(days), "on_time": on_time,
                      "on_time_pct": round(100 * on_time / len(days), 1),
                      "interest_received": round(sum(i.interest_paid or 0 for i in invs if i.status == "paid"), 2)})
    return {
        "as_of": today.isoformat(), "invoices": rows, "rejected": rejected,
        "by_submission_age": by_sub, "by_past_due": by_due,
        "total_outstanding": round(sum(r["balance"] for r in rows), 2),
        "total_past_due": round(sum(r["balance"] for r in rows if r["days_past_due"] > 0), 2),
        "interest_accrued": round(sum(r["interest"]["payable"] for r in rows), 2),
        "payment_stats": stats,
    }


@router.get("/aging")
def aging(db: Session = Depends(get_db)):
    return aging_report(db)


# ------------------------------------------------------------------ exports
QBO_HEADERS = ["InvoiceNo", "Customer", "InvoiceDate", "DueDate", "Terms", "Memo", "Item(Product/Service)",
               "ItemDescription", "ItemQuantity", "ItemRate", "ItemAmount"]
GENERIC_INVOICE_HEADERS = ["Invoice number", "Status", "Customer", "Contract number", "Delivery order", "PIEE document type",
                           "Invoice date", "Submitted date", "Acceptance date", "Payment due date", "Paid date", "CLIN",
                           "Description", "Quantity", "Unit", "Unit price", "Line amount", "Invoice total"]
WAVE_HEADERS = ["Date", "Description", "Amount"]
PAYMENT_HEADERS = ["Paid date", "Invoice number", "Customer", "Contract number", "Delivery order", "Invoice total",
                   "Amount paid", "Interest paid", "Submitted date", "Payment due date", "Days to pay"]
LOAN_PAYMENT_HEADERS = ["Date", "Lender", "Program", "Amount", "Principal", "Interest", "Balance after", "Note"]


def _us(s: str) -> str:
    d = _d(s)
    return d.strftime("%m/%d/%Y") if d else ""


def _in_range(s: str, start: str, end: str) -> bool:
    d = _d(s)
    if d is None:
        return not start and not end
    return (not _d(start) or d >= _d(start)) and (not _d(end) or d <= _d(end))


def _csv(headers: list[str], rows: list[list], name: str) -> Response:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    w.writerows(rows)
    return Response(buf.getvalue(), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


def _export_invoices(db: Session, start: str, end: str, include_drafts: bool) -> list[tuple[Invoice, dict]]:
    rates = get_settings(db).prompt_pay_rates or []
    out = []
    for inv in db.scalars(select(Invoice).order_by(Invoice.number)).all():
        if inv.status == "draft" and not include_drafts:
            continue
        if _in_range(inv.invoice_date or inv.submitted_date, start, end):
            out.append((inv, invoice_view(inv, rates)))
    return out


def qbo_invoice_rows(db: Session, start: str = "", end: str = "", include_drafts: bool = False) -> list[list]:
    s = get_settings(db)
    rows = []
    for inv, v in _export_invoices(db, start, end, include_drafts):
        due = v["prompt_payment"]["due_date"] or ((_d(inv.invoice_date) + timedelta(days=30)).isoformat() if _d(inv.invoice_date) else "")
        memo = " ".join(x for x in [f"Contract {inv.contract_number}" if inv.contract_number else "",
                                    f"Order {inv.delivery_order}" if inv.delivery_order else ""] if x)
        for ln in inv.lines or [{"clin": "", "description": "", "quantity": None, "unit_price": None, "amount": 0}]:
            rows.append([inv.number, inv.customer, _us(inv.invoice_date or inv.submitted_date), _us(due), s.payment_terms, memo,
                         f"CLIN {ln.get('clin')}" if ln.get("clin") else "", ln.get("description") or "",
                         "" if ln.get("quantity") is None else ln["quantity"],
                         "" if ln.get("unit_price") is None else ln["unit_price"], f"{float(ln.get('amount') or 0):.2f}"])
    return rows


@router.get("/export/quickbooks-invoices.csv")
def export_qbo(start: str = "", end: str = "", include_drafts: bool = False, db: Session = Depends(get_db)):
    return _csv(QBO_HEADERS, qbo_invoice_rows(db, start, end, include_drafts), "quickbooks-invoices.csv")


@router.get("/export/invoices.csv")
def export_generic_invoices(start: str = "", end: str = "", include_drafts: bool = False, db: Session = Depends(get_db)):
    rows = []
    for inv, v in _export_invoices(db, start, end, include_drafts):
        for ln in inv.lines or [{}]:
            rows.append([inv.number, inv.status, inv.customer, inv.contract_number, inv.delivery_order, v["doc_type_label"],
                         inv.invoice_date, inv.submitted_date, inv.acceptance_date, v["prompt_payment"]["due_date"], inv.paid_date,
                         ln.get("clin", ""), ln.get("description", ""), "" if ln.get("quantity") is None else ln["quantity"],
                         ln.get("unit", ""), "" if ln.get("unit_price") is None else ln["unit_price"],
                         f"{float(ln.get('amount') or 0):.2f}", f"{v['total']:.2f}"])
    return _csv(GENERIC_INVOICE_HEADERS, rows, "invoices.csv")


def _paid(db: Session, start: str, end: str) -> list[tuple[Invoice, dict]]:
    rates = get_settings(db).prompt_pay_rates or []
    return [(i, invoice_view(i, rates)) for i in db.scalars(select(Invoice).where(Invoice.status == "paid").order_by(Invoice.paid_date)).all()
            if _in_range(i.paid_date, start, end)]


@router.get("/export/payments.csv")
def export_payments(start: str = "", end: str = "", db: Session = Depends(get_db)):
    rows = []
    for inv, v in _paid(db, start, end):
        sub, pd = _d(inv.submitted_date), _d(inv.paid_date)
        rows.append([inv.paid_date, inv.number, inv.customer, inv.contract_number, inv.delivery_order, f"{v['total']:.2f}",
                     f"{(inv.amount_paid if inv.amount_paid is not None else v['total']):.2f}",
                     "" if inv.interest_paid is None else f"{inv.interest_paid:.2f}", inv.submitted_date,
                     v["prompt_payment"]["due_date"], (pd - sub).days if sub and pd else ""])
    return _csv(PAYMENT_HEADERS, rows, "payments-received.csv")


@router.get("/export/wave-payments.csv")
def export_wave_payments(start: str = "", end: str = "", db: Session = Depends(get_db)):
    rows = []
    for inv, v in _paid(db, start, end):
        amt = (inv.amount_paid if inv.amount_paid is not None else v["total"]) + (inv.interest_paid or 0)
        desc = f"Payment {inv.customer} invoice {inv.number}" + (f" contract {inv.contract_number}" if inv.contract_number else "")
        rows.append([inv.paid_date, desc.strip(), f"{amt:.2f}"])
    return _csv(WAVE_HEADERS, rows, "wave-payments.csv")


@router.get("/export/loan-payments.csv")
def export_loan_payments(start: str = "", end: str = "", db: Session = Depends(get_db)):
    rows = []
    for loan in db.scalars(select(Loan)).all():
        st = loan_state(db, loan)
        for p in st["payments"]:
            if _in_range(p["date"], start, end):
                rows.append([p["date"], loan.lender, PROGRAM_LABELS.get(loan.program, loan.program), f"{p['amount']:.2f}",
                             f"{p['principal']:.2f}", f"{p['interest']:.2f}", f"{p['balance_after']:.2f}", p["note"]])
    rows.sort(key=lambda r: r[0])
    return _csv(LOAN_PAYMENT_HEADERS, rows, "loan-payments.csv")


# ------------------------------------------------------------------ loans
def amortized_payment(principal: float, annual_rate_pct: float, months: int) -> float:
    """Level monthly payment: P * r / (1 - (1 + r)^-n), r = annual rate / 12."""
    if months <= 0 or principal <= 0:
        return 0.0
    r = (annual_rate_pct or 0) / 100 / 12
    if r == 0:
        return round(principal / months, 2)
    return round(principal * r / (1 - (1 + r) ** -months), 2)


def _first_due(loan: Loan) -> date | None:
    return _d(loan.first_payment_date) or (add_months(_d(loan.start_date), 1) if _d(loan.start_date) else None)


def loan_state(db: Session, loan: Loan, today: date | None = None) -> dict:
    """Balance, payment split (principal and interest) and next payment date for a loan."""
    today = today or date.today()
    draws = db.scalars(select(LoanDraw).where(LoanDraw.loan_id == loan.id).order_by(LoanDraw.date, LoanDraw.id)).all()
    pays = db.scalars(select(LoanPayment).where(LoanPayment.loan_id == loan.id).order_by(LoanPayment.date, LoanPayment.id)).all()
    r = (loan.rate or 0) / 100
    basis = loan.balance_basis if loan.balance_basis in ("principal", "draws") else "principal"
    events: list[tuple[str, int, Any]] = []
    balance = 0.0
    if basis == "principal":
        balance = float(loan.principal or 0)
    else:
        events += [(d.date, 0, d) for d in draws]
    events += [(p.date, 1, p) for p in pays]
    events.sort(key=lambda e: (e[0], e[1]))
    last = _d(loan.start_date)
    accrued = 0.0
    out_pays, int_total, prin_total = [], 0.0, 0.0
    for when, kind, obj in events:
        d = _d(when)
        if loan.day_count in ("actual_365", "actual_360") and d and last and d > last:
            accrued += balance * r * (d - last).days / (365 if loan.day_count == "actual_365" else 360)
        if d and (last is None or d > last):
            last = d
        if kind == 0:
            if last is None:
                last = d
            balance += obj.amount
            continue
        if obj.interest_override is not None or obj.principal_override is not None:
            interest = obj.interest_override if obj.interest_override is not None else obj.amount - (obj.principal_override or 0)
            principal = obj.principal_override if obj.principal_override is not None else obj.amount - interest
            accrued = 0.0
        else:
            due_int = balance * r / 12 if loan.day_count not in ("actual_365", "actual_360") else accrued
            interest = min(obj.amount, round(due_int, 2))
            principal = obj.amount - interest
            accrued = max(accrued - interest, 0.0)
        principal = min(principal, balance)
        balance = round(balance - principal, 2)
        int_total += interest
        prin_total += principal
        out_pays.append({"id": obj.id, "date": obj.date, "amount": round(obj.amount, 2), "interest": round(interest, 2),
                         "principal": round(principal, 2), "balance_after": balance, "note": obj.note,
                         "entered_split": obj.interest_override is not None or obj.principal_override is not None})
    computed = amortized_payment(float(loan.principal or 0), loan.rate or 0, int(loan.term_months or 0)) if loan.amortizing \
        else round(balance * r / 12, 2)
    first = _first_due(loan)
    nxt = None
    if first and loan.status != "paid_off" and not (basis == "principal" and balance <= 0.005 and pays):
        nxt = add_months(first, len(pays))
    disbursed = sum(d.amount for d in draws)
    return {
        "balance": round(balance, 2), "disbursed": round(disbursed, 2),
        "available": round(max(float(loan.principal or 0) - balance, 0), 2) if basis == "draws" else None,
        "principal_paid": round(prin_total, 2), "interest_paid": round(int_total, 2),
        "computed_payment": computed,
        "payment": loan.payment_amount if loan.payment_amount is not None else computed,
        "next_payment_date": nxt.isoformat() if nxt else "",
        "payments": out_pays,
        "draws": [{"id": d.id, "date": d.date, "amount": d.amount, "purpose": d.purpose} for d in draws],
    }


def use_of_funds(draws: list[dict]) -> list[dict]:
    totals: dict[str, float] = defaultdict(float)
    for d in draws:
        key = " ".join((d.get("purpose") or "").split()) or "Unspecified"
        totals[key[:1].upper() + key[1:]] += float(d.get("amount") or 0)
    grand = sum(totals.values()) or 1
    return sorted(({"purpose": k, "amount": round(v, 2), "pct": round(100 * v / grand, 1)} for k, v in totals.items()),
                  key=lambda x: -x["amount"])


def amortization_schedule(loan: Loan, balance_now: float | None = None) -> dict:
    """Full schedule from the original principal (amortizing loans) or interest-only estimate (lines of credit)."""
    r = (loan.rate or 0) / 100 / 12
    first = _first_due(loan) or add_months(date.today(), 1)
    n = int(loan.term_months or 0)
    rows = []
    if not loan.amortizing:
        bal = float(balance_now if balance_now is not None else loan.principal or 0)
        pay = round(bal * r, 2)
        for i in range(min(n or 12, 12)):
            rows.append({"n": i + 1, "date": add_months(first, i).isoformat(), "payment": pay, "interest": pay,
                         "principal": 0.0, "balance": round(bal, 2)})
        return {"payment": pay, "rows": rows, "total_interest": round(pay * len(rows), 2),
                "note": "Interest only on the current balance at the current rate. Principal is repaid per your agreement."}
    bal = float(loan.principal or 0)
    pay = loan.payment_amount if loan.payment_amount is not None else amortized_payment(bal, loan.rate or 0, n)
    total_int = 0.0
    for i in range(n):
        interest = round(bal * r, 2)
        principal = round(min(pay - interest, bal), 2)
        if i == n - 1:
            principal = round(bal, 2)
        payment = round(interest + principal, 2)
        bal = round(bal - principal, 2)
        total_int += interest
        rows.append({"n": i + 1, "date": add_months(first, i).isoformat(), "payment": payment, "interest": interest,
                     "principal": principal, "balance": max(bal, 0.0)})
        if bal <= 0:
            break
    note = "Fixed rate." if loan.rate_type == "fixed" else f"Variable rate ({loan.rate_base or 'base not entered'}); schedule uses today's rate."
    if rows and rows[-1]["payment"] > pay * 1.5:
        note += " The last payment includes a balloon."
    return {"payment": pay, "rows": rows, "total_interest": round(total_int, 2), "note": note}


def _loan_view(db: Session, loan: Loan) -> dict:
    st = loan_state(db, loan)
    return {"id": loan.id, "lender": loan.lender, "program": loan.program, "program_label": PROGRAM_LABELS.get(loan.program, loan.program),
            "principal": loan.principal, "rate_type": loan.rate_type, "rate": loan.rate, "rate_base": loan.rate_base,
            "term_months": loan.term_months, "start_date": loan.start_date, "first_payment_date": loan.first_payment_date,
            "amortizing": loan.amortizing, "payment_amount": loan.payment_amount, "balance_basis": loan.balance_basis,
            "day_count": loan.day_count, "status": loan.status, "notes": loan.notes, **st,
            "use_of_funds": use_of_funds(st["draws"])}


def _get_loan(db: Session, loan_id: int) -> Loan:
    loan = db.get(Loan, loan_id)
    if loan is None:
        raise HTTPException(404, "Loan not found")
    return loan


def _apply_loan(loan: Loan, data: dict) -> None:
    if "program" in data and data["program"] not in LOAN_PROGRAMS:
        raise HTTPException(400, f"Unknown program {data['program']}")
    for k, v in data.items():
        if hasattr(loan, k) and k != "id":
            setattr(loan, k, v)


@router.get("/loans")
def list_loans(db: Session = Depends(get_db)):
    return [_loan_view(db, l) for l in db.scalars(select(Loan).order_by(Loan.id)).all()]


@router.post("/loans")
def create_loan(body: LoanIn, db: Session = Depends(get_db)):
    data = body.model_dump(exclude_none=True)
    loan = Loan(program="sba_7a", rate_type="fixed", term_months=120, amortizing=True, day_count="monthly", status="active")
    if data.get("program") == "line_of_credit":
        loan.amortizing, loan.balance_basis = False, "draws"
    _apply_loan(loan, data)
    loan.balance_basis = loan.balance_basis or "principal"
    db.add(loan)
    db.commit()
    return _loan_view(db, loan)


@router.post("/loans/calc")
def loan_calc(body: CalcIn):
    return {"payment": amortized_payment(body.principal, body.rate, body.term_months)}


@router.get("/loans/use-of-funds")
def all_use_of_funds(db: Session = Depends(get_db)):
    draws = db.scalars(select(LoanDraw)).all()
    return use_of_funds([{"purpose": d.purpose, "amount": d.amount} for d in draws])


@router.get("/loans/{loan_id}")
def get_loan(loan_id: int, db: Session = Depends(get_db)):
    return _loan_view(db, _get_loan(db, loan_id))


@router.put("/loans/{loan_id}")
def update_loan(loan_id: int, body: LoanIn, db: Session = Depends(get_db)):
    loan = _get_loan(db, loan_id)
    data = body.model_dump(exclude_unset=True)
    if "payment_amount" in data and data["payment_amount"] is None:
        loan.payment_amount = None
    _apply_loan(loan, {k: v for k, v in data.items() if v is not None})
    db.commit()
    return _loan_view(db, loan)


@router.delete("/loans/{loan_id}")
def delete_loan(loan_id: int, db: Session = Depends(get_db)):
    loan = _get_loan(db, loan_id)
    for model in (LoanDraw, LoanPayment):
        for row in db.scalars(select(model).where(model.loan_id == loan_id)).all():
            db.delete(row)
    db.delete(loan)
    db.commit()
    return {"ok": True}


@router.get("/loans/{loan_id}/schedule")
def loan_schedule(loan_id: int, db: Session = Depends(get_db)):
    loan = _get_loan(db, loan_id)
    return amortization_schedule(loan, loan_state(db, loan)["balance"])


@router.post("/loans/{loan_id}/draws")
def add_draw(loan_id: int, body: DrawIn, db: Session = Depends(get_db)):
    _get_loan(db, loan_id)
    if not _d(body.date):
        raise HTTPException(400, "Draw date must be YYYY-MM-DD")
    db.add(LoanDraw(loan_id=loan_id, date=body.date, amount=body.amount, purpose=body.purpose.strip()))
    db.commit()
    return _loan_view(db, _get_loan(db, loan_id))


@router.delete("/loans/{loan_id}/draws/{draw_id}")
def delete_draw(loan_id: int, draw_id: int, db: Session = Depends(get_db)):
    row = db.get(LoanDraw, draw_id)
    if row is None or row.loan_id != loan_id:
        raise HTTPException(404, "Draw not found")
    db.delete(row)
    db.commit()
    return _loan_view(db, _get_loan(db, loan_id))


@router.post("/loans/{loan_id}/payments")
def add_loan_payment(loan_id: int, body: LoanPaymentIn, db: Session = Depends(get_db)):
    _get_loan(db, loan_id)
    if not _d(body.date):
        raise HTTPException(400, "Payment date must be YYYY-MM-DD")
    db.add(LoanPayment(loan_id=loan_id, date=body.date, amount=body.amount, principal_override=body.principal_override,
                       interest_override=body.interest_override, note=body.note))
    db.commit()
    return _loan_view(db, _get_loan(db, loan_id))


@router.delete("/loans/{loan_id}/payments/{pay_id}")
def delete_loan_payment(loan_id: int, pay_id: int, db: Session = Depends(get_db)):
    row = db.get(LoanPayment, pay_id)
    if row is None or row.loan_id != loan_id:
        raise HTTPException(404, "Payment not found")
    db.delete(row)
    db.commit()
    return _loan_view(db, _get_loan(db, loan_id))


# ------------------------------------------------------------------ calendar and dashboard hooks
def calendar_items(db: Session) -> list[dict]:
    """Invoice payment due dates (submitted or accepted, unpaid) and the next three payments on each active loan."""
    out = []
    rates = get_settings(db).prompt_pay_rates or []
    for inv in db.scalars(select(Invoice).where(Invoice.status.in_(OPEN_STATUSES))).all():
        v = invoice_view(inv, rates)
        d = _d(v["prompt_payment"]["due_date"])
        if d:
            out.append({"uid": f"fin-inv-{inv.id}@govbid", "date": d,
                        "summary": f"Payment due: invoice {inv.number} (${v['balance']:,.2f})",
                        "description": f"{inv.customer} contract {inv.contract_number or '(none)'}. {v['prompt_payment']['basis']}",
                        "url": f"/finance?tab=invoices&id={inv.id}"})
    for loan in db.scalars(select(Loan).where(Loan.status != "paid_off")).all():
        st = loan_state(db, loan)
        nxt = _d(st["next_payment_date"])
        if not nxt:
            continue
        for i in range(3):
            d = add_months(nxt, i)
            out.append({"uid": f"fin-loan-{loan.id}-{d.isoformat()}@govbid", "date": d,
                        "summary": f"Loan payment: {loan.lender or PROGRAM_LABELS.get(loan.program, 'loan')} (${st['payment']:,.2f})",
                        "description": f"{PROGRAM_LABELS.get(loan.program, loan.program)} loan. Balance ${st['balance']:,.2f}.",
                        "url": f"/finance?tab=loans&id={loan.id}"})
    return out


def dashboard_items(db: Session) -> list[str]:
    """Invoices past their payment due date, rejected invoices to resubmit, loan payments due within 7 days."""
    today = date.today()
    out = []
    rates = get_settings(db).prompt_pay_rates or []
    for inv in db.scalars(select(Invoice).where(Invoice.status.in_(OPEN_STATUSES + ("rejected",)))).all():
        if inv.status == "rejected":
            out.append(f"Invoice {inv.number} was rejected in PIEE" + (f": {inv.rejection_reason[:80]}" if inv.rejection_reason else "")
                       + ". Correct and resubmit.")
            continue
        v = invoice_view(inv, rates, today)
        due = _d(v["prompt_payment"]["due_date"])
        if due and today > next_business_day(due):
            days = (today - due).days
            out.append(f"Invoice {inv.number} (${v['balance']:,.2f}) is {days} day{'s' if days != 1 else ''} past its payment due date {due.isoformat()}"
                       + (f"; about ${v['interest']['payable']:,.2f} interest accrued." if v["interest"]["payable"] else "."))
    for loan in db.scalars(select(Loan).where(Loan.status != "paid_off")).all():
        st = loan_state(db, loan, today)
        nxt = _d(st["next_payment_date"])
        if nxt and nxt <= today + timedelta(days=7):
            name = loan.lender or PROGRAM_LABELS.get(loan.program, "loan")
            when = "is past due" if nxt < today else ("is due today" if nxt == today else f"is due {nxt.isoformat()}")
            out.append(f"Loan payment to {name} (${st['payment']:,.2f}) {when}.")
    return out
