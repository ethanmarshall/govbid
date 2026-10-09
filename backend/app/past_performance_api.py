"""Past performance log: CRUD, keyword matching against solicitations, content library write-ups and a .docx volume draft."""
import re
import tempfile
from datetime import date, datetime
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel, field_validator
from sqlalchemy import String, cast, or_, select
from sqlalchemy.orm import Session

from .db import get_db
from .models import LibraryEntry
from .models_pp import PP_ROLES, PastPerformance

router = APIRouter(prefix="/api/past-performance")

ROLE_LABELS = {
    "prime": "Prime contractor",
    "sub": "Subcontractor",
    "commercial": "Commercial customer",
    "personal_project": "Independent project",
    "employment": "Performed as an employee",
}
FIELDS = [
    "title", "customer", "agency", "contract_number", "role", "contract_type", "value", "start_date", "end_date",
    "naics", "psc", "description", "results", "relevance_keywords", "contact_name", "contact_email", "contact_phone",
    "cpars_rating", "can_use_as_reference", "notes",
]
DATE_RE = re.compile(r"^\d{4}(-\d{2}(-\d{2})?)?$")


class PPIn(BaseModel):
    title: str | None = None
    customer: str | None = None
    agency: str | None = None
    contract_number: str | None = None
    role: str | None = None
    contract_type: str | None = None
    value: float | None = None
    start_date: str | None = None
    end_date: str | None = None
    naics: str | None = None
    psc: str | None = None
    description: str | None = None
    results: str | None = None
    relevance_keywords: list[str] | str | None = None
    contact_name: str | None = None
    contact_email: str | None = None
    contact_phone: str | None = None
    cpars_rating: str | None = None
    can_use_as_reference: bool | None = None
    notes: str | None = None

    @field_validator("relevance_keywords")
    @classmethod
    def _kw(cls, v):
        if v is None:
            return v
        if isinstance(v, str):
            v = v.split(",")
        seen, out = set(), []
        for k in v:
            k = (k or "").strip()
            if k and k.lower() not in seen:
                seen.add(k.lower())
                out.append(k)
        return out

    @field_validator("role")
    @classmethod
    def _role(cls, v):
        if v is not None and v not in PP_ROLES:
            raise ValueError(f"role must be one of {', '.join(PP_ROLES)}")
        return v

    @field_validator("start_date", "end_date")
    @classmethod
    def _date(cls, v):
        v = (v or "").strip() if v is not None else v
        if v and not DATE_RE.match(v):
            raise ValueError("dates must be YYYY-MM-DD")
        return v


def _out(p: PastPerformance) -> dict:
    d = {f: getattr(p, f) for f in FIELDS}
    d["relevance_keywords"] = list(p.relevance_keywords or [])
    d["id"] = p.id
    d["role_label"] = ROLE_LABELS.get(p.role, p.role)
    d["created_at"] = p.created_at.isoformat() if p.created_at else None
    d["updated_at"] = p.updated_at.isoformat() if p.updated_at else None
    return d


def _apply(p: PastPerformance, body: PPIn, partial: bool):
    data = body.model_dump(exclude_unset=partial)
    for k, v in data.items():
        if v is None and k not in ("value",):
            v = [] if k == "relevance_keywords" else (False if k == "can_use_as_reference" else ("prime" if k == "role" else ""))
        setattr(p, k, v)


def _get(db: Session, pp_id: int) -> PastPerformance:
    p = db.get(PastPerformance, pp_id)
    if not p:
        raise HTTPException(404, "Past performance record not found")
    return p


# ------------------------------------------------------------------ matching
_STOP = set("""a an and are as at be by for from has have in into is it its of on or that the this to was were will with
within without shall should may must any all each per other such than then these those their there which who whom
contractor government provide provided provides including include includes item items services service work required
requirement requirements support new use used using under over not only also can our your you we they""".split())


def _tokens(text: str) -> set[str]:
    words = re.findall(r"[a-z0-9][a-z0-9\-/]*[a-z0-9]|[a-z0-9]", (text or "").lower())
    out = set()
    for w in words:
        if len(w) < 3 or w in _STOP or (w.isdigit() and len(w) < 4):
            continue
        out.add(w)
        if w.endswith("s") and len(w) > 4:
            out.add(w[:-1])  # crude plural folding: "panels" also matches "panel"
    return out


def match_past_performance(db: Session, text: str, limit: int = 3) -> list[dict]:
    """Rank past performance records by keyword overlap with `text` (a solicitation title, description or SOW).

    Scoring: each relevance keyword found in the text counts 3 (phrases match as substrings), each overlapping
    title word counts 2, each overlapping description/results word counts 1 (capped at 10). A matching NAICS
    code adds 3. Returns [{id, title, customer, role, score, matched: [...]}], best first, score > 0 only.
    """
    low = (text or "").lower()
    toks = _tokens(text)
    if not toks and not low.strip():
        return []
    results = []
    for p in db.scalars(select(PastPerformance)):
        matched: list[str] = []
        score = 0
        for k in p.relevance_keywords or []:
            kl = k.lower().strip()
            if not kl:
                continue
            hit = kl in low if " " in kl or len(kl) < 4 else bool(_tokens(kl) & toks) or kl in low
            if hit:
                score += 3
                matched.append(k)
        title_hits = (_tokens(p.title) & toks) - {m.lower() for m in matched}
        score += 2 * len(title_hits)
        body_hits = (_tokens(f"{p.description} {p.results}") & toks) - title_hits - {m.lower() for m in matched}
        score += min(len(body_hits), 10)
        if p.naics and p.naics in low:
            score += 3
            matched.append(f"NAICS {p.naics}")
        if score <= 0:
            continue
        matched += sorted(title_hits) + sorted(body_hits)[:8]
        results.append({"id": p.id, "title": p.title, "customer": p.customer, "role": p.role, "score": score, "matched": matched})
    results.sort(key=lambda r: (-r["score"], r["id"]))
    return results[:limit]


# ------------------------------------------------------------------ write-ups
def _money(v):
    return "" if v is None else f"${v:,.0f}" if v >= 100 else f"${v:,.2f}"


def _period(p: PastPerformance) -> str:
    if p.start_date and p.end_date:
        return f"{p.start_date} to {p.end_date}"
    if p.start_date:
        return f"{p.start_date} to present"
    return p.end_date or ""


def _facts(p: PastPerformance) -> list[tuple[str, str]]:
    rows = [
        ("Contract / project", p.title),
        ("Contract number", p.contract_number),
        ("Customer", p.customer),
        ("Agency", p.agency),
        ("Our role", ROLE_LABELS.get(p.role, p.role)),
        ("Contract type", p.contract_type),
        ("Value", _money(p.value)),
        ("Period of performance", _period(p)),
        ("NAICS / PSC", " / ".join(x for x in (p.naics, p.psc) if x)),
        ("CPARS rating", p.cpars_rating),
    ]
    return [(k, v) for k, v in rows if v]


def _reference(p: PastPerformance) -> str:
    if not p.can_use_as_reference:
        return ""
    return ", ".join(x for x in (p.contact_name, p.contact_email, p.contact_phone) if x)


def library_markdown(p: PastPerformance) -> str:
    lines = [f"## {p.title or 'Untitled project'}", ""]
    lines += [f"- **{k}:** {v}" for k, v in _facts(p)]
    ref = _reference(p)
    if ref:
        lines.append(f"- **Reference:** {ref}")
    if p.description:
        lines += ["", "### Description of work", "", p.description.strip()]
    if p.results:
        lines += ["", "### Results", "", p.results.strip()]
    if p.relevance_keywords:
        lines += ["", f"_Relevant to: {', '.join(p.relevance_keywords)}_"]
    if p.role in ("personal_project", "employment", "sub"):
        lines += ["", f"_Note: this work was performed as {ROLE_LABELS[p.role].lower()}, not as a prime contract. Describe it that way in proposals._"]
    return "\n".join(lines).strip() + "\n"


# ------------------------------------------------------------------ endpoints
@router.get("")
def list_pp(q: str = "", db: Session = Depends(get_db)):
    stmt = select(PastPerformance)
    if q.strip():
        like = f"%{q.strip()}%"
        stmt = stmt.where(or_(
            PastPerformance.title.ilike(like), PastPerformance.customer.ilike(like), PastPerformance.agency.ilike(like),
            PastPerformance.description.ilike(like), cast(PastPerformance.relevance_keywords, String).ilike(like),
        ))
    rows = db.scalars(stmt).all()
    rows = sorted(rows, key=lambda p: (p.end_date or "9999", p.start_date or ""), reverse=True)
    return [_out(p) for p in rows]


@router.get("/meta")
def meta():
    return {"roles": [{"value": r, "label": ROLE_LABELS[r]} for r in PP_ROLES]}


@router.post("/match")
def match(body: dict, db: Session = Depends(get_db)):
    return match_past_performance(db, str(body.get("text", "")), int(body.get("limit", 3) or 3))


@router.get("/export.docx")
def export_docx(db: Session = Depends(get_db)):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    rows = sorted(db.scalars(select(PastPerformance)).all(), key=lambda p: (p.end_date or "9999"), reverse=True)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    t = doc.add_heading("Past Performance Volume", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(f"Draft generated {date.today().isoformat()}. Edit to match the solicitation's format and page limits.")
    p.runs[0].italic = True

    doc.add_heading("Summary of relevant experience", level=1)
    if rows:
        tbl = doc.add_table(rows=1, cols=5)
        tbl.style = "Light Grid Accent 1"
        for i, h in enumerate(["#", "Project", "Customer", "Role", "Period"]):
            tbl.rows[0].cells[i].text = h
        for n, r in enumerate(rows, start=1):
            c = tbl.add_row().cells
            c[0].text, c[1].text, c[2].text = str(n), r.title or "", r.customer or r.agency or ""
            c[3].text, c[4].text = ROLE_LABELS.get(r.role, r.role), _period(r)
    else:
        doc.add_paragraph("No past performance records yet.")

    for n, r in enumerate(rows, start=1):
        doc.add_page_break()
        doc.add_heading(f"{n}. {r.title or 'Untitled project'}", level=1)
        tbl = doc.add_table(rows=0, cols=2)
        tbl.style = "Table Grid"
        facts = _facts(r)
        ref = _reference(r)
        if ref:
            facts.append(("Point of contact", ref))
        for k, v in facts:
            cells = tbl.add_row().cells
            cells[0].text = k
            cells[0].paragraphs[0].runs[0].bold = True
            cells[1].text = str(v)
        if r.description:
            doc.add_heading("Description of work", level=2)
            for para in r.description.strip().split("\n\n"):
                doc.add_paragraph(para.strip())
        if r.results:
            doc.add_heading("Results and relevance", level=2)
            for para in r.results.strip().split("\n\n"):
                doc.add_paragraph(para.strip())
        if r.relevance_keywords:
            doc.add_paragraph("Relevant to: " + ", ".join(r.relevance_keywords)).runs[0].italic = True
        if r.role in ("personal_project", "employment", "sub"):
            note = doc.add_paragraph(f"Performed as {ROLE_LABELS[r.role].lower()}, not as a prime contract.")
            note.runs[0].italic = True

    out = Path(tempfile.mkdtemp()) / f"past-performance-{date.today().isoformat()}.docx"
    doc.save(out)
    return FileResponse(out, filename=out.name, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


@router.post("")
def create_pp(body: PPIn, db: Session = Depends(get_db)):
    if not (body.title or "").strip():
        raise HTTPException(422, "title is required")
    p = PastPerformance()
    _apply(p, body, partial=False)
    db.add(p)
    db.commit()
    return _out(p)


@router.get("/{pp_id}")
def get_pp(pp_id: int, db: Session = Depends(get_db)):
    return _out(_get(db, pp_id))


@router.put("/{pp_id}")
def update_pp(pp_id: int, body: PPIn, db: Session = Depends(get_db)):
    p = _get(db, pp_id)
    if body.title is not None and not body.title.strip():
        raise HTTPException(422, "title cannot be empty")
    _apply(p, body, partial=True)
    p.updated_at = datetime.utcnow()
    db.commit()
    return _out(p)


@router.delete("/{pp_id}")
def delete_pp(pp_id: int, db: Session = Depends(get_db)):
    p = _get(db, pp_id)
    db.delete(p)
    db.commit()
    return {"ok": True}


@router.post("/{pp_id}/to-library")
def to_library(pp_id: int, db: Session = Depends(get_db)):
    p = _get(db, pp_id)
    tags = ["past performance", p.role] + [k for k in (p.naics, p.psc) if k] + list(p.relevance_keywords or [])[:8]
    e = LibraryEntry(category="Past performance", title=p.title or "Past performance", content=library_markdown(p), tags=tags)
    db.add(e)
    db.commit()
    return {"id": e.id, "title": e.title, "category": e.category}
