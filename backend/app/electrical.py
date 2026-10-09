"""Electrical quoters: cable and wire harnesses, control panels, and labels / nameplates.

Inputs come from a CSV or XLSX (any sheet), a drawing PDF whose text holds the tables (read with
the layout-preserving extraction from app/extrusion.py so columns stay aligned), or lines typed by
hand in the UI. Rates and placeholder prices live in the shop-rate config under "harness", "panel"
and "labels" (defaults and sources in app/electrical_catalog.py).

Pricing follows the house model used by the other quoters:
  per unit  = material + labor minutes x rate (+ test and packaging)
  per lot   = setup, kitting, test setup, first article, certificate of conformance, freight
  price(q)  = (per unit x q + per lot) x (1 + G&A) x (1 + profit), at least the minimum lot charge

Harness labor is multiplied by an IPC/WHMA-A-620 workmanship class factor (Class 1, 2 or 3). The
factor is an editable assumption, not something the standard defines. Contact size checks use the
MIL-DTL-38999 crimp contact wire ranges in electrical_catalog.CONTACT_SIZE_AWG.
"""
from __future__ import annotations

import csv
import io
import math
import re
from collections import Counter, defaultdict
from pathlib import Path

from . import pricing
from .extrusion import STOP  # title block and notes markers end a table in drawing text
from .electrical_catalog import (CONTACT_SIZE_AWG, LENGTH_PRICED, M39029_DASH_SIZE, PANEL_DEVICE_TYPES,
                                 WORKMANSHIP_CLASSES)

KINDS = ("harness", "panel", "labels")


class ElectricalError(ValueError):
    pass


# ================================================================ small helpers
def _num(v, default: float = 0.0) -> float:
    try:
        return float(v if v not in (None, "") else default)
    except (TypeError, ValueError):
        raise ElectricalError(f"Expected a number, got {v!r}")


def _opt_num(v) -> float | None:
    if v in (None, ""):
        return None
    return _num(v)


def _int_or_none(text) -> int | None | bool:
    """Whole number from a cell; None when blank; False when the cell holds something else."""
    s = str(text if text is not None else "").strip()
    if not s:
        return None
    m = re.fullmatch(r"(\d+)(?:\.0+)?\s*(?:EA|PCS?|X)?", s, re.I)
    return int(m.group(1)) if m else False


def _qtys(quantities) -> list[int]:
    try:
        q = sorted({int(x) for x in (quantities or [1]) if int(x) > 0})
    except (TypeError, ValueError):
        raise ElectricalError("quantities must be positive whole numbers")
    if not q:
        raise ElectricalError("Give at least one quantity")
    return q


def _line(category: str, item: str, cost: float, *, basis: str = "per_unit", hours: float | None = None,
          rate: float | None = None, note: str = "") -> dict:
    return {"category": category, "item": item, "basis": basis, "hours": None if hours is None else round(hours, 3),
            "rate": rate, "cost": round(cost, 2), "note": note}


def _natural(s: str):
    return [int(p) if p.isdigit() else p for p in re.split(r"(\d+)", str(s or "").upper())]


UNIT_IN = {"in": 1.0, "ft": 12.0, "mm": 1 / 25.4, "cm": 1 / 2.54, "m": 1000 / 25.4}


def _unit_from_header(h: str) -> str:
    h = (h or "").upper()
    if re.search(r"\(\s*MM\s*\)|\bMM\b|MILLIM", h):
        return "mm"
    if re.search(r"\(\s*CM\s*\)|\bCM\b", h):
        return "cm"
    if re.search(r"\(\s*FT\s*\)|\bFT\b|FEET|FOOT", h):
        return "ft"
    if re.search(r"\(\s*M\s*\)|\bMETERS?\b|\bMETRES?\b", h):
        return "m"
    return "in"


def length_in(text, unit: str = "in") -> float | None:
    """Inches from a cell: bare numbers use the column unit, otherwise units in the text win (24 in, 2 ft, 610 mm)."""
    s = str(text if text is not None else "").strip()
    if not s:
        return None
    if re.fullmatch(r"\d+(?:\.\d+)?|\.\d+", s):
        return round(float(s) * UNIT_IN.get(unit, 1.0), 4)
    from .extrusion import parse_length

    return parse_length(s, "mm" if unit == "mm" else "in")


# ================================================================ table reading (CSV, XLSX, PDF text)
Cell = tuple[int, str]


def _rows_from_sheet(filename: str, data: bytes) -> list[list[Cell]]:
    name = filename.lower()
    rows: list[list] = []
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ElectricalError(f"Could not read the spreadsheet: {exc}") from exc
        for ws in wb.worksheets:  # every sheet: a wire list and a connector list often sit on separate tabs
            rows += [list(r) for r in ws.iter_rows(values_only=True)]
            rows.append([])
    elif name.endswith((".csv", ".txt", ".tsv")):
        text = data.decode("utf-8-sig", errors="replace")
        first = text.split("\n", 1)[0]
        dialect = csv.excel_tab if name.endswith(".tsv") or ("\t" in first and "," not in first) else csv.excel
        rows = [r for r in csv.reader(io.StringIO(text), dialect)]
    else:
        raise ElectricalError("Upload a drawing PDF, CSV or XLSX file.")

    def fmt(v) -> str:
        if v is None:
            return ""
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v).strip()

    return [[(i, fmt(v)) for i, v in enumerate(r) if fmt(v)] for r in rows]


def _rows_from_text(text: str) -> list[list[Cell]]:
    """Layout-preserved PDF text -> cells (runs separated by 2+ spaces) with their character columns."""
    return [[(m.start(), m.group(0)) for m in re.finditer(r"\S+(?: \S+)*", ln)] for ln in (text or "").splitlines()]


def _map_header(texts: list[str], cols: list[tuple[str, str]]) -> list[str | None]:
    used: set[str] = set()
    out: list[str | None] = []
    for t in texts:
        name = re.sub(r"\s+", " ", str(t or "").strip().upper())
        key = None
        for k, pat in cols:
            if k not in used and re.search(pat, name):
                key = k
                used.add(k)
                break
        out.append(key)
    return out


def _scan(rows: list[list[Cell]], schemas: dict, positional: bool) -> dict[str, list[dict]]:
    """Find header rows for any schema and collect the records under them.

    schemas: name -> (cols, header_ok(keys) -> bool, row_fn(record, header_texts) -> dict | None).
    positional: CSV/XLSX cells align by column index; PDF text cells align to the nearest header column at or left of them.
    """
    found: dict[str, list[dict]] = defaultdict(list)
    cur = None
    for cells in rows:
        if not cells:
            continue
        texts = [t for _, t in cells]
        hit = None
        for name, (cols, ok, _) in schemas.items():
            keys = _map_header(texts, cols)
            if ok({k for k in keys if k}):
                hit = (name, keys)
                break
        if hit:
            name, keys = hit
            cur = {"name": name, "cols": [(pos, k, txt) for (pos, txt), k in zip(cells, keys)], "misses": 0}
            continue
        if cur is None:
            continue
        if not positional and STOP.search(" ".join(texts)):
            cur = None
            continue
        rec: dict[str, str] = {}
        bad = False
        for pos, txt in cells:
            if positional:
                col = next((c for c in cur["cols"] if c[0] == pos), None)
            else:
                col = None
                for c in cur["cols"]:
                    if c[0] <= pos + 2:
                        col = c
                if col is None:
                    bad = True
                    break
            if col is None or col[1] is None:
                continue
            rec[col[1]] = (rec.get(col[1], "") + " " + txt).strip()
        heads = {k: txt for _, k, txt in cur["cols"] if k}
        out = None if bad else schemas[cur["name"]][2](rec, heads)
        if out is None:
            cur["misses"] += 1
            if cur["misses"] >= 4:
                cur = None
            continue
        cur["misses"] = 0
        found[cur["name"]].append(out)
    return found


# ================================================================ harness: parsing
WIRE_COLS = [
    ("from_term", r"^FROM\s*TERM|^TERM(?:INATION)?\s*\(?\s*FROM|^TERM\s*A$"),
    ("to_term", r"^TO\s*TERM|^TERM(?:INATION)?\s*\(?\s*TO\b|^TERM\s*B$"),
    ("from_pin", r"^FROM\s*(?:PIN|CAV|CAVITY|POS|CONTACT)|^PIN\s*\(?\s*FROM|^PIN\s*A$"),
    ("to_pin", r"^TO\s*(?:PIN|CAV|CAVITY|POS|CONTACT)|^PIN\s*\(?\s*TO\b|^PIN\s*B$"),
    ("from_ref", r"^FROM\s*(?:CONN|CONNECTOR|REF|DEVICE|DESIG|LOC)|^CONN(?:ECTOR)?\s*\(?\s*FROM|^END\s*A$"),
    ("to_ref", r"^TO\s*(?:CONN|CONNECTOR|REF|DEVICE|DESIG|LOC)|^CONN(?:ECTOR)?\s*\(?\s*TO\b|^END\s*B$"),
    ("from", r"^FROM\b"),
    ("to", r"^TO\b"),
    ("wire_id", r"^WIRE\s*(?:ID|NO\.?|#|NUMBER|NAME)?$|^CIRCUIT|^ID$|^W/?N$|^NET\b"),
    ("gauge", r"^GAUGE|^AWG|^GA\.?$|^WIRE\s*(?:SIZE|GAUGE)|^SIZE"),
    ("color", r"^COLOU?R|^CLR"),
    ("length", r"^LENGTH|^LEN\b|^CUT\s*LEN"),
    ("spec", r"^WIRE\s*(?:SPEC|TYPE|PART|P/?N)|^SPEC|^TYPE$|^MIL\b|^INSULATION|^PART"),
    ("shield", r"^SHIELD|^TWIST|^SH\s*/\s*TW|^CABLE"),
    ("signal", r"^SIGNAL|^FUNCTION|^DESC"),
    ("notes", r"^NOTE|^REMARK|^COMMENT"),
]

BOM_COLS = [
    ("contact_size", r"^CONTACT\s*SIZE|^CONT\.?\s*SIZE|^SIZE$"),
    ("contacts", r"^CONTACTS?\b|^PINS?\b|^TERMINALS?$"),
    ("backshell", r"^BACK\s*SHELL|^STRAIN"),
    ("ref", r"^REF|^DESIGNATOR|^DESIG|^CONN(?:ECTOR)?\s*(?:REF|ID|DESIG)?$|^LOC|^TAG$|^DEVICE\s*(?:TAG|ID)"),
    ("part", r"^PART|^P/?N\b|^MPN|^MFR\.?\s*P|^MFG\.?\s*P|^MANUFACTURER\s*P|^CAT|^MODEL|^ORDER"),
    ("manufacturer", r"^MFR\.?$|^MFG\.?$|^MANUFACTURER$|^MAKE$|^BRAND|^VENDOR"),
    ("desc", r"^DESC|^NOMENCLATURE|^NAME$"),
    ("qty", r"^QTY|^QUANTITY|^Q'?TY|^COUNT$"),
    ("price", r"^UNIT\s*(?:PRICE|COST)|^PRICE|^COST|^\$"),
    ("length", r"^LENGTH|^LEN\b"),
    ("device_type", r"^DEVICE\s*TYPE|^CATEGORY|^TYPE$|^KIND$"),
    ("item", r"^ITEM|^FIND|^LINE$|^#$|^NO\.?$"),
    ("notes", r"^NOTE|^REMARK|^COMMENT"),
]

REF = r"[A-Z]{1,4}\d{1,4}[A-Z]?"
END_RX = re.compile(rf"^\s*({REF})\s*(?:[-:./]\s*|\s+PIN\s*|\s+)(?:PIN\s*)?([A-Z0-9]{{1,4}})\s*$", re.I)
REF_ONLY = re.compile(rf"^\s*({REF})\s*$", re.I)
SPEC_RX = [
    (re.compile(r"\bM\s*-?\s*22759\s*/\s*(\d{1,3})(?:\s*-\s*(\d{1,2})\s*-\s*(\d{1,4}))?", re.I), "M22759/{}"),
    (re.compile(r"\bM\s*-?\s*16878\s*/\s*(\d{1,2})(?:\s*-?\s*([A-Z]{2,3}))?", re.I), "M16878/{}"),
    (re.compile(r"\bUL\s*-?\s*(\d{4})\b", re.I), "UL{}"),
    (re.compile(r"\bMIL\s*-?\s*W\s*-?\s*76\b", re.I), "MIL-W-76"),
]


def split_end(text: str) -> tuple[str, str]:
    """'P1-3', 'J2:A', 'P1 PIN 3', 'TB1-12' -> (ref, pin). 'E1' -> ('E1', '')."""
    s = str(text or "").strip()
    m = END_RX.match(s)
    if m:
        return m.group(1).upper(), m.group(2).upper()
    m = REF_ONLY.match(s)
    if m:
        return m.group(1).upper(), ""
    return "", ""


def parse_spec(text: str) -> tuple[str, str]:
    """Wire spec text -> (base spec like 'M22759/16', gauge from the part number or '')."""
    t = str(text or "")
    for rx, fmt in SPEC_RX:
        m = rx.search(t)
        if m:
            gauge = ""
            if fmt.startswith("M22759") and m.group(2):
                gauge = str(int(m.group(2)))
            return fmt.format(m.group(1)) if "{}" in fmt else fmt, gauge
    return t.strip().upper(), ""


def parse_gauge(text) -> str:
    m = re.search(r"(?<!\d)#?\s*(\d{1,2})\s*(?:AWG|GA\b)?", str(text or ""), re.I)
    return str(int(m.group(1))) if m and 0 < int(m.group(1)) <= 40 else ""


TERM_WORDS = [("solder", r"SOLDER|\bSLD\b"), ("splice", r"SPLICE"), ("lug", r"\bLUG|RING|SPADE|FORK|TERMINAL"),
              ("open", r"OPEN|FLYING|STRIP|NONE|N/C"), ("crimp", r"CRIMP|CONTACT|PIN|SOCKET")]


def parse_term(text) -> str:
    t = str(text or "").upper()
    for k, pat in TERM_WORDS:
        if re.search(pat, t):
            return k
    return ""


def blank_wire(**kw) -> dict:
    w = {"wire_id": "", "from_ref": "", "from_pin": "", "to_ref": "", "to_pin": "", "gauge": "", "color": "", "length_in": None,
         "spec": "", "shielded": False, "twisted": False, "from_term": "", "to_term": "", "signal": "", "notes": "", "raw": ""}
    w.update(kw)
    return w


def _wire_row(rec: dict, heads: dict) -> dict | None:
    fr, fp = (rec.get("from_ref", "").upper(), rec.get("from_pin", "").upper()) if rec.get("from_ref") else split_end(rec.get("from", ""))
    tr, tp = (rec.get("to_ref", "").upper(), rec.get("to_pin", "").upper()) if rec.get("to_ref") else split_end(rec.get("to", ""))
    if rec.get("from_ref") and not fp and split_end(rec["from_ref"])[1]:
        fr, fp = split_end(rec["from_ref"])
    if rec.get("to_ref") and not tp and split_end(rec["to_ref"])[1]:
        tr, tp = split_end(rec["to_ref"])
    if not (fr and tr):
        return None
    spec, g_from_spec = parse_spec(rec.get("spec", ""))
    gauge = parse_gauge(rec.get("gauge", "")) or g_from_spec
    sh = (rec.get("shield", "") + " " + rec.get("spec", "")).upper()
    return blank_wire(
        wire_id=rec.get("wire_id", ""), from_ref=fr, from_pin=fp, to_ref=tr, to_pin=tp, gauge=gauge, color=rec.get("color", "").upper(),
        length_in=length_in(rec.get("length", ""), _unit_from_header(heads.get("length", ""))),
        spec=spec, shielded=bool(re.search(r"SHIELD|\bSH\b|\bTSP\b|\bSSP\b|\bYES\b|\bY\b", sh) and not re.search(r"UNSHIELD", sh)),
        twisted=bool(re.search(r"TWIST|\bTW\b|\bTP\b|\bTSP\b|PAIR", sh)),
        from_term=parse_term(rec.get("from_term", "")), to_term=parse_term(rec.get("to_term", "")),
        signal=rec.get("signal", ""), notes=rec.get("notes", ""), raw=" | ".join(v for v in rec.values() if v))


def _contact_size(*texts: str) -> str:
    for t in texts:
        t = str(t or "").upper()
        m = re.search(r"M\s*39029\s*/\s*(\d{2}\s*-\s*\d{3})", t)
        if m:
            size = M39029_DASH_SIZE.get(re.sub(r"\s", "", m.group(1)))
            if size:
                return size
        m = re.search(r"(?:\bSIZE\s*#?\s*|#)(22D|22M|22|20|16|12|10)\b|\b(22D|22M)\b", t)
        if m:
            return m.group(1) or m.group(2)
        if re.fullmatch(r"(22D|22M|22|20|16|12|10)", t.strip()):
            return t.strip()
    return ""


def classify_harness_item(part: str, desc: str, ref: str = "") -> str:
    t = f"{part} {desc}".upper()
    if re.search(r"BACK\s*SHELL|M85049|STRAIN\s*RELIEF|CABLE\s*CLAMP", t):
        return "backshell"
    if re.search(r"HEAT\s*SHRINK|SHRINK\s*TUB|SLEEV|\bTIE\b|TIES\b|LACING|LABEL|MARKER|\bBOOT\b|SPLICE|\bLUG\b|GROMMET|TAPE", t):
        return "accessory"
    if re.search(r"D38999|MS\s*-?\s*3\d{3}|MS27\d{3}|M24308|M83723|\bPLUG\b|RECEPTACLE|\bRECPT?\b|HOUSING|\bHEADER\b|D-?SUB", t):
        return "connector"
    if re.search(r"M\s*39029|\bCONTACTS?\b|\bCRIMP\s*(?:PIN|SOCKET)\b|\bTERMINAL\s*PIN", t):
        return "contact"
    if re.search(r"CONNECTOR|\bCONN\b", t) or re.match(r"^[PJ]\d", ref or ""):
        return "connector"
    return "other"


def blank_bom(**kw) -> dict:
    b = {"ref": "", "kind": "other", "part_number": "", "manufacturer": "", "description": "", "qty": 1, "contact_size": "",
         "unit_price": None, "price_source": "", "distributor": "", "notes": "", "raw": ""}
    b.update(kw)
    return b


def _bom_row_harness(rec: dict, heads: dict) -> list[dict] | None:
    part, desc = rec.get("part", ""), rec.get("desc", "")
    if not (part or desc):
        return None
    q = _int_or_none(rec.get("qty"))
    if q is False:
        return None
    ref = rec.get("ref", "").upper()
    kind = classify_harness_item(part, desc, ref)
    contacts = rec.get("contacts", "")
    price = None
    if rec.get("price"):
        try:
            price = float(re.sub(r"[$,\s]", "", rec["price"]))
        except ValueError:
            price = None
    size = _contact_size(rec.get("contact_size", ""), desc, part, contacts)
    raw = " | ".join(v for v in rec.values() if v)
    out = [blank_bom(ref=ref, kind=kind, part_number=part, manufacturer=rec.get("manufacturer", ""), description=desc,
                     qty=q if q is not None else (None if kind == "contact" else 1), contact_size=size, unit_price=price,
                     price_source="manual" if price is not None else "", notes=rec.get("notes", ""), raw=raw)]
    if contacts and not re.fullmatch(r"\d{1,3}", contacts.strip()):  # a contact part number: count comes from the wire list
        out.append(blank_bom(ref=ref, kind="contact", part_number=contacts.strip(), description=f"Contacts for {ref or part}",
                             qty=None, contact_size=_contact_size(contacts), raw=raw))
    if rec.get("backshell") and not re.fullmatch(r"(?i)\s*(?:-|N/?A|NONE|NO)?\s*", rec["backshell"]):
        out.append(blank_bom(ref=ref, kind="backshell", part_number=rec["backshell"].strip(), description=f"Backshell for {ref or part}",
                             qty=q or 1, raw=raw))
    return out


def _wire_header_ok(keys: set) -> bool:
    return bool(({"from", "from_ref"} & keys) and ({"to", "to_ref"} & keys) and len(keys) >= 3)


def _bom_header_ok(keys: set) -> bool:
    return bool(({"part", "desc"} & keys) and ({"qty", "ref"} & keys) and len(keys) >= 2 and not ({"from", "to", "from_ref"} & keys))


def parse_harness_rows(rows: list[list[Cell]], positional: bool) -> dict:
    found = _scan(rows, {"wires": (WIRE_COLS, _wire_header_ok, _wire_row), "bom": (BOM_COLS, _bom_header_ok, _bom_row_harness)}, positional)
    wires = found.get("wires", [])
    bom = [b for group in found.get("bom", []) for b in group]
    for i, w in enumerate(wires, 1):
        w["wire_id"] = w["wire_id"] or f"W{i}"
    return {"wires": wires, "bom": bom}


# ================================================================ panel: parsing
DEVICE_PATTERNS = [
    ("enclosure", r"ENCLOSURE|\bNEMA\s*\d|CABINET|\bJUNCTION\s*BOX"),
    ("back_panel", r"BACK\s*PANEL|SUB\s*-?\s*PANEL|MOUNTING\s*PANEL|\bINNER\s*PANEL"),
    ("din_rail", r"DIN\s*RAIL|\bTS\s*-?\s*35\b|35\s*MM\s*RAIL|\bTOP\s*HAT\s*RAIL"),
    ("wire_duct", r"WIRE\s*DUCT|WIREWAY|\bDUCT\b|RACEWAY|SLOTTED\s*DUCT"),
    ("terminal_block", r"TERMINAL\s*BLOCK|TERM(?:INAL)?\s*BLK|\bTB\b|FEED\s*-?\s*THROUGH\s*TERMINAL|GROUND\s*TERMINAL"),
    ("disconnect", r"DISCONNECT|ROTARY\s*SWITCH|SAFETY\s*SWITCH|LOAD\s*BREAK"),
    ("breaker", r"BREAKER|\bMCB\b|\bMCCB\b|\bCB\b|SUPPLEMENTARY\s*PROTECTOR"),
    ("fuse", r"\bFUSE"),
    ("contactor", r"CONTACTOR|MOTOR\s*STARTER|OVERLOAD|\bMSP\b|MOTOR\s*PROTECT"),
    ("io_module", r"\bI\s*/\s*O\b|INPUT\s*MODULE|OUTPUT\s*MODULE|ANALOG\s*(?:INPUT|OUTPUT|MODULE)|DISCRETE\s*(?:INPUT|OUTPUT)|\bIO\s*MODULE"),
    ("plc", r"\bPLC\b|PROGRAMMABLE\s*(?:LOGIC\s*)?CONTROLLER|\bCPU\b|CONTROLLER"),
    ("hmi", r"\bHMI\b|TOUCH\s*-?\s*SCREEN|OPERATOR\s*(?:INTERFACE|PANEL|TERMINAL)|DISPLAY\s*TERMINAL"),
    ("power_supply", r"POWER\s*SUPPLY|\bPSU\b|\bDC\s*SUPPLY|TRANSFORMER"),
    ("relay", r"RELAY|\bCR\d*\b|TIMER"),
    ("pilot_device", r"PUSH\s*-?\s*BUTTON|\bPB\b|PILOT\s*LIGHT|INDICATOR\s*LIGHT|SELECTOR|E\s*-?\s*STOP|EMERGENCY\s*STOP|\b22\s*MM\b|\b30\s*MM\b|STACK\s*LIGHT"),
    ("meter", r"METER|\bGAUGE\b|INDICATOR|PANEL\s*DISPLAY"),
]


def classify_device(part: str, desc: str, hint: str = "") -> str:
    h = re.sub(r"[\s/]+", "_", (hint or "").strip().lower())
    if h in PANEL_DEVICE_TYPES:
        return h
    t = f"{hint} {desc} {part}".upper()
    for k, pat in DEVICE_PATTERNS:
        if re.search(pat, t):
            return k
    return "other"


def _meters_from_text(text: str) -> float | None:
    m = re.search(r"(?<![\w.])(\d+(?:\.\d+)?)\s*(M|METERS?|METRES?|FT|FEET|FOOT|IN|INCH(?:ES)?)\b(?![\w-])", str(text or ""), re.I)
    if not m:
        return None
    v, u = float(m.group(1)), m.group(2).upper()
    return round(v if u.startswith("M") else v * 0.3048 if u.startswith("F") else v * 0.0254, 4)


def blank_panel_line(**kw) -> dict:
    ln = {"ref": "", "device_type": "other", "part_number": "", "manufacturer": "", "description": "", "qty": 1, "length_m": None,
          "unit_price": None, "price_source": "", "distributor": "", "notes": "", "matched": False, "raw": ""}
    ln.update(kw)
    return ln


def _panel_row(rec: dict, heads: dict) -> dict | None:
    part, desc = rec.get("part", ""), rec.get("desc", "")
    if not (part or desc):
        return None
    q = _int_or_none(rec.get("qty"))
    if q is False:
        try:
            q = float(re.sub(r"[^\d.]", "", rec["qty"]))
        except ValueError:
            return None
    dt = classify_device(part, desc, rec.get("device_type", ""))
    length_m = None
    if rec.get("length"):
        hl = heads.get("length", "")
        unit = _unit_from_header(hl) if (_unit_from_header(hl) != "in" or re.search(r"\bIN\b|INCH", hl, re.I)) else "m"
        li = length_in(rec["length"], unit)
        length_m = round(li * 0.0254, 4) if li else None
    elif dt in LENGTH_PRICED:
        length_m = _meters_from_text(desc)
    price = None
    if rec.get("price"):
        try:
            price = float(re.sub(r"[$,\s]", "", rec["price"]))
        except ValueError:
            price = None
    return blank_panel_line(ref=rec.get("ref", "").upper(), device_type=dt, part_number=part, manufacturer=rec.get("manufacturer", ""),
                            description=desc, qty=q if q is not None else 1, length_m=length_m, unit_price=price,
                            price_source="manual" if price is not None else "", notes=rec.get("notes", ""), matched=dt != "other",
                            raw=" | ".join(v for v in rec.values() if v))


def _panel_header_ok(keys: set) -> bool:
    return bool(({"part", "desc"} & keys) and ({"qty", "ref", "device_type"} & keys) and len(keys) >= 2)


# ================================================================ labels: parsing
LABEL_COLS = [
    ("iuid", r"^IUID|^UID\b|^UII|^DATA\s*MATRIX|^2D"),
    ("chars", r"^CHAR|^#\s*CHAR|^CHARACTERS"),
    ("type", r"^TYPE|^MATERIAL|^KIND|^STYLE"),
    ("width", r"^WIDTH|^W$|^W\s*\("),
    ("height", r"^HEIGHT|^H$|^H\s*\(|^LENGTH"),
    ("size", r"^SIZE|^DIM"),
    ("text", r"^TEXT|^LEGEND|^ENGRAV|^MARKING|^WORDING|^COPY|^CONTENT"),
    ("color", r"^COLOU?R"),
    ("holes", r"^HOLES?|^MOUNT"),
    ("adhesive", r"^ADHESIVE|^BACKING"),
    ("qty", r"^QTY|^QUANTITY|^Q'?TY|^COUNT$"),
    ("item", r"^ITEM|^TAG|^ID$|^#$|^NO\.?$|^PLATE|^MARK\b|^LABEL"),
    ("notes", r"^NOTE|^REMARK|^COMMENT"),
]
LABEL_WORDS = [
    ("heat_shrink_marker", r"HEAT\s*-?\s*SHRINK|SLEEVE|WIRE\s*MARKER"),
    ("photo_anodized_aluminum", r"PHOTO\s*-?\s*ANOD|METALPHOTO"),
    ("anodized_aluminum", r"ANODIZ|ALUMINUM|ALUMINIUM|\bAL\b"),
    ("stainless_steel", r"STAINLESS|\bSS\b|\bCRES\b|304|316"),
    ("engraved_laminate", r"PHENOLIC|LAMINATE|LAMACOID|ACRYLIC|ENGRAV(?:ED)?\s*PLASTIC|RIGID\s*PLASTIC"),
    ("printed_polyester", r"POLYESTER|\bPET\b|MYLAR|PRINTED\s*LABEL"),
    ("printed_vinyl", r"VINYL|\bPVC\b|DECAL"),
]


def classify_label(text: str) -> str:
    t = str(text or "").upper()
    for k, pat in LABEL_WORDS:
        if re.search(pat, t):
            return k
    return ""


SIZE_RX = re.compile(r"(\d+(?:\.\d+)?|\d+/\d+|\d+\s+\d+/\d+)\s*(?:\"|IN\b|MM\b)?\s*[X×]\s*(\d+(?:\.\d+)?|\d+/\d+|\d+\s+\d+/\d+)\s*(\"|IN\b|MM\b)?", re.I)


def parse_size(text: str, unit: str = "in") -> tuple[float | None, float | None]:
    m = SIZE_RX.search(str(text or ""))
    if not m:
        return None, None
    from .extrusion import _frac

    a, b = _frac(m.group(1)), _frac(m.group(2))
    mm = (m.group(3) or "").upper() == "MM" or re.search(r"\bMM\b", str(text), re.I) or unit == "mm"
    f = 1 / 25.4 if mm else 1.0
    return round(a * f, 3), round(b * f, 3)


def count_chars(text: str) -> int:
    return len(re.sub(r"\s|\||/", "", str(text or "")))


def blank_label(**kw) -> dict:
    ln = {"item": "", "type": "", "width_in": None, "height_in": None, "text": "", "chars": None, "color": "", "holes": 0,
          "adhesive": False, "qty": 1, "iuid": False, "notes": "", "raw": ""}
    ln.update(kw)
    return ln


def _yes(v) -> bool:
    return bool(re.match(r"(?i)\s*(Y|YES|TRUE|X|1|REQ)", str(v or "")))


def _label_row(rec: dict, heads: dict) -> dict | None:
    if not any(rec.get(k) for k in ("type", "text", "size", "width")):
        return None
    q = _int_or_none(rec.get("qty"))
    if q is False:
        return None
    t_all = " ".join([rec.get("type", ""), rec.get("notes", "")])
    typ = classify_label(rec.get("type", "")) or classify_label(t_all)
    unit = _unit_from_header(heads.get("width", "") or heads.get("size", ""))
    w = h = None
    if rec.get("width") or rec.get("height"):
        w, h = length_in(rec.get("width"), unit), length_in(rec.get("height"), unit)
    if (w is None or h is None) and rec.get("size"):
        w, h = parse_size(rec["size"], unit)
    chars = _int_or_none(rec.get("chars"))
    hole = _int_or_none(rec.get("holes"))
    iuid = _yes(rec.get("iuid")) or bool(re.search(r"\bIUID\b|\bUII\b|DATA\s*MATRIX|MIL-STD-130", " ".join(rec.values()), re.I))
    return blank_label(item=rec.get("item", ""), type=typ, width_in=w, height_in=h, text=rec.get("text", ""),
                       chars=chars if isinstance(chars, int) else None, color=rec.get("color", ""),
                       holes=hole if isinstance(hole, int) else (4 if _yes(rec.get("holes")) else 0),
                       adhesive=_yes(rec.get("adhesive")) or bool(re.search(r"ADHESIVE|PSA\b", t_all, re.I)),
                       qty=q if q is not None else 1, iuid=iuid, notes=rec.get("notes", ""), raw=" | ".join(v for v in rec.values() if v))


def _label_header_ok(keys: set) -> bool:
    return bool("qty" in keys and ({"type", "text", "size", "width"} & keys) and len(keys) >= 3)


# ================================================================ file entry point
def _pdf(path: Path) -> tuple[str, int, dict, list[str]]:
    from .drawing import DrawingError, parse_text
    from .extrusion import pdf_text

    try:
        text, pages = pdf_text(path)
    except DrawingError as exc:
        raise ElectricalError(str(exc)) from exc
    has_text = len(re.findall(r"[A-Za-z0-9]", text)) >= 20
    warnings: list[str] = []
    drawing: dict = {}
    if has_text:
        read = parse_text(text)
        drawing = {k: read.get(k) for k in ("part_number", "drawing_number", "revision", "title")}
        drawing["export_controlled"] = bool(read.get("export_controlled"))
        drawing["distribution"] = (read.get("distribution") or {}).get("letter") or ""
        warnings += [w for w in read.get("warnings", []) if "Distribution" in w or "Export" in w]
    else:
        warnings.append("The PDF has no text layer (scanned). Enter the lines by hand or import a CSV/XLSX.")
    return text, pages, drawing, warnings


def parse_rows(kind: str, rows: list[list[Cell]], positional: bool) -> dict:
    if kind == "harness":
        return parse_harness_rows(rows, positional)
    if kind == "panel":
        found = _scan(rows, {"bom": (BOM_COLS, _panel_header_ok, _panel_row)}, positional)
        return {"lines": found.get("bom", [])}
    if kind == "labels":
        found = _scan(rows, {"labels": (LABEL_COLS, _label_header_ok, _label_row)}, positional)
        lines = found.get("labels", [])
        for i, ln in enumerate(lines, 1):
            ln["item"] = ln["item"] or str(i)
        return {"lines": lines}
    raise ElectricalError(f"kind must be one of {KINDS}")


def parse_file(kind: str, path: str | Path, filename: str = "", config: dict | None = None) -> dict:
    """Parse an uploaded CSV, XLSX or drawing PDF for one quote kind. Returns the lines plus warnings and source info."""
    path = Path(path)
    filename = filename or path.name
    if filename.lower().endswith(".pdf"):
        text, pages, drawing, warnings = _pdf(path)
        out = parse_rows(kind, _rows_from_text(text), positional=False)
        source = {"filename": filename, "type": "pdf", "pages": pages}
    else:
        out = parse_rows(kind, _rows_from_sheet(filename, path.read_bytes()), positional=True)
        drawing, warnings, source = {}, [], {"filename": filename, "type": "table"}
    empty = not (out.get("wires") or out.get("bom") or out.get("lines"))
    if empty:
        need = {"harness": "a wire list (Wire ID, From, To, Gauge, Length) or a connector list (Ref, Part number, Description, Qty)",
                "panel": "a BOM with Part number or Description and Qty columns",
                "labels": "a plate list with Type or Text, Size (or Width and Height) and Qty columns"}[kind]
        warnings.append(f"No table found. The file needs {need}.")
    if kind == "harness":
        warnings += validate_harness(out["wires"], out["bom"])
    elif kind == "panel":
        warnings += panel_warnings(out["lines"])
    else:
        warnings += label_warnings(out["lines"])
    out.update(warnings=warnings, source=source, drawing=drawing)
    return out


# ================================================================ harness: normalize, validate
def normalize_wires(wires: list[dict]) -> list[dict]:
    out = []
    for i, w in enumerate(wires or []):
        w = blank_wire(**{k: v for k, v in (w or {}).items() if k in blank_wire()})
        for k in ("wire_id", "from_ref", "from_pin", "to_ref", "to_pin", "gauge", "color", "spec", "from_term", "to_term"):
            w[k] = str(w[k] or "").strip().upper() if k not in ("spec",) else str(w[k] or "").strip()
        w["gauge"] = parse_gauge(w["gauge"]) if w["gauge"] else ""
        if w["spec"]:
            base, g = parse_spec(w["spec"])
            w["spec"] = base
            w["gauge"] = w["gauge"] or g
        w["length_in"] = _opt_num(w["length_in"])
        w["shielded"], w["twisted"] = bool(w["shielded"]), bool(w["twisted"])
        for k in ("from_term", "to_term"):
            w[k] = w[k].lower() if w[k].lower() in ("crimp", "solder", "lug", "splice", "open") else ""
        w["wire_id"] = w["wire_id"] or f"W{i + 1}"
        out.append(w)
    return out


def normalize_bom(bom: list[dict]) -> list[dict]:
    out = []
    for b in bom or []:
        b = blank_bom(**{k: v for k, v in (b or {}).items() if k in blank_bom()})
        b["ref"] = str(b["ref"] or "").strip().upper()
        if b["kind"] not in ("connector", "contact", "backshell", "accessory", "other"):
            b["kind"] = "other"
        b["qty"] = None if b["qty"] in (None, "") else _num(b["qty"])
        b["unit_price"] = _opt_num(b["unit_price"])
        b["contact_size"] = str(b["contact_size"] or "").strip().upper()
        out.append(b)
    return out


def refs_of(b: dict) -> list[str]:
    return [r for r in re.split(r"[,;\s]+", b.get("ref") or "") if r]


def auto_term(ref: str, connectors: set[str]) -> str:
    if ref in connectors:
        return "crimp"
    if re.match(r"^SP\d", ref):
        return "splice"
    if re.match(r"^(E|GND|TB|TS|GS)\d", ref):
        return "lug"
    return "crimp"


def _ends(wires: list[dict], connectors: set[str]) -> list[dict]:
    out = []
    for w in wires:
        for side in ("from", "to"):
            ref = w[f"{side}_ref"]
            out.append({"wire": w, "side": side, "ref": ref, "pin": w[f"{side}_pin"], "term": w[f"{side}_term"] or auto_term(ref, connectors)})
    return out


def validate_harness(wires: list[dict], bom: list[dict]) -> list[str]:
    """Warnings for references to missing connectors, pins used twice, contact size vs gauge, missing lengths."""
    wires, bom = normalize_wires(wires), normalize_bom(bom)
    w: list[str] = []
    conn_lines = [b for b in bom if b["kind"] == "connector"]
    connectors = {r for b in conn_lines for r in refs_of(b)}
    ends = _ends(wires, connectors)
    if connectors:
        missing: dict[str, list[str]] = defaultdict(list)
        for e in ends:
            if e["ref"] not in connectors and e["term"] in ("crimp", "solder"):
                missing[e["ref"]].append(e["wire"]["wire_id"])
        for ref, ids in sorted(missing.items(), key=lambda kv: _natural(kv[0])):
            w.append(f"{ref} is not in the connector list (used by {', '.join(sorted(set(ids), key=_natural)[:6])}{'...' if len(set(ids)) > 6 else ''}).")
        used = {e["ref"] for e in ends}
        unused = sorted(connectors - used, key=_natural)
        if unused:
            w.append(f"Connector(s) in the BOM with no wires: {', '.join(unused)}.")
    elif wires:
        w.append("No connector list: add the connectors (Ref, Part number, Qty) so references and contact sizes can be checked.")
    pins: dict[tuple[str, str], list[str]] = defaultdict(list)
    for e in ends:
        if e["pin"] and e["term"] in ("crimp", "solder"):
            pins[(e["ref"], e["pin"])].append(e["wire"]["wire_id"])
    for (ref, pin), ids in sorted(pins.items(), key=lambda kv: (_natural(kv[0][0]), _natural(kv[0][1]))):
        if len(ids) > 1:
            w.append(f"{ref} pin {pin} is used by {len(ids)} wires ({', '.join(ids)}). Check the wire list, or confirm the contact takes two wires.")
    sizes: dict[str, str] = {}
    for b in bom:
        if b["contact_size"] and b["kind"] in ("connector", "contact"):
            for r in refs_of(b):
                sizes.setdefault(r, b["contact_size"])
    for e in ends:
        size = sizes.get(e["ref"])
        g = e["wire"]["gauge"]
        if size and g and size in CONTACT_SIZE_AWG and e["term"] == "crimp":
            big, small = CONTACT_SIZE_AWG[size]
            if not (big <= int(g) <= small):
                w.append(f"{e['wire']['wire_id']} is {g} AWG but {e['ref']} uses size {size} contacts ({big}-{small} AWG).")
    ids = Counter(x["wire_id"] for x in wires)
    dup = [k for k, n in ids.items() if n > 1]
    if dup:
        w.append(f"Duplicate wire IDs: {', '.join(sorted(dup, key=_natural))}.")
    no_len = [x["wire_id"] for x in wires if not x["length_in"]]
    if no_len:
        w.append(f"{len(no_len)} wire(s) have no length ({', '.join(no_len[:5])}{'...' if len(no_len) > 5 else ''}); wire material is not priced for them.")
    no_g = [x["wire_id"] for x in wires if not x["gauge"]]
    if no_g:
        w.append(f"{len(no_g)} wire(s) have no gauge ({', '.join(no_g[:5])}{'...' if len(no_g) > 5 else ''}).")
    return w


# ================================================================ shared totals
def _totals(per_unit: list[dict], per_lot: list[dict], cfg: dict, qtys: list[int], opts: dict, *, lead_days, min_lot: float) -> dict:
    ga = _num(opts.get("ga_rate"), cfg["ga_rate"])
    profit = _num(opts.get("profit_rate"), cfg["profit_rate"])
    unit_cost = sum(l["cost"] for l in per_unit)
    lot_cost = sum(l["cost"] for l in per_lot)
    breaks = []
    for q in qtys:
        cost = unit_cost * q + lot_cost
        price = max(cost * (1 + ga) * (1 + profit), min_lot)
        breaks.append({"quantity": q, "total_cost": round(cost, 2), "unit_cost": round(cost / q, 2), "unit_price": round(price / q, 2),
                       "total_price": round(price, 2), "margin_pct": round((price - cost) / price * 100, 1) if price else 0.0,
                       "lead_time_days": lead_days(q)})
    return {"per_part_lines": per_unit, "per_lot_lines": per_lot, "per_part_cost": round(unit_cost, 2), "per_lot_cost": round(lot_cost, 2),
            "ga_rate": ga, "profit_rate": profit, "price_breaks": breaks}


def _lot_common(cfg: dict, opts: dict, per_unit: list[dict], per_lot: list[dict], fa_hours: float, rate: float, freight_default: float,
                warnings: list[str]) -> int:
    """First article, CoC, packaging, freight. Returns extra lead days."""
    extra = 0
    if opts.get("first_article"):
        per_lot.append(_line("inspection", "first article inspection and report", fa_hours * rate, basis="per_lot", hours=fa_hours, rate=rate))
        extra += cfg["lead_time"]["first_article_days"]
        warnings.append("First article required: the government must approve it before production ships. Allow for that in delivery days.")
    per_lot.append(_line("inspection", "certificate of conformance", cfg["inspection"]["cert_per_lot"], basis="per_lot"))
    level = opts.get("packaging_level") or "commercial"
    pk = cfg["packaging"].get(level)
    if pk is None:
        raise ElectricalError(f"packaging_level must be one of {list(cfg['packaging'])}")
    if pk["per_part"]:
        per_unit.append(_line("packaging", f"{level} packaging", pk["per_part"]))
    if pk["per_lot"]:
        per_lot.append(_line("packaging", f"{level} lot labels and marking", pk["per_lot"], basis="per_lot"))
    per_lot.append(_line("freight", "outbound freight", _num(opts.get("freight_per_lot"), freight_default), basis="per_lot"))
    return extra


# ================================================================ harness: pricing
def _wire_price(h: dict, spec: str, gauge: str) -> tuple[float, bool]:
    """$/ft and whether it came from the generic by-gauge table."""
    table = h["wire_per_ft"].get(spec) or {}
    if gauge in table:
        return float(table[gauge]), False
    return float(h["wire_default_per_ft"].get(gauge, 0.0)), True


def harness_counts(wires: list[dict], bom: list[dict], opts: dict | None = None, cfg_h: dict | None = None) -> dict:
    """Process counts that drive harness labor: ends, crimps, insertions, solder joints, shields, branches, ties."""
    opts = opts or {}
    wires, bom = normalize_wires(wires), normalize_bom(bom)
    connectors = {r for b in bom if b["kind"] == "connector" for r in refs_of(b)}
    ends = _ends(wires, connectors)
    t = Counter(e["term"] for e in ends)
    insert = sum(1 for e in ends if e["term"] == "crimp" and e["ref"] in connectors)
    refs_used = {e["ref"] for e in ends}
    branches = int(_num(opts.get("branches"), 0)) or max(len(refs_used), 1)
    tpb = (cfg_h or {}).get("ties_per_branch", 6.0)
    tie_bom = sum((b["qty"] or 0) for b in bom if b["kind"] == "accessory" and re.search(r"\bTIES?\b|TIE\s*WRAP|STRAP", f"{b['description']} {b['part_number']}", re.I))
    ties = int(_num(opts.get("ties"), 0)) or int(tie_bom) or math.ceil(branches * tpb)
    shields = sum(1 for w in wires if w["shielded"])
    heat = int(_num(opts.get("heat_shrink"), 0)) or (t["solder"] + t["splice"] + 2 * shields)
    backshells = sum((b["qty"] or 0) for b in bom if b["kind"] == "backshell")
    return {"wires": len(wires), "ends": len(ends), "crimp": t["crimp"], "solder": t["solder"], "lug": t["lug"], "splice": t["splice"],
            "open": t["open"], "insert": insert, "shield_ends": 2 * shields, "branches": branches, "ties": ties, "heat_shrink": heat,
            "backshells": backshells, "markers": 2 * len(wires) if opts.get("mark_wires", True) else 0,
            "branches_estimated": not opts.get("branches"), "ties_estimated": not opts.get("ties") and not tie_bom, "ties_in_bom": bool(tie_bom),
            "heat_shrink_estimated": not opts.get("heat_shrink"), "connectors": sorted(connectors, key=_natural)}


def price_harness(wires: list[dict], bom: list[dict], config: dict | None = None, quantities: list[int] | None = None,
                  options: dict | None = None) -> dict:
    cfg = pricing.merged_config(config)
    h = cfg["harness"]
    opts = dict(options or {})
    qtys = _qtys(quantities)
    wires, bom = normalize_wires(wires), normalize_bom(bom)
    if not wires and not bom:
        raise ElectricalError("Add at least one wire or BOM line.")
    warnings = validate_harness(wires, bom)
    assumptions: list[str] = []
    cls = opts.get("workmanship") or "class_3"
    if cls not in h["workmanship_multiplier"]:
        raise ElectricalError(f"workmanship must be one of {list(h['workmanship_multiplier'])}")
    mult = float(h["workmanship_multiplier"][cls])
    rate, trate = float(h["labor_rate"]), float(h["test_rate"])
    c = harness_counts(wires, bom, opts, h)
    per: list[dict] = []
    lot: list[dict] = []

    # ---- wire
    waste = float(h["waste_factor"])
    groups: dict[tuple[str, str], dict] = {}
    generic = set()
    for w in wires:
        if not w["length_in"]:
            continue
        ft = w["length_in"] / 12 * (1 + waste)
        ppf, gen = _wire_price(h, w["spec"], w["gauge"])
        if gen:
            generic.add(f"{w['spec'] or 'no spec'} {w['gauge'] or '?'} AWG")
        adder = (h["shield_adder_per_ft"] if w["shielded"] else 0) + (h["twist_adder_per_ft"] if w["twisted"] else 0)
        g = groups.setdefault((w["spec"] or "unspecified", w["gauge"] or "?"), {"ft": 0.0, "cost": 0.0, "ppf": ppf})
        g["ft"] += ft
        g["cost"] += ft * (ppf + adder)
    for (spec, gauge), g in sorted(groups.items()):
        per.append(_line("wire", f"{spec} {gauge} AWG, {g['ft']:.1f} ft", g["cost"], rate=g["ppf"], note=f"includes {waste:.0%} waste"))
    if generic:
        assumptions.append(f"No price for {', '.join(sorted(generic))} in the wire table: used the generic by-gauge placeholder.")

    # ---- BOM
    placeholders = 0
    has_contacts = any(b["kind"] == "contact" for b in bom)
    conn_set = set(c["connectors"])
    ends = _ends(wires, conn_set)
    for b in bom:
        qty = b["qty"]
        if b["kind"] == "contact" and qty is None:
            refs = set(refs_of(b))
            qty = sum(1 for e in ends if e["term"] == "crimp" and (not refs or e["ref"] in refs))
            note_q = "count from the wire list"
        else:
            qty = qty if qty is not None else 1
            note_q = ""
        price = b["unit_price"]
        if price is None:
            price = float(h["placeholder_prices"].get(b["kind"], h["placeholder_prices"]["other"]))
            placeholders += 1
            src = "PLACEHOLDER price"
        else:
            src = f"live {b['distributor']}".strip() if b["price_source"] == "live" else "your price"
        label = " ".join(x for x in (b["ref"], b["part_number"] or b["description"]) if x)
        per.append(_line(b["kind"] if b["kind"] in ("connector", "contact", "backshell") else "parts", f"{label} x {qty:g}", qty * price, rate=price,
                         note="; ".join(x for x in (src, note_q) if x)))
    if not has_contacts and c["crimp"]:
        p = float(h["placeholder_prices"]["contact"])
        per.append(_line("contact", f"crimp contacts (not in BOM) x {c['crimp']}", c["crimp"] * p, rate=p, note="PLACEHOLDER price, one per crimp end"))
        placeholders += 1
        assumptions.append("The BOM lists no contacts: priced one placeholder contact per crimp end.")

    # ---- consumables
    cons = h["consumables"]
    sleeving_ft, lacing_ft = _num(opts.get("sleeving_ft")), _num(opts.get("lacing_ft"))
    cons_cost = (c["markers"] * cons["wire_marker_each"] + c["heat_shrink"] * cons["heat_shrink_each"] + (0 if c["ties_in_bom"] else c["ties"]) * cons["cable_tie_each"]
                 + sleeving_ft * cons["sleeving_per_ft"] + lacing_ft * cons["lacing_per_ft"])
    per.append(_line("consumables", "markers, heat shrink, ties, sleeving, lacing", cons_cost,
                     note=f"{c['markers']} markers, {c['heat_shrink']} heat shrink, {c['ties']} ties, {sleeving_ft:g} ft sleeving, {lacing_ft:g} ft lacing"))

    # ---- labor (minutes x class multiplier)
    m = h["minutes"]
    steps = [
        ("cut and strip", c["ends"], m["cut_strip_per_end"], "per wire end"),
        ("crimp contacts", c["crimp"], m["crimp_per_contact"], "per crimp"),
        ("insert contacts", c["insert"], m["insert_per_contact"], "per contact in a connector"),
        ("solder joints", c["solder"], m["solder_per_joint"], "per joint"),
        ("lugs", c["lug"], m["lug_per_end"], "per lug"),
        ("splices", c["splice"], m["splice_per_end"], "per splice end"),
        ("backshell assembly", c["backshells"], m["backshell_per_connector"], "per backshell"),
        ("shield terminations", c["shield_ends"], m["shield_term_per_end"], "per shield end"),
        ("wire marking", c["wires"] if c["markers"] else 0, m["mark_per_wire"], "per wire, both ends"),
        ("heat shrink", c["heat_shrink"], m["heat_shrink_each"], "each"),
        ("sleeving", sleeving_ft, m["sleeving_per_ft"], "per ft"),
        ("layout and lacing", 1, m["layout_base"] + m["layout_per_branch"] * c["branches"], f"base + {c['branches']} branch(es)"),
        ("ties", c["ties"], m["tie_each"], "each"),
    ]
    labor_min = 0.0
    for name, n, each, basis in steps:
        if not n:
            continue
        mins = n * each * mult
        labor_min += mins
        per.append(_line("labor", f"{name}: {n:g} x {each:g} min" + (f" x {mult:g}" if mult != 1 else ""), mins / 60 * rate,
                         hours=mins / 60, rate=rate, note=basis))
    # ---- test
    t = h["test"]
    circuits = c["wires"]
    test_min = circuits * t["continuity_per_circuit_minutes"]
    hipot = opts.get("hipot", True)
    if hipot:
        test_min += circuits * t["hipot_per_circuit_minutes"] + t["hipot_per_harness_minutes"]
    per.append(_line("test", f"continuity{' and hipot / insulation resistance' if hipot else ''}, {circuits} circuit(s)", test_min / 60 * trate,
                     hours=test_min / 60, rate=trate))

    # ---- per lot
    fb = float(h["formboard_hours_per_lot"])
    lot.append(_line("setup", "formboard / layout fixture", fb * rate, basis="per_lot", hours=fb, rate=rate))
    kit = float(h["kitting_hours_per_lot"])
    lot.append(_line("setup", "kitting", kit * rate, basis="per_lot", hours=kit, rate=rate))
    ts = t["continuity_setup_minutes"] + (t["hipot_setup_minutes"] if hipot else 0)
    lot.append(_line("test", "test setup and program", ts / 60 * trate, basis="per_lot", hours=ts / 60, rate=trate))
    extra = _lot_common(cfg, opts, per, lot, float(h["first_article_hours"]), cfg["rates"]["inspection"], cfg["default_freight_per_lot"], warnings)

    if placeholders:
        warnings.append(f"{placeholders} BOM line(s) use placeholder prices. Enter prices or fetch live prices before quoting.")
    assumptions.insert(0, f"Workmanship {WORKMANSHIP_CLASSES.get(cls, cls)}. Labor x {mult:g} (your assumption).")
    if c["branches_estimated"]:
        assumptions.append(f"Branches estimated as one per connector or termination point ({c['branches']}); enter a count to override.")
    if c["ties_estimated"]:
        assumptions.append(f"Cable ties estimated at {h['ties_per_branch']:g} per branch ({c['ties']}).")
    if c["heat_shrink_estimated"] and c["heat_shrink"]:
        assumptions.append(f"Heat shrink estimated as one per solder joint, splice and shield end ({c['heat_shrink']}).")
    assumptions.append("Wire and component prices are placeholders until you enter yours or fetch live distributor prices.")

    def lead(q: int) -> int:
        return int(cfg["lead_time"]["base_days"] + h["supplier_lead_days"] + extra + math.ceil(q / max(h["harnesses_per_day"], 1)))

    out = _totals(per, lot, cfg, qtys, opts, lead_days=lead, min_lot=cfg["min_lot_charge"])
    out.update(kind="harness", part={k: opts.get(k) for k in ("name", "part_number", "nsn") if opts.get(k)}, material="Cable / wire harness",
               part_weight_lb=0.0, stock_volume_in3=0.0, workmanship=cls, workmanship_multiplier=mult, counts=c,
               labor_minutes=round(labor_min, 2), test_minutes=round(test_min, 2), assumptions=assumptions, warnings=warnings,
               config_note=h.get("note", ""), cut_list=cut_list(wires, h), pinout=pinout(wires, bom))
    return out


def cut_list(wires: list[dict], h: dict | None = None, builds: int = 1) -> list[dict]:
    rows = []
    for w in sorted(normalize_wires(wires), key=lambda x: _natural(x["wire_id"])):
        L = w["length_in"]
        rows.append({"wire_id": w["wire_id"], "spec": w["spec"], "gauge": w["gauge"], "color": w["color"],
                     "length_in": round(L, 2) if L else None, "length_ft": round(L / 12, 2) if L else None,
                     "shield_twist": " ".join(x for x in ("shielded" if w["shielded"] else "", "twisted" if w["twisted"] else "") if x),
                     "from": f"{w['from_ref']}-{w['from_pin']}" if w["from_pin"] else w["from_ref"],
                     "to": f"{w['to_ref']}-{w['to_pin']}" if w["to_pin"] else w["to_ref"],
                     "marker": w["wire_id"], "qty_per_harness": 1, "qty_total": builds, "signal": w["signal"]})
    return rows


def pinout(wires: list[dict], bom: list[dict]) -> list[dict]:
    wires, bom = normalize_wires(wires), normalize_bom(bom)
    connectors = {r for b in bom if b["kind"] == "connector" for r in refs_of(b)}
    rows = []
    for e in _ends(wires, connectors):
        w = e["wire"]
        other = "to" if e["side"] == "from" else "from"
        o_ref, o_pin = w[f"{other}_ref"], w[f"{other}_pin"]
        rows.append({"ref": e["ref"], "pin": e["pin"], "wire_id": w["wire_id"], "gauge": w["gauge"], "color": w["color"], "spec": w["spec"],
                     "termination": e["term"], "mates_to": f"{o_ref}-{o_pin}" if o_pin else o_ref, "signal": w["signal"]})
    rows.sort(key=lambda r: (_natural(r["ref"]), _natural(r["pin"]), _natural(r["wire_id"])))
    return rows


# ================================================================ panel: normalize, estimate wires, price
def normalize_panel(lines: list[dict]) -> list[dict]:
    out = []
    for i, ln in enumerate(lines or []):
        ln = blank_panel_line(**{k: v for k, v in (ln or {}).items() if k in blank_panel_line()})
        if ln["device_type"] not in PANEL_DEVICE_TYPES:
            ln["device_type"] = "other"
        ln["qty"] = _num(ln["qty"], 0)
        ln["length_m"] = _opt_num(ln["length_m"])
        ln["unit_price"] = _opt_num(ln["unit_price"])
        out.append(ln)
    return out


def panel_counts(lines: list[dict], opts: dict | None, p: dict) -> dict:
    """Device counts, rail and duct meters, wire count (given or estimated with the assumption spelled out), I/O points, cutouts."""
    opts = opts or {}
    lines = normalize_panel(lines)
    qty = Counter()
    meters = Counter()
    for ln in lines:
        if ln["device_type"] in LENGTH_PRICED:
            meters[ln["device_type"]] += ln["qty"] * ln["length_m"] if ln["length_m"] else ln["qty"]
        else:
            qty[ln["device_type"]] += ln["qty"]
    wr = p["wiring"]
    io_given = _opt_num(opts.get("io_points"))
    io_points = io_given if io_given is not None else qty["io_module"] * wr["io_points_per_module"]
    wire_given = _opt_num(opts.get("wire_count"))
    term = p["device_terminations"]
    if wire_given is not None:
        wires = int(wire_given)
        basis = "wire count given"
    elif io_given is not None:
        power = sum(term.get(t, 0) * n for t, n in qty.items() if t not in ("io_module", "terminal_block", "plc"))
        wires = math.ceil(io_given * wr["wires_per_io_point"] + power / 2)
        basis = (f"estimated: {io_given:g} I/O points x {wr['wires_per_io_point']:g} wires per point "
                 f"+ {power:g} power and control landings / 2")
    else:
        landings = sum(term.get(t, 0) * n for t, n in qty.items())
        wires = math.ceil(landings / 2)
        basis = f"estimated: {landings:g} device and terminal landings / 2 (each wire lands twice); enter a wire or I/O count to override"
    cut = dict(p["cutout_minutes"])
    cutouts = {"operator_device": qty["pilot_device"], "hmi": qty["hmi"], "meter": qty["meter"], "round_hole": 0, "rect_cutout": 0}
    for k, v in (opts.get("cutouts") or {}).items():
        if k in cut and v not in (None, ""):
            cutouts[k] = _num(v)
    tags = sum(n for t, n in qty.items() if t not in ("enclosure", "back_panel", "terminal_block"))
    return {"devices": dict(qty), "meters": {k: round(v, 3) for k, v in meters.items()}, "wires": wires, "wire_basis": basis,
            "io_points": io_points, "io_estimated": io_given is None, "cutouts": cutouts, "device_tags": tags,
            "legend_plates": qty["pilot_device"], "nameplates": int(_num(opts.get("nameplates"), 1))}


def panel_warnings(lines: list[dict]) -> list[str]:
    lines = normalize_panel(lines)
    w = []
    other = [ln for ln in lines if ln["device_type"] == "other"]
    if other:
        w.append(f"{len(other)} line(s) did not match a device type. Pick a type or enter a price.")
    if lines and not any(ln["device_type"] == "enclosure" for ln in lines):
        w.append("No enclosure in the BOM. Add one, or confirm the customer furnishes it.")
    nolen = [ln for ln in lines if ln["device_type"] in LENGTH_PRICED and not ln["length_m"]]
    if nolen:
        w.append(f"{len(nolen)} DIN rail / wire duct line(s) have no length: the quantity is read as meters.")
    return w


UL508A_NOTE = ("UL 508A: only a shop enrolled in UL's Industrial Control Panels program (UL 508A training done and a qualified "
               "manufacturer technical representative on staff) may apply the UL mark, and UL inspects enrolled shops. If you are not "
               "enrolled, do not quote a UL 508A labeled panel yourself: subcontract the build to a UL 508A shop or ask whether the label is required.")


def price_panel(lines: list[dict], config: dict | None = None, quantities: list[int] | None = None, options: dict | None = None) -> dict:
    cfg = pricing.merged_config(config)
    p = cfg["panel"]
    opts = dict(options or {})
    qtys = _qtys(quantities)
    lines = normalize_panel(lines)
    if not lines:
        raise ElectricalError("Add at least one BOM line.")
    warnings = panel_warnings(lines)
    assumptions: list[str] = []
    rate, trate, drate = float(p["labor_rate"]), float(p["test_rate"]), float(p["drafting_rate"])
    c = panel_counts(lines, opts, p)
    per: list[dict] = []
    lot: list[dict] = []
    placeholders = 0
    for ln in lines:
        t = ln["device_type"]
        price = ln["unit_price"]
        if price is None:
            price = float(p["placeholder_prices"].get(t, p["placeholder_prices"]["other"]))
            placeholders += 1
            src = "PLACEHOLDER price"
        else:
            src = f"live {ln['distributor']}".strip() if ln["price_source"] == "live" else "your price"
        label = " ".join(x for x in (ln["ref"], ln["part_number"] or ln["description"]) if x) or PANEL_DEVICE_TYPES[t]
        if t in LENGTH_PRICED:
            m = ln["qty"] * ln["length_m"] if ln["length_m"] else ln["qty"]
            per.append(_line("components", f"{label}, {m:g} m", m * price, rate=price, note=f"{src}, per meter"))
        else:
            per.append(_line("components", f"{label} x {ln['qty']:g}", ln["qty"] * price, rate=price, note=src))
    # ---- fabrication and mounting
    mins = []
    for t, n in sorted(c["devices"].items()):
        each = p["mount_minutes"].get(t, p["mount_minutes"]["other"])
        if n and each:
            mins.append((f"mount {PANEL_DEVICE_TYPES[t].lower()}: {n:g} x {each:g} min", n * each))
    for t, key in (("din_rail", "din_rail_minutes_per_m"), ("wire_duct", "wire_duct_minutes_per_m")):
        mtr = c["meters"].get(t, 0)
        if mtr:
            mins.append((f"cut and mount {PANEL_DEVICE_TYPES[t].split(' (')[0].lower()}: {mtr:g} m x {p[key]:g} min", mtr * p[key]))
    if c["devices"].get("back_panel") or c["devices"].get("enclosure"):
        mins.append(("back panel layout and drilling", p["back_panel_layout_minutes"]))
    for k, n in c["cutouts"].items():
        if n:
            mins.append((f"enclosure cutouts, {k.replace('_', ' ')}: {n:g} x {p['cutout_minutes'][k]:g} min", n * p["cutout_minutes"][k]))
    wr = p["wiring"]
    mins.append((f"wiring: {c['wires']} wires x {wr['minutes_per_wire']:g} min", c["wires"] * wr["minutes_per_wire"]))
    lb = p["labeling"]
    mins.append((f"wire markers: {2 * c['wires']} x {lb['marker_minutes']:g} min", 2 * c["wires"] * lb["marker_minutes"]))
    mins.append((f"device tags: {c['device_tags']:g} x {lb['device_tag_minutes']:g} min", c["device_tags"] * lb["device_tag_minutes"]))
    if c["nameplates"]:
        mins.append((f"nameplates: {c['nameplates']} x {lb['nameplate_minutes']:g} min", c["nameplates"] * lb["nameplate_minutes"]))
    for name, m in mins:
        if m:
            per.append(_line("labor", name, m / 60 * rate, hours=m / 60, rate=rate))
    per.append(_line("consumables", f"wire, ferrules, lugs for {c['wires']} wires", c["wires"] * wr["consumables_per_wire"], rate=wr["consumables_per_wire"]))
    lab_cost = 2 * c["wires"] * lb["wire_marker_each"] + c["device_tags"] * lb["device_tag_each"] + c["legend_plates"] * lb["legend_plate_each"] + c["nameplates"] * lb["nameplate_each"]
    per.append(_line("labels", "wire markers, device tags, legend plates, nameplates", lab_cost,
                     note=f"{2 * c['wires']} markers, {c['device_tags']:g} tags, {c['legend_plates']:g} legend plates, {c['nameplates']} nameplate(s)"))
    # ---- test
    ts = p["testing"]
    tm = c["wires"] * ts["point_to_point_minutes_per_wire"] + ts["power_up_minutes"] + c["io_points"] * ts["io_check_minutes_per_point"]
    per.append(_line("test", f"point-to-point ({c['wires']} wires), power-up, I/O check ({c['io_points']:g} points)", tm / 60 * trate, hours=tm / 60, rate=trate))
    # ---- UL 508A
    if opts.get("ul508a"):
        u = p["ul508a"]
        per.append(_line("ul508a", "UL 508A label", u["label_per_panel"], note="placeholder"))
        per.append(_line("ul508a", "UL 508A construction review and inspection", u["inspection_minutes_per_panel"] / 60 * trate,
                         hours=u["inspection_minutes_per_panel"] / 60, rate=trate))
        lot.append(_line("ul508a", "UL program cost share", u["program_cost_per_lot"], basis="per_lot", note="placeholder"))
        lot.append(_line("ul508a", "UL 508A design review (SCCR, wire sizing, spacing)", u["design_review_hours"] * drate, basis="per_lot",
                         hours=u["design_review_hours"], rate=drate))
        warnings.append(UL508A_NOTE)
    if opts.get("crate", True):
        per.append(_line("packaging", "crate", p["crate_per_panel"]))
    # ---- lot
    if opts.get("documentation", True):
        dh = float(p["documentation_hours"])
        lot.append(_line("engineering", "as-built drawings and documentation", dh * drate, basis="per_lot", hours=dh, rate=drate))
    kit = float(p["kitting_hours_per_lot"])
    lot.append(_line("setup", "kitting", kit * rate, basis="per_lot", hours=kit, rate=rate))
    extra = _lot_common(cfg, opts, per, lot, float(p["first_article_hours"]), cfg["rates"]["inspection"], p["freight_per_lot"], warnings)
    if placeholders:
        warnings.append(f"{placeholders} line(s) use placeholder prices. Enter prices or fetch live prices before quoting.")
    assumptions.append(f"Wire count {c['wires']}: {c['wire_basis']}.")
    if c["io_estimated"] and c["io_points"]:
        assumptions.append(f"I/O points estimated at {p['wiring']['io_points_per_module']:g} per I/O module ({c['io_points']:g}).")
    if c["cutouts"]["operator_device"] or c["cutouts"]["hmi"] or c["cutouts"]["meter"]:
        assumptions.append("Door cutouts assumed for every pilot device, HMI and meter; override the counts if some mount inside.")
    assumptions.append("Component prices are placeholders until you enter yours or fetch live distributor prices.")

    def lead(q: int) -> int:
        return int(cfg["lead_time"]["base_days"] + p["supplier_lead_days"] + extra + math.ceil(q / max(p["panels_per_day"], 1)))

    out = _totals(per, lot, cfg, qtys, opts, lead_days=lead, min_lot=cfg["min_lot_charge"])
    out.update(kind="panel", part={k: opts.get(k) for k in ("name", "part_number", "nsn") if opts.get(k)}, material="Control panel",
               part_weight_lb=0.0, stock_volume_in3=0.0, counts=c, assumptions=assumptions, warnings=warnings, config_note=p.get("note", ""))
    return out


# ================================================================ labels: normalize, price
def normalize_labels(lines: list[dict]) -> list[dict]:
    out = []
    for i, ln in enumerate(lines or []):
        ln = blank_label(**{k: v for k, v in (ln or {}).items() if k in blank_label()})
        ln["width_in"], ln["height_in"] = _opt_num(ln["width_in"]), _opt_num(ln["height_in"])
        ln["qty"] = _num(ln["qty"], 0)
        ln["holes"] = int(_num(ln["holes"], 0))
        ln["chars"] = None if ln["chars"] in (None, "") else int(_num(ln["chars"]))
        ln["adhesive"], ln["iuid"] = bool(ln["adhesive"]), bool(ln["iuid"])
        ln["item"] = str(ln["item"] or i + 1)
        out.append(ln)
    return out


def label_warnings(lines: list[dict]) -> list[str]:
    lines = normalize_labels(lines)
    w = []
    bad = [l["item"] for l in lines if not l["type"]]
    if bad:
        w.append(f"Label(s) {', '.join(bad[:8])} have no recognized type. Pick one.")
    blank = [l["item"] for l in lines if not l["text"].strip() and not l["chars"] and not l["iuid"]]
    if blank:
        w.append(f"Label(s) {', '.join(blank[:8])} have no text or character count; marking time is zero.")
    nosize = [l["item"] for l in lines if l["type"] != "heat_shrink_marker" and not (l["width_in"] and l["height_in"])]
    if nosize:
        w.append(f"Label(s) {', '.join(nosize[:8])} have no size; material is priced at the minimum.")
    if any(l["iuid"] for l in lines):
        w.append("IUID marks: DFARS 252.211-7003 calls for MIL-STD-130 marking, Data Matrix (ISO/IEC 16022) marks verified as machine "
                 "readable per MIL-STD-130 Appendix A (commonly grade B or better), and UII data reported to the IUID Registry "
                 "(through the WAWF receiving report for end items). Confirm you have a verifier.")
    return w


def price_labels(lines: list[dict], config: dict | None = None, quantities: list[int] | None = None, options: dict | None = None) -> dict:
    cfg = pricing.merged_config(config)
    L = cfg["labels"]
    opts = dict(options or {})
    qtys = _qtys(quantities)
    lines = normalize_labels(lines)
    if not lines:
        raise ElectricalError("Add at least one label or plate.")
    warnings = label_warnings(lines)
    laser, labor = float(L["laser_rate"]), float(L["labor_rate"])
    detail = []
    sums = Counter()
    for ln in lines:
        mat = L["materials"].get(ln["type"])
        if mat is None:
            mat = L["materials"]["engraved_laminate"]
        area = (ln["width_in"] or 0) * (ln["height_in"] or 0)
        chars = ln["chars"] if ln["chars"] is not None else count_chars(ln["text"])
        m_cost = max(area * mat["price_per_sqin"], mat["min_material"])
        mark_s = chars * mat["seconds_per_char"] + area * mat["seconds_per_sqin"] + (L["iuid"]["mark_seconds"] if ln["iuid"] else 0.0)
        fin_min = mat["handling_seconds"] / 60 + ln["holes"] * L["hole_minutes"] + (L["adhesive_minutes"] if ln["adhesive"] else 0)
        adh = area * L["adhesive_per_sqin"] if ln["adhesive"] else 0.0
        iu = L["iuid"]
        iuid_min = (iu["generate_minutes"] + iu["verify_minutes"] + iu["registry_minutes"]) if ln["iuid"] else 0.0
        each = m_cost + adh + mark_s / 3600 * laser + (fin_min + iuid_min) / 60 * labor
        detail.append({"item": ln["item"], "type": ln["type"], "area_sqin": round(area, 3), "chars": chars, "qty": ln["qty"],
                       "material_each": round(m_cost + adh, 3), "marking_seconds_each": round(mark_s, 1), "labor_minutes_each": round(fin_min + iuid_min, 2),
                       "cost_each": round(each, 3), "cost_per_set": round(each * ln["qty"], 2)})
        q = ln["qty"]
        sums["material"] += m_cost * q
        sums["adhesive"] += adh * q
        sums["mark_s"] += mark_s * q
        sums["fin_min"] += fin_min * q
        sums["iuid_min"] += iuid_min * q
        sums["iuid_marks"] += q if ln["iuid"] else 0
        sums["pieces"] += q
    per: list[dict] = [
        _line("material", f"plate and label stock, {sums['pieces']:g} piece(s)", sums["material"], note="by area with a per-piece minimum (placeholder prices)"),
    ]
    if sums["adhesive"]:
        per.append(_line("material", "adhesive backing", sums["adhesive"]))
    h = sums["mark_s"] / 3600
    if h:
        per.append(_line("marking", "laser engraving / printing time", h * laser, hours=h, rate=laser, note="seconds per character and per square inch of fill"))
    h = sums["fin_min"] / 60
    per.append(_line("labor", "handling, holes, adhesive application", h * labor, hours=h, rate=labor))
    if sums["iuid_min"]:
        h = sums["iuid_min"] / 60
        per.append(_line("iuid", f"IUID: generate, verify and report {sums['iuid_marks']:g} mark(s)", h * labor, hours=h, rate=labor))
    lot: list[dict] = []
    designs = len(lines)
    su = designs * L["setup_minutes_per_design"] / 60
    lot.append(_line("setup", f"artwork and setup, {designs} unique design(s)", su * labor, basis="per_lot", hours=su, rate=labor))
    if sums["iuid_marks"]:
        lot.append(_line("iuid", "verifier setup and verification report", L["iuid"]["verifier_setup_per_lot"], basis="per_lot"))
    opts.setdefault("packaging_level", "commercial")
    extra = _lot_common(cfg, opts, per, lot, 1.0, cfg["rates"]["inspection"], cfg["default_freight_per_lot"], warnings)

    def lead(q: int) -> int:
        return int(L["lead_days"] + extra + math.ceil(q * sums["pieces"] / max(L["pieces_per_day"], 1)))

    out = _totals(per, lot, cfg, qtys, opts, lead_days=lead, min_lot=float(L["min_lot_charge"]))
    out.update(kind="labels", part={k: opts.get(k) for k in ("name", "part_number", "nsn") if opts.get(k)}, material="Labels and nameplates",
               part_weight_lb=0.0, stock_volume_in3=0.0, line_detail=detail, pieces_per_set=sums["pieces"],
               assumptions=["Prices are per set (every line at its quantity). Material, laser and IUID times are placeholders until you enter yours."],
               warnings=warnings, config_note=L.get("note", ""))
    return out


# ================================================================ dispatch, saved specs, sheets
def price(kind: str, spec: dict, config: dict | None = None) -> dict:
    opts = dict(spec.get("options") or {})
    for k in ("name", "part_number", "nsn"):
        if spec.get(k):
            opts.setdefault(k, spec[k])
    q = spec.get("quantities") or [1]
    if kind == "harness":
        return price_harness(spec.get("wires") or [], spec.get("bom") or [], config, q, opts)
    if kind == "panel":
        return price_panel(spec.get("lines") or [], config, q, opts)
    if kind == "labels":
        return price_labels(spec.get("lines") or [], config, q, opts)
    raise ElectricalError(f"kind must be one of {KINDS}")


def estimate_spec(spec: dict, config: dict | None = None) -> dict:
    """Price a saved harness / panel / labels spec (used by quotes.save_quote)."""
    try:
        return price(spec.get("kind") or "", spec, config)
    except ElectricalError as exc:
        raise pricing.SpecError(str(exc)) from exc


def build_spec(kind: str, body: dict) -> dict:
    """The PartQuote spec for an electrical quote. Reopening the quote restores these lines and options."""
    names = {"harness": "Cable harness", "panel": "Control panel", "labels": "Labels and nameplates"}
    spec = {"kind": kind, "name": body.get("name") or names[kind], "part_number": body.get("part_number") or "", "nsn": body.get("nsn") or "",
            "quantities": body.get("quantities") or [1], "options": body.get("options") or {}, "source": body.get("source") or {}}
    if kind == "harness":
        spec["wires"] = normalize_wires(body.get("wires") or [])
        spec["bom"] = normalize_bom(body.get("bom") or [])
    elif kind == "panel":
        spec["lines"] = normalize_panel(body.get("lines") or [])
    else:
        spec["lines"] = normalize_labels(body.get("lines") or [])
    return spec


def _xlsx(sheets, title: str = "") -> bytes:
    from .extrusion import _xlsx as xl

    return xl(sheets, title)


def harness_xlsx(wires: list[dict], bom: list[dict], config: dict | None = None, builds: int = 1, title: str = "") -> bytes:
    """Workbook: labeled wire cut list, connector pin-out and BOM for `builds` harnesses."""
    cfg = pricing.merged_config(config)
    h = cfg["harness"]
    cl = cut_list(wires, h, builds)
    waste = h["waste_factor"]
    cut_rows = [[r["marker"], r["spec"], r["gauge"], r["color"], r["length_in"], r["length_ft"], r["shield_twist"], r["from"], r["to"],
                 r["qty_total"], round((r["length_in"] or 0) / 12 * builds * (1 + waste), 2), r["signal"]] for r in cl]
    pin_rows = [[r["ref"], r["pin"], r["wire_id"], r["gauge"], r["color"], r["spec"], r["termination"], r["mates_to"], r["signal"]] for r in pinout(wires, bom)]
    nb = normalize_bom(bom)
    ends = _ends(normalize_wires(wires), {r for b in nb if b["kind"] == "connector" for r in refs_of(b)})
    bom_rows = []
    for b in nb:
        q = b["qty"]
        if b["kind"] == "contact" and q is None:
            refs = set(refs_of(b))
            q = sum(1 for e in ends if e["term"] == "crimp" and (not refs or e["ref"] in refs))
        q = q if q is not None else 1
        price = b["unit_price"] if b["unit_price"] is not None else h["placeholder_prices"].get(b["kind"], h["placeholder_prices"]["other"])
        bom_rows.append([b["ref"], b["kind"], b["part_number"], b["manufacturer"], b["description"], q, q * builds, price,
                         "placeholder" if b["unit_price"] is None else (b["price_source"] or "manual"), round(q * builds * price, 2)])
    head = f"{title + ': ' if title else ''}{builds} harness(es). Lengths in inches. Wire to buy includes {waste:.0%} waste."
    return _xlsx([
        ("Cut list", ["Wire / marker", "Spec", "AWG", "Color", "Length (in)", "Length (ft)", "Shield / twist", "From", "To",
                      f"Pieces for {builds}", "Wire to buy (ft)", "Signal"], cut_rows),
        ("Pin-out", ["Connector", "Pin", "Wire", "AWG", "Color", "Spec", "Termination", "Mates to", "Signal"], pin_rows),
        ("BOM", ["Ref", "Kind", "Part number", "Manufacturer", "Description", "Qty per harness", f"Qty for {builds}", "Unit price", "Price source", "Extended"], bom_rows),
    ], head)


def panel_xlsx(lines: list[dict], config: dict | None = None, builds: int = 1, title: str = "") -> bytes:
    cfg = pricing.merged_config(config)
    p = cfg["panel"]
    rows = []
    for ln in normalize_panel(lines):
        t = ln["device_type"]
        price = ln["unit_price"] if ln["unit_price"] is not None else p["placeholder_prices"].get(t, 0)
        q = (ln["qty"] * ln["length_m"] if ln["length_m"] else ln["qty"]) if t in LENGTH_PRICED else ln["qty"]
        rows.append([ln["ref"], PANEL_DEVICE_TYPES[t], ln["part_number"], ln["manufacturer"], ln["description"], q,
                     "m" if t in LENGTH_PRICED else "each", q * builds, price, "placeholder" if ln["unit_price"] is None else (ln["price_source"] or "manual"),
                     round(q * builds * price, 2)])
    return _xlsx([("Panel BOM", ["Ref", "Type", "Part number", "Manufacturer", "Description", "Qty per panel", "Unit", f"Qty for {builds}",
                                 "Unit price", "Price source", "Extended"], rows)], f"{title + ': ' if title else ''}purchase list for {builds} panel(s)")


def labels_xlsx(lines: list[dict], config: dict | None = None, builds: int = 1, title: str = "") -> bytes:
    rows = [[l["item"], l["type"], l["width_in"], l["height_in"], l["text"], l["chars"] if l["chars"] is not None else count_chars(l["text"]),
             l["color"], l["holes"], "yes" if l["adhesive"] else "", "yes" if l["iuid"] else "", l["qty"], l["qty"] * builds, l["notes"]]
            for l in normalize_labels(lines)]
    return _xlsx([("Plate schedule", ["Item", "Type", "Width (in)", "Height (in)", "Text", "Characters", "Color", "Holes", "Adhesive", "IUID",
                                      "Qty per set", f"Qty for {builds}", "Notes"], rows)], f"{title + ': ' if title else ''}label and plate schedule for {builds} set(s)")


def sheet(kind: str, spec: dict, config: dict | None, builds: int, title: str) -> tuple[bytes, str]:
    if kind == "harness":
        return harness_xlsx(spec.get("wires") or [], spec.get("bom") or [], config, builds, title), "cut-list-pinout"
    if kind == "panel":
        return panel_xlsx(spec.get("lines") or [], config, builds, title), "panel-bom"
    if kind == "labels":
        return labels_xlsx(spec.get("lines") or [], config, builds, title), "plate-schedule"
    raise ElectricalError(f"kind must be one of {KINDS}")
