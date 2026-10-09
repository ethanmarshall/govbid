"""Post-award job tracking: jobs, travelers (work orders), vendor purchase orders, quality records,
Certificates of Conformance, shipping details and delivery metrics.

Certificate of Conformance content follows the elements of FAR 52.246-15 (Certificate of Conformance):
date furnished, contractor name, contract number, carrier and shipping document, a statement that the
supplies are of the quality specified and conform in all respects with the contract requirements
(specifications, drawings, preservation, packaging, packing, marking, physical item identification),
quantity, and signature, title and date. Under that clause a C of C is used in place of source inspection
only when the contract allows it and the Contract Administration Office authorizes it in writing.
Source: https://www.acquisition.gov/far/52.246-15 (checked 2026-10-08).
"""
from __future__ import annotations

import re
import shutil
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel, ConfigDict
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import config
from .db import get_db
from .models import Opportunity, PartQuote, PipelineEntry
from .models_crm import Organization
from .models_jobs import JOB_STATUSES, OPERATION_STATUSES, RECORD_TYPES, Job, JobOperation, JobPurchase, JobRecord

router = APIRouter(prefix="/api/jobs", tags=["jobs"])

OPEN_STATUSES = ["awarded", "in_work", "inspection"]  # work not yet delivered
DONE_STATUSES = ["shipped", "invoiced", "paid", "closed"]

# Traveler step names for quote operation types (pricing spec "operations"[].type)
OP_NAMES = {
    "cnc_mill": "CNC milling",
    "cnc_lathe": "CNC turning",
    "manual_machining": "Manual machining",
    "laser_cut": "Laser cutting",
    "waterjet": "Waterjet cutting",
    "press_brake": "Forming (press brake)",
    "weld": "Welding",
    "additive": "3D printing",
    "hardware_insert": "Install hardware inserts",
    "assembly": "Assembly",
    "deburr": "Deburr and edge break",
    "fabrication": "Fabrication",
}

STANDARD_TRAVELER = [
    ("contract_review", "Contract review (drawing rev, specs, quantity, delivery, packaging, clauses)", "In-house"),
    ("purchasing", "Order material and outside services", "In-house"),
    ("receiving", "Receiving inspection of material and certs", "In-house"),
    ("fabrication", "Fabrication / build", "In-house"),
    ("inspection", "Final inspection", "In-house"),
    ("coc", "Certificate of Conformance and records review", "In-house"),
    ("packaging", "Packaging and marking", "In-house"),
    ("ship", "Ship and submit receiving report", "In-house"),
]


# ------------------------------------------------------------------ schemas
class JobIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    title: str | None = None
    customer: str | None = None
    contract_number: str | None = None
    delivery_order: str | None = None
    cage_ship_to: str | None = None
    clins: list[dict[str, Any]] | None = None
    award_date: str | None = None
    due_date: str | None = None
    value: float | None = None
    status: str | None = None
    fob: str | None = None
    inspection_acceptance: str | None = None
    shipped_date: str | None = None
    notes: str | None = None
    carrier: str | None = None
    tracking_number: str | None = None
    packaging_level: str | None = None
    packaging_notes: str | None = None
    iuid_required: bool | None = None
    opportunity_id: int | None = None
    part_quote_id: int | None = None


class FromOpportunityIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    contract_number: str = ""
    award_date: str = ""
    due_date: str = ""
    part_quote_id: int | None = None
    build_traveler: bool = True


class OperationIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    name: str | None = None
    op_type: str | None = None
    work_center: str | None = None
    planned_date: str | None = None
    status: str | None = None
    signoff_initials: str | None = None
    signoff_date: str | None = None
    qty_good: int | None = None
    qty_rejected: int | None = None
    notes: str | None = None
    position: int | None = None


class TemplateIn(BaseModel):
    template: str = "quote"  # quote | standard | blank
    replace: bool = True


class ReorderIn(BaseModel):
    ids: list[int]


class PurchaseIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    organization_id: int | None = None
    vendor_name: str | None = None
    po_number: str | None = None
    description: str | None = None
    quantity: float | None = None
    unit_price: float | None = None
    ordered_date: str | None = None
    promised_date: str | None = None
    received_date: str | None = None
    qty_accepted: float | None = None
    qty_rejected: float | None = None
    certs_received: bool | None = None
    notes: str | None = None


# ------------------------------------------------------------------ helpers
def _check(value: str | None, allowed: list[str], field: str) -> None:
    if value is not None and value not in allowed:
        raise HTTPException(422, f"{field} must be one of: {', '.join(allowed)}")


def _check_date(value: str | None, field: str) -> None:
    if value:
        try:
            date.fromisoformat(value)
        except ValueError:
            raise HTTPException(422, f"{field} must be YYYY-MM-DD")


def _get_job(db: Session, job_id: int) -> Job:
    j = db.get(Job, job_id)
    if not j:
        raise HTTPException(404, "Job not found")
    return j


def _num(v) -> float | None:
    try:
        return None if v in (None, "") else float(v)
    except (TypeError, ValueError):
        return None


def clin_total(clins: list[dict]) -> float:
    """Sum of quantity x unit price over the CLINs (blank values count as zero)."""
    return round(sum((_num(c.get("quantity")) or 0) * (_num(c.get("unit_price")) or 0) for c in clins or []), 2)


def _clean_clins(clins: list[dict]) -> list[dict]:
    out = []
    for i, c in enumerate(clins or []):
        row = {
            "clin": str(c.get("clin") or f"{i + 1:04d}"),
            "description": str(c.get("description") or ""),
            "nsn": str(c.get("nsn") or ""),
            "part_number": str(c.get("part_number") or ""),
            "quantity": _num(c.get("quantity")),
            "unit": str(c.get("unit") or "EA"),
            "unit_price": _num(c.get("unit_price")),
            "due_date": str(c.get("due_date") or ""),
        }
        if row["quantity"] is not None and row["quantity"] == int(row["quantity"]):
            row["quantity"] = int(row["quantity"])
        _check_date(row["due_date"], f"CLIN {row['clin']} due_date")
        out.append(row)
    return out


def is_late(j: Job, today: date | None = None) -> bool:
    """Not yet shipped and past its due date."""
    today = today or date.today()
    return bool(j.due_date) and not j.shipped_date and j.status in OPEN_STATUSES and j.due_date < today.isoformat()


def job_dict(j: Job, db: Session | None = None, full: bool = False) -> dict:
    d = {
        "id": j.id, "title": j.title, "customer": j.customer, "contract_number": j.contract_number,
        "delivery_order": j.delivery_order, "cage_ship_to": j.cage_ship_to, "clins": j.clins or [],
        "award_date": j.award_date, "due_date": j.due_date, "value": j.value, "clin_total": clin_total(j.clins or []),
        "status": j.status, "fob": j.fob, "inspection_acceptance": j.inspection_acceptance, "shipped_date": j.shipped_date,
        "notes": j.notes, "carrier": j.carrier, "tracking_number": j.tracking_number, "packaging_level": j.packaging_level,
        "packaging_notes": j.packaging_notes, "iuid_required": bool(j.iuid_required),
        "opportunity_id": j.opportunity_id, "part_quote_id": j.part_quote_id,
        "late": is_late(j),
        "on_time": (j.shipped_date <= j.due_date) if (j.shipped_date and j.due_date) else None,
        "created_at": j.created_at.isoformat() if j.created_at else None,
        "updated_at": j.updated_at.isoformat() if j.updated_at else None,
    }
    if db is not None:
        ops = _ops(db, j.id)
        d["traveler_progress"] = {"done": sum(1 for o in ops if o.status == "done"), "total": len(ops)}
        if full:
            d["operations"] = [op_dict(o) for o in ops]
            d["purchases"] = [purchase_dict(p) for p in db.scalars(select(JobPurchase).where(JobPurchase.job_id == j.id).order_by(JobPurchase.id)).all()]
            d["records"] = [record_dict(r) for r in db.scalars(select(JobRecord).where(JobRecord.job_id == j.id).order_by(JobRecord.uploaded_at.desc(), JobRecord.id.desc())).all()]
            opp = db.get(Opportunity, j.opportunity_id) if j.opportunity_id else None
            d["opportunity"] = {"id": opp.id, "title": opp.title, "solicitation_number": opp.solicitation_number} if opp else None
    return d


def _ops(db: Session, job_id: int) -> list[JobOperation]:
    return list(db.scalars(select(JobOperation).where(JobOperation.job_id == job_id).order_by(JobOperation.position, JobOperation.id)).all())


def op_dict(o: JobOperation) -> dict:
    return {"id": o.id, "job_id": o.job_id, "position": o.position, "name": o.name, "op_type": o.op_type,
            "work_center": o.work_center, "planned_date": o.planned_date, "status": o.status,
            "signoff_initials": o.signoff_initials, "signoff_date": o.signoff_date,
            "qty_good": o.qty_good, "qty_rejected": o.qty_rejected, "notes": o.notes}


def purchase_dict(p: JobPurchase) -> dict:
    today = date.today().isoformat()
    return {
        "id": p.id, "job_id": p.job_id, "organization_id": p.organization_id, "vendor_name": p.vendor_name,
        "po_number": p.po_number, "description": p.description, "quantity": p.quantity, "unit_price": p.unit_price,
        "extended": round((p.quantity or 0) * (p.unit_price or 0), 2) if p.quantity and p.unit_price else None,
        "ordered_date": p.ordered_date, "promised_date": p.promised_date, "received_date": p.received_date,
        "qty_accepted": p.qty_accepted, "qty_rejected": p.qty_rejected, "certs_received": bool(p.certs_received),
        "notes": p.notes,
        "on_time": (p.received_date <= p.promised_date) if (p.received_date and p.promised_date) else None,
        "overdue": bool(p.promised_date and not p.received_date and p.promised_date < today),
    }


def record_dict(r: JobRecord) -> dict:
    return {"id": r.id, "job_id": r.job_id, "doc_type": r.doc_type, "clin": r.clin, "filename": r.filename,
            "notes": r.notes, "uploaded_at": r.uploaded_at.isoformat() if r.uploaded_at else None,
            "url": f"/api/jobs/{r.job_id}/records/{r.id}/file"}


def _apply_job(j: Job, data: dict) -> None:
    _check(data.get("status"), JOB_STATUSES, "status")
    for f in ("award_date", "due_date", "shipped_date"):
        _check_date(data.get(f), f)
    for f in ("fob", "inspection_acceptance"):
        if data.get(f):
            _check(data[f], ["origin", "destination"], f)
    if "clins" in data and data["clins"] is not None:
        data["clins"] = _clean_clins(data["clins"])
    for k, v in data.items():
        if k == "value" or v is not None:
            setattr(j, k, v)
    if "clins" in data and "value" not in data and j.clins:  # value follows the CLINs unless you set it
        j.value = clin_total(j.clins) or j.value
    if not j.due_date:  # earliest CLIN due date
        dues = sorted(c["due_date"] for c in (j.clins or []) if c.get("due_date"))
        if dues:
            j.due_date = dues[0]
    if data.get("shipped_date") and j.status in OPEN_STATUSES and "status" not in data:
        j.status = "shipped"


def _job_dir(job_id: int) -> Path:
    d = Path(config.UPLOAD_DIR) / "jobs" / str(job_id)
    d.mkdir(parents=True, exist_ok=True)
    return d


def _safe_name(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", Path(name or "file").name)[:120] or "file"


# ------------------------------------------------------------------ traveler templates
def traveler_from_spec(spec: dict | None) -> list[dict]:
    """Traveler steps for a part quote spec: contract review, material, each quote operation,
    outside finishes, inspection, C of C, packaging and shipping."""
    spec = spec or {}
    steps: list[dict] = [
        {"op_type": "contract_review", "name": "Contract review (drawing rev, specs, quantity, delivery, packaging, clauses)", "work_center": "In-house"},
    ]
    material = spec.get("material") or ""
    ops = spec.get("operations") or []
    additive = any(o.get("type") == "additive" for o in ops)
    if not additive:
        steps.append({"op_type": "purchasing", "name": f"Order material{': ' + material if material else ''}", "work_center": ""})
        steps.append({"op_type": "receiving", "name": "Receiving inspection: material, mill certs, quantity", "work_center": "In-house"})
    for o in ops:
        t = o.get("type") or ""
        name = OP_NAMES.get(t, t.replace("_", " ").capitalize() or "Operation")
        if t == "weld" and o.get("process"):
            name += f" ({str(o['process']).upper()})"
        if t == "additive" and o.get("technology"):
            name += f" ({str(o['technology']).upper()})"
        steps.append({"op_type": t, "name": name, "work_center": "In-house"})
    if ops and not additive and not any(o.get("type") == "deburr" for o in ops):
        steps.append({"op_type": "deburr", "name": "Deburr and edge break", "work_center": "In-house"})
    for f in spec.get("finishes") or []:
        ft = f.get("type") if isinstance(f, dict) else str(f)
        if ft:
            steps.append({"op_type": "finishing", "name": f"Finishing: {ft} (outside service)", "work_center": ""})
            steps.append({"op_type": "receiving", "name": f"Receive and inspect after {ft}, file finish cert", "work_center": "In-house"})
    if (spec.get("inspection") or {}).get("first_article"):
        steps.append({"op_type": "inspection", "name": "First article inspection and report", "work_center": "In-house"})
    steps.append({"op_type": "inspection", "name": "Final inspection (dimensions, finish, marking)", "work_center": "In-house"})
    steps.append({"op_type": "coc", "name": "Certificate of Conformance and records review", "work_center": "In-house"})
    level = (spec.get("packaging") or {}).get("level") or "commercial"
    steps.append({"op_type": "packaging", "name": f"Packaging and marking ({level})", "work_center": "In-house"})
    steps.append({"op_type": "ship", "name": "Ship and submit receiving report", "work_center": "In-house"})
    return steps


def _build_traveler(db: Session, j: Job, template: str, replace: bool = True) -> list[JobOperation]:
    if template not in ("quote", "standard", "blank"):
        raise HTTPException(422, "template must be quote, standard or blank")
    if replace:
        for o in _ops(db, j.id):
            db.delete(o)
        db.flush()
        start = 0
    else:
        existing = _ops(db, j.id)
        start = (max((o.position for o in existing), default=-1) + 1)
    if template == "quote":
        q = db.get(PartQuote, j.part_quote_id) if j.part_quote_id else None
        if not q:
            raise HTTPException(422, "This job has no linked part quote. Use the standard or blank template.")
        steps = traveler_from_spec(q.spec)
    elif template == "standard":
        steps = [{"op_type": t, "name": n, "work_center": w} for t, n, w in STANDARD_TRAVELER]
    else:
        steps = []
    rows = []
    for i, s in enumerate(steps):
        o = JobOperation(job_id=j.id, position=start + i, status="not_started", **s)
        db.add(o)
        rows.append(o)
    db.commit()
    return rows


# ------------------------------------------------------------------ metrics and hooks
def compute_metrics(jobs: list[Job], today: date | None = None) -> dict:
    """On-time delivery (shipped_date <= due_date over shipped jobs with both dates), jobs by status,
    open value (awarded, in work or in inspection) and late count."""
    today = today or date.today()
    by_status = {s: 0 for s in JOB_STATUSES}
    for j in jobs:
        by_status[j.status] = by_status.get(j.status, 0) + 1
    measured = [j for j in jobs if j.shipped_date and j.due_date and j.status != "cancelled"]
    on_time = sum(1 for j in measured if j.shipped_date <= j.due_date)
    open_jobs = [j for j in jobs if j.status in OPEN_STATUSES]
    return {
        "by_status": by_status,
        "total": len(jobs),
        "shipped_measured": len(measured),
        "on_time": on_time,
        "on_time_pct": round(100.0 * on_time / len(measured), 1) if measured else None,
        "open_count": len(open_jobs),
        "open_value": round(sum((j.value if j.value is not None else clin_total(j.clins or [])) for j in open_jobs), 2),
        "late_count": sum(1 for j in jobs if is_late(j, today)),
    }


def calendar_items(db: Session) -> list[dict]:
    """Job due dates (undelivered jobs) and vendor PO promised dates (not yet received)."""
    out = []
    for j in db.scalars(select(Job).where(Job.status.in_(OPEN_STATUSES), Job.due_date != "")).all():
        try:
            d = date.fromisoformat(j.due_date)
        except ValueError:
            continue
        out.append({"uid": f"job-due-{j.id}@govbid", "date": d, "summary": f"Delivery due: {j.title or 'Job ' + str(j.id)}",
                    "description": f"Contract {j.contract_number or '(not entered)'} for {j.customer or 'customer'}. Status: {j.status}.",
                    "url": f"/jobs?id={j.id}"})
    rows = db.execute(select(JobPurchase, Job.title).join(Job, Job.id == JobPurchase.job_id)
                      .where(JobPurchase.promised_date != "", JobPurchase.received_date == "")).all()
    for p, title in rows:
        try:
            d = date.fromisoformat(p.promised_date)
        except ValueError:
            continue
        out.append({"uid": f"job-po-{p.id}@govbid", "date": d, "summary": f"PO due from {p.vendor_name or 'vendor'}: {p.po_number or p.description[:40]}",
                    "description": f"For job: {title}. {p.description}".strip(), "url": f"/jobs?id={p.job_id}&tab=purchases"})
    return out


def dashboard_items(db: Session) -> list[str]:
    """Late jobs and overdue vendor POs."""
    today = date.today()
    out = []
    for j in db.scalars(select(Job).where(Job.status.in_(OPEN_STATUSES))).all():
        if is_late(j, today):
            days = (today - date.fromisoformat(j.due_date)).days
            out.append(f"Job late: {j.title or 'Job ' + str(j.id)} was due {j.due_date} ({days} day{'s' if days != 1 else ''} ago).")
    for p in db.scalars(select(JobPurchase).where(JobPurchase.promised_date != "", JobPurchase.received_date == "")).all():
        if p.promised_date < today.isoformat():
            out.append(f"PO overdue: {p.vendor_name or 'vendor'} {p.po_number or ''} promised {p.promised_date}, not received.".replace("  ", " "))
    return out


# ------------------------------------------------------------------ jobs
@router.get("/meta")
def meta():
    return {"statuses": JOB_STATUSES, "open_statuses": OPEN_STATUSES, "operation_statuses": OPERATION_STATUSES,
            "record_types": RECORD_TYPES, "templates": ["quote", "standard", "blank"]}


@router.get("/metrics")
def metrics(db: Session = Depends(get_db)):
    return compute_metrics(list(db.scalars(select(Job)).all()))


@router.get("")
def list_jobs(status: str = "", q: str = "", db: Session = Depends(get_db)):
    stmt = select(Job).order_by(Job.due_date == "", Job.due_date, Job.id.desc())
    if status:
        stmt = stmt.where(Job.status == status)
    if q:
        like = f"%{q}%"
        stmt = stmt.where(Job.title.ilike(like) | Job.contract_number.ilike(like) | Job.customer.ilike(like) | Job.notes.ilike(like))
    return [job_dict(j, db) for j in db.scalars(stmt).all()]


@router.post("")
def create_job(body: JobIn, db: Session = Depends(get_db)):
    data = body.model_dump(exclude_unset=True)
    j = Job(title=data.get("title") or "Untitled job", clins=[], status="awarded")
    _apply_job(j, data)
    db.add(j)
    db.commit()
    return job_dict(j, db, full=True)


@router.get("/purchases")
def list_purchases(organization_id: int | None = None, open_only: bool = False, db: Session = Depends(get_db)):
    stmt = select(JobPurchase).order_by(JobPurchase.ordered_date.desc(), JobPurchase.id.desc())
    if organization_id:
        stmt = stmt.where(JobPurchase.organization_id == organization_id)
    if open_only:
        stmt = stmt.where(JobPurchase.received_date == "")
    return [purchase_dict(p) for p in db.scalars(stmt).all()]


@router.post("/from-opportunity/{opp_id}")
def from_opportunity(opp_id: int, body: FromOpportunityIn | None = None, db: Session = Depends(get_db)):
    """Create a job from a won opportunity and its latest linked part quote; marks the pipeline entry and quote won."""
    body = body or FromOpportunityIn()
    opp = db.get(Opportunity, opp_id)
    if not opp:
        raise HTTPException(404, "Opportunity not found")
    existing = db.scalars(select(Job).where(Job.opportunity_id == opp_id)).first()
    if existing:
        raise HTTPException(409, f"Job {existing.id} already exists for this opportunity")
    for f in ("award_date", "due_date"):
        _check_date(getattr(body, f), f)
    if body.part_quote_id:
        q = db.get(PartQuote, body.part_quote_id)
        if not q:
            raise HTTPException(404, "Part quote not found")
    else:
        q = db.scalars(select(PartQuote).where(PartQuote.opportunity_id == opp_id)
                       .order_by(PartQuote.updated_at.desc(), PartQuote.id.desc())).first()
    award = body.award_date or date.today().isoformat()
    try:
        opp_qty = int(float(re.sub(r"[^\d.]", "", opp.quantity or "") or 0)) or None
    except ValueError:
        opp_qty = None
    qty = (q.quoted_quantity if q and q.quoted_quantity else None) or opp_qty
    price = q.quoted_unit_price if q else None
    due = body.due_date
    due_note = ""
    if not due and q:
        brk = next((b for b in (q.result or {}).get("price_breaks", []) if b.get("quantity") == qty), None)
        if brk and brk.get("lead_time_days"):
            due = (date.fromisoformat(award) + timedelta(days=int(brk["lead_time_days"]))).isoformat()
            due_note = f"Due date estimated from the quote lead time ({brk['lead_time_days']} days after award). Confirm it against the award's delivery schedule."
    clin = {"clin": "0001", "description": (q.name if q and q.name else opp.title)[:300],
            "nsn": (q.nsn if q and q.nsn else opp.nsn) or "", "part_number": (q.part_number if q else "") or "",
            "quantity": qty, "unit": "EA", "unit_price": price, "due_date": due}
    level = ((q.spec or {}).get("packaging") or {}).get("level", "") if q else ""
    j = Job(
        title=opp.title or (q.name if q else "") or "Awarded job", customer=opp.agency or "", contract_number=body.contract_number or "",
        clins=_clean_clins([clin]), award_date=award, due_date=due or "", status="awarded",
        opportunity_id=opp.id, part_quote_id=q.id if q else None, packaging_level=level,
        notes="\n".join(x for x in [f"Solicitation {opp.solicitation_number}" if opp.solicitation_number else "", due_note] if x),
    )
    j.value = clin_total(j.clins) or None
    db.add(j)
    # pipeline -> won
    p = db.scalars(select(PipelineEntry).where(PipelineEntry.opportunity_id == opp.id)).first()
    if not p:
        p = PipelineEntry(opportunity_id=opp.id)
        db.add(p)
    p.stage = "won"
    if price and qty and p.bid_amount is None:
        p.bid_amount = round(price * qty, 2)
    if q:
        q.status = "won"
    db.commit()
    if q and q.nsn and q.quoted_unit_price:
        try:  # keep the NSN price history current, the same way quotes.save_quote does
            from .nsn_history import record_quote_outcome
            record_quote_outcome(db, q)
        except Exception:  # noqa: BLE001
            db.rollback()
    if body.build_traveler:
        _build_traveler(db, j, "quote" if q else "standard")
    return job_dict(j, db, full=True)


@router.get("/{job_id}")
def get_job(job_id: int, db: Session = Depends(get_db)):
    return job_dict(_get_job(db, job_id), db, full=True)


@router.put("/{job_id}")
def update_job(job_id: int, body: JobIn, db: Session = Depends(get_db)):
    j = _get_job(db, job_id)
    _apply_job(j, body.model_dump(exclude_unset=True))
    db.commit()
    return job_dict(j, db, full=True)


@router.delete("/{job_id}")
def delete_job(job_id: int, db: Session = Depends(get_db)):
    j = _get_job(db, job_id)
    for model in (JobOperation, JobPurchase, JobRecord):
        db.query(model).filter(model.job_id == job_id).delete()
    try:
        from .models_quality import NonconformanceReport
        db.query(NonconformanceReport).filter(NonconformanceReport.job_id == job_id).update({"job_id": None, "purchase_id": None})
    except ImportError:
        pass
    db.delete(j)
    db.commit()
    shutil.rmtree(Path(config.UPLOAD_DIR) / "jobs" / str(job_id), ignore_errors=True)
    return {"ok": True}


# ------------------------------------------------------------------ traveler
@router.post("/{job_id}/traveler/template")
def apply_template(job_id: int, body: TemplateIn, db: Session = Depends(get_db)):
    j = _get_job(db, job_id)
    _build_traveler(db, j, body.template, body.replace)
    return [op_dict(o) for o in _ops(db, job_id)]


def _apply_op(o: JobOperation, data: dict) -> None:
    _check(data.get("status"), OPERATION_STATUSES, "status")
    for f in ("planned_date", "signoff_date"):
        _check_date(data.get(f), f)
    for f in ("qty_good", "qty_rejected"):
        if data.get(f) is not None and data[f] < 0:
            raise HTTPException(422, f"{f} cannot be negative")
    for k, v in data.items():
        if v is not None or k in ("qty_good", "qty_rejected"):
            setattr(o, k, v)
    if data.get("signoff_initials") and not o.signoff_date:
        o.signoff_date = date.today().isoformat()
    if data.get("signoff_initials") and "status" not in data:
        o.status = "done"


@router.post("/{job_id}/operations")
def add_operation(job_id: int, body: OperationIn, db: Session = Depends(get_db)):
    _get_job(db, job_id)
    data = body.model_dump(exclude_unset=True)
    pos = data.pop("position", None)
    if pos is None:
        pos = max((o.position for o in _ops(db, job_id)), default=-1) + 1
    o = JobOperation(job_id=job_id, position=pos, name=data.get("name") or "New step", status="not_started")
    _apply_op(o, data)
    db.add(o)
    db.commit()
    return op_dict(o)


@router.put("/{job_id}/operations/{op_id}")
def update_operation(job_id: int, op_id: int, body: OperationIn, db: Session = Depends(get_db)):
    o = db.get(JobOperation, op_id)
    if not o or o.job_id != job_id:
        raise HTTPException(404, "Operation not found")
    _apply_op(o, body.model_dump(exclude_unset=True))
    db.commit()
    return op_dict(o)


@router.delete("/{job_id}/operations/{op_id}")
def delete_operation(job_id: int, op_id: int, db: Session = Depends(get_db)):
    o = db.get(JobOperation, op_id)
    if not o or o.job_id != job_id:
        raise HTTPException(404, "Operation not found")
    db.delete(o)
    db.commit()
    return {"ok": True}


@router.post("/{job_id}/operations/reorder")
def reorder_operations(job_id: int, body: ReorderIn, db: Session = Depends(get_db)):
    ops = {o.id: o for o in _ops(db, job_id)}
    if set(body.ids) != set(ops):
        raise HTTPException(422, "ids must list every operation on this job exactly once")
    for i, oid in enumerate(body.ids):
        ops[oid].position = i
    db.commit()
    return [op_dict(o) for o in _ops(db, job_id)]


# ------------------------------------------------------------------ purchases
def _apply_purchase(db: Session, p: JobPurchase, data: dict) -> None:
    for f in ("ordered_date", "promised_date", "received_date"):
        _check_date(data.get(f), f)
    for f in ("quantity", "unit_price", "qty_accepted", "qty_rejected"):
        if data.get(f) is not None and data[f] < 0:
            raise HTTPException(422, f"{f} cannot be negative")
    if data.get("organization_id"):
        org = db.get(Organization, data["organization_id"])
        if not org:
            raise HTTPException(404, "Vendor organization not found")
        if not data.get("vendor_name"):
            data["vendor_name"] = org.name
    for k, v in data.items():
        if v is not None or k in ("organization_id", "quantity", "unit_price", "qty_accepted", "qty_rejected"):
            setattr(p, k, v)
    # a receipt with no accept/reject counts accepts the full quantity
    if p.received_date and p.qty_accepted is None and p.quantity is not None:
        p.qty_accepted = max(p.quantity - (p.qty_rejected or 0), 0)


@router.post("/{job_id}/purchases")
def add_purchase(job_id: int, body: PurchaseIn, db: Session = Depends(get_db)):
    _get_job(db, job_id)
    data = body.model_dump(exclude_unset=True)
    if not data.get("organization_id") and not data.get("vendor_name"):
        raise HTTPException(422, "Give a vendor (organization_id or vendor_name)")
    p = JobPurchase(job_id=job_id, ordered_date=date.today().isoformat())
    _apply_purchase(db, p, data)
    db.add(p)
    db.commit()
    return purchase_dict(p)


@router.put("/{job_id}/purchases/{pid}")
def update_purchase(job_id: int, pid: int, body: PurchaseIn, db: Session = Depends(get_db)):
    p = db.get(JobPurchase, pid)
    if not p or p.job_id != job_id:
        raise HTTPException(404, "Purchase order not found")
    _apply_purchase(db, p, body.model_dump(exclude_unset=True))
    db.commit()
    return purchase_dict(p)


@router.delete("/{job_id}/purchases/{pid}")
def delete_purchase(job_id: int, pid: int, db: Session = Depends(get_db)):
    p = db.get(JobPurchase, pid)
    if not p or p.job_id != job_id:
        raise HTTPException(404, "Purchase order not found")
    db.delete(p)
    db.commit()
    return {"ok": True}


# ------------------------------------------------------------------ records
@router.post("/{job_id}/records")
async def upload_record(job_id: int, file: UploadFile = File(...), doc_type: str = Form("other"), clin: str = Form(""),
                        notes: str = Form(""), db: Session = Depends(get_db)):
    _get_job(db, job_id)
    _check(doc_type, RECORD_TYPES, "doc_type")
    name = _safe_name(file.filename or "upload")
    stored = f"{datetime.utcnow():%Y%m%d%H%M%S%f}_{name}"
    (_job_dir(job_id) / stored).write_bytes(await file.read())
    r = JobRecord(job_id=job_id, doc_type=doc_type, clin=clin, filename=name, stored_name=stored, notes=notes)
    db.add(r)
    db.commit()
    return record_dict(r)


@router.get("/{job_id}/records/{rid}/file")
def record_file(job_id: int, rid: int, db: Session = Depends(get_db)):
    r = db.get(JobRecord, rid)
    if not r or r.job_id != job_id:
        raise HTTPException(404, "Record not found")
    path = _job_dir(job_id) / r.stored_name
    if not path.exists():
        raise HTTPException(404, "File is missing on disk")
    return FileResponse(path, filename=r.filename)


@router.delete("/{job_id}/records/{rid}")
def delete_record(job_id: int, rid: int, db: Session = Depends(get_db)):
    r = db.get(JobRecord, rid)
    if not r or r.job_id != job_id:
        raise HTTPException(404, "Record not found")
    (_job_dir(job_id) / r.stored_name).unlink(missing_ok=True)
    db.delete(r)
    db.commit()
    return {"ok": True}


# ------------------------------------------------------------------ certificate of conformance
def coc_fields(db: Session, j: Job, clin_id: str = "") -> dict:
    """Everything printed on a C of C for one CLIN (or the first CLIN when none is given)."""
    from .services import get_profile

    prof = get_profile(db)
    clins = j.clins or []
    clin = next((c for c in clins if c.get("clin") == clin_id), None) if clin_id else (clins[0] if clins else None)
    if clin_id and not clin:
        raise HTTPException(404, f"CLIN {clin_id} not found on this job")
    clin = clin or {}
    company = prof.name or "[Company name]"
    shipped = j.shipped_date or "____________"
    contract = j.contract_number or "____________"
    if j.delivery_order:
        contract += f", Order {j.delivery_order}"
    ship_doc = " ".join(x for x in [j.carrier, j.tracking_number] if x) or "____________"
    qty = clin.get("quantity")
    statement = (
        f"I certify that on {shipped}, {company} furnished the supplies called for by Contract No. {contract} "
        f"via carrier and shipping document {ship_doc}, and that the supplies are of the quality specified and conform "
        "in all respects with the contract requirements, including specifications, drawings, preservation, packaging, "
        "packing, marking requirements, and physical item identification (part number), and are in the quantity shown "
        "on this certificate or on the attached acceptance document."
    )
    return {
        "company": company, "cage": prof.cage or "", "uei": prof.uei or "", "contract": contract, "customer": j.customer,
        "clin": clin.get("clin", ""), "nsn": clin.get("nsn", ""), "part_number": clin.get("part_number", ""),
        "description": clin.get("description", "") or j.title, "quantity": "" if qty is None else f"{qty:g}" if isinstance(qty, float) else str(qty),
        "unit": clin.get("unit", "EA"), "shipped_date": j.shipped_date, "ship_doc": ship_doc, "statement": statement,
    }


def _coc_rows(f: dict) -> list[tuple[str, str]]:
    return [("Contractor", f["company"]), ("CAGE", f["cage"]), ("UEI", f["uei"]), ("Customer", f["customer"]),
            ("Contract / order", f["contract"]), ("CLIN", f["clin"]), ("NSN", f["nsn"]), ("Part number", f["part_number"]),
            ("Description", f["description"]), ("Quantity", f"{f['quantity']} {f['unit']}".strip()),
            ("Date shipped", f["shipped_date"]), ("Carrier / shipping document", f["ship_doc"])]


def build_coc_docx(f: dict, path: Path) -> Path:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    d.add_heading("Certificate of Conformance", level=1)
    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in _coc_rows(f):
        cells = t.add_row().cells
        cells[0].text, cells[1].text = k, v or ""
        cells[0].paragraphs[0].runs[0].bold = True
    d.add_paragraph()
    d.add_paragraph(f["statement"])
    d.add_paragraph()
    for line in ("Signature: ________________________________", "Name: ________________________________",
                 "Title: ________________________________", "Date: ________________"):
        d.add_paragraph(line)
    d.save(str(path))
    return path


def build_coc_pdf(f: dict, path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    ss = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=letter, leftMargin=0.9 * inch, rightMargin=0.9 * inch, topMargin=0.8 * inch)
    data = [[Paragraph(f"<b>{escape(k)}</b>", ss["Normal"]), Paragraph(escape(v or ""), ss["Normal"])] for k, v in _coc_rows(f)]
    table = Table(data, colWidths=[2.0 * inch, 4.6 * inch])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story = [Paragraph("Certificate of Conformance", ss["Title"]), table, Spacer(1, 16), Paragraph(escape(f["statement"]), ss["Normal"]),
             Spacer(1, 30)]
    for line in ("Signature: ________________________________", "Name: ________________________________",
                 "Title: ________________________________", "Date: ________________"):
        story += [Paragraph(line, ss["Normal"]), Spacer(1, 14)]
    doc.build(story)
    return path


@router.get("/{job_id}/coc")
def certificate_of_conformance(job_id: int, clin: str = "", fmt: str = Query("docx", pattern="^(docx|pdf)$"), save: bool = False,
                               db: Session = Depends(get_db)):
    """Download a Certificate of Conformance for a CLIN as DOCX or PDF. save=true also files it under the job's records."""
    j = _get_job(db, job_id)
    f = coc_fields(db, j, clin)
    stem = _safe_name(f"CofC_{j.contract_number or 'job' + str(j.id)}_{f['clin'] or 'all'}")
    out_dir = _job_dir(job_id) if save else Path(config.UPLOAD_DIR) / "jobs" / "_generated"
    out_dir.mkdir(parents=True, exist_ok=True)
    stored = f"{datetime.utcnow():%Y%m%d%H%M%S%f}_{stem}.{fmt}" if save else f"{stem}.{fmt}"
    path = out_dir / stored
    (build_coc_pdf if fmt == "pdf" else build_coc_docx)(f, path)
    if save:
        db.add(JobRecord(job_id=job_id, doc_type="coc", clin=f["clin"], filename=f"{stem}.{fmt}", stored_name=stored,
                         notes="Generated, unsigned"))
        db.commit()
    media = "application/pdf" if fmt == "pdf" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return FileResponse(path, filename=f"{stem}.{fmt}", media_type=media)
