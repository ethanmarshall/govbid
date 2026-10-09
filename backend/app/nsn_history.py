"""NSN award price history: what a part sold for before, as a reference for quoting.

DIBBS blocks automated access, so this app never scrapes it. History comes from:
  - DIBBS award search results you save as CSV/XLSX and import (DIBBS lists awardee CAGE,
    contract number, award date and total contract price, but no unit price or quantity),
  - rows you enter by hand, for example from the procurement history on a DIBBS RFQ,
  - your own quotes: a won quote is an award at your price, a lost quote records your losing price.
"""
from __future__ import annotations

import csv
import io
import re
import statistics
from datetime import date, datetime
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import CompanyProfile, PartQuote
from .models_nsn import AWARD_SOURCES, AwardRecord

FIELDS = ["nsn", "part_number", "nomenclature", "cage", "awardee", "contract_number", "award_date", "quantity",
          "unit_price", "total", "unit_of_issue", "source", "notes"]


# ---------------------------------------------------------------- NSN
def normalize_nsn(value: str) -> dict:
    """Parse an NSN or NIIN.

    Accepts 13 digits with or without dashes/spaces ("5340-01-480-5627", "5340014805627") or a
    9-digit NIIN ("01-480-5627", "014805627"). Returns {"nsn", "niin", "fsc"}: nsn is 4-2-3-4 for a
    full NSN, or 2-3-4 for a NIIN with the FSC unknown. Raises ValueError otherwise.
    """
    raw = (value or "").strip().upper()
    raw = re.sub(r"^NSN[:#\s]*", "", raw)
    if re.search(r"[\-\s]", raw) and not re.fullmatch(r"\d{4}[\s\-]*\d{2}[\s\-]*\d{3}[\s\-]*\d{4}|\d{2}[\s\-]*\d{3}[\s\-]*\d{4}", raw):
        raise ValueError(f"'{value}' is not an NSN. Use 13 digits (5340-01-480-5627) or a 9-digit NIIN.")
    digits = re.sub(r"[\s\-]", "", raw)
    if re.fullmatch(r"\d{13}", digits):
        fsc, niin = digits[:4], digits[4:]
        return {"nsn": f"{fsc}-{niin[:2]}-{niin[2:5]}-{niin[5:]}", "niin": niin, "fsc": fsc}
    if re.fullmatch(r"\d{9}", digits):
        return {"nsn": f"{digits[:2]}-{digits[2:5]}-{digits[5:]}", "niin": digits, "fsc": ""}
    raise ValueError(f"'{value}' is not an NSN. Use 13 digits (5340-01-480-5627) or a 9-digit NIIN.")


def try_nsn(value: str) -> dict | None:
    try:
        return normalize_nsn(value)
    except ValueError:
        return None


# ---------------------------------------------------------------- value parsing
def parse_date(v) -> str:
    """Many date spellings -> YYYY-MM-DD, or '' when unknown."""
    if v is None or v == "":
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        y, mo, d = map(int, m.groups())
        return _safe_date(y, mo, d)
    m = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})$", s)  # DIBBS uses MM-DD-YYYY
    if m:
        mo, d, y = map(int, m.groups())
        if y < 100:
            y += 2000 if y < 70 else 1900
        return _safe_date(y, mo, d)
    for fmt in ("%b %d, %Y", "%B %d, %Y", "%d %b %Y", "%d-%b-%Y", "%d-%b-%y", "%Y%m%d"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return ""


def _safe_date(y: int, mo: int, d: int) -> str:
    try:
        return date(y, mo, d).isoformat()
    except ValueError:
        return ""


def parse_money(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("$", "").strip()
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return float(m.group(0)) if m else None


def parse_int(v) -> int | None:
    f = parse_money(v)
    return int(round(f)) if f is not None else None


# ---------------------------------------------------------------- records
def to_dict(r: AwardRecord) -> dict:
    return {
        "id": r.id, "nsn": r.nsn, "niin": r.niin, "part_number": r.part_number, "nomenclature": r.nomenclature,
        "cage": r.cage, "awardee": r.awardee, "contract_number": r.contract_number, "award_date": r.award_date,
        "quantity": r.quantity, "unit_price": r.unit_price, "total": r.total, "unit_of_issue": r.unit_of_issue,
        "source": r.source, "quote_id": r.quote_id, "notes": r.notes,
        "created_at": r.created_at.isoformat() if r.created_at else None,
    }


def _find_duplicate(db: Session, niin: str, contract_number: str, award_date: str, unit_price: float | None) -> AwardRecord | None:
    """Same NIIN, contract number, award date and unit price counts as the same award."""
    rows = db.scalars(select(AwardRecord).where(AwardRecord.niin == niin, AwardRecord.contract_number == contract_number,
                                                AwardRecord.award_date == award_date)).all()
    for r in rows:
        if (r.unit_price is None and unit_price is None) or (r.unit_price is not None and unit_price is not None and abs(r.unit_price - unit_price) < 0.005):
            return r
    return None


def add_record(db: Session, data: dict, commit: bool = True) -> tuple[AwardRecord, bool]:
    """Add one award record. Returns (record, created); created is False for a duplicate."""
    n = normalize_nsn(str(data.get("nsn") or ""))
    source = data.get("source") or "manual"
    if source not in AWARD_SOURCES:
        raise ValueError(f"source must be one of {AWARD_SOURCES}")
    unit = parse_money(data.get("unit_price"))
    qty = parse_int(data.get("quantity"))
    total = parse_money(data.get("total"))
    if unit is None and total is not None and qty:
        unit = round(total / qty, 4)
    if total is None and unit is not None and qty:
        total = round(unit * qty, 2)
    if unit is not None and unit < 0 or (qty is not None and qty < 0):
        raise ValueError("Prices and quantities cannot be negative.")
    award_date = parse_date(data.get("award_date"))
    contract = str(data.get("contract_number") or "").strip().upper()
    dup = _find_duplicate(db, n["niin"], contract, award_date, unit)
    if dup:
        return dup, False
    r = AwardRecord(
        nsn=n["nsn"], niin=n["niin"], part_number=str(data.get("part_number") or "").strip()[:80],
        nomenclature=str(data.get("nomenclature") or "").strip()[:200],
        cage=str(data.get("cage") or "").strip().upper()[:10], awardee=str(data.get("awardee") or "").strip()[:300],
        contract_number=contract[:80], award_date=award_date, quantity=qty, unit_price=unit, total=total,
        unit_of_issue=str(data.get("unit_of_issue") or "").strip().upper()[:10], source=source,
        quote_id=data.get("quote_id"), notes=str(data.get("notes") or "").strip(),
    )
    db.add(r)
    if commit:
        db.commit()
        db.refresh(r)
    else:
        db.flush()
    return r, True


def delete_record(db: Session, record_id: int) -> bool:
    r = db.get(AwardRecord, record_id)
    if not r:
        return False
    db.delete(r)
    db.commit()
    return True


def _records(db: Session, niin: str) -> list[AwardRecord]:
    rows = db.scalars(select(AwardRecord).where(AwardRecord.niin == niin)).all()
    return sorted(rows, key=lambda r: (r.award_date or "", r.created_at or datetime.min), reverse=True)


def history(db: Session, nsn: str) -> dict:
    """All records for an NSN (matched on NIIN, so NIIN-only records count), newest first, with price stats.

    Stats use award prices only: a quote_lost record is your losing price, not an award, so it is
    reported separately as our_last_losing_price.
    """
    n = normalize_nsn(nsn)
    sync_quote_outcomes(db, n["niin"])
    rows = _records(db, n["niin"])
    awards = [r for r in rows if r.source != "quote_lost"]
    priced = [r for r in awards if r.unit_price is not None]
    prices = [r.unit_price for r in priced]
    last = awards[0] if awards else None
    last_priced = priced[0] if priced else None
    lost = [r for r in rows if r.source == "quote_lost" and r.unit_price is not None]
    stats = {
        "count": len(rows),
        "award_count": len(awards),
        "priced_count": len(priced),
        "last_unit_price": last_priced.unit_price if last_priced else None,
        "last_price_date": last_priced.award_date if last_priced else None,
        "last_award_date": last.award_date if last else None,
        "last_awardee": (last.awardee or last.cage) if last else None,
        "last_priced_awardee": (last_priced.awardee or last_priced.cage) if last_priced else None,
        "median_unit_price": round(statistics.median(prices), 4) if prices else None,
        "min": min(prices) if prices else None,
        "max": max(prices) if prices else None,
        "our_last_losing_price": lost[0].unit_price if lost else None,
    }
    return {"nsn": n["nsn"], "niin": n["niin"], "fsc": n["fsc"], "records": [to_dict(r) for r in rows], "stats": stats}


def reference_price(db: Session, nsn: str) -> float | None:
    """Most recent award unit price for the NSN (your losing quotes are not awards), or None."""
    try:
        return history(db, nsn)["stats"]["last_unit_price"]
    except ValueError:
        return None


# ---------------------------------------------------------------- our own quotes
def record_quote_outcome(db: Session, part_quote: PartQuote, commit: bool = True) -> AwardRecord | None:
    """Mirror a won or lost PartQuote into the award history.

    won  -> source quote_won at our quoted unit price (we are the awardee).
    lost -> source quote_lost at our price, noted "our losing price".
    Any other status removes a record this quote created earlier. Returns the record or None.
    """
    existing = db.scalars(select(AwardRecord).where(AwardRecord.quote_id == part_quote.id)).first()
    n = try_nsn(part_quote.nsn or "")
    if part_quote.status not in ("won", "lost") or not n or part_quote.quoted_unit_price is None:
        if existing:
            db.delete(existing)
            if commit:
                db.commit()
        return None
    won = part_quote.status == "won"
    profile = db.get(CompanyProfile, 1) or db.scalars(select(CompanyProfile)).first()
    when = (part_quote.updated_at or part_quote.created_at or datetime.utcnow()).date().isoformat()
    sol = part_quote.opportunity.solicitation_number if getattr(part_quote, "opportunity", None) else ""
    values = {
        "nsn": n["nsn"], "niin": n["niin"], "part_number": (part_quote.part_number or "")[:80],
        "nomenclature": (part_quote.name or "")[:200],
        "cage": (profile.cage if (won and profile) else "") or "",
        "awardee": (profile.name if (won and profile and profile.name) else "Us") if won else "",
        "contract_number": "", "quantity": part_quote.quoted_quantity, "unit_price": float(part_quote.quoted_unit_price),
        "total": round(float(part_quote.quoted_unit_price) * part_quote.quoted_quantity, 2) if part_quote.quoted_quantity else None,
        "source": "quote_won" if won else "quote_lost",
        "notes": (f"Our winning price, quote #{part_quote.id}" if won else f"Our losing price, quote #{part_quote.id}") + (f" ({sol})" if sol else ""),
    }
    if existing:
        values["award_date"] = existing.award_date if (existing.source == values["source"] and existing.award_date) else when
        for k, v in values.items():
            setattr(existing, k, v)
        rec = existing
    else:
        rec = AwardRecord(award_date=when, quote_id=part_quote.id, unit_of_issue="", **values)
        db.add(rec)
    if commit:
        db.commit()
    return rec


def sync_quote_outcomes(db: Session, niin: str | None = None) -> int:
    """Bring award history in line with every won/lost quote (optionally only one NIIN). Returns records touched."""
    touched = 0
    quote_ids_with_records = set(db.scalars(select(AwardRecord.quote_id).where(AwardRecord.quote_id.is_not(None))).all())
    for q in db.scalars(select(PartQuote)).all():
        n = try_nsn(q.nsn or "")
        if niin and (not n or n["niin"] != niin) and q.id not in quote_ids_with_records:
            continue
        if q.status in ("won", "lost") or q.id in quote_ids_with_records:
            record_quote_outcome(db, q, commit=False)
            touched += 1
    db.commit()
    return touched


# ---------------------------------------------------------------- import
def _norm_header(h) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9/#$ ]+", " ", str(h or "").lower())).strip()


def classify_header(h) -> str | None:
    """Map a spreadsheet column header to an AwardRecord field, or None to ignore it."""
    s = _norm_header(h)
    if not s:
        return None
    if any(w in s for w in ("posted", "last mod", "counter", "purchase request", "solicitation", "package")):
        return None
    if "nsn" in s or "national stock" in s or "niin" in s or "stock number" in s or "stock no" in s:
        return "nsn"
    if "cage" in s:
        return "cage"
    if s in ("p/n", "pn", "part #") or "part number" in s or "part no" in s:
        return "part_number"
    if any(w in s for w in ("awardee", "vendor", "contractor", "supplier", "company", "manufacturer")):
        return "awardee"
    if any(w in s for w in ("unit price", "unit cost", "price each", "unit $", "each price")) or s in ("price", "unit_price", "cost"):
        return "unit_price"
    if any(w in s for w in ("total", "extended", "amount", "contract price", "contact price", "award value", "obligated")):
        return "total"
    if "delivery order" in s or s in ("order number", "order #", "order no", "do number", "call number"):
        return "order_number"
    if any(w in s for w in ("contract", "piid", "award number", "award/basic", "basic number", "award #", "award no")):
        return "contract_number"
    if "award date" in s or "date awarded" in s or "awd date" in s or s in ("date", "awarded"):
        return "award_date"
    if "qty" in s or "quantity" in s:
        return "quantity"
    if s in ("u/i", "ui", "unit of issue", "unit", "uom", "u/m", "unit of measure"):
        return "unit_of_issue"
    if "nomenclature" in s or "description" in s or "item name" in s:
        return "nomenclature"
    if "note" in s or "remark" in s or "comment" in s:
        return "notes"
    return None


def _rows_from_bytes(data: bytes, filename: str) -> list[list]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        return [list(row) for row in ws.iter_rows(values_only=True)]
    text = data.decode("utf-8-sig", errors="replace")
    first = text.splitlines()[0] if text.strip() else ""
    dialect = csv.excel_tab if "\t" in first else csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def parse_award_table(data: bytes, filename: str) -> list[dict]:
    """Rows of a DIBBS award search sheet or a generic award sheet -> record dicts (unvalidated)."""
    rows = _rows_from_bytes(data, filename)
    header_idx, mapping = None, {}
    for i, row in enumerate(rows[:20]):
        m = {j: f for j, h in enumerate(row) if (f := classify_header(h))}
        fields = set(m.values())
        if len(fields) >= 2 and fields & {"nsn", "contract_number", "unit_price", "award_date"}:
            header_idx, mapping = i, m
            break
    if header_idx is None:
        return []
    out = []
    for row in rows[header_idx + 1:]:
        if not any(str(c or "").strip() for c in row):
            continue
        rec: dict = {}
        for j, f in mapping.items():
            if j < len(row) and row[j] not in (None, "") and f not in rec:
                rec[f] = row[j] if isinstance(row[j], (int, float, datetime, date)) else str(row[j]).strip()
        order = str(rec.pop("order_number", "") or "").strip()
        if order:
            rec["contract_number"] = f"{rec.get('contract_number', '')}/{order}".strip("/")
        # DIBBS puts part numbers in the NSN/Part Number column when there is no NSN
        if rec.get("nsn") and not try_nsn(str(rec["nsn"])):
            rec.setdefault("part_number", str(rec["nsn"]))
            rec["nsn"] = ""
        out.append(rec)
    return out


def import_awards(db: Session, data: bytes, filename: str, default_nsn: str = "", source: str = "dibbs_import") -> dict:
    """Import award rows from CSV/XLSX. Rows without an NSN use default_nsn, if given."""
    if source not in AWARD_SOURCES:
        raise ValueError(f"source must be one of {AWARD_SOURCES}")
    default = normalize_nsn(default_nsn)["nsn"] if default_nsn else ""
    rows = parse_award_table(data, filename)
    if not rows:
        raise ValueError("No award table found. The sheet needs a header row with columns such as NSN, Contract Number, Award Date, Unit Price.")
    imported = duplicates = skipped = 0
    errors: list[str] = []
    nsns: set[str] = set()
    for i, rec in enumerate(rows, 1):
        rec["nsn"] = rec.get("nsn") or default
        if not rec["nsn"]:
            skipped += 1
            continue
        rec["source"] = source
        try:
            r, created = add_record(db, rec, commit=False)
        except ValueError as exc:
            errors.append(f"Row {i}: {exc}")
            skipped += 1
            continue
        nsns.add(r.nsn)
        if created:
            imported += 1
        else:
            duplicates += 1
    db.commit()
    return {"rows": len(rows), "imported": imported, "duplicates": duplicates, "skipped": skipped, "nsns": sorted(nsns), "errors": errors[:20]}


# ---------------------------------------------------------------- lookup links (verified 2026-10-08)
def lookup_links(nsn: str) -> dict:
    """Pages for researching an NSN by hand. Each URL was checked to exist; DIBBS search links need a full 13-digit NSN."""
    n = normalize_nsn(nsn)
    links = []
    if n["fsc"]:
        digits = n["fsc"] + n["niin"]
        links += [
            {"label": "DIBBS award history for this NSN", "url": f"https://www.dibbs.bsm.dla.mil/Awards/AwdRecs.aspx?Category=nsn&TypeSrch=cq&Value={digits}",
             "note": "Lists awardee CAGE, contract number, award date and total contract price. Save the results to a spreadsheet and import them here."},
            {"label": "DIBBS open RFQs for this NSN", "url": f"https://www.dibbs.bsm.dla.mil/RFQ/RfqRecs.aspx?category=nsn&value={digits}&scope=all",
             "note": "The RFQ document usually carries a procurement history with past unit prices you can enter by hand."},
        ]
    links += [
        {"label": "SAM.gov contract opportunities search", "url": "https://sam.gov/search/?index=opp",
         "note": f"Search the keyword {n['nsn'] if n['fsc'] else n['niin']} (also try it without dashes)."},
        {"label": "PUB LOG (free FLIS data from DLA)", "url": "https://www.dla.mil/Information-Operations/Services/Applications/PUB-LOG/",
         "note": "Download for item names, part numbers and CAGE codes behind an NSN. WebFLIS itself needs a CAC."},
    ]
    return {
        "nsn": n["nsn"],
        "links": links,
        "notes": [
            "DIBBS shows a DoD consent banner first. Click OK, then open the link again: the banner drops the search terms.",
            "DIBBS blocks automated access, so this app does not pull award history itself. Import DIBBS results, enter awards by hand, or let your won and lost quotes fill it in.",
        ] + ([] if n["fsc"] else ["DIBBS searches need the full 13-digit NSN, not the NIIN alone."]),
    }


def import_file(db: Session, path: Path, default_nsn: str = "", source: str = "dibbs_import") -> dict:
    return import_awards(db, path.read_bytes(), path.name, default_nsn, source)
