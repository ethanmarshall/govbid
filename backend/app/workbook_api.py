"""Proposal pricing workbook for service, engineering and design-and-build bids.

Cost build-up (per CLIN, each base configurable on the rates):
  direct labor (DL)      = sum(hours x direct hourly rate)
  fringe                 = fringe % x DL
  overhead               = overhead % x (DL + fringe)        or x DL            (overhead_base)
  materials              = sum(material costs); material handling = MH % x materials
  subcontracts, ODCs     = sums
  total cost input (TCI) = DL + fringe + overhead + materials + handling + subcontracts + ODC
  G&A                    = G&A % x TCI                        (total_cost_input)
                           or x (TCI - materials - handling - subcontracts)    (value_added)
                           or x (DL + fringe + overhead)      (labor_overhead)
  total cost             = TCI + G&A
  profit                 = profit % x total cost              (total_cost)
                           or x (total cost - materials - handling - subcontracts)  (cost_less_pass_through)
  price                  = total cost + profit + fixed-price part lines (already priced, no further burden)

Limitations on subcontracting: FAR 52.219-14 (Oct 2022), read at
https://www.acquisition.gov/far/52.219-14 on 2026-10-08. Paragraph (e)(1), services (except
construction): the prime will not pay more than 50 percent of the amount paid by the Government for
contract performance to subcontractors that are not similarly situated entities. Paragraph (e)(2),
supplies (other than from a nonmanufacturer): not more than 50 percent of the amount paid by the
Government, excluding the cost of materials. Work a similarly situated entity further subcontracts
counts toward the prime's 50 percent. For mixed contracts each limit applies only to its portion.

GSA CALC+ (labor rate comparison): https://buy.gsa.gov/pricing/ (no API calls; compare by hand).
"""
from __future__ import annotations

import io
import re
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from .db import get_db
from .models import Opportunity, PartQuote
from .models_workbook import IndirectRates, LaborCategory, PriceBuild

router = APIRouter(prefix="/api/workbook")

CALC_URL = "https://buy.gsa.gov/pricing/"
LOS_URL = "https://www.acquisition.gov/far/52.219-14"
LOS_LIMIT = 50.0
HOURS_PER_YEAR = 2080

RATE_KEYS = ["fringe_pct", "overhead_pct", "ga_pct", "profit_pct", "material_handling_pct", "overhead_base", "ga_base", "fee_base"]
DEFAULT_RATES = {"fringe_pct": 30.0, "overhead_pct": 40.0, "ga_pct": 12.0, "profit_pct": 8.0, "material_handling_pct": 0.0,
                 "overhead_base": "labor_fringe", "ga_base": "total_cost_input", "fee_base": "total_cost"}
BASES = {
    "overhead_base": {"labor_fringe": "Direct labor + fringe", "labor": "Direct labor only"},
    "ga_base": {"total_cost_input": "Total cost input (all costs before G&A)",
                "value_added": "Value added (total cost input minus materials, handling and subcontracts)",
                "labor_overhead": "Direct labor + fringe + overhead"},
    "fee_base": {"total_cost": "Total cost", "cost_less_pass_through": "Total cost minus materials, handling and subcontracts"},
}

EXPLANATIONS = {
    "fringe_pct": "Fringe is what each direct labor dollar costs you on top of the paycheck: payroll taxes, health insurance, "
                  "retirement match, paid time off and workers' comp. Fringe % = annual fringe costs / annual direct labor dollars.",
    "overhead_pct": "Overhead is the cost of running the place where the work gets done that you cannot charge to one job: rent, "
                    "shop equipment, tools, software, utilities, and supervisors or engineers when they are not on a billable task. "
                    "It is usually spread over direct labor plus fringe.",
    "ga_pct": "General and administrative (G&A) is the cost of running the company as a whole: accounting, legal, insurance, "
              "business development and proposal time, your own time on admin, and office supplies. It is usually spread over "
              "total cost input, meaning every cost of the job before G&A.",
    "profit_pct": "Profit (called fee on cost-type contracts) is what you add on top of total cost. It pays for risk and growth "
                  "and is not a cost. Lower it to sharpen a price; when it goes below zero you are losing money on the bid.",
    "material_handling_pct": "Material handling is an optional rate on purchased materials that covers buying, receiving, "
                             "inspecting and stocking them. Leave it at 0 if those costs are already in overhead.",
    "fully_burdened": "The fully burdened rate is what one hour of a labor category really costs once fringe, overhead and G&A "
                      "are added. The billing rate adds profit. Compare billing rates with GSA CALC+ to see if you are in range.",
    "calculator": "New companies have no history, so estimate next year: the direct labor dollars you expect to bill to jobs, and "
                  "what you expect to spend on fringe, overhead and G&A. The calculator divides each pool by its base to get a "
                  "starting rate. Revisit it once you have real numbers.",
}


# ---------------------------------------------------------------- helpers
def _f(v, default=0.0) -> float:
    try:
        if v is None or v == "":
            return default
        return float(v)
    except (TypeError, ValueError):
        return default


def get_rates_row(db: Session) -> IndirectRates:
    r = db.get(IndirectRates, 1)
    if not r:
        r = IndirectRates(id=1, **DEFAULT_RATES, calc_inputs={})
        db.add(r)
        db.commit()
        db.refresh(r)
    return r


def rates_dict(r: IndirectRates) -> dict:
    return {k: getattr(r, k) for k in RATE_KEYS}


def merged_rates(rates: dict | None) -> dict:
    out = dict(DEFAULT_RATES)
    for k, v in (rates or {}).items():
        if k in out and v is not None and v != "":
            out[k] = v if k.endswith("_base") else _f(v)
    for k, opts in BASES.items():
        if out[k] not in opts:
            out[k] = DEFAULT_RATES[k]
    return out


def category_rate(c: LaborCategory) -> float:
    if c.annual_salary:
        return round(c.annual_salary / HOURS_PER_YEAR, 4)
    return float(c.hourly_rate or 0)


def cat_dict(c: LaborCategory) -> dict:
    return {"id": c.id, "name": c.name, "description": c.description, "hourly_rate": category_rate(c),
            "annual_salary": c.annual_salary, "calc_category": c.calc_category, "calc_note": c.calc_note}


def calculate_rates(direct_labor: float, fringe_costs: float, overhead_costs: float, ga_costs: float,
                    materials: float = 0.0, subcontracts: float = 0.0, other_direct: float = 0.0,
                    overhead_base: str = "labor_fringe", ga_base: str = "total_cost_input") -> dict:
    """Starting indirect rates from the owner's expected annual numbers. Percentages, rounded to 0.01."""
    dl = _f(direct_labor)
    if dl <= 0:
        raise ValueError("Enter the direct labor dollars you expect to charge to jobs in a year.")
    fr, oh, ga = _f(fringe_costs), _f(overhead_costs), _f(ga_costs)
    mat, sub, odc = _f(materials), _f(subcontracts), _f(other_direct)
    fringe_pct = fr / dl * 100
    oh_base_amt = dl + fr if overhead_base != "labor" else dl
    overhead_pct = oh / oh_base_amt * 100
    tci = dl + fr + oh + mat + sub + odc
    ga_base_amt = {"total_cost_input": tci, "value_added": tci - mat - sub, "labor_overhead": dl + fr + oh}.get(ga_base, tci)
    ga_pct = ga / ga_base_amt * 100 if ga_base_amt > 0 else 0.0
    return {
        "fringe_pct": round(fringe_pct, 2), "overhead_pct": round(overhead_pct, 2), "ga_pct": round(ga_pct, 2),
        "steps": [
            f"Fringe % = {fr:,.0f} / {dl:,.0f} direct labor = {fringe_pct:.2f}%",
            f"Overhead % = {oh:,.0f} / {oh_base_amt:,.0f} ({BASES['overhead_base'].get(overhead_base, '')}) = {overhead_pct:.2f}%",
            f"G&A % = {ga:,.0f} / {ga_base_amt:,.0f} ({BASES['ga_base'].get(ga_base, '')}) = {ga_pct:.2f}%",
        ],
        "total_cost_input": round(tci, 2),
    }


def burdened_rate(rate: float, R: dict) -> dict:
    """One labor hour through the build-up."""
    f, oh, ga, p = R["fringe_pct"] / 100, R["overhead_pct"] / 100, R["ga_pct"] / 100, R["profit_pct"] / 100
    fringe = rate * f
    overhead = (rate + fringe) * oh if R["overhead_base"] == "labor_fringe" else rate * oh
    ga_amt = (rate + fringe + overhead) * ga  # every G&A base includes labor and overhead
    cost = rate + fringe + overhead + ga_amt
    profit = cost * p  # every fee base includes labor
    return {"direct": round(rate, 2), "fringe": round(fringe, 2), "overhead": round(overhead, 2), "ga": round(ga_amt, 2),
            "fully_burdened_cost": round(cost, 2), "profit": round(profit, 2), "billing_rate": round(cost + profit, 2),
            "wrap_rate": round((cost / rate) if rate else 0, 4)}


def compute_clin(clin: dict, R: dict, cats: dict[int, LaborCategory]) -> dict:
    labor_lines = []
    dl = 0.0
    for row in clin.get("labor") or []:
        cid = row.get("category_id")
        c = cats.get(int(cid)) if cid not in (None, "") else None
        rate = category_rate(c) if c else _f(row.get("rate"))
        hours = _f(row.get("hours"))
        cost = hours * rate
        dl += cost
        labor_lines.append({"category_id": cid, "category": c.name if c else (row.get("category") or ""), "hours": hours,
                            "rate": rate, "cost": round(cost, 2)})
    materials = sum(_f(m.get("cost")) for m in clin.get("materials") or [])
    subs = clin.get("subcontracts") or []
    sub_total = sum(_f(s.get("cost")) for s in subs)
    non_ss = sum(_f(s.get("cost")) for s in subs if not s.get("similarly_situated"))
    non_ss += sum(_f(s.get("lower_tier_non_ss")) for s in subs if s.get("similarly_situated"))
    odc = sum(_f(o.get("cost")) for o in clin.get("odcs") or [])
    parts = sum(_f(p.get("quantity")) * _f(p.get("unit_price")) for p in clin.get("parts") or [])

    f, oh, ga, pr, mh = (R["fringe_pct"] / 100, R["overhead_pct"] / 100, R["ga_pct"] / 100, R["profit_pct"] / 100,
                         R["material_handling_pct"] / 100)
    fringe = dl * f
    overhead = (dl + fringe) * oh if R["overhead_base"] == "labor_fringe" else dl * oh
    handling = materials * mh
    tci = dl + fringe + overhead + materials + handling + sub_total + odc
    ga_base = {"total_cost_input": tci, "value_added": tci - materials - handling - sub_total,
               "labor_overhead": dl + fringe + overhead}[R["ga_base"]]
    ga_amt = ga_base * ga
    total_cost = tci + ga_amt
    fee_base = total_cost if R["fee_base"] == "total_cost" else total_cost - materials - handling - sub_total
    profit = fee_base * pr
    price = total_cost + profit + parts
    qty = _f(clin.get("quantity"), 1.0) or 1.0
    r2 = lambda x: round(x, 2)  # noqa: E731
    return {
        "clin": clin.get("clin", ""), "description": clin.get("description", ""), "kind": clin.get("kind") or "services",
        "quantity": qty, "unit": clin.get("unit") or "",
        "labor": labor_lines, "hours": round(sum(l["hours"] for l in labor_lines), 2),
        "direct_labor": r2(dl), "fringe": r2(fringe), "overhead": r2(overhead), "materials": r2(materials),
        "material_handling": r2(handling), "subcontracts": r2(sub_total), "odc": r2(odc), "total_cost_input": r2(tci),
        "ga_base": r2(ga_base), "ga": r2(ga_amt), "total_cost": r2(total_cost), "fee_base": r2(fee_base),
        "profit": r2(profit), "parts": r2(parts), "price": r2(price), "unit_price": r2(price / qty),
        "non_ss_subcontracts": r2(non_ss),
    }


def los_check(clins: list[dict], set_aside: bool) -> dict:
    """FAR 52.219-14 percentages per portion (services, supplies) from computed CLINs."""
    out = {"applies": bool(set_aside), "limit_pct": LOS_LIMIT, "source": LOS_URL, "portions": [], "warnings": [],
           "note": ("Services: share of the amount paid by the Government that goes to subcontractors that are not similarly "
                    "situated. Supplies: same, excluding the cost of materials. Work a similarly situated subcontractor passes "
                    "to a lower tier counts against you. A nonmanufacturer supplying another firm's product follows the "
                    "nonmanufacturer rule instead (paragraph (e)(2)).")}
    for kind in ("services", "supplies"):
        rows = [c for c in clins if c["kind"] == kind]
        if not rows:
            continue
        paid = sum(c["price"] for c in rows)
        mats = sum(c["materials"] for c in rows) if kind == "supplies" else 0.0
        base = paid - mats
        nonss = sum(c["non_ss_subcontracts"] for c in rows)
        pct = (nonss / base * 100) if base > 0 else 0.0
        over = pct > LOS_LIMIT
        out["portions"].append({"kind": kind, "amount_paid": round(paid, 2), "cost_of_materials": round(mats, 2),
                                "base": round(base, 2), "non_ss_subcontracts": round(nonss, 2), "pct": round(pct, 2),
                                "over": over, "paragraph": "(e)(1)" if kind == "services" else "(e)(2)"})
        if over and set_aside:
            out["warnings"].append(f"{kind.title()} portion: {pct:.1f}% goes to subcontractors that are not similarly situated, "
                                   f"over the {LOS_LIMIT:.0f}% limit in FAR 52.219-14{'(e)(1)' if kind == 'services' else '(e)(2)'}.")
    return out


def backsolve(total_cost: float, parts: float, fee_base: float, target_price: float) -> dict:
    """Profit % needed to land exactly on the target price."""
    profit = target_price - parts - total_cost
    pct = (profit / fee_base * 100) if fee_base else 0.0
    warn = ""
    if profit < 0:
        warn = f"At this price you lose ${-profit:,.2f}: profit is {pct:.2f}%."
    return {"target_price": round(target_price, 2), "profit": round(profit, 2), "profit_pct": round(pct, 2),
            "negative": profit < 0, "warning": warn}


def compute_build(build: dict, rates: dict | None, cats: dict[int, LaborCategory]) -> dict:
    R = merged_rates(rates)
    clins = [compute_clin(c, R, cats) for c in build.get("clins") or []]
    keys = ["hours", "direct_labor", "fringe", "overhead", "materials", "material_handling", "subcontracts", "odc",
            "total_cost_input", "ga", "total_cost", "fee_base", "profit", "parts", "price", "non_ss_subcontracts"]
    totals = {k: round(sum(c[k] for c in clins), 2) for k in keys}
    by_clin = {c["clin"]: c for c in clins}
    comps = []
    for comp in build.get("competitors") or []:
        prices = {k: _f(v, None) for k, v in (comp.get("prices") or {}).items() if _f(v, None) is not None}
        total = _f(comp.get("total"), None)
        ours_total = totals["price"]
        if total is None and prices:
            # Only some CLINs priced: compare like with like
            total = sum(prices.values())
            ours_total = round(sum(by_clin[k]["price"] for k in prices if k in by_clin), 2)
        gaps = {}
        for k, v in prices.items():
            if k in by_clin:
                ours = by_clin[k]["price"]
                gaps[k] = {"theirs": round(v, 2), "ours": ours, "gap": round(ours - v, 2),
                           "gap_pct": round((ours - v) / v * 100, 2) if v else None}
        tgap = None
        if total is not None:
            tgap = {"theirs": round(total, 2), "ours": ours_total, "gap": round(ours_total - total, 2),
                    "gap_pct": round((ours_total - total) / total * 100, 2) if total else None}
        comps.append({**comp, "total": total, "gaps": gaps, "total_gap": tgap})
    cats_used = {}
    for c in clins:
        for l in c["labor"]:
            key = l["category"] or f"rate {l['rate']}"
            if key not in cats_used:
                cats_used[key] = {"category": key, "category_id": l["category_id"], **burdened_rate(l["rate"], R)}
    target = build.get("target_price")
    bs = backsolve(totals["total_cost"], totals["parts"], totals["fee_base"], _f(target)) if target not in (None, "") else None
    return {
        "rates": R, "clins": clins, "totals": totals,
        "effective_profit_pct": round(totals["profit"] / totals["total_cost"] * 100, 2) if totals["total_cost"] else 0.0,
        "burdened_rates": list(cats_used.values()), "competitors": comps,
        "los": los_check(clins, bool(build.get("set_aside"))), "backsolve": bs,
    }


def _cats(db: Session) -> dict[int, LaborCategory]:
    return {c.id: c for c in db.scalars(select(LaborCategory)).all()}


def build_dict(b: PriceBuild) -> dict:
    return {"id": b.id, "name": b.name, "opportunity_id": b.opportunity_id, "rates": b.rates or {}, "set_aside": b.set_aside,
            "clins": b.clins or [], "competitors": b.competitors or [], "target_price": b.target_price, "status": b.status,
            "notes": b.notes, "created_at": b.created_at.isoformat() if b.created_at else None,
            "updated_at": b.updated_at.isoformat() if b.updated_at else None}


def _opp_label(db: Session, opp_id: int | None) -> dict | None:
    if not opp_id:
        return None
    o = db.get(Opportunity, opp_id)
    if not o:
        return None
    return {"id": o.id, "solicitation_number": o.solicitation_number, "title": o.title, "set_aside_code": o.set_aside_code,
            "set_aside_desc": o.set_aside_desc, "nsn": o.nsn, "quantity": o.quantity}


def full_build(db: Session, b: PriceBuild) -> dict:
    d = build_dict(b)
    d["computed"] = compute_build(d, d["rates"], _cats(db))
    d["opportunity"] = _opp_label(db, b.opportunity_id)
    return d


def build_summary(db: Session, b: PriceBuild) -> dict:
    """Short summary for the opportunity page."""
    c = compute_build(build_dict(b), b.rates, _cats(db))
    return {"id": b.id, "name": b.name, "status": b.status, "total_price": c["totals"]["price"],
            "total_cost": c["totals"]["total_cost"], "profit": c["totals"]["profit"], "profit_pct": c["rates"]["profit_pct"],
            "clins": len(c["clins"]), "los_warnings": c["los"]["warnings"],
            "updated_at": b.updated_at.isoformat() if b.updated_at else None}


# ---------------------------------------------------------------- meta, rates, labor
@router.get("/meta")
def meta():
    return {"explanations": EXPLANATIONS, "bases": BASES, "defaults": DEFAULT_RATES, "calc_url": CALC_URL,
            "los_url": LOS_URL, "los_limit_pct": LOS_LIMIT, "hours_per_year": HOURS_PER_YEAR}


@router.get("/rates")
def get_rates(db: Session = Depends(get_db)):
    r = get_rates_row(db)
    return {**rates_dict(r), "calc_inputs": r.calc_inputs or {}, "notes": r.notes}


class RatesIn(BaseModel):
    fringe_pct: float | None = None
    overhead_pct: float | None = None
    ga_pct: float | None = None
    profit_pct: float | None = None
    material_handling_pct: float | None = None
    overhead_base: str | None = None
    ga_base: str | None = None
    fee_base: str | None = None
    calc_inputs: dict | None = None
    notes: str | None = None


@router.put("/rates")
def put_rates(body: RatesIn, db: Session = Depends(get_db)):
    r = get_rates_row(db)
    for k, v in body.model_dump(exclude_none=True).items():
        if k.endswith("_base") and v not in BASES[k]:
            raise HTTPException(400, f"{k} must be one of {', '.join(BASES[k])}")
        setattr(r, k, v)
    db.commit()
    return get_rates(db)


class CalcIn(BaseModel):
    direct_labor: float = 0
    fringe_costs: float = 0
    overhead_costs: float = 0
    ga_costs: float = 0
    materials: float = 0
    subcontracts: float = 0
    other_direct: float = 0
    overhead_base: str = "labor_fringe"
    ga_base: str = "total_cost_input"


@router.post("/rates/calculate")
def rates_calculate(body: CalcIn):
    try:
        return calculate_rates(**body.model_dump())
    except ValueError as exc:
        raise HTTPException(400, str(exc))


@router.get("/labor")
def list_labor(db: Session = Depends(get_db)):
    R = merged_rates(rates_dict(get_rates_row(db)))
    out = []
    for c in db.scalars(select(LaborCategory).order_by(LaborCategory.name)).all():
        d = cat_dict(c)
        d["burdened"] = burdened_rate(d["hourly_rate"], R)
        out.append(d)
    return {"categories": out, "calc_url": CALC_URL}


class LaborIn(BaseModel):
    name: str = Field(min_length=1)
    description: str = ""
    hourly_rate: float = 0
    annual_salary: float | None = None
    calc_category: str = ""
    calc_note: str = ""


@router.post("/labor")
def create_labor(body: LaborIn, db: Session = Depends(get_db)):
    c = LaborCategory(**body.model_dump())
    if c.annual_salary:
        c.hourly_rate = round(c.annual_salary / HOURS_PER_YEAR, 4)
    db.add(c)
    db.commit()
    return cat_dict(c)


@router.put("/labor/{cat_id}")
def update_labor(cat_id: int, body: LaborIn, db: Session = Depends(get_db)):
    c = db.get(LaborCategory, cat_id)
    if not c:
        raise HTTPException(404, "Labor category not found")
    for k, v in body.model_dump().items():
        setattr(c, k, v)
    if c.annual_salary:
        c.hourly_rate = round(c.annual_salary / HOURS_PER_YEAR, 4)
    db.commit()
    return cat_dict(c)


@router.delete("/labor/{cat_id}")
def delete_labor(cat_id: int, db: Session = Depends(get_db)):
    c = db.get(LaborCategory, cat_id)
    if c:
        db.delete(c)
        db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- pickers
@router.get("/part-quotes")
def part_quote_options(db: Session = Depends(get_db)):
    rows = db.scalars(select(PartQuote).order_by(PartQuote.id.desc()).limit(200)).all()
    return {"quotes": [{"id": q.id, "name": q.name, "nsn": q.nsn, "part_number": q.part_number, "status": q.status,
                        "quoted_unit_price": q.quoted_unit_price, "quoted_quantity": q.quoted_quantity} for q in rows]}


@router.get("/opportunities")
def opportunity_options(q: str = "", db: Session = Depends(get_db)):
    stmt = select(Opportunity).order_by(Opportunity.id.desc()).limit(50)
    if q:
        like = f"%{q}%"
        stmt = stmt.where(or_(Opportunity.title.ilike(like), Opportunity.solicitation_number.ilike(like), Opportunity.nsn.ilike(like)))
    return {"opportunities": [_opp_label(db, o.id) for o in db.scalars(stmt).all()]}


# ---------------------------------------------------------------- builds
class BuildIn(BaseModel):
    name: str = ""
    opportunity_id: int | None = None
    rates: dict | None = None
    set_aside: bool | None = None
    clins: list[dict] | None = None
    competitors: list[dict] | None = None
    target_price: float | None = None
    status: str | None = None
    notes: str | None = None


@router.get("/builds")
def list_builds(db: Session = Depends(get_db)):
    rows = db.scalars(select(PriceBuild).order_by(PriceBuild.updated_at.desc())).all()
    out = []
    for b in rows:
        s = build_summary(db, b)
        s["opportunity"] = _opp_label(db, b.opportunity_id)
        out.append(s)
    return {"builds": out}


@router.post("/builds")
def create_build(body: BuildIn, db: Session = Depends(get_db)):
    data = body.model_dump(exclude_unset=True)
    b = PriceBuild(name=data.get("name") or "New price build", opportunity_id=data.get("opportunity_id"),
                   rates=merged_rates(data.get("rates") or rates_dict(get_rates_row(db))),
                   clins=data.get("clins") or [], competitors=data.get("competitors") or [],
                   target_price=data.get("target_price"), status=data.get("status") or "draft", notes=data.get("notes") or "")
    opp = db.get(Opportunity, b.opportunity_id) if b.opportunity_id else None
    if b.opportunity_id and not opp:
        raise HTTPException(404, "Opportunity not found")
    b.set_aside = data["set_aside"] if data.get("set_aside") is not None else bool(opp and opp.set_aside_code)
    if opp and not data.get("name"):
        b.name = f"{opp.solicitation_number or 'Opportunity'} price build"
    db.add(b)
    db.commit()
    return full_build(db, b)


def _get_build(db: Session, build_id: int) -> PriceBuild:
    b = db.get(PriceBuild, build_id)
    if not b:
        raise HTTPException(404, "Price build not found")
    return b


@router.get("/builds/{build_id}")
def get_build(build_id: int, db: Session = Depends(get_db)):
    return full_build(db, _get_build(db, build_id))


@router.put("/builds/{build_id}")
def update_build(build_id: int, body: BuildIn, db: Session = Depends(get_db)):
    b = _get_build(db, build_id)
    data = body.model_dump(exclude_unset=True)
    if "rates" in data and data["rates"] is not None:
        data["rates"] = merged_rates(data["rates"])
    if data.get("opportunity_id") and not db.get(Opportunity, data["opportunity_id"]):
        raise HTTPException(404, "Opportunity not found")
    for k, v in data.items():
        if k in ("clins", "competitors") and v is None:
            v = []
        setattr(b, k, v)
    b.updated_at = datetime.utcnow()
    db.commit()
    return full_build(db, b)


@router.delete("/builds/{build_id}")
def delete_build(build_id: int, db: Session = Depends(get_db)):
    b = db.get(PriceBuild, build_id)
    if b:
        db.delete(b)
        db.commit()
    return {"ok": True}


@router.post("/builds/{build_id}/reload-rates")
def reload_rates(build_id: int, db: Session = Depends(get_db)):
    b = _get_build(db, build_id)
    b.rates = merged_rates(rates_dict(get_rates_row(db)))
    db.commit()
    return full_build(db, b)


class ComputeIn(BaseModel):
    rates: dict | None = None
    set_aside: bool = False
    clins: list[dict] = Field(default_factory=list)
    competitors: list[dict] = Field(default_factory=list)
    target_price: float | None = None


@router.post("/compute")
def compute(body: ComputeIn, db: Session = Depends(get_db)):
    """Compute an unsaved build (live editing)."""
    rates = body.rates if body.rates else rates_dict(get_rates_row(db))
    return compute_build(body.model_dump(), rates, _cats(db))


class BacksolveIn(BaseModel):
    target_price: float


@router.post("/builds/{build_id}/backsolve")
def backsolve_build(build_id: int, body: BacksolveIn, db: Session = Depends(get_db)):
    b = _get_build(db, build_id)
    c = compute_build(build_dict(b), b.rates, _cats(db))
    t = c["totals"]
    return backsolve(t["total_cost"], t["parts"], t["fee_base"], body.target_price)


@router.post("/builds/{build_id}/reference-prices")
def pull_reference_prices(build_id: int, db: Session = Depends(get_db)):
    """Add or refresh a competitor row with the last award unit price (x quantity) for each CLIN that has an NSN."""
    from .nsn_history import reference_price

    b = _get_build(db, build_id)
    prices, notes = {}, []
    for c in b.clins or []:
        nsn = (c.get("nsn") or "").strip()
        if not nsn:
            continue
        up = reference_price(db, nsn)
        if up is None:
            notes.append(f"{c.get('clin')}: no award history for {nsn}")
            continue
        qty = _f(c.get("quantity"), 1.0) or 1.0
        prices[c.get("clin", "")] = round(up * qty, 2)
        notes.append(f"{c.get('clin')}: last award ${up:,.2f} each x {qty:g}")
    comps = [x for x in (b.competitors or []) if x.get("source") != "nsn_history"]
    if prices:
        comps.append({"name": "Last award (NSN history)", "source": "nsn_history", "prices": prices, "total": None,
                      "note": "; ".join(notes)})
    b.competitors = comps
    db.commit()
    out = full_build(db, b)
    out["reference_notes"] = notes
    return out


@router.get("/by-opportunity/{opp_id}")
def by_opportunity(opp_id: int, db: Session = Depends(get_db)):
    """Price builds linked to an opportunity, newest first, with totals (for the opportunity page)."""
    rows = db.scalars(select(PriceBuild).where(PriceBuild.opportunity_id == opp_id).order_by(PriceBuild.updated_at.desc())).all()
    builds = [build_summary(db, b) for b in rows]
    return {"opportunity_id": opp_id, "builds": builds, "latest": builds[0] if builds else None}


# ---------------------------------------------------------------- xlsx export
def _sheet_title(s: str, used: set) -> str:
    t = re.sub(r"[\[\]\*\?/\\:]", "-", s)[:31] or "Sheet"
    base, n = t, 2
    while t in used:
        t = f"{base[:28]}-{n}"
        n += 1
    used.add(t)
    return t


def export_xlsx(build: dict, rates: dict, cats: dict[int, LaborCategory]) -> bytes:
    """Price volume workbook with live formulas: Rates, Labor, one sheet per CLIN, Summary."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter, quote_sheetname

    R = merged_rates(rates)
    wb = Workbook()
    bold = Font(bold=True)
    head = PatternFill("solid", fgColor="DDE6F0")
    money = '"$"#,##0.00'
    pct = "0.00%"
    used = {"Summary", "Rates", "Labor"}

    # Rates
    rs = wb.active
    rs.title = "Rates"
    rs.append(["Rate", "Value", "Applied to"])
    rows = [("Fringe", R["fringe_pct"], "Direct labor"),
            ("Overhead", R["overhead_pct"], BASES["overhead_base"][R["overhead_base"]]),
            ("G&A", R["ga_pct"], BASES["ga_base"][R["ga_base"]]),
            ("Profit / fee", R["profit_pct"], BASES["fee_base"][R["fee_base"]]),
            ("Material handling", R["material_handling_pct"], "Materials")]
    for name, v, base in rows:
        rs.append([name, v / 100, base])
    for r in range(2, 7):
        rs.cell(r, 2).number_format = pct
    rs.append([])
    rs.append(["Edit the blue values; every other sheet recalculates. The bases above were set when this file was exported."])
    for c in rs[1]:
        c.font, c.fill = bold, head
    for r in range(2, 7):
        rs.cell(r, 2).font = Font(color="1F4E99", bold=True)
    rs.column_dimensions["A"].width, rs.column_dimensions["B"].width, rs.column_dimensions["C"].width = 22, 12, 60
    FR, OH, GA, PF, MH = "Rates!$B$2", "Rates!$B$3", "Rates!$B$4", "Rates!$B$5", "Rates!$B$6"

    def oh_formula(dl, fr):
        return f"=({dl}+{fr})*{OH}" if R["overhead_base"] == "labor_fringe" else f"={dl}*{OH}"

    # Labor
    ls = wb.create_sheet("Labor")
    ls.append(["Labor category", "Direct rate", "Fringe", "Overhead", "G&A", "Fully burdened cost", "Profit", "Billing rate", "GSA CALC+ category"])
    cat_row: dict = {}
    used_cats: list = []
    for clin in build.get("clins") or []:
        for row in clin.get("labor") or []:
            cid = row.get("category_id")
            key = int(cid) if cid not in (None, "") and int(cid) in cats else ("rate", _f(row.get("rate")), row.get("category") or "")
            if key not in used_cats:
                used_cats.append(key)
    for i, key in enumerate(used_cats, start=2):
        if isinstance(key, int):
            c = cats[key]
            name, rate, calc = c.name, category_rate(c), c.calc_category
        else:
            name, rate, calc = key[2] or f"Direct rate {key[1]:g}", key[1], ""
        ls.append([name, rate, f"=B{i}*{FR}", oh_formula(f"B{i}", f"C{i}")[0:] , f"=(B{i}+C{i}+D{i})*{GA}",
                   f"=B{i}+C{i}+D{i}+E{i}", f"=F{i}*{PF}", f"=F{i}+G{i}", calc])
        for col in range(2, 9):
            ls.cell(i, col).number_format = money
        cat_row[key] = i
    for c in ls[1]:
        c.font, c.fill = bold, head
    ls.column_dimensions["A"].width, ls.column_dimensions["I"].width = 30, 30
    for col in "BCDEFGH":
        ls.column_dimensions[col].width = 14

    # CLIN sheets
    refs = []
    for clin in build.get("clins") or []:
        title = _sheet_title(f"CLIN {clin.get('clin') or len(refs) + 1}", used)
        ws = wb.create_sheet(title)
        q = quote_sheetname(title)
        ws.append([f"CLIN {clin.get('clin', '')}: {clin.get('description', '')}"])
        ws["A1"].font = Font(bold=True, size=12)
        r = 3

        def section(label, cols):
            nonlocal r
            ws.cell(r, 1, label).font = bold
            for j, h in enumerate(cols, start=2):
                ws.cell(r, j, h).font = bold
            ws.cell(r, 1).fill = head
            r += 1

        def rng(start, end):
            return (start, end) if end >= start else None

        section("Labor category", ["Hours", "Rate", "Cost"])
        s = r
        for row in clin.get("labor") or []:
            cid = row.get("category_id")
            key = int(cid) if cid not in (None, "") and int(cid) in cats else ("rate", _f(row.get("rate")), row.get("category") or "")
            lr = cat_row[key]
            ws.cell(r, 1, f"=Labor!A{lr}")
            ws.cell(r, 2, _f(row.get("hours")))
            ws.cell(r, 3, f"=Labor!B{lr}").number_format = money
            ws.cell(r, 4, f"=B{r}*C{r}").number_format = money
            r += 1
        lab = rng(s, r - 1)
        r += 1
        section("Materials", ["", "", "Cost"])
        s = r
        for m in clin.get("materials") or []:
            ws.cell(r, 1, m.get("description") or "")
            ws.cell(r, 4, _f(m.get("cost"))).number_format = money
            r += 1
        mat = rng(s, r - 1)
        r += 1
        section("Subcontracts", ["Similarly situated", "Lower-tier non-SS", "Cost"])
        s = r
        for sc in clin.get("subcontracts") or []:
            ws.cell(r, 1, sc.get("name") or "")
            ws.cell(r, 2, "Yes" if sc.get("similarly_situated") else "No")
            ws.cell(r, 3, _f(sc.get("lower_tier_non_ss"))).number_format = money
            ws.cell(r, 4, _f(sc.get("cost"))).number_format = money
            r += 1
        sub = rng(s, r - 1)
        r += 1
        section("ODCs and travel", ["Type", "", "Cost"])
        s = r
        for o in clin.get("odcs") or []:
            ws.cell(r, 1, o.get("description") or "")
            ws.cell(r, 2, o.get("kind") or "odc")
            ws.cell(r, 4, _f(o.get("cost"))).number_format = money
            r += 1
        odc = rng(s, r - 1)
        r += 1
        section("Fixed-price part lines", ["Quantity", "Unit price", "Extended"])
        s = r
        for p in clin.get("parts") or []:
            ws.cell(r, 1, p.get("description") or (f"Part quote #{p.get('part_quote_id')}" if p.get("part_quote_id") else ""))
            ws.cell(r, 2, _f(p.get("quantity")))
            ws.cell(r, 3, _f(p.get("unit_price"))).number_format = money
            ws.cell(r, 4, f"=B{r}*C{r}").number_format = money
            r += 1
        prt = rng(s, r - 1)
        r += 1

        def sum_of(rg, col="D"):
            return f"=SUM({col}{rg[0]}:{col}{rg[1]})" if rg else 0

        ws.cell(r, 1, "Cost build-up").font = bold
        ws.cell(r, 1).fill = head
        r += 1
        cells = {}

        def line(key, label, formula):
            nonlocal r
            ws.cell(r, 1, label)
            ws.cell(r, 4, formula).number_format = money
            cells[key] = f"D{r}"
            r += 1

        line("direct_labor", "Direct labor", sum_of(lab))
        line("fringe", "Fringe", f"={cells['direct_labor']}*{FR}")
        line("overhead", "Overhead", oh_formula(cells["direct_labor"], cells["fringe"]))
        line("materials", "Materials", sum_of(mat))
        line("material_handling", "Material handling", f"={cells['materials']}*{MH}")
        line("subcontracts", "Subcontracts", sum_of(sub))
        line("odc", "ODCs and travel", sum_of(odc))
        C = cells
        line("total_cost_input", "Total cost input",
             f"={C['direct_labor']}+{C['fringe']}+{C['overhead']}+{C['materials']}+{C['material_handling']}+{C['subcontracts']}+{C['odc']}")
        ga_base = {"total_cost_input": f"{C['total_cost_input']}",
                   "value_added": f"({C['total_cost_input']}-{C['materials']}-{C['material_handling']}-{C['subcontracts']})",
                   "labor_overhead": f"({C['direct_labor']}+{C['fringe']}+{C['overhead']})"}[R["ga_base"]]
        line("ga", "G&A", f"={ga_base}*{GA}")
        line("total_cost", "Total cost", f"={C['total_cost_input']}+{C['ga']}")
        fee_base = C["total_cost"] if R["fee_base"] == "total_cost" else f"({C['total_cost']}-{C['materials']}-{C['material_handling']}-{C['subcontracts']})"
        line("profit", "Profit", f"={fee_base}*{PF}")
        line("parts", "Fixed-price part lines", sum_of(prt))
        line("price", "CLIN price", f"={C['total_cost']}+{C['profit']}+{C['parts']}")
        ws.cell(r - 1, 1).font = bold
        ws.cell(r - 1, 4).font = bold
        ws.cell(r, 1, "Quantity")
        ws.cell(r, 4, _f(clin.get("quantity"), 1.0) or 1.0)
        cells["quantity"] = f"D{r}"
        r += 1
        line("unit_price", "Unit price", f"=IF({C['quantity']}=0,0,{C['price']}/{C['quantity']})")
        nonss = "0"
        if sub:
            nonss = f'SUMIF(B{sub[0]}:B{sub[1]},"No",D{sub[0]}:D{sub[1]})+SUMIF(B{sub[0]}:B{sub[1]},"Yes",C{sub[0]}:C{sub[1]})'
        line("non_ss", "Paid to subcontractors not similarly situated", f"={nonss}")
        ws.column_dimensions["A"].width = 44
        for col in "BCD":
            ws.column_dimensions[col].width = 16
        refs.append({"clin": clin.get("clin", ""), "description": clin.get("description", ""),
                     "kind": clin.get("kind") or "services", "q": q, "cells": dict(cells)})

    # Summary
    sm = wb.create_sheet("Summary", 0)
    cols = [("direct_labor", "Direct labor"), ("fringe", "Fringe"), ("overhead", "Overhead"), ("materials", "Materials"),
            ("material_handling", "Handling"), ("subcontracts", "Subcontracts"), ("odc", "ODC/travel"), ("ga", "G&A"),
            ("total_cost", "Total cost"), ("profit", "Profit"), ("parts", "Part lines"), ("price", "Price"),
            ("unit_price", "Unit price"), ("non_ss", "Non-SS subs")]
    comps = [c for c in (build.get("competitors") or []) if c.get("prices") or c.get("total") not in (None, "")]
    header = ["CLIN", "Description", "Type"] + [h for _, h in cols]
    for c in comps:
        header += [c.get("name") or "Competitor", f"Gap vs {c.get('name') or 'competitor'}"]
    sm.append([build.get("name") or "Price build"])
    sm["A1"].font = Font(bold=True, size=13)
    sm.append([])
    sm.append(header)
    for c in sm[3]:
        c.font, c.fill = bold, head
    first = 4
    price_col = get_column_letter(3 + [k for k, _ in cols].index("price") + 1)
    for i, ref in enumerate(refs):
        rr = first + i
        row = [ref["clin"], ref["description"], ref["kind"]]
        row += [f"={ref['q']}!{ref['cells'][k]}" for k, _ in cols]
        for comp in comps:
            v = (comp.get("prices") or {}).get(ref["clin"])
            ccol = get_column_letter(len(row) + 1)
            row += [_f(v) if v not in (None, "") else None, f'=IF({ccol}{rr}="","",{price_col}{rr}-{ccol}{rr})']
        sm.append(row)
    last = first + len(refs) - 1
    tr = last + 1
    total_row = ["Total", "", ""]
    for j, (k, _) in enumerate(cols):
        col = get_column_letter(4 + j)
        total_row.append("" if k == "unit_price" else (f"=SUM({col}{first}:{col}{last})" if refs else 0))
    for ci, comp in enumerate(comps):
        ccol = get_column_letter(4 + len(cols) + ci * 2)
        tot = comp.get("total")
        if tot not in (None, ""):
            total_row.append(_f(tot))
            total_row.append(f"={price_col}{tr}-{ccol}{tr}")
        elif refs:
            # Competitor priced only some CLINs: compare against our price for those CLINs
            total_row.append(f"=SUM({ccol}{first}:{ccol}{last})")
            total_row.append(f'=SUMIF({ccol}{first}:{ccol}{last},"<>",{price_col}{first}:{price_col}{last})-{ccol}{tr}')
        else:
            total_row += [0, 0]
    sm.append(total_row)
    for c in sm[tr]:
        c.font = bold
    for rr in range(first, tr + 1):
        for j in range(4, len(header) + 1):
            sm.cell(rr, j).number_format = money
    # Limitations on subcontracting
    kcol = "C"
    mcol = get_column_letter(4 + [k for k, _ in cols].index("materials"))
    ncol = get_column_letter(4 + [k for k, _ in cols].index("non_ss"))
    r = tr + 2
    sm.cell(r, 1, "Limitations on subcontracting (FAR 52.219-14)").font = bold
    sm.cell(r, 4, LOS_URL)
    r += 1
    for j, h in enumerate(["Portion", "", "", "Amount paid", "Cost of materials", "Base", "Non-SS subs", "Share", "Limit", "Status"], start=1):
        sm.cell(r, j, h).font = bold
    r += 1
    if refs:
        for kind in ("services", "supplies"):
            paid = f'SUMIF({kcol}{first}:{kcol}{last},"{kind}",{price_col}{first}:{price_col}{last})'
            mats = f'SUMIF({kcol}{first}:{kcol}{last},"{kind}",{mcol}{first}:{mcol}{last})' if kind == "supplies" else "0"
            sm.cell(r, 1, kind.title())
            sm.cell(r, 4, f"={paid}").number_format = money
            sm.cell(r, 5, f"={mats}").number_format = money
            sm.cell(r, 6, f"=D{r}-E{r}").number_format = money
            sm.cell(r, 7, f'=SUMIF({kcol}{first}:{kcol}{last},"{kind}",{ncol}{first}:{ncol}{last})').number_format = money
            sm.cell(r, 8, f"=IF(F{r}<=0,0,G{r}/F{r})").number_format = pct
            sm.cell(r, 9, LOS_LIMIT / 100).number_format = pct
            sm.cell(r, 10, f'=IF(F{r}<=0,"n/a",IF(H{r}>I{r},"OVER LIMIT","OK"))')
            r += 1
    sm.column_dimensions["A"].width, sm.column_dimensions["B"].width = 12, 36
    for j in range(3, len(header) + 1):
        sm.column_dimensions[get_column_letter(j)].width = 14
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/builds/{build_id}/export.xlsx")
def export_build(build_id: int, db: Session = Depends(get_db)):
    b = _get_build(db, build_id)
    data = export_xlsx(build_dict(b), b.rates, _cats(db))
    name = re.sub(r"[^A-Za-z0-9_-]+", "_", b.name or f"price_build_{b.id}")[:60] + ".xlsx"
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})
