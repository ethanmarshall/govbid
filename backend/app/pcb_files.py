"""Read PCB fabrication and assembly files to fill in a board for the box build quoter.

Accepted inputs (one or several files, or a zip holding any of them, as KiCad, Altium and Eagle export):
  - Gerber RS-274X layers. Copper layers are found from the X2 file attribute
    (%TF.FileFunction,Copper,L1,Top*%) when present, otherwise from the file name
    (KiCad F_Cu / B_Cu / In1_Cu, Protel .GTL .GBL .G1 .G2 ..., Eagle .cmp .sol). The board size comes
    from the outline layer (FileFunction Profile, Edge_Cuts, .GKO, .GM1, .GML, .OLN), using the
    %FS format and %MO units of that file. Arcs are taken at their end points, so a rounded outline
    may read slightly small.
  - Excellon drill files (.drl .xln .exc .drd, or any file starting with M48). Plated holes from
    0.7 mm to 3.0 mm are counted as through-hole component pins (smaller ones are vias, larger ones
    mounting holes). That is an estimate, shown as one.
  - Pick-and-place (centroid) files: KiCad .pos (text or CSV) and Altium/Eagle CSV. Placements per side,
    fine pitch and BGA parts by package name. Fiducials, test points and mounting holes are skipped.
  - A BOM (CSV/XLSX): one line per part with references, quantity, value, footprint, MPN and
    manufacturer. Through-hole parts are found from the footprint name; pins per part come from the
    footprint (1x04, DIP-8, TO-220 ...) or default to 2.

Nothing is sent anywhere: the files are read in memory and discarded.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from pathlib import PurePosixPath

MAX_MEMBER = 60 * 1024 * 1024  # largest single file read out of a zip
MAX_TOTAL = 200 * 1024 * 1024  # total bytes read out of one zip
MAX_MEMBERS = 400


class PcbFileError(ValueError):
    pass


# ================================================================ helpers
def _text(data: bytes) -> str:
    return data.decode("utf-8-sig", errors="replace")


def _is_gerber(name: str, text: str) -> bool:
    n = name.lower()
    if re.search(r"\.(gbr|ger|pho|art|gtl|gbl|gto|gbo|gts|gbs|gtp|gbp|gko|gm\d*|gml|g\d+|gp\d+|gd\d*|cmp|sol|stc|sts|plc|pls|crc|crs|oln|bor|ly\d+)$", n):
        return "%FS" in text[:4000] or "%MO" in text[:4000] or "D10" in text or "G04" in text[:4000]
    return text.lstrip().startswith(("G04", "%FS", "%TF", "%MO")) and ("D01" in text or "D02" in text or "D03" in text)


def _is_drill(name: str, text: str) -> bool:
    n = name.lower()
    head = text.lstrip()[:200].upper()
    if re.search(r"\.(drl|xln|exc|drd|txt|tap|nc)$", n) or "drill" in n:
        return head.startswith(("M48", ";", "%")) and bool(re.search(r"^T\d+", text, re.M))
    return head.startswith("M48")


# ================================================================ gerber
COPPER_NAME = [
    (r"(?:^|[-_.])f[._]?cu\b|\.gtl$|\.cmp$|top[._ -]?(?:copper|layer)|copper[._ -]?top|\.toplayer", "top"),
    (r"(?:^|[-_.])b[._]?cu\b|\.gbl$|\.sol$|bot(?:tom)?[._ -]?(?:copper|layer)|copper[._ -]?bot", "bottom"),
    (r"(?:^|[-_.])in(\d+)[._]?cu\b|\.g(\d+)$|\.gp(\d+)$|\.ly(\d+)$|inner[._ -]?(\d+)|(?:mid|internal)[._ -]?(?:layer)?[._ -]?(\d+)", "inner"),
]
OUTLINE_NAME = r"edge[._]?cuts|\.gko$|\.gm1$|\.gml$|\.gm$|\.oln$|\.bor$|outline|profile|board[._ -]?edge|mechanical[._ -]?1\b"


def gerber_role(name: str, text: str) -> tuple[str, str]:
    """('copper', 'L1'|'top'|'bottom'|'inner:3') or ('outline', '') or ('other', '')."""
    m = re.search(r"%TF\.FileFunction,([^,*]+)(?:,([^,*]+))?(?:,([^,*]+))?", text[:6000])
    if m:
        func = m.group(1).strip().lower()
        if func == "copper":
            return "copper", (m.group(2) or "").strip()
        if func == "profile":
            return "outline", ""
        return "other", func
    base = PurePosixPath(name).name.lower()
    if re.search(OUTLINE_NAME, base):
        return "outline", ""
    for pat, side in COPPER_NAME:
        mm = re.search(pat, base)
        if mm:
            if side == "inner":
                n = next((g for g in mm.groups() if g), "")
                return "copper", f"inner:{n}"
            return "copper", side
    return "other", ""


def gerber_extents(text: str) -> tuple[float, float, float, float] | None:
    """(xmin, ymin, xmax, ymax) in inches from D01/D02/D03 coordinates."""
    fs = re.search(r"%FS([LTD]?)([AI]?)X(\d)(\d)Y(\d)(\d)\*%", text)
    if not fs:
        return None
    trailing = fs.group(1) == "T"
    xi, xd, yi, yd = (int(fs.group(i)) for i in (3, 4, 5, 6))
    unit_mm = bool(re.search(r"%MOMM\*%", text)) or (not re.search(r"%MOIN\*%", text) and "G71" in text)
    scale = 1 / 25.4 if unit_mm else 1.0

    def val(s: str, ints: int, decs: int) -> float:
        neg = s.startswith("-")
        s = s.lstrip("+-")
        if trailing:
            s = s.ljust(ints + decs, "0")
        v = int(s) / (10 ** decs) if s else 0.0
        return -v if neg else v

    x = y = 0.0
    xs: list[float] = []
    ys: list[float] = []
    for m in re.finditer(r"(?:X([+-]?\d+))?(?:Y([+-]?\d+))?(?:I[+-]?\d+)?(?:J[+-]?\d+)?D0?([123])\*", text):
        if m.group(1) is None and m.group(2) is None:
            continue
        if m.group(1) is not None:
            x = val(m.group(1), xi, xd)
        if m.group(2) is not None:
            y = val(m.group(2), yi, yd)
        xs.append(x * scale)
        ys.append(y * scale)
    if len(xs) < 2:
        return None
    return min(xs), min(ys), max(xs), max(ys)


# ================================================================ excellon
def drill_holes(text: str) -> dict:
    """Hole counts by plating and size band from an Excellon file."""
    metric = bool(re.search(r"^\s*(METRIC|M71)\b", text, re.M | re.I))
    if re.search(r"^\s*(INCH|M72)\b", text, re.M | re.I):
        metric = False
    npth_file = bool(re.search(r"NON[_ -]?PLATED|NPTH", text[:3000], re.I))
    tools: dict[str, float] = {}
    for m in re.finditer(r"^T(\d+)(?:[FS][\d.]+)*C([\d.]+)", text, re.M):
        d = float(m.group(2))
        tools[str(int(m.group(1)))] = d if metric else d * 25.4  # mm
    counts: dict[str, int] = {}
    cur = None
    in_header = True
    for ln in text.splitlines():
        s = ln.strip()
        if s in ("%", "M95"):
            in_header = False
            continue
        t = re.match(r"^T(\d+)\s*$", s)
        if t:
            cur = str(int(t.group(1)))
            in_header = False
            continue
        if not in_header and cur and re.match(r"^[XY][+-]?[\d.]", s):
            counts[cur] = counts.get(cur, 0) + 1
    out = {"holes": 0, "vias": 0, "component": 0, "mounting": 0, "plated": not npth_file, "tools": len(tools)}
    for t, n in counts.items():
        d = tools.get(t)
        out["holes"] += n
        if d is None:
            continue
        if npth_file or d > 3.0:
            out["mounting"] += n
        elif d >= 0.7:
            out["component"] += n
        else:
            out["vias"] += n
    return out


# ================================================================ pick and place
FINE_PITCH = r"QFN|DFN|QFP|TQFP|LQFP|TSSOP|SSOP|MSOP|VSSOP|0201|01005|LGA|CSP|WLCSP|SON\b|SC-?70|SOT-?563|SOT-?363|UDFN|WSON|VQFN"
BGA = r"BGA|FBGA|UBGA|CABGA"
NOT_PLACED_REF = r"^(FID|FD|FIDUCIAL|TP|MH|H\d|MP|MTG|LOGO|REF\*|KIBUZZARD|G\*\*\*)"


def _sniff_rows(text: str) -> list[list[str]]:
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if not lines:
        return []
    # KiCad text .pos: "# Ref Val Package PosX PosY Rot Side" then whitespace columns
    hdr = next((ln for ln in lines[:15] if ln.lstrip().startswith("#") and re.search(r"\bRef\b", ln, re.I) and re.search(r"Side|Layer", ln, re.I)), None)
    if hdr:
        out = [hdr.lstrip("# ").split()]
        out += [ln.split() for ln in lines if not ln.lstrip().startswith("#")]
        return out
    sample = "\n".join(lines[:20])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [r for r in csv.reader(io.StringIO("\n".join(lines)), dialect)]


def _header_index(rows: list[list[str]], want: dict[str, str], need: set[str]) -> tuple[int, dict[str, int]] | None:
    for i, r in enumerate(rows[:40]):
        idx: dict[str, int] = {}
        for j, cell in enumerate(r):
            c = re.sub(r"\s+", " ", str(cell or "").strip().strip('"').upper())
            for k, pat in want.items():
                if k not in idx and re.search(pat, c):
                    idx[k] = j
                    break
        if need <= set(idx):
            return i, idx
    return None


CPL_COLS = {"ref": r"^REF|^DESIGNATOR|^REFDES|^PART ?REF|^COMPONENT$|^NAME$", "package": r"^PACKAGE|^FOOTPRINT|^PATTERN|^CASE",
            "value": r"^VAL|^COMMENT|^PART$|^DESCRIPTION", "side": r"^SIDE|^LAYER|^TB$|^T/B|^MIRROR",
            "x": r"^POS ?X|^MID ?X|^CENTER-?X|^X$|^X ?\(|^REF ?X|^CENTER ?X", "y": r"^POS ?Y|^MID ?Y|^CENTER-?Y|^Y$|^Y ?\(|^REF ?Y"}


def parse_cpl(text: str) -> dict | None:
    rows = _sniff_rows(text)
    hit = _header_index(rows, CPL_COLS, {"ref", "x", "y"})
    if not hit:
        return None
    hi, idx = hit
    top = bot = fine = bga = skipped = 0
    for r in rows[hi + 1:]:
        if len(r) <= max(idx.values()):
            continue
        ref = r[idx["ref"]].strip().strip('"').upper()
        if not ref:
            continue
        if re.match(NOT_PLACED_REF, ref):
            skipped += 1
            continue
        pkg = f"{r[idx['package']] if 'package' in idx else ''} {r[idx['value']] if 'value' in idx else ''}".upper()
        side = (r[idx["side"]] if "side" in idx else "top").strip().strip('"').lower()
        if side.startswith("b"):  # bottom, BottomLayer, B
            bot += 1
        else:
            top += 1
        if re.search(BGA, pkg):
            bga += 1
        elif re.search(FINE_PITCH, pkg):
            fine += 1
    if top + bot == 0:
        return None
    return {"placements": top + bot, "top": top, "bottom": bot, "fine_pitch": fine, "bga": bga, "skipped": skipped}


# ================================================================ BOM
BOM_COLS = {"refs": r"^REF|^DESIGNATOR|^REFDES|^PARTS?$|^REFERENCES?|^LOCATION", "qty": r"^QTY|^QUANTITY|^Q'?TY|^COUNT$",
            "value": r"^VAL|^COMMENT", "footprint": r"^FOOTPRINT|^PACKAGE|^PATTERN|^CASE",
            "mpn": r"^MPN|^MFR\.? ?P|^MFG\.? ?P|^MANUFACTURER ?P|^MANUFACTURER_?PART|^PART ?(?:NUMBER|NO|#)|^P/?N$|^MFR ?PART",
            "manufacturer": r"^MFR\.?$|^MFG\.?$|^MANUFACTURER$|^MANUFACTURER ?NAME|^MAKE$|^BRAND",
            "desc": r"^DESC|^DESCRIPTION", "dnp": r"^DNP|^DNI|^POPULATE|^FITTED|^EXCLUDE",
            "price": r"^UNIT ?(?:PRICE|COST)|^PRICE|^COST"}
THT = (r"THT|(?:^|[_:-])TH(?:$|[_-])|\bDIP\b|DIP-?\d|PDIP|PINHEADER|PIN_?HEADER|PINSOCKET|PIN_?SOCKET|TO-?220|TO-?92|TO-?247|TO-?3P|\bAXIAL|_AXIAL|RADIAL|"
       r"CP_RADIAL|D_DO-?\d|DO-?41|DO-?15|DO-?201|TERMINALBLOCK|TERMINAL_?BLOCK|BARRELJACK|BARREL_?JACK|RELAY_THT|\bSIP|IDC|DSUB|D-?SUB|"
       r"MOLEX.*VERTICAL|THROUGH ?HOLE|BUZZER|CRYSTAL_HC49|HC-?49|POTENTIOMETER_|TRANSFORMER_THT|FUSEHOLDER|BATTERYHOLDER")


def tht_pins(fp: str) -> int:
    f = fp.upper()
    for pat, fn in ((r"(\d+)X(\d+)", lambda m: int(m.group(1)) * int(m.group(2))), (r"DIP-?(\d+)", lambda m: int(m.group(1))),
                    (r"SIP-?(\d+)", lambda m: int(m.group(1))), (r"DSUB-?(\d+)|DE-?(\d+)|DB-?(\d+)", lambda m: int(next(g for g in m.groups() if g))),
                    (r"TO-?(?:220|92|247|3P)", lambda m: 3), (r"RELAY", lambda m: 5), (r"BARREL", lambda m: 3), (r"POTENTIOMETER", lambda m: 3)):
        m = re.search(pat, f)
        if m:
            try:
                return max(int(fn(m)), 1)
            except (StopIteration, ValueError):
                continue
    return 2


def expand_refs(text: str) -> list[str]:
    out: list[str] = []
    for tok in re.split(r"[,;\s]+", text or ""):
        tok = tok.strip().upper()
        if not tok:
            continue
        m = re.match(r"^([A-Z]+)(\d+)\s*-\s*(?:[A-Z]+)?(\d+)$", tok)
        if m and int(m.group(3)) >= int(m.group(2)) and int(m.group(3)) - int(m.group(2)) < 500:
            out += [f"{m.group(1)}{i}" for i in range(int(m.group(2)), int(m.group(3)) + 1)]
        else:
            out.append(tok)
    return out


def parse_bom_rows(rows: list[list[str]]) -> dict | None:
    hit = _header_index(rows, BOM_COLS, {"refs"}) or _header_index(rows, BOM_COLS, {"qty", "mpn"}) or _header_index(rows, BOM_COLS, {"qty", "value"})
    if not hit:
        return None
    hi, idx = hit
    if len(idx) < 2:
        return None
    lines = []
    smt = tht_parts = tht_joints = fine = bga = 0
    smt_unique = 0
    for r in rows[hi + 1:]:
        get = lambda k: (str(r[idx[k]]).strip() if k in idx and idx[k] < len(r) and r[idx[k]] is not None else "")  # noqa: E731
        refs = [x for x in expand_refs(get("refs")) if not re.match(NOT_PLACED_REF, x)]
        if not any(get(k) for k in ("refs", "value", "mpn", "desc")):
            continue
        if "dnp" in idx:
            dnp = get("dnp").upper()
            positive = bool(re.search(r"POPULATE|FITTED", str(rows[hi][idx["dnp"]]).upper()))  # "Populate: No" vs "DNP: Yes"
            if (positive and dnp in ("NO", "N", "FALSE", "0", "DNP", "NOT FITTED")) or (not positive and dnp in ("DNP", "DNI", "YES", "Y", "TRUE", "1", "X", "EXCLUDED")):
                continue
        if "refs" in idx and not refs and get("refs"):
            continue  # only fiducials / test points / mounting holes on this line
        try:
            qty = float(re.sub(r"[^\d.]", "", get("qty"))) if get("qty") else float(len(refs) or 1)
        except ValueError:
            qty = float(len(refs) or 1)
        if qty <= 0:
            continue
        fp = get("footprint")
        text = f"{fp} {get('desc')} {get('value')}".upper()
        is_tht = bool(re.search(THT, text))
        if is_tht:
            tht_parts += qty
            tht_joints += qty * tht_pins(fp or text)
        else:
            smt += qty
            smt_unique += 1
            if re.search(BGA, text):
                bga += qty
            elif re.search(FINE_PITCH, text):
                fine += qty
        price = None
        if get("price"):
            try:
                price = float(re.sub(r"[$,\s]", "", get("price")))
            except ValueError:
                price = None
        lines.append({"ref": ",".join(refs[:12]) + ("..." if len(refs) > 12 else ""), "mpn": get("mpn") or "", "manufacturer": get("manufacturer"),
                      "description": " ".join(x for x in (get("value"), get("desc") if get("desc") != get("value") else "", fp) if x)[:160],
                      "qty": qty, "tht": is_tht, "unit_price": price, "price_source": "manual" if price is not None else "", "distributor": ""})
    if not lines:
        return None
    return {"lines": lines, "smt_placements": int(smt), "smt_unique": smt_unique, "tht_parts": int(tht_parts), "tht_joints": int(tht_joints),
            "fine_pitch": int(fine), "bga": int(bga)}


def _sheet_rows(name: str, data: bytes) -> list[list[str]]:
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise PcbFileError(f"Could not read {name}: {exc}") from exc
        rows: list[list[str]] = []
        for ws in wb.worksheets:
            rows += [["" if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)) for v in r] for r in ws.iter_rows(values_only=True)]
        return rows
    return _sniff_rows(_text(data))


# ================================================================ main entry
def _members(files: list[tuple[str, bytes]]) -> list[tuple[str, bytes]]:
    out: list[tuple[str, bytes]] = []
    total = 0
    for name, data in files:
        if name.lower().endswith(".zip"):
            try:
                z = zipfile.ZipFile(io.BytesIO(data))
            except zipfile.BadZipFile as exc:
                raise PcbFileError(f"{name} is not a readable zip file.") from exc
            infos = [i for i in z.infolist() if not i.is_dir() and not PurePosixPath(i.filename).name.startswith((".", "__MACOSX"))
                     and "__MACOSX" not in i.filename]
            if len(infos) > MAX_MEMBERS:
                raise PcbFileError(f"{name} holds more than {MAX_MEMBERS} files.")
            for i in infos:
                if i.file_size > MAX_MEMBER:
                    continue
                total += i.file_size
                if total > MAX_TOTAL:
                    raise PcbFileError(f"{name} unpacks to more than {MAX_TOTAL // 1024 // 1024} MB.")
                inner = z.read(i)
                if i.filename.lower().endswith(".zip"):
                    continue  # one level only
                out.append((i.filename, inner))
        else:
            out.append((name, data))
    return out


def parse_files(files: list[tuple[str, bytes]]) -> dict:
    """Read Gerbers, drill, pick-and-place and BOM files. Returns the board fields found, BOM lines,
    evidence (what each value came from) and warnings."""
    members = _members(files)
    if not members:
        raise PcbFileError("No files to read.")
    copper: dict[str, str] = {}
    outline = None
    copper_ext = None
    drills = {"holes": 0, "vias": 0, "component": 0, "mounting": 0}
    drill_files = 0
    cpl = None
    bom = None
    evidence: list[str] = []
    warnings: list[str] = []
    used: list[str] = []
    for name, data in members:
        base = PurePosixPath(name).name
        low = base.lower()
        if low.endswith((".pdf", ".png", ".jpg", ".jpeg", ".step", ".stp", ".kicad_pcb", ".brd", ".pcbdoc", ".schdoc", ".kicad_sch", ".sch")):
            continue
        if low.endswith((".xlsx", ".xlsm")):
            rows = _sheet_rows(base, data)
            got = parse_bom_rows(rows) if bom is None else None
            if got:
                bom = got
                used.append(base)
                evidence.append(f"BOM from {base}: {len(got['lines'])} lines.")
            else:
                c = parse_cpl("\n".join(",".join(r) for r in rows))
                if c and cpl is None:
                    cpl = c
                    used.append(base)
            continue
        text = _text(data)
        if _is_drill(base, text):
            d = drill_holes(text)
            drill_files += 1
            for k in drills:
                drills[k] += d[k]
            used.append(base)
            continue
        if _is_gerber(base, text):
            role, which = gerber_role(base, text)
            if role == "copper":
                key = which or base
                copper[key] = base
                if copper_ext is None and (which.lower() in ("l1", "top") or "top" in which.lower()):
                    copper_ext = gerber_extents(text)
            elif role == "outline" and outline is None:
                ext = gerber_extents(text)
                if ext:
                    outline = (base, ext)
            used.append(base)
            continue
        if low.endswith((".csv", ".pos", ".txt", ".tsv", ".xy", ".mnt", ".mnb")):
            rows = _sniff_rows(text)
            c = parse_cpl(text)
            if c and cpl is None and not (_header_index(rows, BOM_COLS, {"refs", "qty"}) and not _header_index(rows, CPL_COLS, {"ref", "x", "y"})):
                cpl = c
                used.append(base)
                evidence.append(f"Pick-and-place from {base}: {c['top']} top, {c['bottom']} bottom placements"
                                f"{f', {c['skipped']} fiducials / test points skipped' if c['skipped'] else ''}.")
                continue
            got = parse_bom_rows(rows) if bom is None else None
            if got:
                bom = got
                used.append(base)
                evidence.append(f"BOM from {base}: {len(got['lines'])} lines.")
    board: dict = {}
    if copper:
        board["layers"] = len(copper)
        evidence.append(f"{len(copper)} copper layer(s): {', '.join(sorted(copper.values()))}.")
        if len(copper) not in (1, 2, 4, 6, 8, 10, 12, 14, 16):
            warnings.append(f"Found {len(copper)} copper layers, an unusual count. Check that every copper Gerber is included.")
    ext = outline[1] if outline else copper_ext
    if ext:
        w, h = ext[2] - ext[0], ext[3] - ext[1]
        if 0.1 < w < 60 and 0.1 < h < 60:
            board["width_in"], board["height_in"] = round(max(w, h), 3), round(min(w, h), 3)
            evidence.append(f"Board {board['width_in']} x {board['height_in']} in from {'outline ' + outline[0] if outline else 'the top copper extents (no outline layer)'}.")
            if not outline:
                warnings.append("No board outline layer found, so the size comes from the copper extents and may read small. Include Edge_Cuts / GKO.")
        else:
            warnings.append("The outline size did not make sense (check the Gerber units). Enter the board size by hand.")
    elif copper:
        warnings.append("Could not read the board size from the Gerbers. Enter it by hand.")
    if drill_files:
        evidence.append(f"Drill: {drills['holes']} holes ({drills['vias']} vias under 0.7 mm, {drills['component']} component holes 0.7 to 3.0 mm, "
                        f"{drills['mounting']} mounting or unplated).")
    if cpl:
        board.update(smt_placements=cpl["placements"], sides=2 if cpl["bottom"] else 1, fine_pitch=cpl["fine_pitch"], bga=cpl["bga"])
    if bom:
        if not cpl:
            board.update(smt_placements=bom["smt_placements"], fine_pitch=bom["fine_pitch"], bga=bom["bga"])
            evidence.append(f"SMT placements {bom['smt_placements']} counted from the BOM (no pick-and-place file).")
        board.update(smt_unique=bom["smt_unique"], tht_parts=bom["tht_parts"], tht_joints=bom["tht_joints"])
        evidence.append(f"{bom['smt_unique']} unique SMT parts; {bom['tht_parts']} through-hole parts with about {bom['tht_joints']} solder joints "
                        "(pins from the footprint names, 2 when not shown).")
        missing = sum(1 for ln in bom["lines"] if not ln["mpn"])
        if missing:
            warnings.append(f"{missing} BOM line(s) have no MPN, so live pricing cannot find them.")
    elif drill_files and drills["component"]:
        board.update(tht_joints=drills["component"])
        evidence.append(f"Through-hole joints estimated at {drills['component']} from the component-size plated holes (no BOM).")
    if not (copper or cpl or bom or drill_files):
        raise PcbFileError("No Gerber, drill, pick-and-place or BOM data found in the files.")
    if not bom:
        warnings.append("No BOM found: component cost is not known. Add the BOM to price parts.")
    return {"board": board, "bom_lines": (bom or {}).get("lines", []), "evidence": evidence, "warnings": warnings,
            "files_used": used, "found": {"gerbers": len(copper) + (1 if outline else 0), "drill": drill_files, "cpl": bool(cpl), "bom": bool(bom)}}


def area_sqin(board: dict) -> float:
    return max(float(board.get("width_in") or 0) * float(board.get("height_in") or 0), 0.0)


__all__ = ["PcbFileError", "parse_files", "parse_cpl", "parse_bom_rows", "drill_holes", "gerber_extents", "gerber_role", "area_sqin"]
