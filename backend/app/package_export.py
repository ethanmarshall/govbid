"""Word and zip export for proposal and technical data packages."""
from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def safe_name(s: str, limit: int = 80) -> str:
    s = re.sub(r"[^A-Za-z0-9 ._()-]+", "_", s or "").strip(" ._")
    return (s or "untitled")[:limit]


def words(text: str) -> int:
    return len(re.findall(r"\b\w+\b", text or ""))


# ------------------------------------------------------------------ low-level docx helpers
def _field(paragraph, instr: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, code, sep, text, end):
        run._r.append(el)


def _add_inline(paragraph, text: str) -> None:
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _table(doc, header: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        _shade(c, "D9E2EC")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            _add_inline(cells[i].paragraphs[0], str(v or ""))
    # repeat header row on each page
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    tr_pr.append(hdr)
    if widths:
        t.autofit = False
        for i, w in enumerate(widths):
            t.columns[i].width = Inches(w)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()


def markdown_to_docx(doc, md: str, base_level: int = 3) -> None:
    """Small markdown subset: headings, bullets, numbered lists, pipe tables, bold/italic/code, blank-line paragraphs."""
    lines = (md or "").replace("\r\n", "\n").split("\n")
    i = 0
    para_buf: list[str] = []

    def flush():
        if para_buf:
            p = doc.add_paragraph()
            _add_inline(p, " ".join(s.strip() for s in para_buf))
            para_buf.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush()
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if m:
            flush()
            level = min(base_level + len(m.group(1)) - 1, 9)
            doc.add_heading(m.group(2).strip(), level=level)
            i += 1
            continue
        if stripped.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                width = max(len(r) for r in rows)
                rows = [r + [""] * (width - len(r)) for r in rows]
                _table(doc, rows[0], rows[1:])
            continue
        m = re.match(r"^[-*+]\s+(.*)", stripped)
        if m:
            flush()
            indent = (len(line) - len(line.lstrip())) // 2
            style = "List Bullet 2" if indent >= 1 else "List Bullet"
            _add_inline(doc.add_paragraph(style=style), m.group(1))
            i += 1
            continue
        m = re.match(r"^\d+[.)]\s+(.*)", stripped)
        if m:
            flush()
            _add_inline(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue
        para_buf.append(stripped)
        i += 1
    flush()


# ------------------------------------------------------------------ document assembly
def _setup(doc, font: str, size: float) -> None:
    st = doc.styles["Normal"]
    st.font.name = font
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    st.font.size = Pt(size)
    st.paragraph_format.space_after = Pt(6)
    for lvl, sz in ((1, 16), (2, 13), (3, 12), (4, 11), (5, 11)):
        h = doc.styles[f"Heading {lvl}"]
        h.font.name = font
        h.font.size = Pt(sz)
        h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        rfonts = h.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), font)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
    # Ask Word to refresh the table of contents and page fields on open
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    doc.settings.element.append(upd)


def _header_footer(section, left: str, right: str) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.text = ""
    hp.add_run(left).font.size = Pt(9)
    if right:
        hp.add_run("\t\t" + right).font.size = Pt(9)
    fp = section.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Page ")
    r.font.size = Pt(9)
    _field(fp, "PAGE", "1")
    if left:
        fp.add_run(f"    {left}").font.size = Pt(9)


def build_docx(pkg: dict, profile: dict, matrix: list[dict], out_path: Path) -> Path:
    """pkg: dict with name, kind, cover, sections (list of dicts), items (list of dicts)."""
    cover = pkg.get("cover") or {}
    font = cover.get("font") or "Times New Roman"
    size = float(cover.get("font_size") or 12)
    doc = Document()
    _setup(doc, font, size)

    company = profile.get("name") or "Company name"
    sol = cover.get("solicitation_number") or ""
    is_tdp = pkg.get("kind") == "tdp"

    # Cover page
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(pkg["name"])
    r.bold = True
    r.font.size = Pt(22)
    sub = "Technical Data Package" if is_tdp else (cover.get("volume_title") or "Proposal")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(sub).font.size = Pt(14)
    doc.add_paragraph()
    lines = [
        ("Solicitation", sol),
        ("Contract", pkg.get("contract_number") or ""),
        ("Issuing agency", cover.get("agency") or ""),
        ("Submitted by", company),
        ("UEI / CAGE", " / ".join(x for x in (profile.get("uei"), profile.get("cage")) if x)),
        ("Business status", cover.get("business_status") or ""),
        ("Point of contact", ", ".join(x for x in (cover.get("poc_name"), cover.get("poc_email"), cover.get("poc_phone")) if x)),
        ("Offer valid for", f"{cover['validity_days']} days" if cover.get("validity_days") else ""),
        ("Date", cover.get("date") or date.today().strftime("%B %d, %Y").replace(" 0", " ")),
    ]
    for label, value in lines:
        if not value:
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)
    if cover.get("restriction_notice", True) and not is_tdp:
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(
            "This proposal includes data that shall not be disclosed outside the Government and shall not be duplicated, "
            "used, or disclosed, in whole or in part, for any purpose other than to evaluate this proposal (FAR 52.215-1(e))."
        )
        r.font.size = Pt(9)
        r.italic = True
    if is_tdp and cover.get("distribution_statement"):
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(cover["distribution_statement"])
        r.font.size = Pt(9)
        r.bold = True

    # Body in a new section so the cover has no header
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    _header_footer(body, company, sol or pkg.get("contract_number") or "")

    doc.add_heading("Table of Contents", level=1)
    toc = doc.add_paragraph()
    _field(toc, 'TOC \\o "1-2" \\h \\z \\u', "Right-click and choose Update Field to build the table of contents.")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    current_volume = None
    first = True
    for s in pkg["sections"]:
        vol = s.get("volume") or ""
        if vol and vol != current_volume:
            if not first:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_heading(vol, level=1)
            current_volume = vol
        elif s.get("page_break_before") and not first:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        first = False
        title = " ".join(x for x in (s.get("number"), s.get("title")) if x)
        doc.add_heading(title, level=2)
        if s.get("content", "").strip():
            markdown_to_docx(doc, s["content"], base_level=3)
        else:
            p = doc.add_paragraph()
            r = p.add_run(f"[Section not yet written. {s.get('guidance') or ''}]".strip())
            r.italic = True
            r.font.color.rgb = RGBColor(0xA8, 0x40, 0x2B)

    items = pkg.get("items") or []
    if items:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading("Deliverables index" if is_tdp else "Attachments", level=1)
        header = ["CDRL", "DID", "Title", "Status", "Files"] if is_tdp else ["#", "Attachment", "Status", "Files"]
        rows = []
        for n, it in enumerate(items, 1):
            files = ", ".join(f["name"] for f in it.get("files") or []) or "none"
            status = (it.get("status") or "").replace("_", " ")
            rows.append([it.get("cdrl") or "", it.get("did") or "", it.get("title") or "", status, files] if is_tdp else [str(n), it.get("title") or "", status, files])
        _table(doc, header, rows)

    if matrix and not is_tdp:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading("Compliance cross-reference", level=1)
        rows = []
        for r in matrix:
            req = r.get("requirement") or ""
            rows.append([str(r.get("id")), (req[:220] + "...") if len(req) > 220 else req, r.get("reference") or "", r.get("response_location") or "Not yet addressed"])
        _table(doc, ["#", "Requirement", "Solicitation ref", "Where addressed"], rows, [0.4, 3.7, 1.0, 1.4])

    doc.save(out_path)
    return out_path


def build_zip(pkg: dict, docx_path: Path, item_dir, matrix_xlsx: Path | None, out_path: Path) -> Path:
    root = safe_name(pkg["name"])
    index_rows = [["Folder", "CDRL", "DID", "Title", "Status", "File"]]
    with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(docx_path, f"{root}/{docx_path.name}")
        if matrix_xlsx and matrix_xlsx.exists():
            z.write(matrix_xlsx, f"{root}/Compliance_Matrix.xlsx")
        for n, it in enumerate(pkg.get("items") or [], 1):
            label = safe_name(" ".join(x for x in (it.get("cdrl") or f"{n:02d}", it.get("title")) if x), 60)
            folder = f"{root}/{'Deliverables' if pkg.get('kind') == 'tdp' else 'Attachments'}/{label}"
            files = it.get("files") or []
            if not files:
                index_rows.append([label, it.get("cdrl", ""), it.get("did", ""), it.get("title", ""), it.get("status", ""), ""])
            for f in files:
                src = item_dir(it["id"]) / f["name"]
                if src.exists():
                    z.write(src, f"{folder}/{f['name']}")
                index_rows.append([label, it.get("cdrl", ""), it.get("did", ""), it.get("title", ""), it.get("status", ""), f["name"]])
        buf = io.StringIO()
        csv.writer(buf).writerows(index_rows)
        z.writestr(f"{root}/index.csv", buf.getvalue())
    return out_path
