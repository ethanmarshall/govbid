"""A bill of materials sent with (or without) a 3D model, as CSV, TSV or Excel.

The BOM tells us what the model can't: material, finish, how each part is made, which parts are bought, and
parts that aren't modeled at all (screws, nuts, labels). Rows are matched to the bodies in the STEP file by
part number or name; anything left over becomes its own line.
"""
from __future__ import annotations

import csv
import io
import math
import re
from pathlib import Path

TABLE_EXTS = (".csv", ".tsv", ".xlsx", ".xlsm")

# header patterns, checked in this order for each cell (the first that matches wins)
COLS = [
    ("item", r"^ITEM|^FIND|^NO\.?$|^#$|^LINE$|^BALLOON|^POS(ITION)?$"),
    ("buy", r"MAKE ?(/|OR|-) ?BUY|^BUY|^PROCURE|^SOURCING|^PART TYPE$|^TYPE$"),
    ("profile", r"^PROFILE|^EXTRUSION$|^SERIES$"),
    ("part", r"PART ?(NO|NUM|#)|^P/?N\.?$|^PART$|^NUMBER$|^DWG|^DRAWING|^MPN|MFR\.? ?(PART|P/?N)|MANUFACTURER ?(PART|P/?N)|"
             r"^CATALOG|^SKU|VENDOR ?(PART|P/?N|#)|SUPPLIER ?(PART|P/?N|#)|^ORDER ?(NO|NUM|#)"),
    ("name", r"^NAME$|^PART ?NAME|^COMPONENT|^TITLE$|^FILE ?NAME|^COMPONENT NAME"),
    ("desc", r"^DESC|^DESCRIPTION"),
    ("total_qty", r"^TOTAL ?(QTY|QUANTITY|REQ)"),  # wins over "Qty per unit" when a sheet has both
    ("qty", r"^QTY|^QUANTITY|^Q'?TY|^COUNT$|^AMOUNT$|^REQ|^PCS$|^EA$"),
    ("material", r"^MAT(ERIA)?L?\.?$|^MATERIAL|^STOCK$|^ALLOY"),
    ("finish", r"^FINISH|^COATING|^PLATING|^TREATMENT|^SURFACE"),
    ("process", r"^PROCESS|^METHOD|^(MFG|MANUFACTURING) ?(METHOD|PROCESS)|^OPERATION"),
    ("vendor", r"^VENDOR|^SUPPLIER|^MFR\.?$|^MFG\.?$|^MANUFACTURER$|^SOURCE$|^BRAND|^MAKE$"),
    ("notes", r"^NOTE|^REMARK|^COMMENT"),
]

# signs a sheet belongs to a circuit board (BOM with designators, or pick-and-place)
PCB_HEAD = re.compile(r"^(REF ?DES|DESIGNATOR|REFERENCE DESIGNATOR|FOOTPRINT|PACKAGE|PATTERN|MID ?X|MID ?Y|CENTER ?X|POS ?X|POSX|"
                      r"ROTATION|ROT|LAYER|SIDE|DNP|POPULATE)", re.I)
DESIGNATOR = re.compile(r"^(R|C|L|D|U|Q|J|P|Y|X|F|K|T|SW|LED|TP|FB|IC|CN|JP)\d+([,\s-]|$)", re.I)

BUY_WORDS = re.compile(r"buy|bought|purchas|cots|catalog|off.?the.?shelf|vendor|oem|stock item|procure|hardware", re.I)
MAKE_WORDS = re.compile(r"make|made|fab|machin|manufactur|in.?house|custom|print", re.I)
PROCESS_WORDS = [
    ("cnc_lathe", r"\bturn|lathe"),
    ("sheet_metal", r"sheet|brake|bend|form(ed)?\b|press ?brake"),
    ("3d_print", r"print|additive|\bfdm\b|\bsla\b|\bsls\b|\bmjf\b|dmls"),
    ("laser_cut", r"laser"),
    ("waterjet", r"water ?jet"),
    ("plasma", r"plasma"),
    ("router", r"router|routed"),
    ("cnc_mill", r"mill|machin|\bcnc\b"),
]
MCMASTER = re.compile(r"^\d{4,5}[A-Z]\d{1,4}$", re.I)
OWN = re.compile(r"valley power|in.?house|\bvps\b|^self$|^us$|^n/?a$", re.I)


class BomError(ValueError):
    pass


# ---------------------------------------------------------------- reading
def read_sheets(name: str, data: bytes) -> list[tuple[str, list[list]]]:
    """[(sheet name, rows)]: every sheet of a workbook (a BOM is often split into cut list, panels and hardware)."""
    low = name.lower()
    if low.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise BomError(f"Could not read the spreadsheet: {exc}") from exc
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True, max_row=5000)]) for ws in wb.worksheets[:12]]
    return [("", read_rows(name, data))]


def read_rows(name: str, data: bytes) -> list[list]:
    """The rows of the first sheet with a header we know (the first sheet if none has one)."""
    low = name.lower()
    if low.endswith((".xlsx", ".xlsm")):
        sheets = read_sheets(name, data)
        return next((rows for _, rows in sheets if header(rows) is not None), sheets[0][1] if sheets else [])
    text = data.decode("utf-8-sig", errors="replace")
    first = text.split("\n", 1)[0]
    tab = low.endswith(".tsv") or ("\t" in first and "," not in first)
    if not tab and first.count(";") > first.count(","):
        return list(csv.reader(io.StringIO(text), delimiter=";"))
    return list(csv.reader(io.StringIO(text), csv.excel_tab if tab else csv.excel))


def _map(cells: list) -> dict[str, int]:
    found: dict[str, int] = {}
    for i, c in enumerate(cells):
        name = re.sub(r"\s+", " ", str(c or "").strip().upper()).strip(" :")
        if not name or len(name) > 40:
            continue
        for key, pat in COLS:
            if key not in found and re.search(pat, name):
                found[key] = i
                break
    return found


def _texty(rows: list[list], col: int) -> bool:
    """A column of words ("Door stiles"), not item numbers."""
    vals = [str(r[col]).strip() for r in rows[:40] if col < len(r) and r[col] not in (None, "")]
    return len(vals) >= 1 and sum(1 for v in vals if re.search(r"[A-Za-z]{3}", v)) >= 0.6 * len(vals)


def header(rows: list[list]) -> tuple[int, dict[str, int]] | None:
    for i, r in enumerate(rows[:30]):
        c = _map(r)
        if not ({"qty", "total_qty"} & set(c)) or len(c) < 2:
            continue
        if not ({"name", "desc"} & set(c)) and "item" in c and _texty(rows[i + 1:], c["item"]):
            c["name"] = c.pop("item")  # "Item" holding the part names, as many hand-made BOMs do
        if {"part", "name", "desc"} & set(c):
            return i, c
        if "profile" in c:  # an extrusion cut list: the profile is the part
            c["part"] = c["profile"]
            return i, c
    return None


def sniff(name: str, data: bytes) -> str | None:
    """'pcb' for circuit board BOMs and pick-and-place files, 'bom' for a mechanical BOM, None if we can't tell."""
    try:
        rows = read_rows(name, data)
    except Exception:  # noqa: BLE001
        return None
    for r in rows[:30]:
        heads = [str(c or "").strip() for c in r]
        if sum(1 for h in heads if h and PCB_HEAD.match(h)) >= 1 and any(re.match(r"^(REF|DESIGNATOR)", h, re.I) for h in heads if h):
            return "pcb"
        if sum(1 for h in heads if h and PCB_HEAD.match(h)) >= 2:
            return "pcb"
    hit = header(rows)
    if not hit:
        return None
    hi, cols = hit
    # a column full of R1, C12, U3 is a board BOM even when the header says "Part"
    for key in ("part", "item", "name"):
        if key in cols:
            vals = [str(r[cols[key]] or "").strip() for r in rows[hi + 1: hi + 40] if cols[key] < len(r) and r[cols[key]]]
            if len(vals) >= 3 and sum(1 for v in vals if DESIGNATOR.match(v)) >= 0.6 * len(vals):
                return "pcb"
    return "bom"


# ---------------------------------------------------------------- parsing
def _cell(r: list, cols: dict, key: str) -> str:
    i = cols.get(key)
    if i is None or i >= len(r) or r[i] is None:
        return ""
    v = r[i]
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def map_material(raw: str, cfg: dict | None) -> str | None:
    if not raw:
        return None
    from . import drawing

    if cfg:
        pool = list(cfg.get("materials") or {}) + list((cfg.get("additive") or {}).get("materials") or {})
        exact = next((m for m in pool if m.lower() == raw.lower()), None)
        if exact:
            return exact
        # printed plastics by name: "PETG", "Nylon 12"
        for m in (cfg.get("additive") or {}).get("materials") or {}:
            if re.search(rf"\b{re.escape(m.split(' (')[0])}\b", raw, re.I):
                return m
    return drawing.map_material(raw)


# plain words people put in part names ("Pedestal PETG", "Bearing Cup Aluminum", "Guard - Stainless")
NAME_WORDS = [
    (r"\b(aluminum|aluminium|alum)\b", "6061-T6 aluminum"),
    (r"\b(stainless|ss)\b", "304 stainless"),
    (r"\b(mild steel|carbon steel|steel)\b", "A36 / 1018 steel"),
    (r"\btitanium\b", "titanium Ti-6Al-4V"),
    (r"\b(polycarbonate)\b", "PC (polycarbonate)"),
    (r"\btpu\b", "TPU 95A (flexible)"),
    (r"\b(carbon fiber nylon|nylon cf|pa ?cf)\b", "Nylon CF (carbon fiber)"),
    (r"\bresin\b", "Standard resin"),
]
PART_NUMBERISH = re.compile(r"^\d{3,4}$")


def material_from_name(name: str, cfg: dict | None = None) -> str | None:
    """A material written in a part's name, for models whose CAD export carries no material (Onshape's STEP export
    leaves it out). Strict: a bare 304 or 1018 in a name is more often a part number than an alloy."""
    from . import drawing

    n = re.sub(r"[_]+", " ", str(name or "")).strip()
    if not n:
        return None
    pool = list((cfg or {}).get("materials") or {}) + list(((cfg or {}).get("additive") or {}).get("materials") or {})
    exact = next((m for m in pool if re.search(rf"(^|[\s(-]){re.escape(m)}($|[\s)-])", n, re.I)), None)
    if exact:
        return exact
    for m in ((cfg or {}).get("additive") or {}).get("materials") or {}:
        short = m.split(" (")[0]
        if len(short) >= 3 and not PART_NUMBERISH.match(short) and re.search(rf"\b{re.escape(short)}\b", n, re.I):
            return m
    hit = drawing.map_material(n, strict=True)
    if hit:
        return hit
    m = re.search(r"\b(6061|7075|5052)\b", n)  # aluminum alloys, written bare
    if m:
        return {"6061": "6061-T6 aluminum", "7075": "7075-T6 aluminum", "5052": "5052-H32 aluminum"}[m.group(1)]
    for pat, mat in NAME_WORDS:
        if re.search(pat, n, re.I) and (not cfg or mat in pool):
            return mat
    return None


def infill_from_name(name: str) -> int | None:
    """'Pedestal PETG 80% infill' -> 80."""
    m = re.search(r"\b(\d{2,3})\s*%\s*(infill|fill)\b", str(name or ""), re.I)
    return int(m.group(1)) if m and 10 <= int(m.group(1)) <= 100 else None


def map_finish(raw: str) -> str | None:
    if not raw:
        return None
    if re.fullmatch(r"(none|n/?a|-+|as machined|mill finish|bare|raw)", raw.strip(), re.I):
        return "none"
    from . import drawing

    hit = [f["mapped"] for f in drawing._find_finishes(raw) if f.get("mapped")]
    return hit[0] if hit else None


def map_process(raw: str) -> str | None:
    for key, pat in PROCESS_WORDS:
        if raw and re.search(pat, raw, re.I):
            return key
    return None


def _buy(e: dict) -> bool | None:
    """True bought, False made here, None the BOM doesn't say."""
    from . import catalog_8020

    b = e["buy_raw"]
    if b:
        if BUY_WORDS.search(b) and not re.search(r"\bmake\b", b, re.I):
            return True
        if MAKE_WORDS.search(b):
            return False
    if e["process"]:
        return False
    v = e["vendor"]
    if v and not OWN.search(v):
        return True
    pn = e["part"]
    if pn and (MCMASTER.match(pn) or catalog_8020.match_part(pn)):
        return True
    if e["material"]:
        return False
    return None


TBD = re.compile(r"^(tbd|tba|n/?a|none|-+|\?+|x+)$", re.I)


def parse(name: str, data: bytes, cfg: dict | None = None) -> dict:
    """Rows of a BOM, from every sheet with a header: [{item, part, name, desc, qty, material, finish, process,
    vendor, buy, profile, label, sheet, raw}]."""
    out, warnings, columns = [], [], set()
    for sheet, rows in read_sheets(name, data):
        hit = header(rows)
        if hit:
            columns |= set(hit[1])
            out += _rows(sheet, rows, hit, cfg, warnings)
    if not columns:
        raise BomError("No header row found. A BOM needs a Qty column and a Part number, Name or Description column.")
    if not out:
        raise BomError("The BOM has a header row but no parts under it.")
    return {"rows": out, "warnings": warnings, "columns": sorted(columns)}


def _rows(sheet: str, rows: list[list], hit: tuple, cfg: dict | None, warnings: list[str]) -> list[dict]:
    hi, cols = hit
    out = []
    for r in rows[hi + 1:]:
        if not any(str(x or "").strip() for x in r):
            continue
        part, nm, desc = _cell(r, cols, "part"), _cell(r, cols, "name"), _cell(r, cols, "desc")
        part = "" if TBD.match(part) else part
        if not (part or nm or desc):
            continue
        if re.match(r"^((sub|grand) ?)?totals?\b", f"{nm or desc or part}", re.I):
            continue
        q = _cell(r, cols, "total_qty") or _cell(r, cols, "qty")
        if not q and len(nm or desc) > 50:
            continue  # a note under the table, not a part
        try:
            qf = float(re.sub(r"[^\d.]", "", q)) if q else 1.0
        except ValueError:
            qf = 1.0
        if q and not re.sub(r"[^\d.]", "", q):
            warnings.append(f"Quantity '{q}' for '{part or nm or desc}' is not a number; used 1.")
        if qf <= 0:
            continue
        mat_raw, fin_raw, proc_raw = _cell(r, cols, "material"), _cell(r, cols, "finish"), _cell(r, cols, "process")
        e = {"item": _cell(r, cols, "item"), "part": part, "name": nm, "desc": desc, "qty": max(1, math.ceil(qf - 1e-9)),
             "material_raw": mat_raw, "material": map_material(mat_raw, cfg), "finish_raw": fin_raw, "finish": map_finish(fin_raw),
             "process_raw": proc_raw, "process": map_process(proc_raw) or map_process(_cell(r, cols, "buy")),
             "vendor": _cell(r, cols, "vendor"), "buy_raw": _cell(r, cols, "buy"), "notes": _cell(r, cols, "notes"),
             "profile": _cell(r, cols, "profile"), "sheet": sheet,
             "raw": " | ".join(str(x) for x in r if x not in (None, ""))[:240]}
        e["buy"] = _buy(e)
        e["label"] = (nm or desc or part)[:80]
        out.append(e)
    return out


# ---------------------------------------------------------------- matching to the model
def norm(s: str) -> str:
    s = str(s or "")
    s = re.sub(r"\.(step|stp|sldprt|sldasm|ipt|iam|prt|par|f3d|x_t)$", "", s, flags=re.I)
    s = re.sub(r"<\d+>|:\d+$|\s*\(\d+\)$", "", s)
    s = re.sub(r"[_\-./\\,\s]+", " ", s.lower()).strip()
    return s


def _forms(s: str) -> set[str]:
    n = norm(s)
    if not n:
        return set()
    n = " ".join(w[:-1] if len(w) > 4 and w.endswith("s") and not w.endswith("ss") else w for w in n.split())  # plurals
    out = {n}
    stripped = re.sub(r"^\d{1,3} ", "", n)  # "04 guide vane" -> "guide vane"
    if stripped and stripped != n:
        out.add(stripped)
    return out


def _score(e: dict, texts: list[str]) -> int:
    mine = set()
    for t in texts:
        mine |= _forms(t)
    if not mine:
        return 0
    best = 0
    words = lambda t: {w for w in re.findall(r"[a-z0-9]{4,}", t)}  # noqa: E731
    mine_w = set().union(*(words(m) for m in mine))
    for key, exact in (("part", 4), ("name", 4), ("desc", 3)):
        for f in _forms(e.get(key) or ""):
            if f in mine:
                best = max(best, exact)
            elif key == "part" and len(f) >= 4 and any(m.startswith(f + " ") for m in mine):
                best = max(best, 2)  # "VP-101" and a body named "VP-101 base plate"
            elif len(words(f)) >= 2 and len(words(f) & mine_w) >= max(2, round(len(words(f)) * 2 / 3)):
                best = max(best, 1)  # "Upright mounting plates" and "CCDP Mounting Plate"
            elif len(f) >= 5 and any(re.search(rf"(^| ){re.escape(f)}( |$)", m) or (len(m) >= 5 and re.search(rf"(^| ){re.escape(m)}( |$)", f))
                                     for m in mine):
                best = max(best, 1)
    return best


class Bom:
    """BOM rows shared across the models in one request; each row is used once."""

    def __init__(self, files: list[dict]):
        self.files = files  # [{"name", "rows", "warnings"}]
        self.rows = [dict(r, src=f["name"], idx=i) for f in files for i, r in enumerate(f["rows"])]
        self.used: set[int] = set()

    def __bool__(self) -> bool:
        return bool(self.rows)

    @property
    def names(self) -> str:
        return ", ".join(f["name"] for f in self.files)

    def left(self) -> list[dict]:
        return [r for n, r in enumerate(self.rows) if n not in self.used]

    def match_groups(self, groups: list[dict]) -> dict[int, dict]:
        """{group number: BOM row}. Best matches first; a row can cover several groups with the same name."""
        cands = []
        for gi, g in enumerate(groups):
            texts = list(g.get("raw_names") or []) + [g.get("name") or ""]
            for n, e in enumerate(self.rows):
                if n in self.used:
                    continue
                s = _score(e, texts)
                if s:
                    cands.append((-s, gi, n))
        cands.sort()
        out: dict[int, dict] = {}
        took: dict[int, int] = {}  # row -> score it was taken with
        for neg, gi, n in cands:
            if gi in out:
                continue
            if n in took and (took[n] != -neg or -neg < 3):
                continue
            out[gi] = self.rows[n]
            took[n] = -neg
        self.used |= set(took)
        counts: dict[int, int] = {}
        for e in out.values():
            counts[id(e)] = counts.get(id(e), 0) + 1
        for gi, e in out.items():
            groups[gi]["_bom_shared"] = counts[id(e)] > 1
        return out

    def match_name(self, *names: str) -> dict | None:
        best, at = 0, None
        for n, e in enumerate(self.rows):
            if n in self.used:
                continue
            s = _score(e, list(names))
            if s > best:
                best, at = s, n
        if at is None or best < 3:
            return None
        self.used.add(at)
        return self.rows[at]

    def skip_top(self, *names: str) -> None:
        """The assembly's own row ("Cart assembly, qty 1") isn't a part to price."""
        for n, e in enumerate(self.rows):
            if n not in self.used and _score(e, list(names)) >= 3:
                self.used.add(n)


FASTENER_ROW = re.compile(r"\b(fasteners?|screws?|bolts?|t-?nuts?|washers?|shcs|bhcs|bhscs|fhcs|fbhscs|hardware)\b", re.I)
PANEL_ROW = re.compile(r"\b(panels?|shelf|shelves|gaskets?|retainers?)\b", re.I)


def frame_hardware(bom: "Bom") -> tuple[dict, list[str], list[str]]:
    """BOM rows a T-slot frame quote covers, marked used. A row with an 80/20 part number sets the count of that
    hardware (the BOM is what gets ordered). Profiles, panels and other hardware are already in the frame quote
    from the model. Returns ({(base, 80/20 part number): qty}, covered row labels, notes)."""
    from collections import Counter

    from . import catalog_8020, tslot
    from .assembly import is_bought

    counts: Counter = Counter()
    covered, notes = [], []
    for n, e in enumerate(bom.rows):
        if n in bom.used:
            continue
        label = e["label"]
        hit = catalog_8020.match_part(e["part"]) if e["part"] else None
        if hit:
            counts[(hit[0].rsplit("-", 1)[0], hit[1])] += e["qty"]
            bom.used.add(n)
            continue
        if e.get("profile") or catalog_8020.profile(e["part"]) or PANEL_ROW.search(label) or FASTENER_ROW.search(label) \
                or is_bought(label) or tslot.hardware_for([label])[0] or e["buy"] is True:
            covered.append(label if not e["part"] else f"{label} ({e['part']})")
            bom.used.add(n)
    if counts:
        notes.append("Hardware counts from your BOM: " + ", ".join(f"{q} x {pn}" for (_b, pn), q in counts.items()))
    return dict(counts), covered, notes


def apply(g: dict, e: dict, defaults: dict) -> tuple[dict, list[dict], list[str]]:
    """BOM row -> (defaults for this part, facts, staff notes). Changes g in place (bought, qty, raw_names)."""
    from . import portal_read

    d = dict(defaults)
    facts, notes = [], []
    if e.get("buy") is True:
        g["bought"] = True
    elif e.get("buy") is False:
        g["bought"] = False
    names = [x for x in (e.get("part"), e.get("desc"), e.get("name")) if x]
    g["raw_names"] = names + [x for x in (g.get("raw_names") or []) if x not in names]
    mq = int(g.get("qty", 1) or 1)
    if e.get("qty") and e["qty"] != mq and not g.pop("_bom_shared", False):
        notes.append(f"{g.get('name')}: your BOM says {e['qty']} per assembly, the model has {mq}. We used the BOM.")
        g["qty"] = e["qty"]
    if d.get("material_source") != "chosen" and e.get("material"):
        d["material"], d["material_source"] = e["material"], "bom"
    elif e.get("material_raw") and not e.get("material") and not TBD.match(e["material_raw"]):
        notes.append(f"{g.get('name')}: BOM material '{e['material_raw']}' isn't in our list; pick the closest in review.")
    if d.get("finish_source") != "chosen" and e.get("finish"):
        d["finishes"], d["finish_source"] = ([] if e["finish"] == "none" else [e["finish"]]), "bom"
    elif e.get("finish_raw") and not e.get("finish") and not TBD.match(e["finish_raw"]):
        notes.append(f"{g.get('name')}: BOM finish '{e['finish_raw']}' needs a look.")
    if e.get("process"):
        d["process"] = e["process"]
    bits = [f"item {e['item']}" if e.get("item") else "", e.get("part") or "", f"qty {e['qty']}"]
    if e.get("vendor"):
        bits.append(e["vendor"])
    facts.append(portal_read.row("Your BOM", ", ".join(b for b in bits if b)))
    return d, facts, notes
