"""CMMC Level 1 and Level 2 (NIST SP 800-171) compliance tracker: control status, SPRS score, POA&M and export."""
import re
import tempfile
from datetime import date, datetime
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import cmmc_controls as cc
from .db import get_db
from .models_cmmc import CONTROL_STATUSES, CmmcControl, CmmcSettings

router = APIRouter(prefix="/api/cmmc")

DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
OPEN_STATUSES = {"not_started", "planned", "partial"}


# ------------------------------------------------------------------ helpers
def _ensure_rows(db: Session) -> dict[tuple[str, str], CmmcControl]:
    """Create a status row for every known control the first time it is needed."""
    rows = {(r.framework, r.control_id): r for r in db.scalars(select(CmmcControl))}
    added = False
    for key in cc.CONTROL_INDEX:
        if key not in rows:
            r = CmmcControl(framework=key[0], control_id=key[1], status="not_started")
            db.add(r)
            rows[key] = r
            added = True
    if added:
        db.commit()
    return rows


def _settings(db: Session) -> CmmcSettings:
    s = db.scalar(select(CmmcSettings).limit(1))
    if not s:
        s = CmmcSettings()
        db.add(s)
        db.commit()
    return s


def _check_date(v: str | None, field: str) -> str:
    v = (v or "").strip()
    if v and not DATE_RE.match(v):
        raise HTTPException(422, f"{field} must be YYYY-MM-DD")
    if v:
        try:
            date.fromisoformat(v)
        except ValueError:
            raise HTTPException(422, f"{field} is not a valid date")
    return v


def _out(meta: dict, r: CmmcControl) -> dict:
    d = dict(meta)
    d.update(
        id=r.id,
        status=r.status,
        evidence=r.evidence or "",
        owner=r.owner or "",
        poam_due=r.poam_due or "",
        notes=r.notes or "",
        updated_at=r.updated_at.isoformat() if r.updated_at else None,
    )
    if meta["framework"] == "L2":
        d["deduction"] = cc.deduction(meta["control_id"], r.status)
        d["partial_credit"] = cc.PARTIAL_CREDIT.get(meta["control_id"])
    return d


def compute_summary(db: Session) -> dict:
    """Score and progress numbers for the dashboard and the Compliance page."""
    rows = _ensure_rows(db)
    s = _settings(db)
    l1 = [rows[("L1", c["control_id"])] for c in cc.LEVEL1_CONTROLS]
    l2 = {c["control_id"]: rows[("L2", c["control_id"])].status for c in cc.LEVEL2_CONTROLS}
    score = cc.sprs_score(l2)

    open_items = [r for r in rows.values() if r.status in OPEN_STATUSES and r.poam_due]
    dues = sorted(r.poam_due for r in open_items)
    warnings = []
    if l2.get("3.12.4") not in ("implemented",):
        warnings.append("No system security plan (3.12.4) marked implemented. Without an SSP a Level 2 assessment cannot be completed.")
    blocked = [
        r.control_id for r in open_items
        if r.framework == "L2" and not cc.CONTROL_INDEX[("L2", r.control_id)]["poam_allowed"]
        and not (r.control_id == "3.13.11" and r.status == "partial")
    ]
    if blocked:
        warnings.append("These requirements cannot be closed out on a POA&M under 32 CFR 170.21 and must be met before a Level 2 assessment: " + ", ".join(blocked) + ".")
    l1_poam = [r.control_id for r in open_items if r.framework == "L1"]
    if l1_poam:
        warnings.append("Level 1 does not allow POA&Ms. Every Level 1 requirement must be met before you affirm in SPRS.")
    if score < cc.CONDITIONAL_MIN_SCORE:
        warnings.append(f"Score is below {cc.CONDITIONAL_MIN_SCORE} (80% of 110), the minimum for Conditional Level 2 status.")
    today = date.today().isoformat()
    return {
        "level1_done": sum(1 for r in l1 if r.status in ("implemented", "not_applicable")),
        "level1_total": cc.LEVEL1_TOTAL,
        "level2_done": sum(1 for v in l2.values() if v in ("implemented", "not_applicable")),
        "level2_total": len(cc.LEVEL2_CONTROLS),
        "sprs_score": score,
        "sprs_max": cc.SPRS_MAX,
        "sprs_min": cc.SPRS_MAX - sum(cc.WEIGHTS.values()),
        "conditional_min_score": cc.CONDITIONAL_MIN_SCORE,
        "open_poam": len(open_items),
        "overdue_poam": sum(1 for d in dues if d < today),
        "next_poam_due": dues[0] if dues else None,
        "last_sprs_submission_date": s.last_sprs_submission_date or None,
        "affirmation_date": s.affirmation_date or None,
        "warnings": warnings,
    }


# ------------------------------------------------------------------ endpoints
@router.get("/meta")
def meta():
    return {
        "statuses": CONTROL_STATUSES,
        "families": [{"code": k, "number": v[0], "name": v[1]} for k, v in cc.FAMILIES.items()],
        "partial_credit": cc.PARTIAL_CREDIT,
        "no_poam": sorted(cc.NO_POAM),
        "sprs_max": cc.SPRS_MAX,
        "conditional_min_score": cc.CONDITIONAL_MIN_SCORE,
    }


@router.get("/controls")
def list_controls(framework: str = "", family: str = "", status: str = "", db: Session = Depends(get_db)):
    rows = _ensure_rows(db)
    out = []
    for c in cc.ALL_CONTROLS:
        if framework and c["framework"] != framework.upper():
            continue
        if family and c["family"] != family.upper():
            continue
        r = rows[(c["framework"], c["control_id"])]
        if status and r.status not in status.split(","):
            continue
        out.append(_out(c, r))
    return out


class ControlUpdate(BaseModel):
    status: str | None = None
    evidence: str | None = None
    owner: str | None = None
    poam_due: str | None = None
    notes: str | None = None


@router.put("/controls/{framework}/{control_id}")
def update_control(framework: str, control_id: str, body: ControlUpdate, db: Session = Depends(get_db)):
    key = (framework.upper(), control_id)
    if key not in cc.CONTROL_INDEX:
        raise HTTPException(404, "Unknown control")
    rows = _ensure_rows(db)
    r = rows[key]
    data = body.model_dump(exclude_unset=True)
    if "status" in data:
        if data["status"] not in CONTROL_STATUSES:
            raise HTTPException(422, f"status must be one of {', '.join(CONTROL_STATUSES)}")
        r.status = data["status"]
    if "poam_due" in data:
        r.poam_due = _check_date(data["poam_due"], "poam_due")
    for f in ("evidence", "owner", "notes"):
        if f in data:
            setattr(r, f, data[f] or "")
    r.updated_at = datetime.utcnow()
    db.commit()
    return _out(cc.CONTROL_INDEX[key], r)


@router.get("/summary")
def summary(db: Session = Depends(get_db)):
    return compute_summary(db)


class SettingsUpdate(BaseModel):
    last_sprs_submission_date: str | None = None
    affirmation_date: str | None = None
    scope_notes: str | None = None


def _settings_out(s: CmmcSettings) -> dict:
    return {
        "last_sprs_submission_date": s.last_sprs_submission_date or "",
        "affirmation_date": s.affirmation_date or "",
        "scope_notes": s.scope_notes or "",
    }


@router.get("/settings")
def get_settings(db: Session = Depends(get_db)):
    return _settings_out(_settings(db))


@router.put("/settings")
def put_settings(body: SettingsUpdate, db: Session = Depends(get_db)):
    s = _settings(db)
    data = body.model_dump(exclude_unset=True)
    for f in ("last_sprs_submission_date", "affirmation_date"):
        if f in data:
            setattr(s, f, _check_date(data[f], f))
    if "scope_notes" in data:
        s.scope_notes = data["scope_notes"] or ""
    db.commit()
    return _settings_out(s)


@router.get("/export.xlsx")
def export_xlsx(db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    rows = _ensure_rows(db)
    summ = compute_summary(db)
    s = _settings(db)
    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDE4EE")

    def header(ws, cols, widths):
        ws.append(cols)
        for i, c in enumerate(ws[ws.max_row], start=0):
            c.font = bold
            c.fill = head_fill
            ws.column_dimensions[c.column_letter].width = widths[i]
        ws.freeze_panes = f"A{ws.max_row + 1}"

    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 80
    for k, v in [
        ("Generated", date.today().isoformat()),
        ("SPRS score (NIST SP 800-171 DoD Assessment Methodology v1.2.1)", f"{summ['sprs_score']} of {summ['sprs_max']}"),
        ("Level 2 requirements met", f"{summ['level2_done']} of {summ['level2_total']}"),
        ("Level 1 requirements met (FAR 52.204-21)", f"{summ['level1_done']} of {summ['level1_total']}"),
        ("Open POA&M items", summ["open_poam"]),
        ("Next POA&M due", summ["next_poam_due"] or ""),
        ("Last SPRS submission", s.last_sprs_submission_date or ""),
        ("Affirmation date", s.affirmation_date or ""),
        ("Scope (systems holding FCI/CUI)", s.scope_notes or ""),
    ]:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = bold
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    for w in summ["warnings"]:
        ws.append(["Warning", w])
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True)

    ws = wb.create_sheet("Status")
    header(ws, ["Framework", "Control", "CMMC ID", "Family", "Requirement", "Points", "Status", "Deduction", "Evidence", "Owner", "POA&M due", "Notes"],
           [10, 14, 18, 30, 70, 8, 15, 10, 50, 16, 12, 40])
    for c in cc.ALL_CONTROLS:
        r = rows[(c["framework"], c["control_id"])]
        ws.append([
            c["framework"], c["control_id"], c.get("cmmc_id") or c["control_id"], c["family_name"], c["text"],
            c["weight"] if c["weight"] is not None else "", r.status,
            cc.deduction(c["control_id"], r.status) if c["framework"] == "L2" else "",
            r.evidence or "", r.owner or "", r.poam_due or "", r.notes or "",
        ])
        for col in (5, 9, 12):
            ws.cell(row=ws.max_row, column=col).alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("POA&M")
    header(ws, ["Control", "Requirement", "Points", "Status", "Deduction now", "POA&M allowed", "Planned action / notes", "Owner", "Due", "Evidence so far"],
           [14, 70, 8, 12, 12, 14, 50, 16, 12, 40])
    open_l2 = [c for c in cc.LEVEL2_CONTROLS if rows[("L2", c["control_id"])].status in OPEN_STATUSES]
    open_l2.sort(key=lambda c: (rows[("L2", c["control_id"])].poam_due or "9999", -(c["weight"] or 0)))
    for c in open_l2:
        r = rows[("L2", c["control_id"])]
        allowed = c["poam_allowed"] or (c["control_id"] == "3.13.11" and r.status == "partial")
        ws.append([c["control_id"], c["text"], c["weight"], r.status, cc.deduction(c["control_id"], r.status),
                   "yes" if allowed else "no", r.notes or "", r.owner or "", r.poam_due or "", r.evidence or ""])
        for col in (2, 7, 10):
            ws.cell(row=ws.max_row, column=col).alignment = Alignment(wrap_text=True, vertical="top")

    out = Path(tempfile.mkdtemp()) / f"cmmc-status-{date.today().isoformat()}.xlsx"
    wb.save(out)
    return FileResponse(out, filename=out.name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
