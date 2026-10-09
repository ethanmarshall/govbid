"""Lightweight quality system for a one-person shop: nonconformance reports (NCR), corrective actions (CAR),
calibration log, editable quality documents (DOCX download) and supplier approval with scorecards
computed from job purchase orders. Regulatory sources for the starter documents are listed in
app/quality_templates.py.
"""
from __future__ import annotations

import re
import shutil
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel, ConfigDict
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from . import config
from .db import get_db
from .models_crm import Organization
from .models_jobs import Job, JobPurchase
from .models_quality import (APPROVAL_BASES, APPROVAL_STATUSES, CAR_STATUSES, INSTRUMENT_STATUSES, NCR_DISPOSITIONS, NCR_SOURCES,
                             NCR_STATUSES, SUPPLIER_CERT_TYPES, CorrectiveAction, Instrument, NonconformanceReport, QualityDocument,
                             SupplierApproval, SupplierCert)
from .quality_templates import COMPANY, STARTER_DOCUMENTS

router = APIRouter(prefix="/api/quality", tags=["quality"])

INSTRUMENT_KINDS = ["Calipers", "Micrometer", "Height gauge", "Indicator", "Gauge blocks", "Pin gauges", "Thread gauges",
                    "Torque wrench", "Multimeter", "Megohmmeter", "Hipot tester", "Crimp tool", "Pull tester", "Scale", "Other"]
DUE_SOON_DAYS = 30
CERT_SOON_DAYS = 60


# ------------------------------------------------------------------ schemas
class NcrIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    job_id: int | None = None
    purchase_id: int | None = None
    source: str | None = None
    part: str | None = None
    quantity: float | None = None
    description: str | None = None
    containment: str | None = None
    disposition: str | None = None
    disposition_by: str | None = None
    root_cause: str | None = None
    status: str | None = None
    opened_date: str | None = None
    closed_date: str | None = None
    notes: str | None = None


class CarIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    ncr_id: int | None = None
    problem: str | None = None
    root_cause: str | None = None
    action: str | None = None
    owner: str | None = None
    opened_date: str | None = None
    due_date: str | None = None
    effectiveness_check: str | None = None
    effectiveness_date: str | None = None
    effective: bool | None = None
    status: str | None = None
    closed_date: str | None = None


class InstrumentIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    name: str | None = None
    kind: str | None = None
    asset_id: str | None = None
    serial: str | None = None
    manufacturer: str | None = None
    range_resolution: str | None = None
    location: str | None = None
    interval_days: int | None = None
    last_cal_date: str | None = None
    cal_source: str | None = None
    status: str | None = None
    notes: str | None = None


class DocIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    title: str | None = None
    doc_number: str | None = None
    content: str | None = None
    effective_date: str | None = None
    version: str | None = None
    new_revision: bool = False
    change_note: str = ""


class ApprovalIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    status: str | None = None
    approval_date: str | None = None
    basis: list[str] | None = None
    scope: str | None = None
    review_due: str | None = None
    notes: str | None = None


# ------------------------------------------------------------------ helpers
def _check(value: str | None, allowed: list[str], field: str) -> None:
    if value is not None and value not in allowed:
        raise HTTPException(422, f"{field} must be one of: {', '.join(a or '(blank)' for a in allowed)}")


def _check_date(value: str | None, field: str) -> None:
    if value:
        try:
            date.fromisoformat(value)
        except ValueError:
            raise HTTPException(422, f"{field} must be YYYY-MM-DD")


def _safe_name(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", Path(name or "file").name)[:120] or "file"


def _qdir(*parts) -> Path:
    d = Path(config.UPLOAD_DIR) / "quality"
    for p in parts:
        d = d / str(p)
    d.mkdir(parents=True, exist_ok=True)
    return d


def next_number(db: Session, model, prefix: str, year: int | None = None) -> str:
    """Next sequential number for the year: PREFIX-YYYY-NNN (NCR-2026-001, NCR-2026-002, ...)."""
    year = year or date.today().year
    stem = f"{prefix}-{year}-"
    seqs = []
    for n in db.scalars(select(model.number).where(model.number.like(stem + "%"))).all():
        m = re.fullmatch(re.escape(stem) + r"(\d+)", n or "")
        if m:
            seqs.append(int(m.group(1)))
    return f"{stem}{(max(seqs) + 1) if seqs else 1:03d}"


def next_due(last_cal_date: str, interval_days: int | None) -> str:
    """Next calibration due date: last calibration date plus the interval (blank if either is missing)."""
    if not last_cal_date or not interval_days:
        return ""
    return (date.fromisoformat(last_cal_date) + timedelta(days=int(interval_days))).isoformat()


def cal_status(inst: Instrument, today: date | None = None) -> str:
    """overdue | due_soon (within 30 days) | ok | no_record | out_of_service | retired."""
    today = today or date.today()
    if inst.status != "active":
        return inst.status
    due = next_due(inst.last_cal_date, inst.interval_days)
    if not due:
        return "no_record"
    if due < today.isoformat():
        return "overdue"
    if due <= (today + timedelta(days=DUE_SOON_DAYS)).isoformat():
        return "due_soon"
    return "ok"


def _org_purchases(db: Session, org: Organization) -> list[JobPurchase]:
    return list(db.scalars(select(JobPurchase).where(or_(
        JobPurchase.organization_id == org.id,
        (JobPurchase.organization_id.is_(None)) & (func.lower(JobPurchase.vendor_name) == (org.name or "").lower()),
    ))).all())


def scorecard(purchases: list[JobPurchase], today: date | None = None) -> dict:
    """Supplier performance from purchase orders.

    on_time_pct: received orders with a promised date where received_date <= promised_date, over those orders.
    acceptance_pct: quantity accepted / (accepted + rejected) over received orders.
    """
    today_s = (today or date.today()).isoformat()
    received = [p for p in purchases if p.received_date]
    timed = [p for p in received if p.promised_date]
    on_time = sum(1 for p in timed if p.received_date <= p.promised_date)
    acc = sum((p.qty_accepted or 0) for p in received)
    rej = sum((p.qty_rejected or 0) for p in received)
    dates = [p.ordered_date for p in purchases if p.ordered_date]
    return {
        "orders": len(purchases),
        "received": len(received),
        "open": len(purchases) - len(received),
        "late_open": sum(1 for p in purchases if not p.received_date and p.promised_date and p.promised_date < today_s),
        "on_time": on_time,
        "on_time_measured": len(timed),
        "on_time_pct": round(100.0 * on_time / len(timed), 1) if timed else None,
        "qty_accepted": acc,
        "qty_rejected": rej,
        "acceptance_pct": round(100.0 * acc / (acc + rej), 1) if (acc + rej) else None,
        "certs_missing": sum(1 for p in received if not p.certs_received),
        "spend": round(sum((p.quantity or 0) * (p.unit_price or 0) for p in purchases), 2),
        "last_order": max(dates) if dates else "",
    }


def _cert_state(c: SupplierCert, today: date | None = None) -> str:
    today = today or date.today()
    if not c.expiration_date:
        return "no_expiration"
    if c.expiration_date < today.isoformat():
        return "expired"
    if c.expiration_date <= (today + timedelta(days=CERT_SOON_DAYS)).isoformat():
        return "expiring_soon"
    return "current"


# ------------------------------------------------------------------ dicts
def ncr_dict(n: NonconformanceReport, db: Session) -> dict:
    job = db.get(Job, n.job_id) if n.job_id else None
    cars = db.scalars(select(CorrectiveAction).where(CorrectiveAction.ncr_id == n.id)).all()
    return {"id": n.id, "number": n.number, "job_id": n.job_id, "job_title": job.title if job else None, "purchase_id": n.purchase_id,
            "source": n.source, "part": n.part, "quantity": n.quantity, "description": n.description, "containment": n.containment,
            "disposition": n.disposition, "disposition_by": n.disposition_by, "root_cause": n.root_cause, "status": n.status,
            "opened_date": n.opened_date, "closed_date": n.closed_date, "notes": n.notes,
            "cars": [{"id": c.id, "number": c.number, "status": c.status} for c in cars]}


def car_dict(c: CorrectiveAction, db: Session) -> dict:
    ncr = db.get(NonconformanceReport, c.ncr_id) if c.ncr_id else None
    today = date.today().isoformat()
    return {"id": c.id, "number": c.number, "ncr_id": c.ncr_id, "ncr_number": ncr.number if ncr else None, "problem": c.problem,
            "root_cause": c.root_cause, "action": c.action, "owner": c.owner, "opened_date": c.opened_date, "due_date": c.due_date,
            "effectiveness_check": c.effectiveness_check, "effectiveness_date": c.effectiveness_date, "effective": c.effective,
            "status": c.status, "closed_date": c.closed_date,
            "overdue": bool(c.status != "closed" and c.due_date and c.due_date < today)}


def instrument_dict(i: Instrument) -> dict:
    hist = []
    for idx, h in enumerate(i.history or []):
        h = dict(h)
        h["index"] = idx
        h["url"] = f"/api/quality/instruments/{i.id}/calibrations/{idx}/file" if h.get("stored_name") else None
        h.pop("stored_name", None)
        hist.append(h)
    return {"id": i.id, "name": i.name, "kind": i.kind, "asset_id": i.asset_id, "serial": i.serial, "manufacturer": i.manufacturer,
            "range_resolution": i.range_resolution, "location": i.location, "interval_days": i.interval_days,
            "last_cal_date": i.last_cal_date, "cal_source": i.cal_source, "status": i.status, "notes": i.notes,
            "next_due": next_due(i.last_cal_date, i.interval_days), "cal_status": cal_status(i),
            "overdue": cal_status(i) == "overdue", "history": list(reversed(hist))}


def doc_dict(d: QualityDocument, full: bool = True) -> dict:
    out = {"id": d.id, "key": d.key, "doc_number": d.doc_number, "title": d.title, "version": d.version,
           "effective_date": d.effective_date, "revisions": len(d.history or []),
           "updated_at": d.updated_at.isoformat() if d.updated_at else None}
    if full:
        out["content"] = d.content
        out["history"] = [{"index": i, **h} for i, h in enumerate(d.history or [])]
    return out


def cert_dict(c: SupplierCert) -> dict:
    return {"id": c.id, "organization_id": c.organization_id, "cert_type": c.cert_type, "number": c.number, "issuer": c.issuer,
            "expiration_date": c.expiration_date, "filename": c.filename, "notes": c.notes, "state": _cert_state(c),
            "url": f"/api/quality/suppliers/{c.organization_id}/certs/{c.id}/file" if c.stored_name else None}


def supplier_dict(db: Session, org: Organization) -> dict:
    a = db.scalars(select(SupplierApproval).where(SupplierApproval.organization_id == org.id)).first()
    certs = db.scalars(select(SupplierCert).where(SupplierCert.organization_id == org.id).order_by(SupplierCert.expiration_date)).all()
    return {
        "organization_id": org.id, "name": org.name, "kind": org.kind, "cage": org.cage, "uei": org.uei, "city": org.city,
        "state": org.state, "is_manufacturer": org.is_manufacturer, "capabilities": org.capabilities,
        "approval": {"status": a.status if a else "pending", "approval_date": a.approval_date if a else "", "basis": (a.basis or []) if a else [],
                     "scope": a.scope if a else "", "review_due": a.review_due if a else "", "notes": a.notes if a else ""},
        "certs": [cert_dict(c) for c in certs],
        "expired_certs": sum(1 for c in certs if _cert_state(c) == "expired"),
        "scorecard": scorecard(_org_purchases(db, org)),
    }


def _supplier_orgs(db: Session) -> list[Organization]:
    ids = set(db.scalars(select(Organization.id).where(Organization.kind == "vendor")).all())
    ids |= set(db.scalars(select(SupplierApproval.organization_id)).all())
    ids |= {i for i in db.scalars(select(JobPurchase.organization_id)).all() if i}
    if not ids:
        return []
    return list(db.scalars(select(Organization).where(Organization.id.in_(ids)).order_by(Organization.name)).all())


def _get_org(db: Session, org_id: int) -> Organization:
    org = db.get(Organization, org_id)
    if not org:
        raise HTTPException(404, "Organization not found")
    return org


# ------------------------------------------------------------------ calendar and dashboard hooks
def calendar_items(db: Session) -> list[dict]:
    """Calibration due dates, open CAR due dates and supplier certificate expirations."""
    out = []
    for i in db.scalars(select(Instrument).where(Instrument.status == "active")).all():
        due = next_due(i.last_cal_date, i.interval_days)
        if due:
            out.append({"uid": f"cal-{i.id}-{due}@govbid", "date": date.fromisoformat(due),
                        "summary": f"Calibration due: {i.name or i.kind} {i.asset_id}".strip(),
                        "description": f"Serial {i.serial or 'n/a'}. Last calibrated {i.last_cal_date} by {i.cal_source or 'unknown'}.",
                        "url": "/quality?tab=calibration"})
    for c in db.scalars(select(CorrectiveAction).where(CorrectiveAction.status != "closed", CorrectiveAction.due_date != "")).all():
        out.append({"uid": f"car-{c.id}@govbid", "date": date.fromisoformat(c.due_date), "summary": f"{c.number} due",
                    "description": (c.problem or "")[:300], "url": "/quality?tab=car"})
    rows = db.execute(select(SupplierCert, Organization.name).join(Organization, Organization.id == SupplierCert.organization_id)
                      .where(SupplierCert.expiration_date != "")).all()
    for c, name in rows:
        out.append({"uid": f"supcert-{c.id}@govbid", "date": date.fromisoformat(c.expiration_date),
                    "summary": f"{name}: {c.cert_type} expires", "description": f"Certificate {c.number}".strip(),
                    "url": "/quality?tab=suppliers"})
    return out


def dashboard_items(db: Session) -> list[str]:
    """Overdue calibration, overdue corrective actions and expired certifications for suppliers still in use."""
    out = []
    today = date.today()
    for i in db.scalars(select(Instrument).where(Instrument.status == "active")).all():
        if cal_status(i, today) == "overdue":
            out.append(f"Calibration overdue: {i.name or i.kind} {i.asset_id} was due {next_due(i.last_cal_date, i.interval_days)}.".replace("  ", " "))
    for c in db.scalars(select(CorrectiveAction).where(CorrectiveAction.status != "closed", CorrectiveAction.due_date != "")).all():
        if c.due_date < today.isoformat():
            out.append(f"Corrective action overdue: {c.number} was due {c.due_date}.")
    disapproved = set(db.scalars(select(SupplierApproval.organization_id).where(SupplierApproval.status == "disapproved")).all())
    rows = db.execute(select(SupplierCert, Organization.name).join(Organization, Organization.id == SupplierCert.organization_id)
                      .where(SupplierCert.expiration_date != "", SupplierCert.expiration_date < today.isoformat())).all()
    for c, name in rows:
        if c.organization_id not in disapproved:
            out.append(f"Supplier cert expired: {name} {c.cert_type} expired {c.expiration_date}.")
    return out


# ------------------------------------------------------------------ meta and summary
@router.get("/meta")
def meta():
    return {"ncr_dispositions": NCR_DISPOSITIONS, "ncr_statuses": NCR_STATUSES, "ncr_sources": NCR_SOURCES, "car_statuses": CAR_STATUSES,
            "instrument_statuses": INSTRUMENT_STATUSES, "instrument_kinds": INSTRUMENT_KINDS, "approval_statuses": APPROVAL_STATUSES,
            "approval_bases": APPROVAL_BASES, "supplier_cert_types": SUPPLIER_CERT_TYPES}


@router.get("/summary")
def summary(db: Session = Depends(get_db)):
    insts = db.scalars(select(Instrument)).all()
    cars = [car_dict(c, db) for c in db.scalars(select(CorrectiveAction)).all()]
    statuses = [s["approval"]["status"] for s in (supplier_dict(db, o) for o in _supplier_orgs(db))]
    return {
        "open_ncrs": db.scalar(select(func.count()).select_from(NonconformanceReport).where(NonconformanceReport.status != "closed")) or 0,
        "open_cars": sum(1 for c in cars if c["status"] != "closed"),
        "overdue_cars": sum(1 for c in cars if c["overdue"]),
        "cal_overdue": sum(1 for i in insts if cal_status(i) == "overdue"),
        "cal_due_soon": sum(1 for i in insts if cal_status(i) == "due_soon"),
        "suppliers": len(statuses),
        "suppliers_approved": sum(1 for s in statuses if s in ("approved", "conditional")),
        "alerts": dashboard_items(db),
    }


# ------------------------------------------------------------------ NCRs
def _apply_ncr(db: Session, n: NonconformanceReport, data: dict) -> None:
    _check(data.get("disposition"), NCR_DISPOSITIONS, "disposition")
    _check(data.get("status"), NCR_STATUSES, "status")
    _check(data.get("source"), NCR_SOURCES, "source")
    for f in ("opened_date", "closed_date"):
        _check_date(data.get(f), f)
    if data.get("job_id") and not db.get(Job, data["job_id"]):
        raise HTTPException(404, "Job not found")
    if data.get("purchase_id"):
        p = db.get(JobPurchase, data["purchase_id"])
        if not p:
            raise HTTPException(404, "Purchase order not found")
        data.setdefault("job_id", p.job_id)
    for k, v in data.items():
        if v is not None or k in ("job_id", "purchase_id", "quantity"):
            setattr(n, k, v)
    if n.status == "closed" and not n.closed_date:
        n.closed_date = date.today().isoformat()
    if n.status != "closed" and "status" in data:
        n.closed_date = data.get("closed_date") or ""
    if n.disposition and n.status == "open" and "status" not in data:
        n.status = "dispositioned"


@router.get("/ncrs")
def list_ncrs(status: str = "", job_id: int | None = None, db: Session = Depends(get_db)):
    stmt = select(NonconformanceReport).order_by(NonconformanceReport.id.desc())
    if status:
        stmt = stmt.where(NonconformanceReport.status == status)
    if job_id:
        stmt = stmt.where(NonconformanceReport.job_id == job_id)
    return [ncr_dict(n, db) for n in db.scalars(stmt).all()]


@router.post("/ncrs")
def create_ncr(body: NcrIn, db: Session = Depends(get_db)):
    data = body.model_dump(exclude_unset=True)
    opened = data.get("opened_date") or date.today().isoformat()
    _check_date(opened, "opened_date")
    n = NonconformanceReport(number=next_number(db, NonconformanceReport, "NCR", int(opened[:4])), opened_date=opened, status="open")
    _apply_ncr(db, n, data)
    db.add(n)
    db.commit()
    return ncr_dict(n, db)


@router.get("/ncrs/{ncr_id}")
def get_ncr(ncr_id: int, db: Session = Depends(get_db)):
    n = db.get(NonconformanceReport, ncr_id)
    if not n:
        raise HTTPException(404, "NCR not found")
    return ncr_dict(n, db)


@router.put("/ncrs/{ncr_id}")
def update_ncr(ncr_id: int, body: NcrIn, db: Session = Depends(get_db)):
    n = db.get(NonconformanceReport, ncr_id)
    if not n:
        raise HTTPException(404, "NCR not found")
    _apply_ncr(db, n, body.model_dump(exclude_unset=True))
    db.commit()
    return ncr_dict(n, db)


@router.delete("/ncrs/{ncr_id}")
def delete_ncr(ncr_id: int, db: Session = Depends(get_db)):
    n = db.get(NonconformanceReport, ncr_id)
    if not n:
        raise HTTPException(404, "NCR not found")
    db.query(CorrectiveAction).filter(CorrectiveAction.ncr_id == ncr_id).update({"ncr_id": None})
    db.delete(n)
    db.commit()
    return {"ok": True}


# ------------------------------------------------------------------ CARs
def _apply_car(db: Session, c: CorrectiveAction, data: dict) -> None:
    _check(data.get("status"), CAR_STATUSES, "status")
    for f in ("opened_date", "due_date", "effectiveness_date", "closed_date"):
        _check_date(data.get(f), f)
    if data.get("ncr_id") and not db.get(NonconformanceReport, data["ncr_id"]):
        raise HTTPException(404, "NCR not found")
    for k, v in data.items():
        if v is not None or k in ("ncr_id", "effective"):
            setattr(c, k, v)
    if c.status == "closed":
        if c.effective is False:
            raise HTTPException(422, "The effectiveness check failed. Revise the action instead of closing this CAR.")
        c.closed_date = c.closed_date or date.today().isoformat()
    elif "status" in data:
        c.closed_date = data.get("closed_date") or ""


@router.get("/cars")
def list_cars(status: str = "", db: Session = Depends(get_db)):
    stmt = select(CorrectiveAction).order_by(CorrectiveAction.id.desc())
    if status:
        stmt = stmt.where(CorrectiveAction.status == status)
    return [car_dict(c, db) for c in db.scalars(stmt).all()]


@router.post("/cars")
def create_car(body: CarIn, db: Session = Depends(get_db)):
    data = body.model_dump(exclude_unset=True)
    opened = data.get("opened_date") or date.today().isoformat()
    _check_date(opened, "opened_date")
    c = CorrectiveAction(number=next_number(db, CorrectiveAction, "CAR", int(opened[:4])), opened_date=opened, status="open")
    if data.get("ncr_id") and not data.get("problem"):
        ncr = db.get(NonconformanceReport, data["ncr_id"])
        if ncr:
            data["problem"] = f"{ncr.number}: {ncr.description}".strip()
            data.setdefault("root_cause", ncr.root_cause)
    _apply_car(db, c, data)
    db.add(c)
    db.commit()
    return car_dict(c, db)


@router.put("/cars/{car_id}")
def update_car(car_id: int, body: CarIn, db: Session = Depends(get_db)):
    c = db.get(CorrectiveAction, car_id)
    if not c:
        raise HTTPException(404, "Corrective action not found")
    _apply_car(db, c, body.model_dump(exclude_unset=True))
    db.commit()
    return car_dict(c, db)


@router.delete("/cars/{car_id}")
def delete_car(car_id: int, db: Session = Depends(get_db)):
    c = db.get(CorrectiveAction, car_id)
    if not c:
        raise HTTPException(404, "Corrective action not found")
    db.delete(c)
    db.commit()
    return {"ok": True}


# ------------------------------------------------------------------ calibration
def _apply_inst(i: Instrument, data: dict) -> None:
    _check(data.get("status"), INSTRUMENT_STATUSES, "status")
    _check_date(data.get("last_cal_date"), "last_cal_date")
    if data.get("interval_days") is not None and data["interval_days"] <= 0:
        raise HTTPException(422, "interval_days must be above zero")
    for k, v in data.items():
        if v is not None:
            setattr(i, k, v)


@router.get("/instruments")
def list_instruments(db: Session = Depends(get_db)):
    rows = [instrument_dict(i) for i in db.scalars(select(Instrument)).all()]
    order = {"overdue": 0, "due_soon": 1, "no_record": 2, "ok": 3, "out_of_service": 4, "retired": 5}
    return sorted(rows, key=lambda r: (order.get(r["cal_status"], 9), r["next_due"] or "9999", r["name"]))


@router.post("/instruments")
def create_instrument(body: InstrumentIn, db: Session = Depends(get_db)):
    data = body.model_dump(exclude_unset=True)
    i = Instrument(name=data.get("name") or data.get("kind") or "Instrument", interval_days=365, status="active", history=[])
    _apply_inst(i, data)
    db.add(i)
    db.commit()
    return instrument_dict(i)


@router.put("/instruments/{inst_id}")
def update_instrument(inst_id: int, body: InstrumentIn, db: Session = Depends(get_db)):
    i = db.get(Instrument, inst_id)
    if not i:
        raise HTTPException(404, "Instrument not found")
    _apply_inst(i, body.model_dump(exclude_unset=True))
    db.commit()
    return instrument_dict(i)


@router.delete("/instruments/{inst_id}")
def delete_instrument(inst_id: int, db: Session = Depends(get_db)):
    i = db.get(Instrument, inst_id)
    if not i:
        raise HTTPException(404, "Instrument not found")
    db.delete(i)
    db.commit()
    shutil.rmtree(Path(config.UPLOAD_DIR) / "quality" / "cal" / str(inst_id), ignore_errors=True)
    return {"ok": True}


@router.post("/instruments/{inst_id}/calibrations")
async def record_calibration(inst_id: int, cal_date: str = Form(...), source: str = Form(""), result: str = Form("pass"),
                             notes: str = Form(""), file: UploadFile | None = File(None), db: Session = Depends(get_db)):
    """Log a calibration (optionally with its certificate) and move the instrument's last calibration date forward."""
    i = db.get(Instrument, inst_id)
    if not i:
        raise HTTPException(404, "Instrument not found")
    _check_date(cal_date, "cal_date")
    _check(result, ["pass", "adjusted", "fail"], "result")
    entry: dict[str, Any] = {"date": cal_date, "source": source, "result": result, "notes": notes, "filename": "", "stored_name": ""}
    if file is not None and file.filename:
        name = _safe_name(file.filename)
        stored = f"{datetime.utcnow():%Y%m%d%H%M%S%f}_{name}"
        (_qdir("cal", inst_id) / stored).write_bytes(await file.read())
        entry["filename"], entry["stored_name"] = name, stored
    i.history = list(i.history or []) + [entry]
    if not i.last_cal_date or cal_date >= i.last_cal_date:
        i.last_cal_date = cal_date
        if source:
            i.cal_source = source
    if result == "fail":
        i.status = "out_of_service"
    db.commit()
    return instrument_dict(i)


@router.get("/instruments/{inst_id}/calibrations/{idx}/file")
def calibration_file(inst_id: int, idx: int, db: Session = Depends(get_db)):
    i = db.get(Instrument, inst_id)
    hist = (i.history or []) if i else []
    if not i or idx < 0 or idx >= len(hist) or not hist[idx].get("stored_name"):
        raise HTTPException(404, "Certificate not found")
    path = _qdir("cal", inst_id) / hist[idx]["stored_name"]
    if not path.exists():
        raise HTTPException(404, "File is missing on disk")
    return FileResponse(path, filename=hist[idx]["filename"])


# ------------------------------------------------------------------ quality documents
def ensure_starter_documents(db: Session) -> int:
    """Add any starter template that is not in the database yet. Returns how many were added."""
    have = set(db.scalars(select(QualityDocument.key)).all())
    added = 0
    for t in STARTER_DOCUMENTS:
        if t["key"] not in have:
            db.add(QualityDocument(key=t["key"], doc_number=t["doc_number"], title=t["title"], content=t["content"].strip() + "\n",
                                   version="A", effective_date="", history=[]))
            added += 1
    if added:
        db.commit()
    return added


def next_revision(v: str) -> str:
    """A -> B, Z -> AA, 1 -> 2, 1.0 -> 1.1, anything else gets '.1' appended."""
    v = (v or "").strip()
    if re.fullmatch(r"[A-Z]+", v):
        chars = list(v)
        i = len(chars) - 1
        while i >= 0:
            if chars[i] != "Z":
                chars[i] = chr(ord(chars[i]) + 1)
                return "".join(chars)
            chars[i] = "A"
            i -= 1
        return "A" + "".join(chars)
    m = re.fullmatch(r"(.*?)(\d+)", v)
    if m:
        return f"{m.group(1)}{int(m.group(2)) + 1}"
    return (v + ".1") if v else "A"


@router.get("/documents")
def list_documents(db: Session = Depends(get_db)):
    ensure_starter_documents(db)
    return [doc_dict(d, full=False) for d in db.scalars(select(QualityDocument).order_by(QualityDocument.doc_number, QualityDocument.id)).all()]


@router.post("/documents/seed")
def seed_documents(db: Session = Depends(get_db)):
    return {"added": ensure_starter_documents(db)}


@router.post("/documents")
def create_document(body: DocIn, db: Session = Depends(get_db)):
    _check_date(body.effective_date, "effective_date")
    d = QualityDocument(key="", title=body.title or "New document", doc_number=body.doc_number or "", content=body.content or "",
                        version=body.version or "A", effective_date=body.effective_date or "", history=[])
    db.add(d)
    db.commit()
    return doc_dict(d)


@router.get("/documents/{doc_id}")
def get_document(doc_id: int, db: Session = Depends(get_db)):
    d = db.get(QualityDocument, doc_id)
    if not d:
        raise HTTPException(404, "Document not found")
    return doc_dict(d)


@router.put("/documents/{doc_id}")
def update_document(doc_id: int, body: DocIn, db: Session = Depends(get_db)):
    """Edit a document. new_revision=true files the current text in the history and bumps the version."""
    d = db.get(QualityDocument, doc_id)
    if not d:
        raise HTTPException(404, "Document not found")
    _check_date(body.effective_date, "effective_date")
    if body.new_revision:
        d.history = list(d.history or []) + [{"version": d.version, "effective_date": d.effective_date, "title": d.title,
                                              "content": d.content, "note": body.change_note, "replaced_at": datetime.utcnow().isoformat(timespec="seconds")}]
        d.version = body.version or next_revision(d.version)
        d.effective_date = body.effective_date or date.today().isoformat()
    else:
        if body.version is not None:
            d.version = body.version
        if body.effective_date is not None:
            d.effective_date = body.effective_date
    for f in ("title", "doc_number", "content"):
        v = getattr(body, f)
        if v is not None:
            setattr(d, f, v)
    db.commit()
    return doc_dict(d)


@router.delete("/documents/{doc_id}")
def delete_document(doc_id: int, db: Session = Depends(get_db)):
    d = db.get(QualityDocument, doc_id)
    if not d:
        raise HTTPException(404, "Document not found")
    if d.key:
        raise HTTPException(422, "Starter documents cannot be deleted. Edit them instead.")
    db.delete(d)
    db.commit()
    return {"ok": True}


def _add_runs(p, text: str) -> None:
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part:
            p.add_run(part).bold = (i % 2 == 1)


def markdown_to_docx(md: str, path: Path, header: list[str] | None = None) -> Path:
    """Small markdown subset to DOCX: #/##/### headings, - bullets, 1. numbered lists, **bold**, paragraphs."""
    import docx
    from docx.shared import Pt

    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    for line in header or []:
        p = d.add_paragraph()
        p.add_run(line).italic = True
    para: list[str] = []

    def flush():
        if para:
            _add_runs(d.add_paragraph(), " ".join(para))
            para.clear()

    for raw in (md or "").splitlines():
        line = raw.rstrip()
        s = line.strip()
        m = re.match(r"^(#{1,3})\s+(.*)", s)
        if not s:
            flush()
        elif m:
            flush()
            d.add_heading(m.group(2), level=len(m.group(1)))
        elif re.match(r"^[-*]\s+", s):
            flush()
            style = "List Bullet 2" if line.startswith("   ") else "List Bullet"
            _add_runs(d.add_paragraph(style=style), re.sub(r"^[-*]\s+", "", s))
        elif re.match(r"^\d+\.\s+", s):
            flush()
            _add_runs(d.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s+", "", s))
        else:
            para.append(s)
    flush()
    d.save(str(path))
    return path


@router.get("/documents/{doc_id}/docx")
def document_docx(doc_id: int, db: Session = Depends(get_db)):
    from .services import get_profile

    d = db.get(QualityDocument, doc_id)
    if not d:
        raise HTTPException(404, "Document not found")
    company = get_profile(db).name
    content = d.content.replace(COMPANY, company) if company else d.content
    header = [f"{d.doc_number}  Revision {d.version}  Effective {d.effective_date or '(not yet effective)'}".strip()]
    stem = _safe_name(f"{d.doc_number}_{d.title}_Rev{d.version}")
    path = markdown_to_docx(content, _qdir("generated") / f"{stem}.docx", header)
    return FileResponse(path, filename=f"{stem}.docx",
                        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


# ------------------------------------------------------------------ suppliers
@router.get("/suppliers")
def list_suppliers(db: Session = Depends(get_db)):
    return [supplier_dict(db, o) for o in _supplier_orgs(db)]


@router.get("/suppliers/asl.xlsx")
def approved_supplier_list(all: bool = False, db: Session = Depends(get_db)):
    """Approved Supplier List: approved and conditional suppliers (all=true includes every status)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    rows = [s for s in (supplier_dict(db, o) for o in _supplier_orgs(db))
            if all or s["approval"]["status"] in ("approved", "conditional")]
    wb = Workbook()
    ws = wb.active
    ws.title = "Approved Suppliers"
    cols = ["Supplier", "CAGE", "UEI", "Location", "Status", "Approval date", "Basis", "Approved for", "Review due",
            "Certifications on file", "Orders", "On-time %", "Acceptance %", "Last order", "Notes"]
    widths = [32, 9, 15, 18, 13, 13, 26, 34, 12, 40, 8, 10, 12, 12, 40]
    ws.append(cols)
    for i, c in enumerate(ws[1]):
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE4EE")
        ws.column_dimensions[c.column_letter].width = widths[i]
    ws.freeze_panes = "A2"
    for s in rows:
        a, sc = s["approval"], s["scorecard"]
        certs = "; ".join(f"{c['cert_type']}{' exp ' + c['expiration_date'] if c['expiration_date'] else ''}{' (EXPIRED)' if c['state'] == 'expired' else ''}"
                          for c in s["certs"])
        ws.append([s["name"], s["cage"], s["uei"], ", ".join(x for x in [s["city"], s["state"]] if x), a["status"], a["approval_date"],
                   ", ".join(b.replace("_", " ") for b in a["basis"]), a["scope"], a["review_due"], certs, sc["orders"],
                   sc["on_time_pct"], sc["acceptance_pct"], sc["last_order"], a["notes"]])
        for c in ws[ws.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws2 = wb.create_sheet("Notes")
    ws2.column_dimensions["A"].width = 100
    for line in [f"Generated {date.today().isoformat()} from GovBid Pro.",
                 "On-time % = received orders with a promised date that arrived on or before it.",
                 "Acceptance % = quantity accepted / (accepted + rejected) on received orders."]:
        ws2.append([line])
    out = _qdir("generated") / f"Approved_Supplier_List_{date.today().isoformat()}.xlsx"
    wb.save(out)
    return FileResponse(out, filename=out.name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@router.get("/suppliers/{org_id}")
def get_supplier(org_id: int, db: Session = Depends(get_db)):
    return supplier_dict(db, _get_org(db, org_id))


@router.put("/suppliers/{org_id}")
def update_supplier(org_id: int, body: ApprovalIn, db: Session = Depends(get_db)):
    org = _get_org(db, org_id)
    data = body.model_dump(exclude_unset=True)
    _check(data.get("status"), APPROVAL_STATUSES, "status")
    for b in data.get("basis") or []:
        _check(b, APPROVAL_BASES, "basis")
    for f in ("approval_date", "review_due"):
        _check_date(data.get(f), f)
    a = db.scalars(select(SupplierApproval).where(SupplierApproval.organization_id == org_id)).first()
    if not a:
        a = SupplierApproval(organization_id=org_id, status="pending", basis=[])
        db.add(a)
    for k, v in data.items():
        if v is not None:
            setattr(a, k, v)
    if a.status in ("approved", "conditional") and not a.approval_date and "approval_date" not in data:
        a.approval_date = date.today().isoformat()
    db.commit()
    return supplier_dict(db, org)


@router.post("/suppliers/{org_id}/certs")
async def add_supplier_cert(org_id: int, cert_type: str = Form(...), number: str = Form(""), issuer: str = Form(""),
                            expiration_date: str = Form(""), notes: str = Form(""), file: UploadFile | None = File(None),
                            db: Session = Depends(get_db)):
    _get_org(db, org_id)
    _check_date(expiration_date, "expiration_date")
    c = SupplierCert(organization_id=org_id, cert_type=cert_type, number=number, issuer=issuer, expiration_date=expiration_date, notes=notes)
    if file is not None and file.filename:
        name = _safe_name(file.filename)
        stored = f"{datetime.utcnow():%Y%m%d%H%M%S%f}_{name}"
        (_qdir("suppliers", org_id) / stored).write_bytes(await file.read())
        c.filename, c.stored_name = name, stored
    db.add(c)
    db.commit()
    return cert_dict(c)


@router.get("/suppliers/{org_id}/certs/{cert_id}/file")
def supplier_cert_file(org_id: int, cert_id: int, db: Session = Depends(get_db)):
    c = db.get(SupplierCert, cert_id)
    if not c or c.organization_id != org_id or not c.stored_name:
        raise HTTPException(404, "Certificate file not found")
    path = _qdir("suppliers", org_id) / c.stored_name
    if not path.exists():
        raise HTTPException(404, "File is missing on disk")
    return FileResponse(path, filename=c.filename)


@router.delete("/suppliers/{org_id}/certs/{cert_id}")
def delete_supplier_cert(org_id: int, cert_id: int, db: Session = Depends(get_db)):
    c = db.get(SupplierCert, cert_id)
    if not c or c.organization_id != org_id:
        raise HTTPException(404, "Certificate not found")
    if c.stored_name:
        (_qdir("suppliers", org_id) / c.stored_name).unlink(missing_ok=True)
    db.delete(c)
    db.commit()
    return {"ok": True}
