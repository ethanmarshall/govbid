import io

import docx
from fastapi.testclient import TestClient

from app.main import app

TRAINER = {
    "title": "Electrical training equipment design",
    "customer": "Training equipment program",
    "agency": "",
    "role": "employment",
    "start_date": "2022-03-01",
    "end_date": "2025-06-30",
    "naics": "333318",
    "description": "Designed electrical trainer panels with power distribution, motor controls and RTD temperature monitoring.",
    "results": "Delivered panels to multiple schools on schedule.",
    "relevance_keywords": "power distribution, trainer panels, RTD",
    "contact_name": "Jane Roe",
    "contact_email": "jane@example.com",
    "can_use_as_reference": True,
    "value": 250000,
}


def test_crud_and_search():
    with TestClient(app) as c:
        r = c.post("/api/past-performance", json=TRAINER)
        assert r.status_code == 200, r.text
        rec = r.json()
        assert rec["relevance_keywords"] == ["power distribution", "trainer panels", "RTD"]
        assert rec["role_label"] == "Performed as an employee"
        pid = rec["id"]

        assert c.post("/api/past-performance", json={"title": ""}).status_code == 422
        assert c.post("/api/past-performance", json={"title": "x", "role": "boss"}).status_code == 422

        r = c.put(f"/api/past-performance/{pid}", json={"cpars_rating": "Exceptional", "value": None})
        assert r.json()["cpars_rating"] == "Exceptional" and r.json()["value"] is None
        assert r.json()["title"] == TRAINER["title"]  # untouched by partial update

        assert any(x["id"] == pid for x in c.get("/api/past-performance?q=trainer").json())
        assert any(x["id"] == pid for x in c.get("/api/past-performance?q=RTD").json())
        assert not any(x["id"] == pid for x in c.get("/api/past-performance?q=zzzqqq").json())
        assert c.get(f"/api/past-performance/{pid}").json()["id"] == pid

        assert c.delete(f"/api/past-performance/{pid}").json() == {"ok": True}
        assert c.get(f"/api/past-performance/{pid}").status_code == 404


def test_to_library():
    with TestClient(app) as c:
        pid = c.post("/api/past-performance", json=TRAINER).json()["id"]
        r = c.post(f"/api/past-performance/{pid}/to-library")
        assert r.status_code == 200
        lid = r.json()["id"]
        from app.db import SessionLocal
        from app.models import LibraryEntry

        with SessionLocal() as db:
            e = db.get(LibraryEntry, lid)
            assert e.category == "Past performance"
            assert e.content.startswith("## Electrical training equipment design")
            assert "**Our role:** Performed as an employee" in e.content
            assert "**Value:** $250,000" in e.content
            assert "**Period of performance:** 2022-03-01 to 2025-06-30" in e.content
            assert "Jane Roe" in e.content and "### Results" in e.content
        c.delete(f"/api/past-performance/{pid}")


def test_matcher():
    from app.db import SessionLocal
    from app.past_performance_api import match_past_performance

    with TestClient(app) as c:
        a = c.post("/api/past-performance", json=TRAINER).json()["id"]
        b = c.post("/api/past-performance", json={"title": "Website redesign", "role": "commercial", "description": "React site for a bakery", "relevance_keywords": ["web development"]}).json()["id"]
        text = "Solicitation: Power distribution trainer panels for Navy electrical school, NAICS 333318."
        with SessionLocal() as db:
            res = match_past_performance(db, text)
            assert res and res[0]["id"] == a and res[0]["score"] > 5
            assert "power distribution" in res[0]["matched"]
            assert all(r["id"] != b for r in res)
            assert match_past_performance(db, "") == []
            assert len(match_past_performance(db, text, limit=1)) == 1
        r = c.post("/api/past-performance/match", json={"text": text})
        assert r.json()[0]["id"] == a
        c.delete(f"/api/past-performance/{a}")
        c.delete(f"/api/past-performance/{b}")


def test_docx_export():
    with TestClient(app) as c:
        pid = c.post("/api/past-performance", json=TRAINER).json()["id"]
        r = c.get("/api/past-performance/export.docx")
        assert r.status_code == 200
        d = docx.Document(io.BytesIO(r.content))
        text = "\n".join(p.text for p in d.paragraphs)
        assert "Past Performance Volume" in text
        assert "Electrical training equipment design" in text
        assert "Description of work" in text
        cells = [cell.text for t in d.tables for row in t.rows for cell in row.cells]
        assert "Point of contact" in cells and "$250,000" in cells
        c.delete(f"/api/past-performance/{pid}")
