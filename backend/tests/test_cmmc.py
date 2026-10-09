import io

import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy import delete

from app import cmmc_controls as cc
from app.main import app


def _reset():
    from app.db import SessionLocal, init_db
    from app.models_cmmc import CmmcControl, CmmcSettings

    init_db()
    with SessionLocal() as db:
        db.execute(delete(CmmcControl))
        db.execute(delete(CmmcSettings))
        db.commit()


# ------------------------------------------------------------------ data and score math
def test_control_data_shape():
    assert len(cc.LEVEL1_CONTROLS) == 15
    assert len(cc.LEVEL2_CONTROLS) == 110
    ids = [c["control_id"] for c in cc.LEVEL2_CONTROLS]
    assert ids[0] == "3.1.1" and ids[-1] == "3.14.7" and len(set(ids)) == 110
    fams = {c["family"] for c in cc.LEVEL2_CONTROLS}
    assert len(fams) == 14
    assert {c["weight"] for c in cc.LEVEL2_CONTROLS} == {0, 1, 3, 5}
    assert cc.WEIGHTS["3.12.4"] == 0  # "NA" in the methodology
    assert cc.WEIGHTS["3.1.1"] == 5 and cc.WEIGHTS["3.1.5"] == 3 and cc.WEIGHTS["3.1.3"] == 1
    assert cc.LEVEL1_CONTROLS[0]["control_id"] == "AC.L1-b.1.i"
    assert cc.LEVEL1_CONTROLS[-1]["control_id"] == "SI.L1-b.1.xv"
    assert cc.LEVEL1_CONTROLS[8]["nist_ids"] == ["3.10.3", "3.10.4", "3.10.5"]
    # POA&M eligibility per 32 CFR 170.21
    idx = {c["control_id"]: c for c in cc.LEVEL2_CONTROLS}
    assert not idx["3.1.20"]["poam_allowed"] and not idx["3.12.4"]["poam_allowed"]
    assert not idx["3.1.1"]["poam_allowed"]  # 5 points
    assert idx["3.1.3"]["poam_allowed"] and idx["3.13.11"]["poam_allowed"]


def test_score_all_not_started_is_strongly_negative():
    score = cc.sprs_score({})
    assert score == 110 - sum(cc.WEIGHTS.values()) == -203


def test_score_all_implemented_is_110():
    assert cc.sprs_score({c["control_id"]: "implemented" for c in cc.LEVEL2_CONTROLS}) == 110
    assert cc.sprs_score({c["control_id"]: "not_applicable" for c in cc.LEVEL2_CONTROLS}) == 110


def test_partial_credit_rules():
    allok = {c["control_id"]: "implemented" for c in cc.LEVEL2_CONTROLS}
    assert cc.sprs_score({**allok, "3.5.3": "partial"}) == 107
    assert cc.sprs_score({**allok, "3.5.3": "not_started"}) == 105
    assert cc.sprs_score({**allok, "3.13.11": "partial"}) == 107
    assert cc.sprs_score({**allok, "3.13.11": "planned"}) == 105
    # partial on any other requirement earns no credit
    assert cc.sprs_score({**allok, "3.1.1": "partial"}) == 105
    assert cc.sprs_score({**allok, "3.1.3": "partial"}) == 109


# ------------------------------------------------------------------ API
def test_controls_list_update_and_summary():
    _reset()
    with TestClient(app) as c:
        rows = c.get("/api/cmmc/controls").json()
        assert len(rows) == 125
        l2 = c.get("/api/cmmc/controls?framework=L2&family=AC").json()
        assert len(l2) == 22 and l2[0]["control_id"] == "3.1.1" and l2[0]["weight"] == 5

        s = c.get("/api/cmmc/summary").json()
        assert s["sprs_score"] == -203 and s["sprs_max"] == 110 and s["level1_done"] == 0 and s["level1_total"] == 15
        assert s["open_poam"] == 0 and s["next_poam_due"] is None

        r = c.put("/api/cmmc/controls/L2/3.1.1", json={"status": "implemented", "evidence": "AD accounts only", "owner": "Ethan"})
        assert r.status_code == 200 and r.json()["status"] == "implemented" and r.json()["deduction"] == 0
        c.put("/api/cmmc/controls/L2/3.5.3", json={"status": "partial", "poam_due": "2027-01-15"})
        c.put("/api/cmmc/controls/L2/3.1.3", json={"status": "planned", "poam_due": "2026-12-01"})
        c.put("/api/cmmc/controls/L1/AC.L1-b.1.i", json={"status": "implemented"})
        c.put("/api/cmmc/controls/L1/PE.L1-b.1.ix", json={"status": "not_applicable"})

        assert c.put("/api/cmmc/controls/L2/3.1.1", json={"status": "done"}).status_code == 422
        assert c.put("/api/cmmc/controls/L2/3.1.1", json={"poam_due": "Jan 5"}).status_code == 422
        assert c.put("/api/cmmc/controls/L2/9.9.9", json={"status": "implemented"}).status_code == 404

        s = c.get("/api/cmmc/summary").json()
        assert s["sprs_score"] == -203 + 5 + 2  # 3.1.1 met, 3.5.3 partial (5 -> 3)
        assert s["level1_done"] == 2
        assert s["open_poam"] == 2 and s["next_poam_due"] == "2026-12-01"
        assert any("system security plan" in w for w in s["warnings"])

        planned = c.get("/api/cmmc/controls?framework=L2&status=planned,partial").json()
        assert {r["control_id"] for r in planned} == {"3.5.3", "3.1.3"}
        # evidence persisted
        one = [r for r in c.get("/api/cmmc/controls?framework=L2&family=AC").json() if r["control_id"] == "3.1.1"][0]
        assert one["evidence"] == "AD accounts only" and one["owner"] == "Ethan"


def test_settings_and_xlsx_export():
    _reset()
    with TestClient(app) as c:
        r = c.put("/api/cmmc/settings", json={"last_sprs_submission_date": "2026-09-30", "scope_notes": "One laptop and a M365 GCC tenant"})
        assert r.status_code == 200 and r.json()["last_sprs_submission_date"] == "2026-09-30"
        assert c.put("/api/cmmc/settings", json={"affirmation_date": "2026-13-40"}).status_code == 422
        assert c.get("/api/cmmc/summary").json()["last_sprs_submission_date"] == "2026-09-30"
        c.put("/api/cmmc/controls/L2/3.13.11", json={"status": "partial", "poam_due": "2027-02-01", "notes": "Move to FIPS mode"})

        r = c.get("/api/cmmc/export.xlsx")
        assert r.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Summary", "Status", "POA&M"]
        assert wb["Status"].max_row == 126  # header + 125 controls
        poam = list(wb["POA&M"].iter_rows(min_row=2, values_only=True))
        assert poam[0][0] == "3.13.11" and poam[0][4] == 3 and poam[0][5] == "yes"
        summ = {row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True)}
        assert "One laptop" in summ["Scope (systems holding FCI/CUI)"]
