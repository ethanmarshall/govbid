import io
import json
import re
import zipfile
from datetime import date, timedelta

from docx import Document
from fastapi.testclient import TestClient

from app import pricing, quote_tools_api
from app.cad_quote import CAD_DIR
from app.db import SessionLocal
from app.drawings_api import DRAWING_DIR
from app.main import app
from app.models import PartQuote

NSN_A = "5340-01-480-5627"


def _opp(c, **kw):
    body = {"title": "Bracket buy", "solicitation_number": kw.pop("sol", "SPE7M1-26-T-0001"), "agency": "DLA Land and Maritime",
            "set_aside_code": "SDVOSBC", "naics": "332710", "description": "Machined bracket.",
            "response_deadline": (date.today() + timedelta(days=20)).isoformat(), **kw}
    return c.post("/api/opportunities", json=body).json()["id"]


def _quote(c, status="draft", quoted_unit_price=None, oid=None, **spec_kw):
    spec = dict(pricing.EXAMPLE_SPEC, quantities=[10, 50, 100], **spec_kw)
    body = {"spec": spec, "status": status, "quoted_quantity": 50, "opportunity_id": oid}
    if quoted_unit_price is not None:
        body["quoted_unit_price"] = quoted_unit_price
    r = c.post("/api/pricing/quotes", json=body)
    assert r.status_code == 200, r.text
    return r.json()


def _vendor(c, name="Good Shop", email="sales@goodshop.example"):
    org = c.post("/api/crm/organizations", json={"name": name, "kind": "vendor"}).json()
    c.put(f"/api/crm/organizations/{org['id']}", json={"business_types": {"SB": True, "SDVOSB": True}, "is_manufacturer": True})
    if email:
        c.post(f"/api/crm/organizations/{org['id']}/contacts", json={"name": "Pat Jones", "email": email})
    return org["id"]


def _store_drawing(did: str, read: dict):
    DRAWING_DIR.mkdir(parents=True, exist_ok=True)
    (DRAWING_DIR / f"{did}.pdf").write_bytes(b"%PDF-1.4\n% test drawing\n")
    (DRAWING_DIR / f"{did}.json").write_text(json.dumps({"filenames": ["bracket.pdf"], "read": read}))


def _store_step(fid: str):
    CAD_DIR.mkdir(parents=True, exist_ok=True)
    (CAD_DIR / f"{fid}.step").write_text("ISO-10303-21;\nEND-ISO-10303-21;\n")
    (CAD_DIR / f"{fid}.json").write_text(json.dumps({"filenames": ["bracket.step"]}))


def test_settings_roundtrip_and_validation():
    with TestClient(app) as c:
        s = c.get("/api/quote-tools/settings").json()
        assert s["validity_days"] == 30 and s["payment_terms"] == "Net 30"
        s2 = c.put("/api/quote-tools/settings", json={"validity_days": 45, "fob": "Origin", "address": "1 Main St\nTown, ST 00000"}).json()
        assert s2["validity_days"] == 45 and s2["fob"] == "Origin"
        assert c.put("/api/quote-tools/settings", json={"validity_days": 0}).status_code == 400
        c.put("/api/quote-tools/settings", json={"validity_days": 30, "fob": "Destination"})


def test_customer_quote_pdf_and_docx():
    with TestClient(app) as c:
        prof = c.get("/api/profile").json()
        prof.update(name="Marshall Engineering LLC", uei="ABCDEFGH1234", cage="1AB23")
        c.put("/api/profile", json=prof)
        c.put("/api/quote-tools/settings", json={"address": "100 Test Rd\nSpringfield, VA 22150"})
        oid = _opp(c, sol="SPE7M1-26-T-0100")
        q = _quote(c, oid=oid, revision="C")
        pq = c.get(f"/api/pricing/quotes/{q['id']}").json()
        breaks = {b["quantity"]: b for b in pq["result"]["price_breaks"]}

        prev = c.get(f"/api/quote-tools/{q['id']}/customer-quote").json()
        assert re.fullmatch(rf"Q-{date.today().year}-\d{{4}}", prev["number"])
        assert prev["available_quantities"] == [10, 50, 100]
        assert prev["reference"]["solicitation_number"] == "SPE7M1-26-T-0100" and prev["reference"]["revision"] == "C"
        assert prev["terms"]["payment_terms"] == "Net 30"
        assert "Military packaging" in prev["terms"]["packaging"]
        for ln in prev["lines"]:
            assert set(ln) == {"quantity", "unit_price", "extended_price", "lead_time"}  # no cost or margin

        r = c.get(f"/api/quote-tools/{q['id']}/customer-quote.pdf",
                  params={"customer_name": "DLA Land and Maritime", "attn": "Buyer", "quantities": "10,100", "notes": "Price excludes FAT."})
        assert r.status_code == 200 and r.content.startswith(b"%PDF")
        assert prev["number"] in r.headers["content-disposition"]
        # the form values are remembered
        again = c.get(f"/api/quote-tools/{q['id']}/customer-quote").json()
        assert again["number"] == prev["number"]
        assert again["options"]["quantities"] == [10, 100] and again["options"]["customer_name"] == "DLA Land and Maritime"
        assert [ln["quantity"] for ln in again["lines"]] == [10, 50, 100]  # the preview lists every quantity

        r = c.get(f"/api/quote-tools/{q['id']}/customer-quote.docx")
        assert r.status_code == 200 and r.content[:2] == b"PK"
        doc = Document(io.BytesIO(r.content))
        text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
        assert "Marshall Engineering LLC" in text and "UEI ABCDEFGH1234" in text and "CAGE 1AB23" in text
        assert "100 Test Rd" in text and prev["number"] in text and "Net 30" in text and "ARO" in text
        assert f"${breaks[100]['unit_price']:,.2f}" in text
        assert "margin" not in text.lower() and "unit cost" not in text.lower()

        q2 = _quote(c)
        n2 = c.get(f"/api/quote-tools/{q2['id']}/customer-quote").json()["number"]
        assert int(n2[-4:]) == int(prev["number"][-4:]) + 1
        assert c.get(f"/api/quote-tools/{q2['id']}/customer-quote.pdf", params={"quantities": "7"}).status_code == 400
        assert c.get("/api/quote-tools/999999/customer-quote.pdf").status_code == 404


def test_rfq_drafts_package_response_and_make_or_buy():
    with TestClient(app) as c:
        did, fid = "a" * 32, "b" * 40
        _store_drawing(did, {"export_controlled": False, "distribution": {"letter": "A"}})
        _store_step(fid)
        q = _quote(c, material_certs_required=True, inspection={"first_article": True},
                   drawing={"drawing_id": did, "filename": "bracket.pdf", "revision": "B"},
                   cad={"file_id": fid, "filename": "bracket.step"},
                   bom=[{"item": 1, "mpn": "91251A540", "qty": 4, "description": "SHCS 1/4-20"}])
        v1, v2 = _vendor(c), _vendor(c, "No Email Machine", email="")
        vendors = c.get("/api/quote-tools/vendors").json()
        good = next(v for v in vendors if v["id"] == v1)
        assert good["is_manufacturer"] and good["business_types"]["SDVOSB"] and good["email"] == "sales@goodshop.example"

        due = (date.today() + timedelta(days=7)).isoformat()
        r = c.post(f"/api/quote-tools/{q['id']}/rfq", json={"vendor_ids": [v1, v2], "due_date": due, "message": "Repeat buy.",
                                                            "include_files": True, "quantities": [10, 100]})
        assert r.status_code == 200, r.text
        d = r.json()
        assert not d["files_blocked"] and "cannot attach" in d["mailto_note"]
        a, b = d["rfqs"]
        assert a["status"] == "sent" and a["vendor_email"] == "sales@goodshop.example" and b["vendor_email"] == ""
        assert "Mounting bracket" in a["subject"] and due in a["subject"]
        for needle in ("Hello Pat,", "Repeat buy.", "Quantities: 10, 100", "Material: 6061-T6 aluminum (material certifications required)",
                       "Finish: anodize (Type II)", "First article", "MIL-STD-2073-1", f"Please reply by: {due}", "bracket.step"):
            assert needle in a["body"], needle
        assert a["mailto"].startswith("mailto:sales@goodshop.example?subject=") and "%20" in a["mailto"]
        assert set(a["package_files"]) == {"RFQ.txt", "bracket.step", "bracket.pdf", "bom.csv"}

        z = c.get(f"/api/quote-tools/rfqs/{a['id']}/package.zip")
        assert z.status_code == 200 and z.headers["content-type"] == "application/zip"
        names = zipfile.ZipFile(io.BytesIO(z.content)).namelist()
        assert set(names) == {"RFQ.txt", "bracket.step", "bracket.pdf", "bom.csv"}
        bom = zipfile.ZipFile(io.BytesIO(z.content)).read("bom.csv").decode()
        assert "91251A540" in bom
        m = c.get(f"/api/quote-tools/rfqs/{a['id']}/mailto").json()
        assert "cannot attach" in m["note"]

        assert len(c.get(f"/api/quote-tools/{q['id']}/rfqs").json()) == 2
        assert c.put(f"/api/quote-tools/rfqs/{a['id']}", json={"status": "bogus"}).status_code == 400
        assert c.put(f"/api/quote-tools/rfqs/{a['id']}", json={"status": "awaiting"}).json()["status"] == "awaiting"

        resp = c.post(f"/api/quote-tools/rfqs/{a['id']}/response",
                      json={"prices": [{"quantity": 10, "unit_price": 40}, {"quantity": 100, "unit_price": 15}], "lead_days": 21, "notes": "Firm 30 days"}).json()
        assert resp["status"] == "responded" and resp["response"]["lead_days"] == 21
        mob = c.get(f"/api/pricing/quotes/{q['id']}/make-or-buy").json()
        assert len(mob["vendor_quotes"]) == 1 and mob["vendor_quotes"][0]["vendor_name"] == "Good Shop"
        # recording again updates the same vendor quote
        c.post(f"/api/quote-tools/rfqs/{a['id']}/response", json={"prices": [{"quantity": 100, "unit_price": 14}], "lead_days": 20})
        mob = c.get(f"/api/pricing/quotes/{q['id']}/make-or-buy").json()
        assert len(mob["vendor_quotes"]) == 1 and mob["vendor_quotes"][0]["prices"] == [{"quantity": 100, "unit_price": 14.0}]
        assert c.post(f"/api/quote-tools/rfqs/{a['id']}/response", json={"prices": []}).status_code == 400

        dec = c.post(f"/api/quote-tools/rfqs/{b['id']}/decline", json={"notes": "No capacity"}).json()
        assert dec["status"] == "declined" and dec["response_notes"] == "No capacity"
        assert c.post(f"/api/quote-tools/{q['id']}/rfq", json={"vendor_ids": []}).status_code == 400
        assert c.post(f"/api/quote-tools/{q['id']}/rfq", json={"vendor_ids": [999999]}).status_code == 404


def test_rfq_blocks_files_for_controlled_drawing():
    with TestClient(app) as c:
        did, fid = "c" * 32, "d" * 40
        _store_drawing(did, {"export_controlled": True, "distribution": {"letter": "D", "meaning": "DoD and U.S. DoD contractors only."}})
        _store_step(fid)
        q = _quote(c, drawing={"drawing_id": did, "filename": "ctl.pdf"}, cad={"file_id": fid, "filename": "ctl.step"})
        chk = c.get(f"/api/quote-tools/{q['id']}/export-check").json()
        assert chk["flagged"] and any("Distribution Statement D" in x for x in chk["reasons"])
        v = _vendor(c, "Controlled Shop", "rfq@controlled.example")
        d = c.post(f"/api/quote-tools/{q['id']}/rfq", json={"vendor_ids": [v], "include_files": True}).json()
        assert d["files_blocked"] and any("JCP" in w for w in d["warnings"])
        rfq = d["rfqs"][0]
        assert rfq["package_files"] == ["RFQ.txt"] and "not attached" in rfq["body"]
        names = zipfile.ZipFile(io.BytesIO(c.get(f"/api/quote-tools/rfqs/{rfq['id']}/package.zip").content)).namelist()
        assert "ctl.step" not in names and "ctl.pdf" not in names and "FILES_NOT_INCLUDED.txt" in names

        # an ITAR-flagged solicitation blocks files too
        oid = _opp(c, sol="ITAR-1", description="Technical data is subject to ITAR. JCP certification required.")
        q2 = _quote(c, oid=oid, cad={"file_id": fid, "filename": "ctl.step"})
        d2 = c.post(f"/api/quote-tools/{q2['id']}/rfq", json={"vendor_ids": [v], "include_files": True}).json()
        assert d2["files_blocked"] and any("solicitation" in w for w in d2["warnings"])


def test_rfq_calendar_and_dashboard_hooks():
    with TestClient(app) as c:
        q = _quote(c)
        v = _vendor(c, "Late Shop", "late@shop.example")
        past = (date.today() - timedelta(days=2)).isoformat()
        rid = c.post(f"/api/quote-tools/{q['id']}/rfq", json={"vendor_ids": [v], "due_date": past, "include_files": False}).json()["rfqs"][0]["id"]
        with SessionLocal() as db:
            items = quote_tools_api.calendar_items(db)
            assert any(i["uid"] == f"vendor-rfq-{rid}@govbid" and i["date"].isoformat() == past and i["url"].endswith(f"id={q['id']}") for i in items)
            assert any("Late Shop" in s for s in quote_tools_api.dashboard_items(db))
        c.post(f"/api/quote-tools/rfqs/{rid}/decline", json={})
        with SessionLocal() as db:
            assert not any("Late Shop" in s for s in quote_tools_api.dashboard_items(db))
        assert c.delete(f"/api/quote-tools/rfqs/{rid}").status_code == 200


def test_win_loss_insights():
    with TestClient(app) as c:
        with SessionLocal() as db:  # isolate from quotes made by other tests
            for pq in db.query(PartQuote).filter(PartQuote.status.in_(("won", "lost", "submitted"))).all():
                pq.status = "draft"
            db.commit()
        base = _quote(c)
        cost50 = next(b for b in base["price_breaks"] if b["quantity"] == 50)["unit_cost"]

        def at_margin(m):  # unit price giving margin m (fraction) over our unit cost
            return round(cost50 / (1 - m), 2)

        nsn_b = "5340-01-111-2222"
        c.post("/api/nsn/records", json={"nsn": nsn_b, "contract_number": "SPE-OTHER-1", "award_date": "2099-01-01",
                                         "quantity": 50, "unit_price": round(at_margin(0.35) * 0.8, 2), "awardee": "Rival"})
        _quote(c, "won", at_margin(0.15), nsn=NSN_A)
        _quote(c, "won", at_margin(0.25), nsn=NSN_A)
        _quote(c, "won", at_margin(0.05), nsn=NSN_A)
        _quote(c, "lost", at_margin(0.35), nsn=nsn_b)
        _quote(c, "lost", at_margin(0.45), nsn=NSN_A)
        _quote(c, "submitted", at_margin(0.2), nsn=NSN_A)
        d = c.get("/api/quote-tools/insights").json()
        o = d["overall"]
        assert o["quotes"] == 6 and o["won"] == 3 and o["lost"] == 2 and o["submitted"] == 1
        assert o["win_rate"] == 60.0
        assert abs(o["avg_price_ratio"] - 1.25) < 0.01 and o["ratio_count"] == 1  # our losing price / the rival's award
        assert abs(o["avg_margin_won"] - 15.0) < 0.2 and abs(o["avg_margin_lost"] - 40.0) < 0.2
        assert o["suggestion"]["enough"] and o["suggestion"]["low"] == 20 and o["suggestion"]["high"] == 30
        assert "5 decided" in o["suggestion"]["text"]
        fsc = {g["value"]: g for g in d["groups"]["fsc"]}
        assert fsc["5340"]["quotes"] == 6
        mat = d["groups"]["material"][0]
        assert mat["value"] == "6061-T6 aluminum"
        small = {g["value"]: g for g in d["groups"]["kind"]}
        assert small["part"]["suggestion"]["enough"]
        lost_b = next(p for p in d["points"] if p["nsn"] == nsn_b)
        assert lost_b["award_price"] is not None and abs(lost_b["price_ratio"] - 1.25) < 0.01
        # under five decided quotes there is not enough data
        with SessionLocal() as db:
            for pq in db.query(PartQuote).filter(PartQuote.status == "won").all():
                pq.status = "submitted"
            db.commit()
        o2 = c.get("/api/quote-tools/insights").json()["overall"]
        assert not o2["suggestion"]["enough"] and "Not enough data" in o2["suggestion"]["text"]
        with SessionLocal() as db:
            for pq in db.query(PartQuote).filter(PartQuote.status.in_(("won", "lost", "submitted"))).all():
                pq.status = "draft"
            db.commit()


def test_deleted_quote_does_not_pass_its_number_to_a_new_quote():
    from fastapi.testclient import TestClient
    from app.main import app
    from app import pricing
    with TestClient(app) as c:
        a = c.post("/api/pricing/quotes", json={"spec": pricing.EXAMPLE_SPEC}).json()
        na = c.get(f"/api/quote-tools/{a['id']}/customer-quote").json()["number"]
        assert c.delete(f"/api/pricing/quotes/{a['id']}").status_code in (200, 204)
        b = c.post("/api/pricing/quotes", json={"spec": pricing.EXAMPLE_SPEC}).json()
        nb = c.get(f"/api/quote-tools/{b['id']}/customer-quote").json()["number"]
        assert int(nb[-4:]) == int(na[-4:]) + 1  # never reissued, even when SQLite reuses the quote id
