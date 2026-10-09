import io

import docx
from fastapi.testclient import TestClient

from app import flowdown_data as fd
from app.flowdown_api import extract_clause_numbers
from app.main import app

REQUIRED = ("52.203-13 52.203-19 52.204-21 52.204-23 52.204-25 52.204-27 52.209-6 52.219-8 52.219-9 52.219-14 52.222-21 "
            "52.222-26 52.222-35 52.222-36 52.222-37 52.222-40 52.222-41 52.222-50 52.222-54 52.223-18 52.224-3 52.225-1 "
            "52.225-13 52.232-40 52.244-6 52.246-2 52.247-64 252.203-7002 252.204-7012 252.204-7020 252.204-7021 "
            "252.211-7003 252.222-7006 252.223-7008 252.225-7009 252.225-7012 252.225-7048 252.226-7001 252.244-7000 "
            "252.246-7007 252.246-7008 252.247-7023").split()


def _fd(number, **sub):
    return fd.evaluate(number, sub)["flow_down"]


def test_table_complete_with_sources():
    for n in REQUIRED:
        c = fd.BY_NUMBER[n]
        assert c["source_url"].startswith("https://www.acquisition.gov/")
        assert c["status"] in ("mandatory", "conditional", "not_required", "check")
        assert c["paragraph"]


def test_thresholds():
    assert _fd("52.222-36", value=20_000) == "no"
    assert _fd("52.222-36", value=20_001) == "yes"
    assert _fd("52.222-35", value=199_999) == "no"
    assert _fd("52.222-35", value=200_000) == "yes"
    assert _fd("52.203-13", value=8_000_000, performance_days=90) == "no"
    assert _fd("52.203-13", value=8_000_000, performance_days=180) == "yes"
    assert _fd("52.209-6", value=45_000) == "no"
    assert _fd("52.209-6", value=45_001) == "yes"
    assert _fd("52.209-6", value=100_000, cots=True) == "no"
    assert _fd("52.222-54", value=3_000, services=True) == "no"
    assert _fd("52.222-54", value=5_000, services=True) == "yes"
    assert _fd("52.222-54", value=5_000, services=False, supplies=True) == "no"
    assert _fd("252.226-7001", value=600_000) == "yes"


def test_conditions():
    assert _fd("252.204-7012", involves_cui=False) == "no"
    assert _fd("252.204-7012", involves_cui=True, commercial=True) == "yes"
    assert _fd("52.204-21", involves_fci=True) == "yes"
    assert _fd("52.204-21", involves_fci=True, cots=True) == "no"
    assert _fd("252.204-7021", involves_cui=True) == "yes"  # CUI implies FCI
    assert _fd("252.246-7008", electronic_parts=True, original_manufacturer=True) == "no"
    assert _fd("252.246-7008", electronic_parts=True) == "yes"
    assert _fd("52.232-40", small_business=True) == "yes"
    assert _fd("52.219-14") == "no"
    assert _fd("52.246-2") == "check"
    assert _fd("52.999-99") == "check"


def test_commercial_filtering():
    # Not on the commercial lists -> not flowed to a commercial subcontract
    assert _fd("52.225-13", commercial=True) == "no"
    assert _fd("252.225-7048", commercial=True) == "no"
    assert _fd("252.225-7048", commercial=False) == "yes"
    # On the list -> still evaluated
    assert _fd("52.204-25", commercial=True) == "yes"
    assert _fd("52.247-64", commercial=True) == "no"
    assert _fd("52.247-64", commercial=True, resale_no_value_added=True) == "yes"


def test_extract_clause_numbers():
    text = ("Clauses: FAR 52.212-4, 52.219-14 Limitations on Subcontracting; DFARS 252.204-7012 and 252.225-7048. "
            "Section 52.204-21. Page 52.204-21 again. Not a clause: 1252.204-70123 or 52.2041-1")
    assert extract_clause_numbers(text) == ["52.212-4", "52.219-14", "252.204-7012", "252.225-7048", "52.204-21"]


def test_api_check_opportunity_and_docx():
    from app.db import SessionLocal
    from app.models import Analysis, Opportunity

    with TestClient(app) as c:
        with SessionLocal() as db:
            o = Opportunity(source="manual", external_id="fd-1", solicitation_number="SPE-FD", title="Flowdown test")
            db.add(o)
            db.flush()
            db.add(Analysis(opportunity_id=o.id, summary="Includes DFARS 252.204-7012.",
                            breakdown={"key_clauses": [{"clause": "52.219-14", "note": "LOS"}, {"clause": "FAR 52.204-25"}]},
                            compliance_matrix=[{"id": 1, "requirement": "Comply with 252.225-7048 export controls and 52.222-36"}]))
            db.commit()
            opp_id = o.id

        assert len(c.get("/api/flowdown/clauses").json()["clauses"]) >= 42
        cl = c.get(f"/api/flowdown/opportunity/{opp_id}/clauses").json()["clauses"]
        assert set(cl) == {"52.219-14", "52.204-25", "252.204-7012", "252.225-7048", "52.222-36"}

        r = c.post("/api/flowdown/check", json={"opportunity_id": opp_id, "subcontract": {"value": 15_000}}).json()
        yes = {x["number"] for x in r["flow_down"]}
        assert yes == {"52.204-25", "252.225-7048"}
        no = {x["number"] for x in r["not_required"]}
        assert {"52.219-14", "252.204-7012", "52.222-36"} <= no

        r = c.post("/api/flowdown/check", json={"clauses": ["52.222-36", "foo"], "clause_text": "and 252.204-7012",
                                                "subcontract": {"value": 50_000, "involves_cui": True}}).json()
        assert {x["number"] for x in r["flow_down"]} == {"52.222-36", "252.204-7012"}
        assert r["unknown"] == ["foo"]

        resp = c.post("/api/flowdown/po-attachment", json={"opportunity_id": opp_id, "vendor": "Acme Machine",
                                                           "po_number": "PO-1001", "subcontract": {"value": 15_000}})
        assert resp.status_code == 200
        d = docx.Document(io.BytesIO(resp.content))
        text = "\n".join(p.text for p in d.paragraphs)
        assert "The following clauses are incorporated by reference" in text
        assert "Acme Machine" in text
        cells = [cell.text for t in d.tables for row in t.rows for cell in row.cells]
        assert "52.204-25" in cells and "252.225-7048" in cells
        assert any("excluding paragraph (b)(2)" in x for x in cells)

        bad = c.post("/api/flowdown/po-attachment", json={"clauses": ["52.219-14"], "subcontract": {}})
        assert bad.status_code == 400
