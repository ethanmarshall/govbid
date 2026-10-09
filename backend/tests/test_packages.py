import io
import sys
import types
import zipfile

import docx
from fastapi.testclient import TestClient

from app import packages_api
from app.main import app

SOL = b"""SECTION L
L.1 The offeror shall submit a technical volume not to exceed 10 pages.
L.2 The offeror shall provide three past performance references.
L.3 The offeror shall submit pricing for all CLINs.
SECTION C
C.1 The contractor shall staff the help desk from 8am to 5pm.
C.2 The contractor shall deliver monthly status reports.
"""


def _opp_with_analysis(c):
    oid = c.post("/api/opportunities", json={"title": "Help Desk", "solicitation_number": "36C10X26Q9999", "agency": "VA"}).json()["id"]
    c.post(f"/api/opportunities/{oid}/documents", files={"files": ("rfq.txt", SOL, "text/plain")})
    a = c.post(f"/api/opportunities/{oid}/analyze").json()["analysis"]
    return oid, a["compliance_matrix"]


def test_proposal_package_flow():
    with TestClient(app) as c:
        c.put("/api/profile", json={"name": "Marshall Tech LLC", "uei": "ABC123DEF456", "cage": "1A2B3", "certifications": {"SB": "certified", "SDVOSB": "pending"}})
        oid, matrix = _opp_with_analysis(c)
        assert len(matrix) >= 5

        p = c.post("/api/packages", json={"kind": "proposal", "opportunity_id": oid}).json()
        assert p["name"] == "Help Desk"
        assert p["cover"]["solicitation_number"] == "36C10X26Q9999"
        assert p["cover"]["business_status"] == "Small Business"  # SDVOSB pending, so not claimed
        titles = [s["title"] for s in p["sections"]]
        assert "Technical approach" in titles and "Past performance references" in titles
        # every requirement assigned exactly once
        assigned = [r for s in p["sections"] for r in s["requirement_ids"]]
        assert sorted(assigned) == sorted(r["id"] for r in matrix)
        assert p["requirements_assigned"] == len(matrix)
        by_title = {s["title"]: s for s in p["sections"]}
        pp_rows = by_title["Past performance references"]["requirement_ids"]
        assert any("past performance" in next(r["requirement"] for r in matrix if r["id"] == rid).lower() for rid in pp_rows)

        # write a section and mark its requirements covered -> syncs back to the opportunity matrix
        tech = by_title["Technical approach"]
        r = c.put(f"/api/sections/{tech['id']}", json={
            "content": "### Help desk staffing\nWe staff **two** technicians 8am to 5pm.\n\n- Tier 1 triage\n- Tier 2 escalation\n\n| Role | Hours |\n|---|---|\n| Tech | 8-5 |",
            "status": "done", "covered_ids": tech["requirement_ids"], "page_limit": 10,
        }).json()
        assert r["words"] > 10 and r["status"] == "done"
        opp = c.get(f"/api/opportunities/{oid}").json()
        rows = {x["id"]: x for x in opp["analysis"]["compliance_matrix"]}
        for rid in tech["requirement_ids"]:
            assert rows[rid]["status"] == "done"
            assert rows[rid]["response_location"] == "Volume I 3.0 Technical approach"

        # uncovering clears the location again
        c.put(f"/api/sections/{tech['id']}", json={"covered_ids": []})
        rows = {x["id"]: x for x in c.get(f"/api/opportunities/{oid}").json()["analysis"]["compliance_matrix"]}
        assert all(rows[rid]["response_location"] == "" for rid in tech["requirement_ids"])

        # moving a requirement to another section removes it from the first
        mgmt = by_title["Management approach and schedule"]
        moved = tech["requirement_ids"][0]
        c.put(f"/api/sections/{mgmt['id']}", json={"requirement_ids": mgmt["requirement_ids"] + [moved]})
        p = c.get(f"/api/packages/{p['id']}").json()
        by_title = {s["title"]: s for s in p["sections"]}
        assert moved not in by_title["Technical approach"]["requirement_ids"]
        assert moved in by_title["Management approach and schedule"]["requirement_ids"]

        # add a section after technical approach, then delete it
        p = c.post(f"/api/packages/{p['id']}/sections", json={"title": "Transition plan", "number": "3.1", "after_id": tech["id"]}).json()
        titles = [s["title"] for s in p["sections"]]
        assert titles[titles.index("Technical approach") + 1] == "Transition plan"
        new_id = next(s["id"] for s in p["sections"] if s["title"] == "Transition plan")
        p = c.delete(f"/api/sections/{new_id}").json()
        assert "Transition plan" not in [s["title"] for s in p["sections"]]

        # attachments with file upload
        item = p["items"][0]
        r = c.post(f"/api/items/{item['id']}/files", files={"files": ("SF1449_signed.pdf", b"%PDF-1.4 fake", "application/pdf")}).json()
        assert r["files"][0]["name"] == "SF1449_signed.pdf" and r["status"] == "in_progress"
        assert c.get(f"/api/items/{item['id']}/files/SF1449_signed.pdf").content.startswith(b"%PDF")

        # Word export
        d = c.get(f"/api/packages/{p['id']}/export.docx")
        assert d.status_code == 200
        doc = docx.Document(io.BytesIO(d.content))
        text = "\n".join(par.text for par in doc.paragraphs)
        assert "Help Desk" in text and "36C10X26Q9999" in text and "Marshall Tech LLC" in text
        assert "3.0 Technical approach" in text and "Help desk staffing" in text
        assert "Compliance cross-reference" in text
        assert any(par.style.name == "List Bullet" and "Tier 1" in par.text for par in doc.paragraphs)
        assert any(t.rows[0].cells[0].text == "Role" for t in doc.tables)
        assert "[Section not yet written." in text
        assert "—" not in text

        # zip export
        z = zipfile.ZipFile(io.BytesIO(c.get(f"/api/packages/{p['id']}/export.zip").content))
        names = z.namelist()
        assert any(n.endswith(".docx") for n in names)
        assert any(n.endswith("Compliance_Matrix.xlsx") for n in names)
        assert any(n.endswith("SF1449_signed.pdf") and "/Attachments/" in n for n in names)
        assert any(n.endswith("index.csv") for n in names)

        # listing and deleting the opportunity keeps the package
        assert c.get("/api/packages").json()[0]["id"] == p["id"]
        c.delete(f"/api/opportunities/{oid}")
        assert c.get(f"/api/packages/{p['id']}").json()["opportunity_id"] is None


def test_tdp_package_and_library():
    with TestClient(app) as c:
        p = c.post("/api/packages", json={"kind": "tdp", "name": "Trainer TDP", "contract_number": "W912XX-26-C-0001"}).json()
        assert p["sections"][0]["title"] == "Transmittal letter"
        assert [i["cdrl"] for i in p["items"]][:4] == ["A001", "A002", "A003", "A004"]
        assert p["items"][0]["did"] == "DI-SESS-81000"
        assert "DISTRIBUTION STATEMENT" in p["cover"]["distribution_statement"]

        it = p["items"][0]
        c.post(f"/api/items/{it['id']}/files", files={"files": ("12345-001_RevA.pdf", b"%PDF", "application/pdf")})
        assert c.put(f"/api/items/{it['id']}", json={"status": "delivered"}).json()["status"] == "delivered"
        assert c.put(f"/api/items/{it['id']}", json={"status": "bogus"}).status_code == 400

        z = zipfile.ZipFile(io.BytesIO(c.get(f"/api/packages/{p['id']}/export.zip").content))
        assert any("/Deliverables/A001 Product drawings" in n and n.endswith("12345-001_RevA.pdf") for n in z.namelist())
        doc = docx.Document(io.BytesIO(c.get(f"/api/packages/{p['id']}/export.docx").content))
        assert any(t.rows[0].cells[0].text == "CDRL" for t in doc.tables)

        # library
        e = c.post("/api/library", json={"category": "Company overview", "title": "Capability statement", "content": "We build training equipment."}).json()
        assert c.get("/api/library", params={"q": "training"}).json()[0]["id"] == e["id"]
        sec = p["sections"][1]
        c.put(f"/api/sections/{sec['id']}", json={"content": "Overview text for reuse."})
        saved = c.post(f"/api/sections/{sec['id']}/save-to-library", json={"category": "Technical"}).json()
        assert saved["content"] == "Overview text for reuse."
        assert c.delete(f"/api/library/{e['id']}").json()["ok"]


def test_ai_draft(monkeypatch):
    captured = {}

    class FakeMessages:
        def create(self, **kw):
            captured.update(kw)
            return types.SimpleNamespace(content=[types.SimpleNamespace(type="text", text="### Approach\nWe will staff the desk — daily.")])

    monkeypatch.setitem(sys.modules, "anthropic", types.SimpleNamespace(Anthropic=lambda api_key: types.SimpleNamespace(messages=FakeMessages())))
    monkeypatch.setattr(packages_api, "ANTHROPIC_API_KEY", "sk-test")
    with TestClient(app) as c:
        oid, _ = _opp_with_analysis(c)
        p = c.post("/api/packages", json={"kind": "proposal", "opportunity_id": oid}).json()
        tech = next(s for s in p["sections"] if s["title"] == "Technical approach")
        lib = c.post("/api/library", json={"title": "Help desk past work", "content": "Ran a 24/7 desk for Army."}).json()
        r = c.post(f"/api/sections/{tech['id']}/draft", json={"instructions": "Keep it short", "library_ids": [lib["id"]]}).json()
        assert "—" not in r["draft"] and "staff the desk" in r["draft"]
        prompt = captured["messages"][0]["content"]
        assert "staff the help desk" in prompt and "Ran a 24/7 desk" in prompt and "Keep it short" in prompt
        assert "no em dashes" in captured["system"]


def test_ai_draft_requires_key(monkeypatch):
    monkeypatch.setattr(packages_api, "ANTHROPIC_API_KEY", "")
    with TestClient(app) as c:
        p = c.post("/api/packages", json={"kind": "proposal"}).json()
        assert c.post(f"/api/sections/{p['sections'][0]['id']}/draft", json={}).status_code == 400
