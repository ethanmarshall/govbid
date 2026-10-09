import csv
import io
from datetime import date, timedelta

from docx import Document
from fastapi.testclient import TestClient

from app import finance_api as fa
from app.main import app

RATES = [{"start": "2026-07-01", "end": "2026-12-31", "rate": 4.75}]


# ------------------------------------------------------------------ prompt payment math
def test_due_date_receipt_later_than_acceptance():
    # delivered Mar 2, accepted Mar 4, invoice received Mar 10 -> 30 days after receipt
    r = fa.payment_due("2026-03-10", ship_date="2026-03-02", acceptance_date="2026-03-04")
    assert r["deemed_receipt"] == "2026-03-10"
    assert r["due_date"] == "2026-04-09"


def test_due_date_acceptance_later_than_receipt():
    # submitted with shipment Mar 2, accepted Mar 6 (within 7 days) -> 30 days after acceptance
    r = fa.payment_due("2026-03-02", ship_date="2026-03-02", acceptance_date="2026-03-06")
    assert r["acceptance_used"] == "2026-03-06"
    assert r["due_date"] == "2026-04-05"


def test_constructive_acceptance_seven_days():
    # no acceptance recorded yet: deemed on day 7 after delivery (estimate)
    r = fa.payment_due("2026-03-02", ship_date="2026-03-02")
    assert r["acceptance_used"] == "2026-03-09" and r["due_date"] == "2026-04-08" and r["estimate"]
    # late actual acceptance (day 20) does not push the due date out for Prompt Payment purposes
    late = fa.payment_due("2026-03-02", ship_date="2026-03-02", acceptance_date="2026-03-22")
    assert late["due_date"] == "2026-04-08"
    # unless disputed, then actual acceptance controls
    disputed = fa.payment_due("2026-03-02", ship_date="2026-03-02", acceptance_date="2026-03-22", disputed=True)
    assert disputed["due_date"] == "2026-04-21"
    # a longer contract acceptance period replaces the 7 days
    longer = fa.payment_due("2026-03-02", ship_date="2026-03-02", acceptance_period_days=15)
    assert longer["acceptance_used"] == "2026-03-17"


def test_override_and_not_submitted():
    assert fa.payment_due("")["due_date"] == ""
    assert fa.payment_due("2026-03-02", due_override="2026-03-20")["due_date"] == "2026-03-20"


def test_weekend_and_holiday_rule():
    assert date(2026, 11, 26) in fa.federal_holidays(2026)  # Thanksgiving
    assert date(2026, 7, 3) in fa.federal_holidays(2026)  # July 4 on Saturday observed Friday
    # due Saturday 2026-10-10 -> next working day is Tuesday 10-13 (Monday 10-12 is Columbus Day)
    assert fa.next_business_day(date(2026, 10, 10)) == date(2026, 10, 13)
    none = fa.late_interest(10000, "2026-10-10", "2026-10-13", RATES)
    assert none["days_late"] == 0 and none["interest"] == 0
    late = fa.late_interest(10000, "2026-10-10", "2026-10-14", RATES)
    assert late["days_late"] == 4


def test_interest_uses_stored_rate_and_compounds():
    # 45 days late at 4.75%: one 30-day compounding period then 15 simple days, 360-day year
    r = fa.late_interest(10000, "2026-08-01", "2026-09-15", RATES)
    expected = round(10000 * (1 + 0.0475 * 30 / 360) * (1 + 0.0475 * 15 / 360) - 10000, 2)
    assert r["days_late"] == 45 and r["rate"] == 4.75 and r["interest"] == expected
    # a different stored rate changes the result
    r2 = fa.late_interest(10000, "2026-08-01", "2026-09-15", [{"start": "2026-07-01", "end": "2026-12-31", "rate": 6.0}])
    assert r2["interest"] > r["interest"]
    # no stored rate: no number, explains why
    r3 = fa.late_interest(10000, "2026-08-01", "2026-09-15", [])
    assert r3["rate"] is None and "fiscal.treasury.gov" in r3["note"]
    # under $1 need not be paid
    tiny = fa.late_interest(100, "2026-08-03", "2026-08-05", RATES)
    assert 0 < tiny["interest"] < 1 and tiny["payable"] == 0


def test_amortization_formula():
    # $100,000, 6%, 120 months: standard textbook payment 1110.21
    assert fa.amortized_payment(100000, 6, 120) == 1110.21
    assert fa.amortized_payment(1200, 0, 12) == 100.0
    p, r, n = 50000, 0.0825 / 12, 84
    assert fa.amortized_payment(50000, 8.25, 84) == round(p * r / (1 - (1 + r) ** -n), 2)


# ------------------------------------------------------------------ API flows
def _job(c):
    from app.db import SessionLocal
    from app.models_jobs import Job
    with SessionLocal() as db:
        j = Job(title="Brackets", customer="DLA Land and Maritime", contract_number="SPE7M126P0001", delivery_order="",
                clins=[{"clin": "0001", "description": "Bracket", "quantity": 10, "unit": "EA", "unit_price": 42.5}],
                status="shipped", shipped_date="2026-03-02")
        db.add(j)
        db.commit()
        return j.id


def _job_status(job_id):
    from app.db import SessionLocal
    from app.models_jobs import Job
    with SessionLocal() as db:
        return db.get(Job, job_id).status


def test_invoice_flow_aging_exports_and_document():
    with TestClient(app) as c:
        c.put("/api/finance/settings", json={"invoice_prefix": "SDV-", "next_invoice_seq": 7})
        job_id = _job(c)
        inv = c.post("/api/finance/invoices", json={"job_id": job_id}).json()
        assert inv["number"] == "SDV-0007"
        assert inv["contract_number"] == "SPE7M126P0001" and inv["total"] == 425.0
        assert inv["lines"][0]["amount"] == 425.0 and inv["ship_date"] == "2026-03-02"

        sub = (date.today() - timedelta(days=75)).isoformat()
        inv = c.post(f"/api/finance/invoices/{inv['id']}/status", json={"status": "submitted", "date": sub}).json()
        assert _job_status(job_id) == "invoiced"
        assert inv["prompt_payment"]["due_date"]

        # a rejected invoice shows on the dashboard
        rej = c.post("/api/finance/invoices", json={"customer": "Prime Co", "lines": [{"clin": "1", "quantity": 1, "unit_price": 100}]}).json()
        c.post(f"/api/finance/invoices/{rej['id']}/status", json={"status": "rejected", "rejection_reason": "Wrong pay DoDAAC"})

        ag = c.get("/api/finance/aging").json()
        row = next(r for r in ag["invoices"] if r["id"] == inv["id"])
        assert row["submission_bucket"] == "61-90"
        assert ag["by_submission_age"]["61-90"]["amount"] == 425.0
        assert row["days_past_due"] > 0 and ag["total_past_due"] == 425.0
        assert any(r["id"] == rej["id"] for r in ag["rejected"])

        from app.db import SessionLocal
        with SessionLocal() as db:
            alerts = fa.dashboard_items(db)
            cal = fa.calendar_items(db)
        assert any("SDV-0007" in a and "past its payment due date" in a for a in alerts)
        assert any("rejected" in a for a in alerts)
        assert any(i["uid"] == f"fin-inv-{inv['id']}@govbid" for i in cal)

        paid_on = date.today().isoformat()
        paid = c.post(f"/api/finance/invoices/{inv['id']}/status", json={"status": "paid", "date": paid_on, "interest_paid": 1.5}).json()
        assert paid["amount_paid"] == 425.0 and _job_status(job_id) == "paid"
        ag = c.get("/api/finance/aging").json()
        assert ag["payment_stats"]["count"] == 1 and ag["payment_stats"]["average_days"] == 75

        # exports
        qbo = c.get("/api/finance/export/quickbooks-invoices.csv").text
        rows = list(csv.reader(io.StringIO(qbo)))
        assert rows[0] == fa.QBO_HEADERS
        assert rows[1][0] == "SDV-0007" and rows[1][-1] == "425.00"
        assert "/" in rows[1][2]  # MM/DD/YYYY
        wave = list(csv.reader(io.StringIO(c.get("/api/finance/export/wave-payments.csv").text)))
        assert wave[0] == ["Date", "Description", "Amount"] and wave[1][2] == "426.50"
        pay = list(csv.reader(io.StringIO(c.get("/api/finance/export/payments.csv").text)))
        assert pay[0] == fa.PAYMENT_HEADERS and pay[1][1] == "SDV-0007"
        gen = list(csv.reader(io.StringIO(c.get("/api/finance/export/invoices.csv").text)))
        assert gen[0] == fa.GENERIC_INVOICE_HEADERS
        # date range filter excludes everything outside it
        empty = list(csv.reader(io.StringIO(c.get("/api/finance/export/payments.csv?start=2000-01-01&end=2000-12-31").text)))
        assert len(empty) == 1

        # invoice document
        r = c.get(f"/api/finance/invoices/{inv['id']}/document.docx")
        assert r.status_code == 200
        doc = Document(io.BytesIO(r.content))
        text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
        assert "INVOICE" in text and "SDV-0007" in text and "SPE7M126P0001" in text and "$425.00" in text

        # duplicate number refused
        assert c.post("/api/finance/invoices", json={"number": "SDV-0007"}).status_code == 409


def test_loans_balance_schedule_and_use_of_funds():
    with TestClient(app) as c:
        loan = c.post("/api/finance/loans", json={"lender": "Main Street Bank", "program": "sba_7a", "principal": 100000,
                                                  "rate": 6, "term_months": 120, "start_date": "2026-01-15",
                                                  "first_payment_date": "2026-02-15"}).json()
        assert loan["computed_payment"] == 1110.21 and loan["next_payment_date"] == "2026-02-15"
        c.post(f"/api/finance/loans/{loan['id']}/draws", json={"date": "2026-01-15", "amount": 60000, "purpose": "Equipment"})
        c.post(f"/api/finance/loans/{loan['id']}/draws", json={"date": "2026-01-15", "amount": 40000, "purpose": "working capital"})
        loan = c.post(f"/api/finance/loans/{loan['id']}/payments", json={"date": "2026-02-15", "amount": 1110.21}).json()
        p = loan["payments"][0]
        assert p["interest"] == 500.0 and p["principal"] == 610.21 and loan["balance"] == 99389.79
        assert loan["next_payment_date"] == "2026-03-15"
        uof = {u["purpose"]: u["pct"] for u in loan["use_of_funds"]}
        assert uof == {"Equipment": 60.0, "Working capital": 40.0}
        sched = c.get(f"/api/finance/loans/{loan['id']}/schedule").json()
        assert len(sched["rows"]) == 120 and sched["rows"][0]["interest"] == 500.0
        assert sched["rows"][-1]["balance"] == 0
        assert abs(sched["rows"][-1]["payment"] - 1110.21) < 1

        loc = c.post("/api/finance/loans", json={"lender": "Credit Union", "program": "line_of_credit", "principal": 50000,
                                                 "rate": 9, "rate_type": "variable", "rate_base": "Prime + 1.5",
                                                 "start_date": "2026-01-01", "first_payment_date": "2026-02-01"}).json()
        assert loc["balance_basis"] == "draws" and not loc["amortizing"]
        loc = c.post(f"/api/finance/loans/{loc['id']}/draws", json={"date": "2026-01-01", "amount": 10000, "purpose": "Inventory and materials"}).json()
        assert loc["balance"] == 10000 and loc["available"] == 40000 and loc["computed_payment"] == 75.0

        exp = list(csv.reader(io.StringIO(c.get("/api/finance/export/loan-payments.csv").text)))
        assert exp[0] == fa.LOAN_PAYMENT_HEADERS and len(exp) == 2
