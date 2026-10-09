import io
import json

import docx
from fastapi.testclient import TestClient
from pypdf import PdfReader

from app import config, writing
from app.main import app

SOL = b"""SECTION L
L.1 The offeror shall submit a technical volume not to exceed 2 pages.
L.2 The offeror shall provide three past performance references.
L.3 The offeror shall submit pricing for all CLINs.
SECTION M
Factor 1: Technical Approach
Factor 2: Past Performance
Factor 3: Price
Award will be made on a best value tradeoff basis.
SECTION C
C.1 The contractor shall calibrate the temperature panels annually.
C.2 The contractor shall deliver monthly status reports.
"""


def _profile(c, **extra):
    body = {"name": "Marshall Tech LLC", "uei": "ABC123DEF456", "cage": "1A2B3", "sam_status": "active",
            "naics_codes": ["335314", "541330"], "psc_codes": ["6110"], "small_under_naics": {"335314": True, "541330": True},
            "certifications": {"SB": "certified", "SDVOSB": "pending", "VOSB": "none"}}
    body.update(extra)
    c.put("/api/profile", json=body)


def _seeded_package(c):
    _profile(c)
    oid = c.post("/api/opportunities", json={"title": "Panel Calibration", "solicitation_number": "N00001-26-Q-0001", "agency": "Navy"}).json()["id"]
    c.post(f"/api/opportunities/{oid}/documents", files={"files": ("rfq.txt", SOL, "text/plain")})
    a = c.post(f"/api/opportunities/{oid}/analyze").json()["analysis"]
    p = c.post("/api/packages", json={"kind": "proposal", "opportunity_id": oid}).json()
    tech = next(s for s in p["sections"] if s["title"] == "Technical approach")
    # Long enough to blow a 2 page limit at 500 words/page, with a leftover placeholder.
    body = ("We calibrate temperature panels annually using traceable standards. " * 160) + "\n\nLead technician: [insert name]. Schedule TBD."
    c.put(f"/api/sections/{tech['id']}", json={"content": body, "page_limit": 2})
    return oid, p["id"], tech["id"], a["compliance_matrix"]


def test_rule_based_review_findings():
    with TestClient(app) as c:
        oid, pid, tech_id, matrix = _seeded_package(c)
        run = c.post(f"/api/writing/review/{pid}").json()
        assert run["method"] == "rules"
        res = run["result"]
        assert res["overall"]["rating"] in writing.RATINGS and res["overall"]["estimate"] is True
        # placeholders flagged against the technical section
        assert any(f["section_id"] == tech_id and "placeholders" in f["action"].lower() and "[insert name]" in f["action"] for f in res["fixes"])
        # page limit exceeded
        assert any("page limit" in i for i in res["page_limit_issues"])
        # nothing marked covered, so requirements are missing or partial
        assert len(res["compliance"]) == len(matrix)
        assert all(r["status"] in ("missing", "partial") for r in res["compliance"])
        assert res["fixes"][0]["priority"] == 1
        assert {f["factor"] for f in res["factors"]} >= {"Factor 1: Technical Approach"}
        # empty sections produce deficiencies
        assert any("is empty" in d for f in res["factors"] for d in f["deficiencies"]) or any("empty" in f["action"] for f in res["fixes"])

        # runs are stored for comparison
        c.post(f"/api/writing/review/{pid}")
        hist = c.get(f"/api/writing/review/{pid}").json()
        assert len(hist["runs"]) == 2 and hist["latest"]["result"]["overall"]["rating"]
        one = c.get(f"/api/writing/review-run/{hist['runs'][1]['id']}").json()
        assert one["result"]["compliance"]


def test_placeholder_detection():
    hits = writing.find_placeholders("Per [L.3.2] we will [insert lead time]. See [COMPANY NAME] and TBD. A [link](http://x).")
    assert "[insert lead time]" in hits and "[COMPANY NAME]" in hits and "TBD" in hits
    assert "[L.3.2]" not in hits and "[link]" not in hits


def test_claude_review_path(monkeypatch):
    with TestClient(app) as c:
        oid, pid, tech_id, matrix = _seeded_package(c)
        monkeypatch.setattr(config, "ANTHROPIC_API_KEY", "sk-test")
        seen = {}

        def fake(system, user, max_tokens=8000):
            seen["user"] = user
            return "```json\n" + json.dumps({
                "overall": {"rating": "marginal", "summary": "Thin technical volume — needs work."},
                "factors": [{"factor": "Technical Approach", "rating": "Marginal", "strengths": ["Clear calibration plan"],
                             "weaknesses": ["No staffing detail"], "deficiencies": [], "risks": [], "section_ids": [tech_id, 99999]}],
                "compliance": [{"requirement_id": matrix[0]["id"], "status": "met", "where": "Technical", "note": ""}],
                "page_limit_issues": [], "fixes": [{"priority": 1, "section_id": tech_id, "action": "Add staffing."}],
            }) + "\n```"

        monkeypatch.setattr(writing, "call_claude", fake)
        run = c.post(f"/api/writing/review/{pid}").json()
        assert run["method"] == "claude"
        res = run["result"]
        assert res["overall"]["rating"] == "Marginal"
        assert "—" not in res["overall"]["summary"]
        assert res["factors"][0]["section_ids"] == [tech_id]  # unknown section id dropped
        assert any(f["action"] == "Add staffing." for f in res["fixes"])
        assert res["page_limit_issues"]  # measured page overage merged in from the rules
        assert "Section" in seen["user"] or "section" in seen["user"]

        monkeypatch.setattr(writing, "call_claude", lambda *a, **k: "Sorry, I cannot produce JSON today.")
        run = c.post(f"/api/writing/review/{pid}").json()
        assert run["method"] == "rules"
        assert "AI review failed" in run["result"]["overall"]["summary"]


def test_sources_sought_detection():
    assert writing.is_sources_sought("Sources Sought", "Valve repair", "")
    assert writing.is_sources_sought("Special Notice", "RFI - Trainer upgrades", "")
    assert writing.is_sources_sought("Presolicitation", "Relays", "Interested firms should submit a capability statement by 1 May.")
    assert writing.is_sources_sought("", "Widgets", "This is a Request for Information for market research purposes only.")
    assert not writing.is_sources_sought("Combined Synopsis/Solicitation", "Relays", "Offerors shall submit pricing.")
    assert not writing.is_sources_sought("Award Notice", "Sources sought award", "")


SS_DESC = """This is a SOURCES SOUGHT notice for market research purposes only.
Interested firms shall provide the following:
1. Company name, UEI, CAGE code, and point of contact.
2. Business size and socioeconomic status under NAICS 335314.
3. Describe your experience manufacturing relay control panels for training devices.
4. What is your typical lead time for a first article?
"""


def _ss_opp(c):
    return c.post("/api/opportunities", json={"title": "Relay control panels", "solicitation_number": "SPE7M1-26-R-0001",
                                              "agency": "DLA", "naics": "335314", "description": SS_DESC}).json()["id"]


def test_sources_sought_template_and_save():
    with TestClient(app) as c:
        _profile(c)
        c.post("/api/past-performance", json={"title": "Relay control panel trainer build", "customer": "Navy NPTU", "role": "sub",
                                              "description": "Built relay control panels for a training device.", "relevance_keywords": ["relay", "control panels"]})
        oid = _ss_opp(c)
        info = c.get(f"/api/writing/sources-sought/{oid}").json()
        assert info["applies"] is True
        assert len(info["requests"]) == 4

        d = c.post(f"/api/writing/sources-sought/{oid}").json()
        md = d["markdown"]
        assert d["method"] == "template"
        assert "VetCert application pending" in md
        assert "SBA-certified" not in md
        assert "ABC123DEF456" in md and "1A2B3" in md
        assert "Small business under NAICS 335314" in md
        assert "is registered in our SAM profile" in md
        assert "Relay control panel trainer build" in md and "Subcontractor" in md
        assert "[Contact name]" in md  # no POC saved yet, left as a blank
        asked = [x for x in d["checklist"] if x["source"] == "notice"]
        assert len(asked) == 4 and all(x["status"] != "missing" for x in asked)
        assert any(x["item"] == "UEI stated" and x["status"] == "covered" for x in d["checklist"])

        # an edited draft that claims a pending certification gets flagged
        chk = c.post(f"/api/writing/sources-sought/{oid}/checklist", json={"markdown": md + "\nWe are a certified SDVOSB."}).json()["checklist"]
        assert any(x["source"] == "accuracy" for x in chk)

        r = c.post(f"/api/writing/sources-sought/{oid}/docx", json={"markdown": md})
        assert r.status_code == 200
        text = "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
        assert "Company information" in text

        saved = c.post(f"/api/writing/sources-sought/{oid}/save", json={"markdown": md}).json()
        pkg = c.get(f"/api/packages/{saved['package_id']}").json()
        assert pkg["name"] == "SPE7M1-26-R-0001 sources sought response"
        assert pkg["kind"] == "proposal" and pkg["opportunity_id"] == oid
        assert len(pkg["sections"]) == 1 and "VetCert application pending" in pkg["sections"][0]["content"]


def test_sources_sought_claude(monkeypatch):
    with TestClient(app) as c:
        _profile(c)
        oid = _ss_opp(c)
        monkeypatch.setattr(config, "ANTHROPIC_API_KEY", "sk-test")
        captured = {}

        def fake(system, user, max_tokens=8000):
            captured["user"] = user
            return "# Response\n\n## Company\nUEI ABC123DEF456, CAGE 1A2B3 — VetCert application pending."

        monkeypatch.setattr(writing, "call_claude", fake)
        d = c.post(f"/api/writing/sources-sought/{oid}").json()
        assert d["method"] == "claude" and "—" not in d["markdown"]
        assert "VetCert application pending (not yet certified as an SDVOSB)" in captured["user"]


def test_capability_statement_pdf_and_docx(tmp_path):
    with TestClient(app) as c:
        _profile(c)
        pp = c.post("/api/past-performance", json={"title": "Temperature monitoring panels", "customer": "Naval lab", "role": "employment",
                                                   "description": "RTD panels with lead length compensation. " * 8}).json()
        s = c.put("/api/writing/capability", json={
            "tagline": "Electrical control panels and test equipment for training and industry",
            "overview": "Veteran-owned engineering and manufacturing firm.",
            "competencies": ["Relay and industrial control panels", "RTD temperature monitoring", "Wire harness assembly", " "],
            "differentiators": ["Nuclear power plant experience"], "past_performance_ids": [pp["id"]],
            "contact_name": "Ethan Marshall", "contact_title": "Owner", "contact_phone": "555-0100", "contact_email": "info@example.com"}).json()
        assert s["competencies"] == ["Relay and industrial control panels", "RTD temperature monitoring", "Wire harness assembly"]
        info = c.get("/api/writing/capability").json()
        assert "335314 Relay and Industrial Control Manufacturing" in info["company"]["naics"]
        assert "SDVOSB: VetCert application pending" in info["company"]["certifications"]

        r = c.get("/api/writing/capability.pdf")
        assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
        reader = PdfReader(io.BytesIO(r.content))
        assert len(reader.pages) == 1
        text = reader.pages[0].extract_text()
        assert "ABC123DEF456" in text and "1A2B3" in text
        assert "Temperature monitoring panels" in text
        assert r.headers["X-Fit-Overflow"] == "0"

        # Overflowing content: shrinks, then warns, but stays on one page
        big = {"competencies": [f"Competency number {i} with a fairly long description of the work" for i in range(60)]}
        r = c.post("/api/writing/capability/pdf", json={"settings": {**{k: s[k] for k in ("tagline", "overview", "differentiators", "past_performance_ids")}, **big}})
        assert r.headers["X-Fit-Shrunk"] == "1" and r.headers["X-Fit-Overflow"] == "1"
        assert len(PdfReader(io.BytesIO(r.content)).pages) == 1

        # Logo upload
        from PIL import Image
        buf = io.BytesIO()
        Image.new("RGB", (200, 80), "navy").save(buf, "PNG")
        assert c.post("/api/writing/capability/logo", files={"file": ("logo.png", buf.getvalue(), "image/png")}).json()["logo_file"] == "logo.png"
        assert c.post("/api/writing/capability/logo", files={"file": ("logo.gif", b"GIF89a", "image/gif")}).status_code == 400
        assert len(PdfReader(io.BytesIO(c.get("/api/writing/capability.pdf").content)).pages) == 1

        r = c.get("/api/writing/capability.docx")
        assert r.status_code == 200
        d = docx.Document(io.BytesIO(r.content))
        all_text = "\n".join(cell.text for t in d.tables for row in t.rows for cell in row.cells)
        assert "ABC123DEF456" in all_text and "CORE COMPETENCIES" in all_text and "VetCert application pending" in all_text
        c.delete("/api/writing/capability/logo")


def test_capability_tailoring():
    with TestClient(app) as c:
        _profile(c)
        c.put("/api/writing/capability", json={"tagline": "Panels", "competencies": ["Machining", "Wire harness assembly", "Relay control panels"]})
        c.post("/api/past-performance", json={"title": "Relay control panel build", "customer": "Navy", "role": "prime",
                                              "relevance_keywords": ["relay", "control panels"]})
        oid = _ss_opp(c)
        t = c.post("/api/writing/capability/tailor", json={"opp_id": oid}).json()
        assert t["method"] == "keywords"
        assert t["competencies"][0] == "Relay control panels"
        assert t["past_performance_ids"]
        assert any(o["id"] == oid for o in c.get("/api/writing/capability/opportunities").json())
        r = c.get(f"/api/writing/capability.pdf?opp_id={oid}")
        assert "Prepared for" in PdfReader(io.BytesIO(r.content)).pages[0].extract_text()
