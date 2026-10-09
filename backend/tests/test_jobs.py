import copy
from datetime import date, timedelta

from fastapi.testclient import TestClient

from app import pricing
from app.main import app


def _d(days: int) -> str:
    return (date.today() + timedelta(days=days)).isoformat()


def _make_opp(**kw) -> int:
    from app.db import SessionLocal
    from app.models import Opportunity

    with SessionLocal() as db:
        o = Opportunity(source="manual", external_id=f"jobs-test-{kw.get('title', 'x')}-{date.today()}-{id(kw)}",
                        solicitation_number=kw.get("sol", "SPE7M1-26-T-0001"), title=kw.get("title", "Bracket, mounting"),
                        agency=kw.get("agency", "DLA Land and Maritime"), nsn=kw.get("nsn", "5340-01-777-1234"), quantity=kw.get("quantity", "40"))
        db.add(o)
        db.commit()
        return o.id


def test_job_crud_and_clins_value():
    with TestClient(app) as c:
        r = c.post("/api/jobs", json={"title": "Cable assembly W1", "customer": "NAVSEA",
                                      "clins": [{"clin": "0001", "quantity": 10, "unit_price": 125.5, "due_date": _d(30)},
                                                {"clin": "0002", "quantity": "2", "unit_price": "40", "due_date": _d(20)}]})
        assert r.status_code == 200, r.text
        j = r.json()
        assert j["value"] == 1335.0 and j["clin_total"] == 1335.0
        assert j["due_date"] == _d(20)  # earliest CLIN due date
        assert j["status"] == "awarded" and j["late"] is False
        jid = j["id"]
        assert c.post("/api/jobs", json={"status": "nope"}).status_code == 422
        assert c.post("/api/jobs", json={"due_date": "1/2/2026"}).status_code == 422
        assert c.put(f"/api/jobs/{jid}", json={"fob": "sideways"}).status_code == 422

        r = c.put(f"/api/jobs/{jid}", json={"contract_number": "N00024-26-P-1234", "fob": "destination", "inspection_acceptance": "origin",
                                            "carrier": "UPS", "tracking_number": "1Z999", "packaging_notes": "MIL-STD-129 marking", "iuid_required": True})
        j = r.json()
        assert j["contract_number"] == "N00024-26-P-1234" and j["iuid_required"] is True and j["fob"] == "destination"
        # shipping moves an open job to shipped
        j = c.put(f"/api/jobs/{jid}", json={"shipped_date": _d(10)}).json()
        assert j["status"] == "shipped" and j["on_time"] is True
        assert any(x["id"] == jid for x in c.get("/api/jobs", params={"q": "N00024"}).json())
        assert not any(x["id"] == jid for x in c.get("/api/jobs", params={"status": "awarded"}).json())
        assert c.delete(f"/api/jobs/{jid}").json()["ok"]
        assert c.get(f"/api/jobs/{jid}").status_code == 404


def test_compute_metrics():
    from app.jobs_api import compute_metrics
    from app.models_jobs import Job

    today = date(2026, 6, 1)
    jobs = [
        Job(status="shipped", due_date="2026-05-10", shipped_date="2026-05-09", value=100.0, clins=[]),  # on time
        Job(status="paid", due_date="2026-05-10", shipped_date="2026-05-10", value=100.0, clins=[]),  # on time (same day)
        Job(status="shipped", due_date="2026-05-10", shipped_date="2026-05-12", value=100.0, clins=[]),  # late
        Job(status="in_work", due_date="2026-05-20", shipped_date="", value=None, clins=[{"quantity": 5, "unit_price": 10}]),  # late, open
        Job(status="awarded", due_date="2026-07-01", shipped_date="", value=1000.0, clins=[]),
        Job(status="cancelled", due_date="2026-05-01", shipped_date="2026-05-30", value=999.0, clins=[]),
    ]
    m = compute_metrics(jobs, today)
    assert m["shipped_measured"] == 3 and m["on_time"] == 2
    assert m["on_time_pct"] == 66.7
    assert m["open_count"] == 2 and m["open_value"] == 1050.0
    assert m["late_count"] == 1
    assert m["by_status"]["shipped"] == 2 and m["by_status"]["cancelled"] == 1
    assert compute_metrics([], today)["on_time_pct"] is None


def test_from_opportunity_flow():
    with TestClient(app) as c:
        oid = _make_opp(title="Bracket, mounting (jobs)", nsn="5340-01-777-1234", quantity="25")
        spec = copy.deepcopy(pricing.EXAMPLE_SPEC)
        spec.update(nsn="5340-01-777-1234", quantities=[25], inspection={"first_article": True})
        q = c.post("/api/pricing/quotes", json={"spec": spec, "opportunity_id": oid, "quoted_quantity": 25, "status": "submitted"}).json()
        assert q["quoted_unit_price"]
        c.put(f"/api/opportunities/{oid}/pipeline", json={"stage": "submitted"})

        r = c.post(f"/api/jobs/from-opportunity/{oid}", json={"award_date": "2026-10-01"})
        assert r.status_code == 200, r.text
        j = r.json()
        assert j["customer"] == "DLA Land and Maritime" and j["contract_number"] == ""
        assert j["part_quote_id"] == q["id"] and j["opportunity_id"] == oid
        clin = j["clins"][0]
        assert clin["clin"] == "0001" and clin["nsn"] == "5340-01-777-1234" and clin["quantity"] == 25
        assert clin["unit_price"] == q["quoted_unit_price"] and clin["part_number"] == "12345-001"
        assert j["value"] == round(25 * q["quoted_unit_price"], 2)
        lead = next(b["lead_time_days"] for b in q["price_breaks"] if b["quantity"] == 25)
        assert j["due_date"] == (date(2026, 10, 1) + timedelta(days=lead)).isoformat()
        assert j["packaging_level"] == "mil_std_2073"
        names = [o["name"] for o in j["operations"]]
        assert names[0].startswith("Contract review") and "CNC milling" in names
        assert any("anodize" in n for n in names) and any(n.startswith("First article") for n in names)
        assert names[-1].startswith("Ship")

        assert c.get(f"/api/opportunities/{oid}").json()["pipeline"]["stage"] == "won"
        assert c.get(f"/api/pricing/quotes/{q['id']}").json()["status"] == "won"
        h = c.get("/api/nsn/5340-01-777-1234/history").json()
        assert any(rec["source"] == "quote_won" for rec in h["records"])

        assert c.post(f"/api/jobs/from-opportunity/{oid}").status_code == 409
        assert c.post("/api/jobs/from-opportunity/999999").status_code == 404

        # no quote: standard traveler, quantity from the opportunity
        oid2 = _make_opp(title="Training panel", quantity="3 EA")
        j2 = c.post(f"/api/jobs/from-opportunity/{oid2}").json()
        assert j2["part_quote_id"] is None and j2["clins"][0]["quantity"] == 3 and j2["due_date"] == ""
        assert len(j2["operations"]) == 8
        assert c.get(f"/api/opportunities/{oid2}").json()["pipeline"]["stage"] == "won"
        c.delete(f"/api/jobs/{j['id']}")
        c.delete(f"/api/jobs/{j2['id']}")
        c.delete(f"/api/pricing/quotes/{q['id']}")


def test_traveler_operations():
    with TestClient(app) as c:
        jid = c.post("/api/jobs", json={"title": "Traveler test"}).json()["id"]
        assert c.post(f"/api/jobs/{jid}/traveler/template", json={"template": "quote"}).status_code == 422  # no quote linked
        ops = c.post(f"/api/jobs/{jid}/traveler/template", json={"template": "standard"}).json()
        assert len(ops) == 8 and all(o["status"] == "not_started" for o in ops)
        ops = c.post(f"/api/jobs/{jid}/traveler/template", json={"template": "blank"}).json()
        assert ops == []
        a = c.post(f"/api/jobs/{jid}/operations", json={"name": "Cut wire", "work_center": "Bench"}).json()
        b = c.post(f"/api/jobs/{jid}/operations", json={"name": "Crimp contacts"}).json()
        assert (a["position"], b["position"]) == (0, 1)
        assert c.put(f"/api/jobs/{jid}/operations/{a['id']}", json={"status": "bogus"}).status_code == 422
        u = c.put(f"/api/jobs/{jid}/operations/{a['id']}", json={"signoff_initials": "EM", "qty_good": 10, "qty_rejected": 1}).json()
        assert u["status"] == "done" and u["signoff_date"] == date.today().isoformat() and u["qty_rejected"] == 1
        h = c.put(f"/api/jobs/{jid}/operations/{b['id']}", json={"status": "hold", "notes": "Waiting on contacts"}).json()
        assert h["status"] == "hold"
        ops = c.post(f"/api/jobs/{jid}/operations/reorder", json={"ids": [b["id"], a["id"]]}).json()
        assert [o["id"] for o in ops] == [b["id"], a["id"]]
        assert c.post(f"/api/jobs/{jid}/operations/reorder", json={"ids": [a["id"]]}).status_code == 422
        j = c.get(f"/api/jobs/{jid}").json()
        assert j["traveler_progress"] == {"done": 1, "total": 2}
        assert c.delete(f"/api/jobs/{jid}/operations/{b['id']}").json()["ok"]
        c.delete(f"/api/jobs/{jid}")


def test_purchases_calendar_dashboard():
    from app.db import SessionLocal
    from app.jobs_api import calendar_items, dashboard_items

    with TestClient(app) as c:
        org = c.post("/api/crm/organizations", json={"name": "Jobs Test Metals", "kind": "vendor"}).json()
        jid = c.post("/api/jobs", json={"title": "Late job test", "due_date": _d(-3)}).json()["id"]
        assert c.post(f"/api/jobs/{jid}/purchases", json={"description": "no vendor"}).status_code == 422
        assert c.post(f"/api/jobs/{jid}/purchases", json={"organization_id": 999999}).status_code == 404
        p = c.post(f"/api/jobs/{jid}/purchases", json={"organization_id": org["id"], "po_number": "PO-1001", "description": "6061 plate",
                                                      "quantity": 5, "unit_price": 20, "promised_date": _d(-1)}).json()
        assert p["vendor_name"] == "Jobs Test Metals" and p["extended"] == 100.0 and p["overdue"] is True
        p2 = c.post(f"/api/jobs/{jid}/purchases", json={"vendor_name": "Anodize Co", "po_number": "PO-1002", "promised_date": _d(5)}).json()
        with SessionLocal() as db:
            dash = dashboard_items(db)
            cal = calendar_items(db)
        assert any("Late job test" in s for s in dash) and any("PO-1001" in s for s in dash)
        assert not any("PO-1002" in s for s in dash)
        uids = {e["uid"] for e in cal}
        assert f"job-due-{jid}@govbid" in uids and f"job-po-{p2['id']}@govbid" in uids
        assert all(isinstance(e["date"], date) for e in cal)

        r = c.put(f"/api/jobs/{jid}/purchases/{p['id']}", json={"received_date": _d(0), "qty_rejected": 1, "certs_received": True}).json()
        assert r["qty_accepted"] == 4 and r["on_time"] is False and r["overdue"] is False
        assert any(x["id"] == p["id"] for x in c.get("/api/jobs/purchases", params={"organization_id": org["id"]}).json())
        assert not any(x["id"] == p["id"] for x in c.get("/api/jobs/purchases", params={"open_only": True}).json())
        assert c.delete(f"/api/jobs/{jid}/purchases/{p2['id']}").json()["ok"]
        assert len(c.get(f"/api/jobs/{jid}").json()["purchases"]) == 1
        c.delete(f"/api/jobs/{jid}")
        c.delete(f"/api/crm/organizations/{org['id']}")


def test_records_and_certificate_of_conformance():
    with TestClient(app) as c:
        from app.db import SessionLocal
        from app.services import get_profile
        with SessionLocal() as db:
            prof = get_profile(db)
            old = (prof.name, prof.cage)
            prof.name, prof.cage = "Marshall Precision LLC", "9ZZ99"
            db.commit()
        jid = c.post("/api/jobs", json={"title": "C of C test", "contract_number": "SPE7M1-26-P-0042", "carrier": "FedEx",
                                        "tracking_number": "7788", "shipped_date": "2026-09-30",
                                        "clins": [{"clin": "0001", "nsn": "5340-01-777-1234", "part_number": "12345-001",
                                                   "description": "Bracket", "quantity": 25, "unit_price": 50}]}).json()["id"]
        r = c.post(f"/api/jobs/{jid}/records", data={"doc_type": "material_cert", "clin": "0001"},
                   files={"file": ("mill cert.pdf", b"%PDF-1.4 test", "application/pdf")})
        assert r.status_code == 200, r.text
        rec = r.json()
        assert rec["filename"] == "mill_cert.pdf" and rec["doc_type"] == "material_cert"
        assert c.get(rec["url"]).content == b"%PDF-1.4 test"
        assert c.post(f"/api/jobs/{jid}/records", data={"doc_type": "nope"}, files={"file": ("a.txt", b"x")}).status_code == 422

        r = c.get(f"/api/jobs/{jid}/coc", params={"clin": "0001", "fmt": "docx"})
        assert r.status_code == 200 and r.content[:2] == b"PK"
        import io
        import docx
        text = "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
        cells = " ".join(cell.text for t in docx.Document(io.BytesIO(r.content)).tables for row in t.rows for cell in row.cells)
        assert "Marshall Precision LLC" in text and "SPE7M1-26-P-0042" in text and "conform in all respects" in text
        assert "9ZZ99" in cells and "5340-01-777-1234" in cells and "12345-001" in cells and "25 EA" in cells
        assert "Signature" in text

        r = c.get(f"/api/jobs/{jid}/coc", params={"fmt": "pdf", "save": True})
        assert r.status_code == 200 and r.content[:4] == b"%PDF"
        recs = c.get(f"/api/jobs/{jid}").json()["records"]
        assert any(x["doc_type"] == "coc" and x["filename"].endswith(".pdf") for x in recs)
        assert c.get(f"/api/jobs/{jid}/coc", params={"clin": "9999"}).status_code == 404
        assert c.get(f"/api/jobs/{jid}/coc", params={"fmt": "rtf"}).status_code == 422

        assert c.delete(f"/api/jobs/{jid}/records/{rec['id']}").json()["ok"]
        assert c.get(rec["url"]).status_code == 404
        c.delete(f"/api/jobs/{jid}")
        with SessionLocal() as db:
            prof = get_profile(db)
            prof.name, prof.cage = old
            db.commit()


def test_meta_and_metrics_endpoint():
    with TestClient(app) as c:
        m = c.get("/api/jobs/meta").json()
        assert "in_work" in m["statuses"] and "hold" in m["operation_statuses"] and "coc" in m["record_types"]
        r = c.get("/api/jobs/metrics").json()
        assert {"by_status", "on_time_pct", "open_value", "late_count"} <= set(r)
