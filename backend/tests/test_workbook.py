import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.workbook_api import backsolve, burdened_rate, calculate_rates, compute_build, los_check, merged_rates

RATES = {"fringe_pct": 30, "overhead_pct": 40, "ga_pct": 12, "profit_pct": 8, "material_handling_pct": 5,
         "overhead_base": "labor_fringe", "ga_base": "total_cost_input", "fee_base": "total_cost"}

CLIN = {
    "clin": "0001", "description": "Trainer design", "kind": "services", "quantity": 1,
    "labor": [{"category": "Engineer", "rate": 50, "hours": 100}],
    "materials": [{"description": "Steel", "cost": 1000}],
    "subcontracts": [{"name": "Paint shop", "cost": 2000, "similarly_situated": False}],
    "odcs": [{"description": "Trip", "cost": 500, "kind": "travel"}],
    "parts": [{"description": "Bracket", "quantity": 10, "unit_price": 25}],
}


def test_build_up_matches_hand_calculation():
    # By hand:
    #   DL 100 h x $50 = 5,000; fringe 30% = 1,500; overhead 40% x 6,500 = 2,600
    #   materials 1,000; handling 5% = 50; subcontracts 2,000; travel 500
    #   TCI = 12,650; G&A 12% = 1,518; total cost 14,168; profit 8% = 1,133.44
    #   part line 10 x 25 = 250; price = 15,551.44
    c = compute_build({"clins": [CLIN]}, RATES, {})["clins"][0]
    assert c["direct_labor"] == 5000
    assert c["fringe"] == 1500
    assert c["overhead"] == 2600
    assert c["material_handling"] == 50
    assert c["total_cost_input"] == 12650
    assert c["ga"] == 1518
    assert c["total_cost"] == 14168
    assert c["profit"] == pytest.approx(1133.44)
    assert c["parts"] == 250
    assert c["price"] == pytest.approx(15551.44)


def test_configurable_bases():
    r = dict(RATES, overhead_base="labor", ga_base="value_added", fee_base="cost_less_pass_through")
    c = compute_build({"clins": [CLIN]}, r, {})["clins"][0]
    assert c["overhead"] == 2000  # 40% x 5,000
    # TCI = 5000+1500+2000+1000+50+2000+500 = 12,050; value added = 12,050-1,000-50-2,000 = 9,000
    assert c["ga"] == 1080
    # fee base = (12,050 + 1,080) - 3,050 = 10,080
    assert c["profit"] == pytest.approx(806.4)


def test_burdened_rate():
    b = burdened_rate(50, merged_rates(RATES))
    # 50 + 15 fringe + 26 OH = 91; G&A 10.92 -> 101.92; profit 8.1536 -> 110.07
    assert b["fully_burdened_cost"] == pytest.approx(101.92)
    assert b["billing_rate"] == pytest.approx(110.07)


def test_rate_calculator():
    r = calculate_rates(direct_labor=200_000, fringe_costs=60_000, overhead_costs=52_000, ga_costs=40_000,
                        materials=88_000)
    assert r["fringe_pct"] == 30.0
    assert r["overhead_pct"] == 20.0  # 52,000 / 260,000
    assert r["ga_pct"] == 10.0  # 40,000 / 400,000 TCI
    with pytest.raises(ValueError):
        calculate_rates(0, 1, 1, 1)


def test_backsolve_profit():
    c = compute_build({"clins": [CLIN]}, RATES, {})["totals"]
    bs = backsolve(c["total_cost"], c["parts"], c["fee_base"], 15_000)
    # profit = 15,000 - 250 - 14,168 = 582 -> 582 / 14,168 = 4.11%
    assert bs["profit"] == 582
    assert bs["profit_pct"] == pytest.approx(4.11, abs=0.01)
    assert not bs["negative"]
    low = backsolve(c["total_cost"], c["parts"], c["fee_base"], 14_000)
    assert low["negative"] and low["profit"] == -418 and "lose" in low["warning"]


def test_los_services_and_supplies():
    svc = {"kind": "services", "price": 10_000, "materials": 0, "non_ss_subcontracts": 6_000}
    sup = {"kind": "supplies", "price": 10_000, "materials": 4_000, "non_ss_subcontracts": 2_500}
    out = los_check([svc, sup], True)
    p = {x["kind"]: x for x in out["portions"]}
    assert p["services"]["pct"] == 60.0 and p["services"]["over"]
    # supplies: 2,500 / (10,000 - 4,000) = 41.67%
    assert p["supplies"]["pct"] == pytest.approx(41.67)
    assert not p["supplies"]["over"]
    assert len(out["warnings"]) == 1 and "52.219-14" in out["warnings"][0]
    assert los_check([svc], False)["warnings"] == []


def test_similarly_situated_and_lower_tier():
    clin = dict(CLIN, subcontracts=[
        {"name": "SDVOSB partner", "cost": 3000, "similarly_situated": True, "lower_tier_non_ss": 1000},
        {"name": "Big shop", "cost": 2000, "similarly_situated": False},
    ])
    c = compute_build({"clins": [clin]}, RATES, {})["clins"][0]
    assert c["non_ss_subcontracts"] == 3000


def test_api_build_export_and_opportunity_summary():
    with TestClient(app) as cl:
        cat = cl.post("/api/workbook/labor", json={"name": "Mechanical Engineer", "annual_salary": 104_000,
                                                   "calc_category": "Engineer III"}).json()
        assert cat["hourly_rate"] == 50.0
        cl.put("/api/workbook/rates", json=RATES)
        from app.db import SessionLocal
        from app.models import Opportunity, PartQuote
        with SessionLocal() as db:
            o = Opportunity(source="manual", external_id="wb-1", solicitation_number="W123-WB", title="Trainer",
                            set_aside_code="SDVOSBC")
            pq = PartQuote(name="Bracket", quoted_unit_price=25.0, quoted_quantity=10)
            db.add_all([o, pq])
            db.commit()
            opp_id, pq_id = o.id, pq.id
        clin = dict(CLIN, labor=[{"category_id": cat["id"], "hours": 100}],
                    parts=[{"part_quote_id": pq_id, "description": "Bracket", "quantity": 10, "unit_price": 25}])
        b = cl.post("/api/workbook/builds", json={"opportunity_id": opp_id, "clins": [clin],
                                                  "competitors": [{"name": "Incumbent", "source": "usaspending", "prices": {"0001": 15000}}]}).json()
        assert b["set_aside"] is True
        assert b["computed"]["totals"]["price"] == pytest.approx(15551.44)
        gap = b["computed"]["competitors"][0]["gaps"]["0001"]
        assert gap["gap"] == pytest.approx(551.44)
        assert b["computed"]["los"]["applies"]

        bs = cl.post(f"/api/workbook/builds/{b['id']}/backsolve", json={"target_price": 15000}).json()
        assert bs["profit"] == 582

        s = cl.get(f"/api/workbook/by-opportunity/{opp_id}").json()
        assert s["latest"]["total_price"] == pytest.approx(15551.44)

        r = cl.get(f"/api/workbook/builds/{b['id']}/export.xlsx")
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert {"Summary", "Rates", "Labor", "CLIN 0001"} <= set(wb.sheetnames)
        ws = wb["CLIN 0001"]
        formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
        assert any("Rates!$B$2" in f for f in formulas)  # fringe uses the rate cell
        assert any("Rates!$B$4" in f for f in formulas)  # G&A
        summary = wb["Summary"]
        sformulas = [c.value for row in summary.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
        assert any(f.startswith("=SUM(") for f in sformulas)
        assert any("SUMIF" in f for f in sformulas)  # LOS check
        assert any("'CLIN 0001'!" in f for f in sformulas)
        assert wb["Labor"]["B2"].value == 50.0 and wb["Labor"]["C2"].value.startswith("=")

        cl.delete(f"/api/workbook/builds/{b['id']}")
        assert cl.get(f"/api/workbook/builds/{b['id']}").status_code == 404
