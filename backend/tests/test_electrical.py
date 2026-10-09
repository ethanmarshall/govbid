"""Electrical quoters: harness parsing and validation, labor math, panel wire estimate, label pricing, save/restore, sheets."""
import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import electrical, pricing
from app.main import app

FIX = Path(__file__).parent / "fixtures" / "electrical"


def _w(wires, wid):
    return next(w for w in wires if w["wire_id"] == wid)


# ---------------------------------------------------------------- harness parsing
def test_split_end_and_spec():
    assert electrical.split_end("P1-3") == ("P1", "3")
    assert electrical.split_end("J2:A") == ("J2", "A")
    assert electrical.split_end("TB1 PIN 12") == ("TB1", "12")
    assert electrical.split_end("E1") == ("E1", "")
    assert electrical.parse_spec("M22759/16-22-9") == ("M22759/16", "22")
    assert electrical.parse_spec("UL 1007")[0] == "UL1007"
    assert electrical.parse_gauge("#20 AWG") == "20"


def test_parse_from_to_csv():
    r = electrical.parse_file("harness", FIX / "harness_wires.csv")
    wires = r["wires"]
    assert len(wires) == 6 and r["bom"] == []
    w4 = _w(wires, "W4")
    assert (w4["from_ref"], w4["from_pin"], w4["to_ref"], w4["to_pin"]) == ("P1", "4", "TB1", "1")
    assert w4["gauge"] == "22" and w4["length_in"] == 24.0 and w4["spec"] == "M22759/16"
    assert w4["shielded"] and w4["twisted"]
    assert _w(wires, "W1")["spec"] == "M22759/16" and _w(wires, "W6")["spec"] == "UL1007"
    assert any("No connector list" in x for x in r["warnings"])


def test_parse_xlsx_two_sheets_and_feet():
    r = electrical.parse_file("harness", FIX / "harness.xlsx")
    assert [w["wire_id"] for w in r["wires"]] == ["101", "102"]
    assert _w(r["wires"], "102")["length_in"] == 30.0  # 2.5 ft column
    assert {(b["ref"], b["kind"], b["unit_price"]) for b in r["bom"]} == {("P1", "connector", 0.42), ("P2", "connector", 0.40)}
    assert r["warnings"] == []


def test_parse_drawing_pdf_and_validation_warnings():
    r = electrical.parse_file("harness", FIX / "harness_drawing.pdf")
    assert len(r["wires"]) == 6
    kinds = {(b["ref"], b["kind"]) for b in r["bom"]}
    assert {("P1", "connector"), ("P1", "contact"), ("P1", "backshell"), ("J2", "connector"), ("J2", "contact"), ("", "accessory")} == kinds
    p1 = next(b for b in r["bom"] if b["ref"] == "P1" and b["kind"] == "connector")
    assert p1["contact_size"] == "22D"
    j2c = next(b for b in r["bom"] if b["ref"] == "J2" and b["kind"] == "contact")
    assert j2c["contact_size"] == "20" and j2c["qty"] is None  # count comes from the wire list
    assert r["drawing"]["revision"] == "B"
    w = " ".join(r["warnings"])
    assert "P9 is not in the connector list (used by W6)" in w
    assert "P1 pin 3 is used by 2 wires (W3, W6)" in w
    assert "W5 is 18 AWG but P1 uses size 22D contacts (22-28 AWG)" in w
    assert "W3 is 20 AWG but P1 uses size 22D contacts" in w
    assert "W1 is" not in w  # 22 AWG fits 22D


def test_gauge_vs_contact_size_table():
    wires = [{"wire_id": "A", "from_ref": "P1", "from_pin": "1", "to_ref": "P2", "to_pin": "1", "gauge": "14"}]
    bom = [{"ref": "P1", "kind": "connector", "contact_size": "12"}, {"ref": "P2", "kind": "connector", "contact_size": "16"}]
    w = electrical.validate_harness(wires, bom)
    assert any("P2 uses size 16 contacts (16-20 AWG)" in x for x in w)
    assert not any("P1 uses" in x for x in w)  # 14 AWG fits size 12 (12-14)


# ---------------------------------------------------------------- harness labor math by hand
def test_tiny_harness_labor_math():
    wires = [
        {"wire_id": "W1", "from_ref": "P1", "from_pin": "1", "to_ref": "P2", "to_pin": "1", "gauge": "22", "spec": "M22759/16", "length_in": 12},
        {"wire_id": "W2", "from_ref": "P1", "from_pin": "2", "to_ref": "P2", "to_pin": "2", "gauge": "22", "spec": "M22759/16", "length_in": 24,
         "to_term": "solder"},
    ]
    bom = [
        {"ref": "P1", "kind": "connector", "part_number": "C1", "qty": 1, "unit_price": 10},
        {"ref": "P2", "kind": "connector", "part_number": "C2", "qty": 1, "unit_price": 8},
        {"ref": "P1", "kind": "contact", "part_number": "K1", "qty": None, "unit_price": 0.5},
    ]
    opts = {"workmanship": "class_2", "branches": 2, "ties": 4, "hipot": False}
    r = electrical.price_harness(wires, bom, None, [1, 10], opts)
    c = r["counts"]
    assert (c["ends"], c["crimp"], c["insert"], c["solder"], c["heat_shrink"], c["markers"]) == (4, 3, 3, 1, 1, 4)
    # 4 ends x 0.35 + 3 crimps x 0.6 + 3 inserts x 0.35 + 1 solder x 1.5 + 2 wires marked x 0.75
    # + 1 heat shrink x 0.5 + layout (15 + 2 x 6) + 4 ties x 0.2 = 35.55 min (class 2 multiplier 1.0)
    assert r["labor_minutes"] == pytest.approx(35.55)
    assert r["test_minutes"] == pytest.approx(0.8)  # 2 circuits x 0.4 continuity, no hipot
    labor = sum(l["cost"] for l in r["per_part_lines"] if l["category"] == "labor")
    assert labor == pytest.approx(35.55 / 60 * 55, abs=0.05)
    wire = next(l for l in r["per_part_lines"] if l["category"] == "wire")
    assert wire["cost"] == pytest.approx(36 / 12 * 1.05 * 0.24, abs=0.01)
    contact = next(l for l in r["per_part_lines"] if l["category"] == "contact")
    assert contact["cost"] == pytest.approx(1.0)  # 2 crimp ends at P1 x $0.50
    expected_unit = 0.756 + 18 + 1.0 + (4 * 0.25 + 1 * 0.30 + 4 * 0.06) + 35.55 / 60 * 55 + 0.8 / 60 * 70 + 1.0
    assert r["per_part_cost"] == pytest.approx(expected_unit, abs=0.05)
    lot = 2 * 55 + 1 * 55 + 20 / 60 * 70 + 25 + 45  # formboard, kitting, continuity setup, CoC, freight
    assert r["per_lot_cost"] == pytest.approx(lot, abs=0.02)
    b10 = r["price_breaks"][1]
    assert b10["total_price"] == pytest.approx((r["per_part_cost"] * 10 + r["per_lot_cost"]) * 1.10 * 1.15, abs=0.05)
    # Class 3 multiplies only the assembly labor
    r3 = electrical.price_harness(wires, bom, None, [1], dict(opts, workmanship="class_3"))
    assert r3["labor_minutes"] == pytest.approx(35.55 * 1.35, abs=0.01)
    assert r3["test_minutes"] == pytest.approx(0.8)
    assert len(r["cut_list"]) == 2 and r["pinout"][0]["ref"] == "P1"


def test_harness_placeholder_and_bad_class():
    with pytest.raises(electrical.ElectricalError):
        electrical.price_harness([{"from_ref": "P1", "to_ref": "P2"}], [], None, [1], {"workmanship": "class_9"})
    r = electrical.price_harness([{"from_ref": "P1", "from_pin": "1", "to_ref": "P2", "to_pin": "1", "gauge": "22", "length_in": 10}],
                                 [{"ref": "P1", "kind": "connector"}, {"ref": "P2", "kind": "connector"}], None, [1], {})
    assert any("placeholder prices" in w for w in r["warnings"])
    assert any(l["note"].startswith("PLACEHOLDER") for l in r["per_part_lines"])


# ---------------------------------------------------------------- panel
def _panel_lines():
    return [
        {"device_type": "terminal_block", "qty": 10, "unit_price": 2},
        {"device_type": "relay", "qty": 2, "unit_price": 20},
        {"device_type": "breaker", "qty": 1, "unit_price": 40},
    ]


def test_panel_wire_count_estimate():
    p = pricing.DEFAULT_CONFIG["panel"]
    c = electrical.panel_counts(_panel_lines(), {}, p)
    assert c["wires"] == 13  # (10 x 1 + 2 x 6 + 1 x 4) / 2 = 13
    assert "26 device and terminal landings / 2" in c["wire_basis"]
    c = electrical.panel_counts(_panel_lines(), {"io_points": 32}, p)
    assert c["wires"] == 56  # 32 x 1.5 + (2 x 6 + 1 x 4) / 2
    c = electrical.panel_counts(_panel_lines(), {"wire_count": 40}, p)
    assert c["wires"] == 40 and c["wire_basis"] == "wire count given"


def test_panel_parse_and_price():
    r = electrical.parse_file("panel", FIX / "panel_bom.csv")
    types = [l["device_type"] for l in r["lines"]]
    assert types[:4] == ["enclosure", "back_panel", "din_rail", "wire_duct"] and types[-1] == "other"
    rail = r["lines"][2]
    assert rail["length_m"] == 0.5 and rail["qty"] == 2
    assert any("did not match a device type" in w for w in r["warnings"])
    q = electrical.price_panel(r["lines"], None, [1, 3], {"ul508a": True})
    assert any("UL 508A" in w and "enrolled" in w for w in q["warnings"])
    assert any(l["category"] == "ul508a" for l in q["per_lot_lines"])
    rail_line = next(l for l in q["per_part_lines"] if "DR-35" in l["item"])
    assert rail_line["cost"] == pytest.approx(1.0 * 8.0)  # 2 x 0.5 m at the placeholder $8/m
    assert q["price_breaks"][0]["unit_price"] > q["price_breaks"][1]["unit_price"]
    assert any(a.startswith("Wire count") for a in q["assumptions"])


# ---------------------------------------------------------------- labels
def test_label_pricing_by_hand():
    line = {"item": "1", "type": "engraved_laminate", "width_in": 2, "height_in": 3, "text": "PUMP 1", "holes": 2, "qty": 3}
    r = electrical.price_labels([line], None, [100], {})
    d = r["line_detail"][0]
    # material max(6 sq in x 0.06, 0.75) = 0.75; 5 chars x 1.5 s at $85/h; (30 s handling + 2 holes x 0.5 min) at $45/h
    each = 0.75 + 7.5 / 3600 * 85 + 1.5 / 60 * 45
    assert d["chars"] == 5 and d["cost_each"] == pytest.approx(each, abs=0.001)
    assert r["per_part_cost"] == pytest.approx(each * 3 + 1.0, abs=0.03)  # + commercial packaging per set
    assert r["per_lot_cost"] == pytest.approx(15 / 60 * 45 + 25 + 45)
    b = r["price_breaks"][0]
    assert b["total_price"] == pytest.approx((r["per_part_cost"] * 100 + r["per_lot_cost"]) * 1.1 * 1.15, abs=0.05)
    small = electrical.price_labels([line], None, [1], {"freight_per_lot": 0})
    assert small["price_breaks"][0]["total_price"] == 75.0  # labels minimum lot charge


def test_label_csv_and_iuid():
    r = electrical.parse_file("labels", FIX / "labels.csv")
    l1, l2, l3 = r["lines"]
    assert (l1["type"], l1["width_in"], l1["height_in"], l1["holes"], l1["qty"]) == ("engraved_laminate", 1.0, 3.0, 2, 2)
    assert l2["type"] == "stainless_steel" and l2["iuid"]
    assert l3["type"] == "printed_polyester" and l3["adhesive"] and l3["width_in"] == pytest.approx(50 / 25.4, abs=0.001)
    assert any("DFARS 252.211-7003" in w and "IUID Registry" in w for w in r["warnings"])
    q = electrical.price_labels(r["lines"], None, [1, 10], {})
    assert any(l["category"] == "iuid" for l in q["per_part_lines"])
    assert any(l["category"] == "iuid" for l in q["per_lot_lines"])


# ---------------------------------------------------------------- API: parse, save, restore, sheets, config
def test_api_save_restore_each_kind():
    harness = electrical.parse_file("harness", FIX / "harness_drawing.pdf")
    bodies = {
        "harness": {"wires": harness["wires"], "bom": harness["bom"], "options": {"workmanship": "class_3"}},
        "panel": {"lines": _panel_lines(), "options": {"io_points": 32}},
        "labels": {"lines": [{"item": "1", "type": "anodized_aluminum", "width_in": 2, "height_in": 4, "text": "SERIAL NO", "qty": 1}]},
    }
    with TestClient(app) as c:
        cat = c.get("/api/electrical/catalog").json()
        assert cat["harness"]["contact_sizes"]["22D"] == [22, 28] and "enclosure" in cat["panel"]["device_types"]
        for kind, body in bodies.items():
            body = dict(body, quantities=[1, 5], name=f"Test {kind}", part_number="PN-1")
            est = c.post(f"/api/electrical/{kind}/quote", json=body)
            assert est.status_code == 200, est.text
            row = est.json()["price_breaks"][1]
            saved = c.post(f"/api/electrical/{kind}/save", json=dict(body, status="draft", quoted_quantity=5, quoted_unit_price=row["unit_price"]))
            assert saved.status_code == 200, saved.text
            s = saved.json()
            assert s["kind"] == kind and s["quoted_quantity"] == 5 and s["result"]["price_breaks"][1]["unit_price"] == pytest.approx(row["unit_price"])
            back = c.get(f"/api/electrical/{kind}/quotes/{s['id']}").json()
            assert back["spec"]["kind"] == kind and back["spec"]["name"] == f"Test {kind}"
            if kind == "harness":
                assert len(back["spec"]["wires"]) == 6 and back["spec"]["options"]["workmanship"] == "class_3"
            else:
                assert len(back["spec"]["lines"]) == len(body["lines"])
            # update in place
            upd = c.post(f"/api/electrical/{kind}/save", json=dict(body, quote_id=s["id"], status="submitted"))
            assert upd.json()["id"] == s["id"] and upd.json()["status"] == "submitted"
            # the wrong tab refuses it
            other = "panel" if kind != "panel" else "labels"
            assert c.get(f"/api/electrical/{other}/quotes/{s['id']}").status_code == 400
            x = c.get(f"/api/electrical/{kind}/sheet.xlsx", params={"quote_id": s["id"]})
            assert x.status_code == 200 and x.content[:2] == b"PK"
        assert c.post("/api/electrical/bogus/quote", json={}).status_code == 404
        listed = {q["kind"] for q in c.get("/api/pricing/quotes").json()}
        assert {"harness", "panel", "labels"} <= listed


def test_api_parse_upload_and_harness_sheet():
    with TestClient(app) as c:
        with open(FIX / "harness_wires.csv", "rb") as f:
            r = c.post("/api/electrical/harness/parse", files={"file": ("wires.csv", f, "text/csv")})
        assert r.status_code == 200 and len(r.json()["wires"]) == 6
        assert c.post("/api/electrical/harness/parse", files={"file": ("x.doc", b"abc", "application/msword")}).status_code == 400
        h = electrical.parse_file("harness", FIX / "harness_drawing.pdf")
        x = c.post("/api/electrical/harness/sheet.xlsx", json={"wires": h["wires"], "bom": h["bom"], "builds": 3, "title": "HX-100"})
        assert x.status_code == 200
        wb = load_workbook(io.BytesIO(x.content))
        assert wb.sheetnames == ["Cut list", "Pin-out", "BOM"]
        cut = list(wb["Cut list"].iter_rows(min_row=4, values_only=True))
        assert cut[0][0] == "W1" and cut[0][9] == 3
        pins = list(wb["Pin-out"].iter_rows(min_row=4, values_only=True))
        assert ("J2", "A", "W1") == pins[0][:3]
        v = c.post("/api/electrical/harness/validate", json={"wires": h["wires"], "bom": h["bom"]}).json()
        assert any("P9" in w for w in v["warnings"])


def test_rates_saved_in_shop_config():
    with TestClient(app) as c:
        r = c.put("/api/pricing/config", json={"changes": {"harness": {"labor_rate": 60.0, "workmanship_multiplier": {"class_3": 1.5}}}})
        assert r.status_code == 200, r.text
        cat = c.get("/api/electrical/catalog").json()
        assert cat["harness"]["config"]["labor_rate"] == 60.0 and cat["harness"]["config"]["workmanship_multiplier"]["class_3"] == 1.5
        assert c.put("/api/pricing/config", json={"changes": {"panel": {"labor_rate": -1}}}).status_code == 400
        c.post("/api/pricing/config/reset")
