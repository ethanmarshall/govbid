"""T-slot extrusion build quotes: parsing, units, nesting, price math, save/restore, sheets, AI refusal."""
import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import extrusion, pricing
from app.extrusion_catalog import catalog, default_config
from app.main import app

FIX = Path(__file__).parent / "fixtures" / "extrusion"


def _by(lines, **kw):
    return [l for l in lines if all(l.get(k) == v for k, v in kw.items())]


# ---------------------------------------------------------------- units and callouts
@pytest.mark.parametrize("text,unit,expected", [
    ("24.000", "in", 24.0),
    ('24"', "in", 24.0),
    ("24 1/2 IN", "in", 24.5),
    ("610 mm", "in", 610 / 25.4),
    ("610", "mm", 610 / 25.4),
    ("2 ft", "in", 24.0),
    ("2'-6\"", "in", 30.0),
    ("1.5 m", "in", 1500 / 25.4),
])
def test_parse_length_units(text, unit, expected):
    assert extrusion.parse_length(text, unit) == pytest.approx(expected, abs=1e-3)


def test_machining_callouts():
    ops = {o["op"]: o["count"] for o in extrusion.parse_machining("END TAP BOTH ENDS, 2X ACCESS HOLE @ 7 1/4 IN, CBORE")}
    assert ops == {"end_tap": 2, "access_hole": 2, "counterbore": 1}
    ops = {o["op"]: o["count"] for o in extrusion.parse_machining("45° MITER BOTH ENDS; (4) DRILL THRU & CBORE")}
    assert ops == {"miter_cut": 2, "drill_thru_cbore": 4}


def test_profile_and_panel_matching():
    cat = catalog(None)
    assert extrusion.match_profile("1530-LS PROFILE", cat)[0] == "1530-LS"
    assert extrusion.match_profile("1515 LS", cat)[0] == "1515-LS"
    assert extrusion.match_profile("40-4040", cat)[0] == "40-4040"
    assert extrusion.match_profile("30-3030 EXTRUSION", cat)[0] == "30-3030"
    assert extrusion.match_profile("2020 PROFILE", cat)[0] == "2020"
    assert extrusion.match_profile("2020 PROFILE", cat, metric_doc=True)[0] == "20-2020"
    line = extrusion.make_line(qty=2, desc="PANEL, POLYCARB 1/4 x 24 x 36", cat=cat)
    assert line["kind"] == "panel" and line["catalog_id"] == "PANEL-PC-250"
    assert (line["width_in"], line["height_in"]) == (24.0, 36.0)


# ---------------------------------------------------------------- parsing files
def test_parse_drawing_pdf_bom_and_cut_list():
    r = extrusion.parse_file(FIX / "frame_bom.pdf")
    lines = r["lines"]
    prof = _by(lines, kind="profile")
    # The cut list replaces the BOM's single 1515 row (8 pieces, no length)
    assert sorted((l["catalog_id"], l["length_in"], l["qty"]) for l in prof) == [("1515", 24.0, 4), ("1515", 36.0, 4)]
    a = next(l for l in prof if l["length_in"] == 24.0)
    assert a["machining"] == [{"op": "end_tap", "count": 2}]
    b = next(l for l in prof if l["length_in"] == 36.0)
    assert {o["op"]: o["count"] for o in b["machining"]} == {"access_hole": 2, "counterbore": 1}
    hw = {l["catalog_id"]: l["qty"] for l in _by(lines, kind="hardware")}
    assert hw == {"TNUT-DROP-15": 32, "BRKT-4H-15": 8, "ANCHOR-15": 4, "FOOT-15": 4}
    panel = _by(lines, kind="panel")[0]
    assert panel["catalog_id"] == "PANEL-PC-250" and panel["qty"] == 2 and panel["width_in"] == 24.0
    un = r["unmatched"]
    assert len(un) == 1 and un[0]["part_number"] == "XYZ-999"
    assert r["drawing"]["title"] == "TEST STAND FRAME"
    assert not any("ALL HARDWARE" in (l["description"] or "") for l in lines)  # notes are not BOM rows


def test_parse_metric_drawing_mm_lengths():
    r = extrusion.parse_file(FIX / "metric_cart.pdf")
    prof = {l["catalog_id"]: l for l in _by(r["lines"], kind="profile")}
    assert prof["40-4040"]["length_in"] == pytest.approx(610 / 25.4, abs=1e-3)
    assert prof["40-4040"]["qty"] == 4
    assert prof["40-4040"]["machining"] == [{"op": "miter_cut", "count": 2}]
    assert prof["40-4080"]["length_in"] == pytest.approx(1000 / 25.4, abs=1e-3)
    hw = {l["catalog_id"] for l in _by(r["lines"], kind="hardware")}
    assert hw == {"BRKT-GUSSET-ML", "CASTER-LOCK-ML"}
    assert r["metric_doc"] is True


def test_parse_csv_and_xlsx():
    for name in ("bom.csv", "bom.xlsx"):
        r = extrusion.parse_bom_table(FIX / name)
        lines = r["lines"]
        assert _by(lines, catalog_id="1010")[0]["length_in"] == 18.0
        assert _by(lines, catalog_id="1010")[0]["machining"] == [{"op": "end_tap", "count": 2}]
        assert _by(lines, catalog_id="1020")[0]["length_in"] == 24.0  # "2 ft"
        assert _by(lines, kind="hardware", catalog_id="TNUT-SLIDE-10")[0]["qty"] == 16
        assert _by(lines, kind="hardware", catalog_id="BRKT-4H-10")[0]["qty"] == 8  # series from the Notes column
        assert _by(lines, kind="panel")[0]["catalog_id"] == "PANEL-ACR-125"
        assert len(r["unmatched"]) == 1


def test_parse_csv_mm_header():
    r = extrusion.parse_bom_table(FIX / "bom_metric.csv")
    lens = sorted(l["length_in"] for l in r["lines"])
    assert lens == pytest.approx([250 / 25.4, 500 / 25.4], abs=1e-3)


# ---------------------------------------------------------------- nesting
def test_nesting_first_fit_decreasing():
    n = extrusion.nest_ffd([20, 60, 30, 50, 40], 100)
    assert n["sticks"] == 2 and n["waste_in"] == 0
    assert [s["pieces"] for s in n["stick_detail"]] == [[60, 40], [50, 30, 20]]
    # With a 1/8 in kerf the 40 no longer fits after the 60
    n = extrusion.nest_ffd([20, 60, 30, 50, 40], 100, kerf=0.125)
    assert n["sticks"] == 3
    assert n["bought_in"] == 300 and n["used_in"] == 200 and n["waste_pct"] == pytest.approx(33.3)
    n = extrusion.nest_ffd([250, 10], 242, end_trim=0.5)
    assert n["oversize"] == [250] and n["sticks"] == 1


# ---------------------------------------------------------------- price math
def _small_config(mode="cut_to_length"):
    ext = default_config()
    ext.update(pricing_mode=mode, cut_charge=2.0, kerf_in=0.0, end_trim_in=0.0, inhouse_cut_minutes=3.0,
               stock_length_in={"fractional": 50.0, "metric": 50.0}, inspection_minutes_per_build=0.0,
               kitting_hours_per_lot=0.0, supplier_order_charge=0.0, crate_per_build=0.0, supplier_lead_days=0, builds_per_day=1,
               assembly_minutes={"per_build": 0.0, "per_joint": 6.0, "per_fastener": 0.0, "per_panel": 0.0, "per_accessory": 0.0})
    ext["machining"]["end_tap"] = 3.0
    ext["prices"]["profiles"]["1515"] = 1.0
    ext["prices"]["hardware"]["BRKT-4H-15"] = 5.0
    return {"extrusion": ext, "rates": {"assembly": 60.0, "fabrication": 40.0}, "packaging": {"commercial": {"per_part": 1.0, "per_lot": 0.0}},
            "default_freight_per_lot": 10.0, "ga_rate": 0.1, "profit_rate": 0.2, "min_lot_charge": 0.0,
            "lead_time": {"base_days": 5, "first_article_days": 10, "parts_per_day": 25}}


SMALL = [
    {"kind": "profile", "catalog_id": "1515", "qty": 2, "length_in": 10, "machining": [{"op": "end_tap", "count": 2}]},
    {"kind": "hardware", "catalog_id": "BRKT-4H-15", "qty": 4},
]


def test_price_math_by_hand_cut_to_length():
    r = extrusion.price_build(SMALL, _small_config(), [1, 2], {"certificate_of_conformance": False})
    costs = {l["category"]: l["cost"] for l in r["per_part_lines"]}
    # material 2 x 10 in x $1 = 20; cuts 2 x $2 = 4; end taps 4 x $3 = 12; brackets 4 x $5 = 20
    # assembly 4 joints x 6 min = 24 min at $60/h = 24; packaging 1  => 81 per build
    assert costs == {"material": 20.0, "cutting": 4.0, "machining": 12.0, "hardware": 20.0, "assembly": 24.0, "inspection": 0.0, "packaging": 1.0}
    assert r["per_part_cost"] == 81.0
    assert r["per_lot_cost"] == 10.0  # freight only
    b1, b2 = r["price_breaks"]
    assert b1["total_cost"] == 91.0 and b1["unit_price"] == pytest.approx(91 * 1.1 * 1.2, abs=0.01)
    assert b2["total_cost"] == 172.0 and b2["unit_price"] == pytest.approx(172 * 1.1 * 1.2 / 2, abs=0.01)
    assert b1["lead_time_days"] == 6 and b2["lead_time_days"] == 7
    assert r["counts"]["joints"] == 4


def test_price_math_stock_mode_nests_the_lot():
    r = extrusion.price_build(SMALL, _small_config("stock"), [1, 3], {"certificate_of_conformance": False})
    # q=1: two 10 in pieces fit one 50 in stick ($50). q=3: six pieces = 60 in -> 2 sticks ($100)
    assert r["nesting_by_quantity"]["1"]["1515"]["sticks"] == 1
    assert r["nesting_by_quantity"]["3"]["1515"]["sticks"] == 2
    cut = next(l for l in r["per_part_lines"] if l["category"] == "cutting")
    assert cut["cost"] == pytest.approx(2 * 3 / 60 * 40, abs=0.01)  # 2 pieces x 3 min at $40/h
    non_mat = 81.0 - 20.0 - 4.0 + 4.0  # same as cut mode minus material, cut charge replaced by saw time (also $4)
    b1, b3 = r["price_breaks"]
    assert b1["total_cost"] == pytest.approx(non_mat + 50 + 10, abs=0.01)
    assert b3["total_cost"] == pytest.approx(non_mat * 3 + 100 + 10, abs=0.01)


def test_unmatched_lines_and_custom_price():
    lines = SMALL + [{"kind": "unmatched", "description": "wire tray", "qty": 1},
                     {"kind": "custom", "description": "label kit", "qty": 2, "unit_price": 7.5}]
    r = extrusion.price_build(lines, _small_config(), [1])
    assert any("not priced" in w for w in r["warnings"])
    assert any(l["item"].startswith("label kit") and l["cost"] == 15.0 for l in r["per_part_lines"])


# ---------------------------------------------------------------- API: quote, save, restore, sheets, config
def test_api_quote_save_restore_and_sheets():
    with TestClient(app) as c:
        cat = c.get("/api/extrusion/catalog").json()
        assert any(p["id"] == "1515" for p in cat["profiles"]) and "placeholder" in cat["note"].lower()
        with open(FIX / "frame_bom.pdf", "rb") as f:
            parsed = c.post("/api/extrusion/parse", files={"file": ("frame_bom.pdf", f, "application/pdf")}).json()
        lines = parsed["lines"]
        q = c.post("/api/extrusion/quote", json={"lines": lines, "quantities": [1, 5]}).json()
        assert [b["quantity"] for b in q["price_breaks"]] == [1, 5]
        assert q["price_breaks"][0]["unit_price"] > q["price_breaks"][1]["unit_price"]

        saved = c.post("/api/extrusion/save", json={"lines": lines, "quantities": [1, 5], "name": "Test stand frame",
                                                    "part_number": "TSF-1001", "status": "draft", "notes": "from drawing",
                                                    "quoted_quantity": 5}).json()
        qid = saved["id"]
        assert saved["quoted_unit_price"] == q["price_breaks"][1]["unit_price"]
        back = c.get(f"/api/extrusion/quotes/{qid}").json()
        assert back["spec"]["kind"] == "extrusion_build"
        assert len(back["spec"]["lines"]) == len(lines)
        assert back["result"]["price_breaks"] == q["price_breaks"]
        # The generic quote endpoints keep working for status changes
        upd = c.put(f"/api/pricing/quotes/{qid}", json={"spec": back["spec"], "status": "submitted"}).json()
        assert upd["status"] == "submitted" and upd["price_breaks"] == q["price_breaks"]
        assert any(r["id"] == qid for r in c.get("/api/pricing/quotes").json())
        # Update in place
        again = c.post("/api/extrusion/save", json={"lines": lines[:3], "quantities": [2], "name": "Test stand frame", "quote_id": qid}).json()
        assert again["id"] == qid and [b["quantity"] for b in again["price_breaks"]] == [2]

        res = c.post("/api/extrusion/cutlist.xlsx", json={"lines": lines, "quantities": [1, 5], "builds": 5, "title": "Test stand"})
        assert res.status_code == 200
        wb = load_workbook(io.BytesIO(res.content))
        ws = wb["Cut list"]
        rows = [r for r in ws.iter_rows(min_row=4, values_only=True) if r[0]]
        assert sorted((r[1], r[2], r[4], r[5]) for r in rows) == [("1515", 24.0, 4, 20), ("1515", 36.0, 4, 20)]
        a = next(r for r in rows if r[2] == 24.0)
        assert a[3] == pytest.approx(609.6) and "End tap" in a[6]
        nest = [r for r in wb["Stick nesting"].iter_rows(min_row=4, values_only=True) if r[0]]
        assert len(nest) == 5  # 40 pieces, 240 in per build: one 242 in stick per build

        res = c.get(f"/api/extrusion/cutlist.xlsx?quote_id={qid}&builds=1")
        assert res.status_code == 200
        res = c.post("/api/extrusion/purchase.xlsx", json={"lines": lines, "builds": 2})
        ws = load_workbook(io.BytesIO(res.content))["Purchase list"]
        vals = [r for r in ws.iter_rows(min_row=4, values_only=True)]
        assert any(r[1] == "HW-101" and r[3] == 64 for r in vals)
        assert c.get(f"/api/extrusion/purchase.xlsx?quote_id={qid}").status_code == 200

        # Non-extrusion quotes are refused by the restore endpoint
        other = c.post("/api/pricing/quotes", json={"spec": pricing.EXAMPLE_SPEC}).json()
        assert c.get(f"/api/extrusion/quotes/{other['id']}").status_code == 400


def test_config_rates_round_trip():
    with TestClient(app) as c:
        r = c.put("/api/pricing/config", json={"changes": {"extrusion": {"pricing_mode": "stock", "prices": {"profiles": {"1515": 0.7}}}}})
        assert r.status_code == 200 and r.json()["config"]["extrusion"]["pricing_mode"] == "stock"
        cat = c.get("/api/extrusion/catalog").json()
        assert next(p for p in cat["profiles"] if p["id"] == "1515")["price_per_in"] == 0.7
        assert c.put("/api/pricing/config", json={"changes": {"extrusion": {"pricing_mode": "bogus"}}}).status_code == 400
        assert c.put("/api/pricing/config", json={"changes": {"extrusion": {"cut_charge": -1}}}).status_code == 400
        c.post("/api/pricing/config/reset")


# ---------------------------------------------------------------- AI read safety
def test_ai_refused_for_export_controlled_drawing(monkeypatch):
    called = []
    monkeypatch.setattr(extrusion, "ANTHROPIC_API_KEY", "test")
    monkeypatch.setattr(extrusion, "_claude_read_bom", lambda data: called.append(1) or {
        "units": "in", "lines": [{"item": "1", "qty": 4, "part_number": "1515", "description": "PROFILE", "length": "30", "machining": "END TAP BOTH ENDS"}]})
    with TestClient(app) as c:
        def post(**form):
            with open(FIX / "itar_frame.pdf", "rb") as f:
                return c.post("/api/extrusion/parse", files={"file": ("itar_frame.pdf", f, "application/pdf")}, data=form)

        r = post(use_ai="true")
        assert r.status_code == 400 and "export-controlled" in r.json()["detail"]
        assert called == []
        r = post()  # no AI requested: the text read just finds no table
        assert r.status_code == 200 and r.json()["lines"] == [] and r.json()["drawing"]["export_controlled"]
        r = post(use_ai="true", force="true")
        assert r.status_code == 200 and called == [1]
        line = r.json()["lines"][0]
        assert line["catalog_id"] == "1515" and line["length_in"] == 30.0 and line["source"] == "ai"
