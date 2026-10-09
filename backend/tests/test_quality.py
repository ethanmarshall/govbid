import io
from datetime import date, timedelta

from fastapi.testclient import TestClient

from app.main import app


def _d(days: int) -> str:
    return (date.today() + timedelta(days=days)).isoformat()


def test_next_due_and_cal_status():
    from app.models_quality import Instrument
    from app.quality_api import cal_status, next_due

    assert next_due("2026-01-15", 365) == "2027-01-15"
    assert next_due("2024-02-29", 180) == "2024-08-27"
    assert next_due("", 365) == "" and next_due("2026-01-01", None) == ""
    today = date(2026, 6, 1)
    assert cal_status(Instrument(status="active", last_cal_date="2025-05-01", interval_days=365), today) == "overdue"
    assert cal_status(Instrument(status="active", last_cal_date="2025-06-15", interval_days=365), today) == "due_soon"
    assert cal_status(Instrument(status="active", last_cal_date="2026-05-01", interval_days=365), today) == "ok"
    assert cal_status(Instrument(status="active", last_cal_date="", interval_days=365), today) == "no_record"
    assert cal_status(Instrument(status="out_of_service", last_cal_date="2020-01-01", interval_days=30), today) == "out_of_service"


def test_scorecard_math():
    from app.models_jobs import JobPurchase
    from app.quality_api import scorecard

    pos = [
        JobPurchase(ordered_date="2026-01-02", promised_date="2026-01-10", received_date="2026-01-09", quantity=10, unit_price=5,
                    qty_accepted=10, qty_rejected=0, certs_received=True),
        JobPurchase(ordered_date="2026-02-02", promised_date="2026-02-10", received_date="2026-02-10", quantity=10, unit_price=5,
                    qty_accepted=8, qty_rejected=2, certs_received=True),
        JobPurchase(ordered_date="2026-03-02", promised_date="2026-03-10", received_date="2026-03-15", quantity=20, unit_price=1,
                    qty_accepted=20, qty_rejected=0, certs_received=False),
        JobPurchase(ordered_date="2026-04-02", promised_date="2026-04-10", received_date="", quantity=5, unit_price=2),
        JobPurchase(ordered_date="2026-04-05", promised_date="", received_date="2026-04-20", quantity=0, qty_accepted=0, qty_rejected=0),
    ]
    s = scorecard(pos, today=date(2026, 5, 1))
    assert s["orders"] == 5 and s["received"] == 4 and s["open"] == 1 and s["late_open"] == 1
    assert s["on_time_measured"] == 3 and s["on_time"] == 2 and s["on_time_pct"] == 66.7
    assert s["qty_accepted"] == 38 and s["qty_rejected"] == 2 and s["acceptance_pct"] == 95.0
    assert s["certs_missing"] == 2 and s["last_order"] == "2026-04-05" and s["spend"] == 130.0
    empty = scorecard([])
    assert empty["on_time_pct"] is None and empty["acceptance_pct"] is None and empty["last_order"] == ""


def test_ncr_numbering_and_car_flow():
    with TestClient(app) as c:
        jid = c.post("/api/jobs", json={"title": "NCR job"}).json()["id"]
        a = c.post("/api/quality/ncrs", json={"part": "Bracket", "quantity": 3, "description": "Hole position out of tolerance",
                                              "job_id": jid, "opened_date": "2031-03-01"}).json()
        b = c.post("/api/quality/ncrs", json={"part": "Bracket", "opened_date": "2031-04-01"}).json()
        x = c.post("/api/quality/ncrs", json={"part": "Panel", "opened_date": "2032-01-05"}).json()
        assert a["number"] == "NCR-2031-001" and b["number"] == "NCR-2031-002" and x["number"] == "NCR-2032-001"
        assert a["job_title"] == "NCR job" and a["status"] == "open"
        c.delete(f"/api/quality/ncrs/{b['id']}")
        assert c.post("/api/quality/ncrs", json={"opened_date": "2031-05-01"}).json()["number"] == "NCR-2031-002"
        assert c.post("/api/quality/ncrs", json={"disposition": "burn"}).status_code == 422
        assert c.post("/api/quality/ncrs", json={"job_id": 999999}).status_code == 404

        u = c.put(f"/api/quality/ncrs/{a['id']}", json={"disposition": "rework", "root_cause": "Wrong work offset"}).json()
        assert u["status"] == "dispositioned"
        car = c.post("/api/quality/cars", json={"ncr_id": a["id"], "due_date": _d(-2), "action": "Probe the part before each run",
                                                "opened_date": "2031-03-02"}).json()
        assert car["number"] == "CAR-2031-001" and car["ncr_number"] == "NCR-2031-001"
        assert car["problem"].startswith("NCR-2031-001") and car["root_cause"] == "Wrong work offset" and car["overdue"] is True
        assert c.get(f"/api/quality/ncrs/{a['id']}").json()["cars"][0]["number"] == "CAR-2031-001"

        from app.db import SessionLocal
        from app.quality_api import calendar_items, dashboard_items
        with SessionLocal() as db:
            assert any("CAR-2031-001" in s for s in dashboard_items(db))
            assert any(e["uid"] == f"car-{car['id']}@govbid" for e in calendar_items(db))

        assert c.put(f"/api/quality/cars/{car['id']}", json={"effective": False, "status": "closed"}).status_code == 422
        closed = c.put(f"/api/quality/cars/{car['id']}", json={"effective": True, "effectiveness_check": "Next 3 lots clean",
                                                               "status": "closed"}).json()
        assert closed["closed_date"] == date.today().isoformat() and closed["overdue"] is False
        n = c.put(f"/api/quality/ncrs/{a['id']}", json={"status": "closed"}).json()
        assert n["closed_date"] == date.today().isoformat()
        assert any(r["id"] == a["id"] for r in c.get("/api/quality/ncrs", params={"job_id": jid}).json())
        assert not any(r["id"] == a["id"] for r in c.get("/api/quality/ncrs", params={"status": "open"}).json())

        # deleting the job keeps the NCR but unlinks it
        c.delete(f"/api/jobs/{jid}")
        assert c.get(f"/api/quality/ncrs/{a['id']}").json()["job_id"] is None
        c.delete(f"/api/quality/ncrs/{a['id']}")
        assert next(x for x in c.get("/api/quality/cars").json() if x["id"] == car["id"])["ncr_id"] is None
        c.delete(f"/api/quality/cars/{car['id']}")


def test_calibration_log():
    with TestClient(app) as c:
        assert c.post("/api/quality/instruments", json={"name": "Bad", "interval_days": 0}).status_code == 422
        i = c.post("/api/quality/instruments", json={"name": "6 in calipers", "kind": "Calipers", "asset_id": "CAL-01",
                                                     "serial": "A123", "interval_days": 365, "last_cal_date": _d(-400)}).json()
        assert i["cal_status"] == "overdue" and i["overdue"] is True and i["next_due"] == _d(-35)
        from app.db import SessionLocal
        from app.quality_api import calendar_items, dashboard_items
        with SessionLocal() as db:
            assert any("CAL-01" in s and "overdue" in s.lower() for s in dashboard_items(db))
        r = c.post(f"/api/quality/instruments/{i['id']}/calibrations",
                   data={"cal_date": _d(0), "source": "Acme Cal Lab", "result": "pass"},
                   files={"file": ("cert.pdf", b"%PDF cal", "application/pdf")})
        assert r.status_code == 200, r.text
        i2 = r.json()
        assert i2["last_cal_date"] == _d(0) and i2["next_due"] == _d(365) and i2["cal_status"] == "ok" and i2["cal_source"] == "Acme Cal Lab"
        assert c.get(i2["history"][0]["url"]).content == b"%PDF cal"
        with SessionLocal() as db:
            assert any(e["uid"].startswith(f"cal-{i['id']}-") and e["date"] == date.today() + timedelta(days=365) for e in calendar_items(db))
        # an older calibration entered later does not move the date back
        i3 = c.post(f"/api/quality/instruments/{i['id']}/calibrations", data={"cal_date": _d(-400), "result": "pass"}).json()
        assert i3["last_cal_date"] == _d(0) and len(i3["history"]) == 2
        # a failed calibration takes it out of service
        i4 = c.post(f"/api/quality/instruments/{i['id']}/calibrations", data={"cal_date": _d(0), "result": "fail"}).json()
        assert i4["status"] == "out_of_service" and i4["cal_status"] == "out_of_service"
        assert c.post(f"/api/quality/instruments/{i['id']}/calibrations", data={"cal_date": "June 1"}).status_code == 422
        assert c.put(f"/api/quality/instruments/{i['id']}", json={"status": "active", "interval_days": 180}).json()["next_due"] == _d(180)
        assert any(x["id"] == i["id"] for x in c.get("/api/quality/instruments").json())
        assert c.delete(f"/api/quality/instruments/{i['id']}").json()["ok"]


def test_quality_documents():
    with TestClient(app) as c:
        docs = c.get("/api/quality/documents").json()
        keys = {d["key"] for d in docs}
        assert {"quality_manual", "receiving_inspection", "nonconforming_material", "calibration", "counterfeit_prevention"} <= keys
        assert c.post("/api/quality/documents/seed").json()["added"] == 0
        qm = next(d for d in docs if d["key"] == "quality_manual")
        full = c.get(f"/api/quality/documents/{qm['id']}").json()
        body = full["content"]
        for topic in ("Contract review", "Purchasing and supplier control", "Receiving inspection", "travelers", "Inspection and test",
                      "nonconforming product", "Corrective action", "Calibration", "Records retention", "Counterfeit", "FOD", "Training"):
            assert topic.lower() in body.lower(), topic
        assert "not certified" in body
        for d in docs:
            assert "—" not in c.get(f"/api/quality/documents/{d['id']}").json()["content"]

        r = c.put(f"/api/quality/documents/{qm['id']}", json={"content": body + "\nLocal change.\n", "new_revision": True,
                                                             "change_note": "Added note"}).json()
        assert r["version"] == "B" and r["effective_date"] == date.today().isoformat()
        assert r["history"][-1]["version"] == "A" and r["history"][-1]["note"] == "Added note" and r["history"][-1]["content"] == body
        r = c.put(f"/api/quality/documents/{qm['id']}", json={"title": "Quality Manual"}).json()
        assert r["version"] == "B" and len(r["history"]) == 1

        f = c.get(f"/api/quality/documents/{qm['id']}/docx")
        assert f.status_code == 200 and f.content[:2] == b"PK"
        import docx
        d = docx.Document(io.BytesIO(f.content))
        text = "\n".join(p.text for p in d.paragraphs)
        assert "Quality Manual" in text and "Revision B" in text and "Local change." in text
        assert c.delete(f"/api/quality/documents/{qm['id']}").status_code == 422

        mine = c.post("/api/quality/documents", json={"title": "Wire harness workmanship", "doc_number": "QP-010",
                                                      "content": "# Workmanship\n\n- Follow **IPC/WHMA-A-620** class in the contract\n1. Inspect"}).json()
        assert mine["version"] == "A"
        assert c.get(f"/api/quality/documents/{mine['id']}/docx").status_code == 200
        assert c.delete(f"/api/quality/documents/{mine['id']}").json()["ok"]


def test_next_revision():
    from app.quality_api import next_revision
    assert next_revision("A") == "B" and next_revision("Z") == "AA" and next_revision("AZ") == "BA"
    assert next_revision("1") == "2" and next_revision("1.9") == "1.10" and next_revision("") == "A"


def test_suppliers_scorecard_certs_and_asl():
    with TestClient(app) as c:
        org = c.post("/api/crm/organizations", json={"name": "Quality Test Plating", "kind": "vendor", "cage": "1QQ11"}).json()
        oid = org["id"]
        jid = c.post("/api/jobs", json={"title": "Supplier job"}).json()["id"]
        c.post(f"/api/jobs/{jid}/purchases", json={"organization_id": oid, "po_number": "P1", "quantity": 10, "ordered_date": "2026-01-01",
                                                  "promised_date": "2026-01-10", "received_date": "2026-01-08", "qty_rejected": 2})
        c.post(f"/api/jobs/{jid}/purchases", json={"organization_id": oid, "po_number": "P2", "quantity": 10, "ordered_date": "2026-02-01",
                                                  "promised_date": "2026-02-10", "received_date": "2026-02-12"})
        # an unlinked PO with the same vendor name counts too
        c.post(f"/api/jobs/{jid}/purchases", json={"vendor_name": "quality test plating", "po_number": "P3", "ordered_date": "2026-03-01"})

        s = next(x for x in c.get("/api/quality/suppliers").json() if x["organization_id"] == oid)
        assert s["approval"]["status"] == "pending"
        sc = s["scorecard"]
        assert sc["orders"] == 3 and sc["on_time_pct"] == 50.0 and sc["acceptance_pct"] == 90.0 and sc["last_order"] == "2026-03-01"

        assert c.put(f"/api/quality/suppliers/{oid}", json={"status": "great"}).status_code == 422
        assert c.put(f"/api/quality/suppliers/{oid}", json={"basis": ["vibes"]}).status_code == 422
        s = c.put(f"/api/quality/suppliers/{oid}", json={"status": "approved", "basis": ["iso_cert", "past_performance"],
                                                         "scope": "Type II anodize"}).json()
        assert s["approval"]["status"] == "approved" and s["approval"]["approval_date"] == date.today().isoformat()

        r = c.post(f"/api/quality/suppliers/{oid}/certs", data={"cert_type": "ISO 9001", "number": "C-1", "expiration_date": _d(-1)},
                   files={"file": ("iso.pdf", b"%PDF iso", "application/pdf")})
        assert r.status_code == 200, r.text
        cert = r.json()
        assert cert["state"] == "expired" and c.get(cert["url"]).content == b"%PDF iso"
        cert2 = c.post(f"/api/quality/suppliers/{oid}/certs", data={"cert_type": "AS9100", "expiration_date": _d(20)}).json()
        assert cert2["state"] == "expiring_soon" and cert2["url"] is None
        assert c.post(f"/api/quality/suppliers/{oid}/certs", data={"cert_type": "X", "expiration_date": "soon"}).status_code == 422

        from app.db import SessionLocal
        from app.quality_api import calendar_items, dashboard_items
        with SessionLocal() as db:
            assert any("Quality Test Plating" in a and "ISO 9001" in a for a in dashboard_items(db))
            assert any(e["uid"] == f"supcert-{cert2['id']}@govbid" for e in calendar_items(db))

        f = c.get("/api/quality/suppliers/asl.xlsx")
        assert f.status_code == 200 and f.content[:2] == b"PK"
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(f.content)).active
        rows = [[cell.value for cell in row] for row in ws.iter_rows()]
        assert rows[0][0] == "Supplier"
        mine = next(r for r in rows if r[0] == "Quality Test Plating")
        assert mine[4] == "approved" and mine[10] == 3 and mine[11] == 50.0 and "EXPIRED" in mine[9]

        c.put(f"/api/quality/suppliers/{oid}", json={"status": "disapproved"})
        rows = [[cell.value for cell in row] for row in load_workbook(io.BytesIO(c.get("/api/quality/suppliers/asl.xlsx").content)).active.iter_rows()]
        assert not any(r[0] == "Quality Test Plating" for r in rows)
        rows = [[cell.value for cell in row] for row in load_workbook(io.BytesIO(c.get("/api/quality/suppliers/asl.xlsx", params={"all": True}).content)).active.iter_rows()]
        assert any(r[0] == "Quality Test Plating" for r in rows)
        with SessionLocal() as db:
            assert not any("Quality Test Plating" in a for a in dashboard_items(db))

        assert c.delete(f"/api/quality/suppliers/{oid}/certs/{cert['id']}").json()["ok"]
        assert c.get(f"/api/quality/suppliers/{oid}").json()["certs"][0]["cert_type"] == "AS9100"
        summ = c.get("/api/quality/summary").json()
        assert {"open_ncrs", "open_cars", "cal_overdue", "suppliers_approved", "alerts"} <= set(summ)
        assert "dispositions" not in c.get("/api/quality/meta").json() and "return_to_vendor" in c.get("/api/quality/meta").json()["ncr_dispositions"]
        c.delete(f"/api/jobs/{jid}")
        c.delete(f"/api/quality/suppliers/{oid}/certs/{cert2['id']}")
        c.delete(f"/api/crm/organizations/{oid}")
