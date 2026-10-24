"""A STEP model sent with a separate BOM: material, finish, make or buy, quantities and parts not in the model."""
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app import portal, portal_bom
from app.main import app

FIX = Path(__file__).parent / "fixtures"

BOM = """Item,Part Number,Description,Qty,Material,Finish,Vendor
1,VP-101,Base plate,1,304 Stainless,Passivate,
2,VP-102,Standoff,4,6061-T6,Black anodize,
3,91290A116,SHCS M6x12,4,,,McMaster-Carr
4,94180A333,Heat-set insert M6,8,,,McMaster-Carr
5,VP-103,Cover,1,Delrin,,
"""


@pytest.fixture
def client():
    portal.QUOTE_LIMIT.clear()
    with TestClient(app) as c:
        assert c.put("/api/portal/settings", json={"enabled": True}).status_code == 200
        yield c
        c.put("/api/portal/settings", json={"enabled": False})


def test_reads_common_bom_layouts():
    from app import pricing

    got = portal_bom.parse("bom.csv", BOM.encode(), pricing.DEFAULT_CONFIG)
    rows = {r["part"]: r for r in got["rows"]}
    assert rows["VP-101"]["material"] == "304 stainless" and rows["VP-101"]["buy"] is False
    assert rows["VP-102"]["finish"] == "anodize (Type II)" and rows["VP-102"]["qty"] == 4
    assert rows["91290A116"]["buy"] is True and rows["VP-103"]["material"] == "delrin (acetal)"
    # SolidWorks style headers, a semicolon CSV, and a make/buy column
    sw = "ITEM NO.;PART NUMBER;DESCRIPTION;QTY.;MAKE/BUY\n1;Bracket;Side bracket;2;Make\n2;6204-2RS;Bearing;2;Buy\n"
    rows = portal_bom.parse("sw.csv", sw.encode())["rows"]
    assert [r["buy"] for r in rows] == [False, True]


def test_tells_board_files_from_a_bom():
    board = "Designator,Value,Footprint,Qty\nR1,10k,0603,1\nC1,1u,0603,1\n"
    pnp = "Designator,Mid X,Mid Y,Rotation,Layer\nR1,10,10,0,Top\n"
    refs = "Part,Value,Qty\nR1,10k,1\nR2,10k,1\nC3,1u,1\nU1,LM317,1\n"
    assert portal_bom.sniff("b.csv", board.encode()) == "pcb"
    assert portal_bom.sniff("p.csv", pnp.encode()) == "pcb"
    assert portal_bom.sniff("r.csv", refs.encode()) == "pcb"
    assert portal_bom.sniff("m.csv", BOM.encode()) == "bom"
    assert portal_bom.sniff("x.csv", b"hello,world\n1,2\n") is None


def test_step_and_bom_quote_together(client, tmp_path):
    bom = tmp_path / "parts list.csv"
    bom.write_text(BOM)
    step = FIX / "cad" / "named_asm.step"
    up = [("files", (p.name, p.read_bytes(), "application/octet-stream")) for p in (step, bom)]
    r = client.post("/api/public/quote", data={"quantity": "2"}, files=up)
    assert r.status_code == 200, r.text
    q = r.json()
    items = {ln["name"]: ln for ln in q["result"]["items"]}
    facts = lambda ln: {f["label"]: f["value"] for f in ln["facts"]}  # noqa: E731

    plate = items["VP-101 base plate"]
    assert plate["material"] == "304 stainless" and "from your BOM" in facts(plate)["Material"]
    assert "passivate" in facts(plate)["Finish"].lower() and "Your BOM" in facts(plate)
    standoff = items["VP-102 standoff"]
    assert standoff["qty_per"] == 4 and standoff["qty"] == 8  # the BOM's 4 per assembly, not the model's 2
    screw = items["SHCS M6x12"]
    assert screw["kind"] == "bought" and screw["qty_per"] == 4 and "91290A116" in facts(screw)["Catalog part"]
    insert = next(ln for ln in q["result"]["items"] if ln["name"].startswith("Heat-set insert"))
    assert insert["kind"] == "bought" and insert["qty"] == 16  # not in the model, still quoted
    cover = items["Cover"]
    assert cover["route"] == "manual" and "not in your model" in cover["note"]
    asm = next(ln for ln in q["result"]["items"] if ln["kind"] == "assembly")
    assert "parts list.csv" in facts(asm)["Your BOM"] and "from your BOM" in facts(asm)["Material"]
    text = json.dumps(q).lower()
    assert "unit_cost" not in text and "internal_reasons" not in text


def test_bom_on_its_own(client, tmp_path):
    bom = tmp_path / "bom.xlsx"
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for row in [["Assembly: test fixture"], [], ["Qty", "Description", "Part No.", "Vendor"],
                [6, "Socket head cap screw 1/4-20 x 1", "91251A540", "McMaster"], [1, "Mounting plate", "MP-1", ""]]:
        ws.append(row)
    wb.save(bom)
    up = [("files", (bom.name, bom.read_bytes(), "application/octet-stream"))]
    q = client.post("/api/public/quote", data={"quantity": "1"}, files=up).json()
    kinds = sorted(ln["kind"] for ln in q["result"]["items"])
    assert kinds == ["bought", "part"]
    assert all(ln["group"] == "Parts from your BOM" for ln in q["result"]["items"])


def test_add_a_bom_after_the_model(client):
    step = FIX / "cad" / "named_asm.step"
    q = client.post("/api/public/quote", data={"quantity": "1"},
                    files=[("files", (step.name, step.read_bytes(), "application/octet-stream"))]).json()
    plate = next(ln for ln in q["result"]["items"] if ln["name"] == "VP-101 base plate")
    assert "from your BOM" not in json.dumps(plate["facts"])
    r = client.post(f"/api/public/quote/{q['ref']}/files", data={"token": q["token"]},
                    files=[("files", ("bom.csv", BOM.encode(), "text/csv"))])
    assert r.status_code == 200, r.text
    plate = next(ln for ln in r.json()["result"]["items"] if ln["name"] == "VP-101 base plate")
    assert plate["material"] == "304 stainless"
    bad = client.post(f"/api/public/quote/{q['ref']}/files", data={"token": "wrong"},
                      files=[("files", ("bom.csv", BOM.encode(), "text/csv"))])
    assert bad.status_code == 400


def _workbook(path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Cut list"
    for r in [["Cart: cut list"], ["Item", "Profile", "Qty", "Cut length (in)"], ["Verticals", "1530", 4, 30.75], ["Door rails", "1515", 8, 13.3125],
              ["Totals", None, 12, None]]:
        ws.append(r)
    ws = wb.create_sheet("Hardware")
    for r in [["Item", "80/20 part no.", "Used on", "Units", "Qty per unit", "Total qty"],
              ["Hinge", "12068", "Doors", 4, 2, 8], ["Fasteners for 12068", "75-3628", "", 8, 4, 32], ["Slide latch", "2053", "", 4, 1, 4],
              ["Caster", "TBD", "", 4, 1, 4], ["Mounting plate screw, M8 x 25 SHCS", "TBD", "", 2, 12, 24]]:
        ws.append(r)
    wb.save(path)


def test_reads_a_hand_made_workbook(tmp_path):
    p = tmp_path / "bom.xlsx"
    _workbook(p)
    rows = portal_bom.parse(p.name, p.read_bytes())["rows"]
    by = {r["label"]: r for r in rows}
    assert set(by) == {"Verticals", "Door rails", "Hinge", "Fasteners for 12068", "Slide latch", "Caster", "Mounting plate screw, M8 x 25 SHCS"}
    assert by["Verticals"]["profile"] == "1530" and by["Hinge"]["qty"] == 8  # the Total qty column, not Qty per unit
    assert by["Caster"]["part"] == ""  # TBD is not a part number


def test_bom_sets_frame_hardware_counts(tmp_path):
    p = tmp_path / "bom.xlsx"
    _workbook(p)
    bom = portal_bom.Bom([{"name": p.name, **portal_bom.parse(p.name, p.read_bytes())}])
    counts, covered, notes = portal_bom.frame_hardware(bom)
    assert counts == {("HINGE", "12068"): 8, ("LATCH", "2053"): 4}
    assert len(covered) == 5 and len(bom.left()) == 0  # profiles, the hinge screws, the caster and the M8 screws: in the frame quote


def test_hardware_counted_by_instance_not_body():
    from app import tslot

    hinges = [{"qty": 16, "body_names": [f"Hinge Subassembly <{i}>" for i in range(1, 9) for _ in range(2)]},
              {"qty": 8, "body_names": [f"Hinge Subassembly <{i}>" for i in range(1, 9)]},
              {"qty": 64, "body_names": [f"Hinge Hardware <{i}>" for i in range(1, 9) for _ in range(8)]}]
    assert tslot.unit_count(hinges) == 8  # 3 bodies per hinge; the hinge screws come with the hinges
    assert tslot.unit_count([{"qty": 4, "body_names": ["Caster"] * 4}]) == 4  # no instance names: one body each
    assert tslot.fastener_base(["Mounting Plate Hardware <3>"], {"length": 1.3, "width": 0.52, "height": 0.52}) == "SCREW-BHSCS"
    assert tslot.fastener_base(["Mounting Plate Hardware <3>"], {"length": 0.63, "width": 0.63, "height": 0.24}) == "TNUT-DROP"


def test_material_and_infill_from_part_names():
    from app import pricing, portal

    c = pricing.DEFAULT_CONFIG
    cases = {"Pedestal PETG": "PETG", "Bearing Cup 6061": "6061-T6 aluminum", "Guard - Stainless": "304 stainless",
             "Spacer Delrin": "delrin (acetal)", "Mount 304 SS": "304 stainless", "Knob ASA": "ASA (UV stable)",
             "Bracket 304": None, "1530 End Cap": None, "Part 1": None, "CirPattern2": None, "PC Board mount": None}
    for name, want in cases.items():
        assert portal_bom.material_from_name(name, c) == want, name
    assert portal_bom.infill_from_name("Cover PLA 80% infill") == 80
    d = portal._named({"material": "6061-T6 aluminum", "material_source": "default", "finishes": [], "finish_source": ""}, "Pedestal PETG 80% infill", c)
    assert d["material"] == "PETG" and d["material_source"] == "name" and d["infill"] == 80
    # a material from the drawing, the BOM or the customer stays
    kept = portal._named({"material": "316 stainless", "material_source": "drawing"}, "Pedestal PETG", c)
    assert kept["material"] == "316 stainless"
